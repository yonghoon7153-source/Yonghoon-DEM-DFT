#!/usr/bin/env python3
r"""10-06 보고 덱 조립기 — 1저자 원본 네 조각 (part1 · part1.5 · part1.75 · part2) 을 한 pptx 로 합치고 손본다.

    python3 docs/report_20261006/build_deck.py \
        --part1 <…part1.pptx> --part15 <…part1.5.pptx> --part175 <…part1.75.zip 또는 .pptx> --part2 <…part2.pptx> \
        --tsv docs/data/ps45_r45_network_webapp_20261005/network_5comp.tsv \
        --out <출력.pptx> [--merge-only]

1. **합치기 (OPC 수준)** — 슬라이드 XML · 관계 · 그림 · OLE (Origin 그래프) · 메모를 그대로 옮긴다.  네 조각은 같은 양식이라
   마스터 · 레이아웃 · 테마 · 메모 마스터가 바이트 동일해야 한다 (다르면 멈춘다).  그림은 내용 해시로 한 번만 담는다.
   `--merge-only` = 손보기 없이 1 · 2 · 3 · 4 · 5 · 6 순서로만 합친 덱 (원본과 장마다 대조하는 용도).
2. **손보기** — 1 장 → 1a (DEM 원리 · 시뮬레이션 조건) · 1b (지배 방정식 · 계산 과정), 6 장 → 6a (접촉 네트워크 개념) ·
   6b (네트워크 해석 결과: 표 결과 칸 · 네이티브 그래프 둘 · 글머리 · 각주).  "10월 계산 예정" 상자 삭제.
   장 번호가 하나씩 밀리므로 본문 · 메모의 "N 장" 참조를 새 번호로 고친다.  원 3 장 (새 4 장) 의 천 단위 콤마를 뺀다 (보고자료 원칙 §6).
   글꼴은 바꾸지 않는다 (새 글은 옆 글의 글꼴을 그대로 쓴다).
값은 모두 `--tsv` (웹앱 케이스 페이지 전사, 1저자 10-05) 에서 읽고, σ_contact-free / σ_full 비만 아래 상수 (같은 웹앱 페이지) 로 둔다.
"""
import argparse
import copy
import csv
import hashlib
import io
import posixpath
import re
import zipfile
from decimal import Decimal, ROUND_HALF_UP

from lxml import etree

# ───────────────────────────── 이름공간 ─────────────────────────────
NS = {
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main',
    'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
    'a14': 'http://schemas.microsoft.com/office/drawing/2010/main',
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
}
P, A, R, MC = (f'{{{NS[k]}}}' for k in ('p', 'a', 'r', 'mc'))
REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
CT_NS = 'http://schemas.openxmlformats.org/package/2006/content-types'
RT = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
RT_SLIDE, RT_NOTES, RT_LAYOUT, RT_NOTESMASTER = RT + 'slide', RT + 'notesSlide', RT + 'slideLayout', RT + 'notesMaster'
CT_SLIDE = 'application/vnd.openxmlformats-officedocument.presentationml.slide+xml'
CT_NOTES = 'application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml'

# 웹앱 케이스 페이지의 σ_contact-free / σ_full (이온 · Hertz) — 1저자 10-05 전사, TSV 에 없는 열
CF_OVER_FULL_HERTZ = {'0:10': 4.0, '3:7': 4.0, '5:5': 4.0, '7:3': 3.9, '10:0': 4.2}

# 양식 부분 — 네 조각에서 바이트 동일해야 한다
TEMPLATE_PREFIXES = ('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/', 'ppt/notesMasters/',
                     'ppt/tableStyles.xml', 'ppt/presProps.xml')


# ───────────────────────────── OPC 도우미 ─────────────────────────────
def load_package(path):
    """pptx (또는 pptx 하나를 담은 zip) → {partname: bytes}."""
    if hasattr(path, 'getvalue'):
        data = path.getvalue()
    else:
        with open(path, 'rb') as fh:
            data = fh.read()
    z = zipfile.ZipFile(io.BytesIO(data))
    names = z.namelist()
    if '[Content_Types].xml' not in names:  # 감싼 zip — 안의 .pptx 하나를 꺼낸다
        inner = [n for n in names if n.lower().endswith('.pptx')]
        if len(inner) != 1:
            raise SystemExit(f'{path}: pptx 가 아니고, 안에 pptx 가 정확히 하나 있지도 않다 ({names})')
        z = zipfile.ZipFile(io.BytesIO(z.read(inner[0])))
        names = z.namelist()
    return {n: z.read(n) for n in names}


def content_types(parts):
    t = etree.fromstring(parts['[Content_Types].xml'])
    defaults = {e.get('Extension').lower(): e.get('ContentType') for e in t if e.tag == f'{{{CT_NS}}}Default'}
    overrides = {e.get('PartName').lstrip('/'): e.get('ContentType') for e in t if e.tag == f'{{{CT_NS}}}Override'}
    return defaults, overrides


def part_ct(name, cts):
    defaults, overrides = cts
    if name in overrides:
        return overrides[name]
    return defaults.get(name.rsplit('.', 1)[-1].lower())


def rels_name(part):
    d, b = posixpath.split(part)
    return posixpath.join(d, '_rels', b + '.rels')


def parse_rels(parts, part):
    rn = rels_name(part)
    if rn not in parts:
        return []
    t = etree.fromstring(parts[rn])
    return [dict(Id=e.get('Id'), Type=e.get('Type'), Target=e.get('Target'), TargetMode=e.get('TargetMode'))
            for e in t]


def resolve(part, target):
    return posixpath.normpath(posixpath.join(posixpath.dirname(part), target))


def rel_target(from_part, to_part):
    return posixpath.relpath(to_part, posixpath.dirname(from_part))


def rels_xml(rels):
    root = etree.Element(f'{{{REL_NS}}}Relationships', nsmap={None: REL_NS})
    for r in rels:
        e = etree.SubElement(root, f'{{{REL_NS}}}Relationship', Id=r['Id'], Type=r['Type'], Target=r['Target'])
        if r.get('TargetMode'):
            e.set('TargetMode', r['TargetMode'])
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def xml_bytes(tree_or_el):
    el = tree_or_el.getroot() if hasattr(tree_or_el, 'getroot') else tree_or_el
    return etree.tostring(el, xml_declaration=True, encoding='UTF-8', standalone=True)


def sha(b):
    return hashlib.sha256(b).hexdigest()


# ───────────────────────────── 합치기 ─────────────────────────────
class Deck:
    """출력 패키지.  양식 부분은 part1 에서, 슬라이드는 add_slide 로 차례로."""

    def __init__(self, base):
        self.base = base
        self.base_cts = content_types(base)
        self.parts, self.cts = {}, {}
        self.by_hash = {}            # 그림 · OLE 내용 해시 → 출력 partname
        self.slides = []             # [(partname, sldId)]
        # 양식 부분 = presentation.xml 에서 닿는 것 (슬라이드 제외) + 패키지 머리
        root_rels = parse_rels(base, '')  # _rels/.rels
        todo = [resolve('', r['Target']) for r in root_rels if not r.get('TargetMode')]
        seen = set()
        while todo:
            part = todo.pop()
            if part in seen or part not in base:
                continue
            seen.add(part)
            self._put(part, base[part], part_ct(part, self.base_cts))
            for r in parse_rels(base, part):
                if r.get('TargetMode') == 'External' or r['Type'] in (RT_SLIDE,):
                    continue
                todo.append(resolve(part, r['Target']))
            if rels_name(part) in base and part != 'ppt/presentation.xml':
                self.parts[rels_name(part)] = base[rels_name(part)]
        self.parts['_rels/.rels'] = base['_rels/.rels']
        for part, data in self.parts.items():
            if part.startswith('ppt/media/'):
                self.by_hash.setdefault(sha(data), part)

    def _put(self, part, data, ct):
        self.parts[part] = data
        if ct:
            self.cts[part] = ct

    def _new_name(self, folder, stem, ext):
        n = 1
        while f'{folder}/{stem}{n}.{ext}' in self.parts:
            n += 1
        return f'{folder}/{stem}{n}.{ext}'

    def add_blob(self, data, src_part, ct):
        """그림 · hdphoto · OLE 바이너리 — 같은 내용이면 한 부분을 같이 쓴다."""
        h = sha(data)
        if h in self.by_hash:
            return self.by_hash[h]
        folder, base = posixpath.split(src_part)
        m = re.match(r'([A-Za-z]+)\d*\.(\w+)$', base)
        stem, ext = (m.group(1), m.group(2)) if m else ('file', base.rsplit('.', 1)[-1])
        name = self._new_name(folder, stem, ext)
        self._put(name, data, ct)
        self.by_hash[h] = name
        return name

    def convert_rels(self, src, src_part, rels):
        """원본 슬라이드 관계 → 출력 관계 (Id 유지).  레이아웃 = 같은 이름 (양식 동일 확인됨), 그림 · OLE = 복사.
        모든 슬라이드가 ppt/slides/ 에 있으므로 상대 경로는 슬라이드 이름과 무관하다.  메모 관계는 자리표 (Target=None)."""
        src_cts = content_types(src)
        out = []
        for r in rels:
            r = dict(r)
            if r.get('TargetMode') == 'External':
                out.append(r)
                continue
            tgt = resolve(src_part, r['Target'])
            if r['Type'] in (RT_LAYOUT, RT_NOTESMASTER):
                if tgt not in self.parts:
                    raise SystemExit(f'양식 부분 없음: {tgt}')
                new = tgt
            elif r['Type'] == RT_NOTES:
                r['Target'] = None  # add_slide 가 채운다
                out.append(r)
                continue
            elif r['Type'] == RT_SLIDE:
                raise SystemExit(f'슬라이드 간 링크는 다루지 않는다: {src_part} {r}')
            else:
                new = self.add_blob(src[tgt], tgt, part_ct(tgt, src_cts))
            r['Target'] = rel_target('ppt/slides/slide0.xml', new)
            out.append(r)
        return out

    def add_slide(self, slide_tree, rels, sld_id, notes_tree=None):
        """slide_tree = lxml 트리 · rels = convert_rels 결과 (+ 덧붙인 관계).  notes_tree 가 있으면 메모 부분을 새로 만든다."""
        part = self._new_name('ppt/slides', 'slide', 'xml')
        rels = [dict(r) for r in rels]
        nrel = [r for r in rels if r['Type'] == RT_NOTES]
        rels = [r for r in rels if r['Type'] != RT_NOTES]
        if notes_tree is not None:
            npart = self._new_name('ppt/notesSlides', 'notesSlide', 'xml')
            nid = nrel[0]['Id'] if nrel else next_rid(rels)
            rels.append(dict(Id=nid, Type=RT_NOTES, Target=rel_target(part, npart)))
            nrels = [dict(Id='rId1', Type=RT_NOTESMASTER, Target=rel_target(npart, 'ppt/notesMasters/notesMaster1.xml')),
                     dict(Id='rId2', Type=RT_SLIDE, Target=rel_target(npart, part))]
            self._put(npart, xml_bytes(notes_tree), CT_NOTES)
            self.parts[rels_name(npart)] = rels_xml(nrels)
        self._put(part, xml_bytes(slide_tree), CT_SLIDE)
        self.parts[rels_name(part)] = rels_xml(rels)
        self.slides.append((part, sld_id))
        return part

    def drop_unused(self):
        """어느 관계도 가리키지 않는 그림 · 내장 부분을 뺀다 (지운 패널 틀 그림 등)."""
        used = set()
        for rn, data in self.parts.items():
            if not rn.endswith('.rels'):
                continue
            src = posixpath.join(posixpath.dirname(posixpath.dirname(rn)), posixpath.basename(rn)[:-5])
            for e in etree.fromstring(data):
                if e.get('TargetMode') != 'External':
                    used.add(resolve(src, e.get('Target')))
        gone = [p for p in self.parts if p.startswith(('ppt/media/', 'ppt/embeddings/')) and p not in used]
        for p in gone:
            del self.parts[p]
            self.cts.pop(p, None)
        return gone

    def finish(self, app_src):
        """presentation.xml · 관계 · 구역 · app.xml · [Content_Types].xml 을 다시 쓴다."""
        self.dropped = self.drop_unused()
        pres = etree.fromstring(self.base['ppt/presentation.xml'])
        prels = [r for r in parse_rels(self.base, 'ppt/presentation.xml') if r['Type'] != RT_SLIDE]
        used = {r['Id'] for r in prels}
        lst = pres.find('p:sldIdLst', NS)
        for e in list(lst):
            lst.remove(e)
        k = 100
        for part, sid in self.slides:
            while f'rId{k}' in used:
                k += 1
            rid = f'rId{k}'
            used.add(rid)
            prels.append(dict(Id=rid, Type=RT_SLIDE, Target=rel_target('ppt/presentation.xml', part)))
            etree.SubElement(lst, f'{P}sldId', id=str(sid), attrib={f'{R}id': rid})
        # 구역: 첫 구역에 모든 장, 나머지 구역은 비운다 (원본과 같은 구조)
        for sec_lst in pres.iter(f'{{{NS["p14"]}}}sectionLst'):
            secs = sec_lst.findall('p14:section', NS)
            for i, sec in enumerate(secs):
                sl = sec.find('p14:sldIdLst', NS)
                for e in list(sl):
                    sl.remove(e)
                if i == 0:
                    for _, sid in self.slides:
                        etree.SubElement(sl, f'{{{NS["p14"]}}}sldId', id=str(sid))
        self.parts['ppt/presentation.xml'] = xml_bytes(pres)
        self.cts['ppt/presentation.xml'] = part_ct('ppt/presentation.xml', self.base_cts)
        self.parts['ppt/_rels/presentation.xml.rels'] = rels_xml(prels)
        # app.xml — 장 수 · 메모 수 · 장 제목 목록을 실제와 맞춘다 (나머지는 part2 의 것 = OLE 서버 "Graph" 포함)
        app = etree.fromstring(app_src)
        ep = '{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}'
        vt = '{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}'
        n_notes = sum(1 for p in self.parts if p.startswith('ppt/notesSlides/') and p.endswith('.xml'))
        for tag, val in (('Slides', len(self.slides)), ('Notes', n_notes)):
            e = app.find(ep + tag)
            if e is not None:
                e.text = str(val)
        hp = app.find(f'{ep}HeadingPairs/{vt}vector')
        tp = app.find(f'{ep}TitlesOfParts/{vt}vector')
        if hp is not None and tp is not None:
            pairs = list(hp)
            counts = []
            for i in range(0, len(pairs), 2):
                counts.append([pairs[i].find(vt + 'lpstr').text, pairs[i + 1].find(vt + 'i4')])
            titles = [e.text for e in tp]
            pos, new_titles = 0, []
            for name, i4 in counts:
                n = int(i4.text)
                chunk = titles[pos:pos + n]
                pos += n
                if name == '슬라이드 제목':
                    chunk = ['PowerPoint 프레젠테이션'] * len(self.slides)
                    i4.text = str(len(self.slides))
                new_titles += chunk
            for e in list(tp):
                tp.remove(e)
            for t in new_titles:
                etree.SubElement(tp, vt + 'lpstr').text = t
            tp.set('size', str(len(new_titles)))
        self.parts['docProps/app.xml'] = xml_bytes(app)
        # [Content_Types].xml
        root = etree.Element(f'{{{CT_NS}}}Types', nsmap={None: CT_NS})
        defaults = {'rels': 'application/vnd.openxmlformats-package.relationships+xml', 'xml': 'application/xml'}
        overrides = {}
        for part in self.parts:
            ct = self.cts.get(part)
            if part.endswith('.rels') or part == '[Content_Types].xml':
                continue
            ext = part.rsplit('.', 1)[-1].lower()
            if ct is None:
                raise SystemExit(f'내용 형식 모름: {part}')
            if ext in ('png', 'jpeg', 'jpg', 'wdp', 'emf', 'bin', 'gif', 'tif', 'tiff', 'svg', 'xlsx', 'wmf'):
                if defaults.setdefault(ext, ct) != ct:
                    overrides[part] = ct
            elif not (ext == 'xml' and ct == 'application/xml'):
                overrides[part] = ct
        for ext, ct in sorted(defaults.items()):
            etree.SubElement(root, f'{{{CT_NS}}}Default', Extension=ext, ContentType=ct)
        for part, ct in sorted(overrides.items()):
            etree.SubElement(root, f'{{{CT_NS}}}Override', PartName='/' + part, ContentType=ct)
        self.parts['[Content_Types].xml'] = xml_bytes(root)

    def save(self, path):
        order = ['[Content_Types].xml', '_rels/.rels', 'ppt/presentation.xml']
        names = order + sorted(p for p in self.parts if p not in order)
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
            for n in names:
                zi = zipfile.ZipInfo(n, date_time=(2026, 10, 6, 0, 0, 0))   # 묶음 시각 고정 (재현용)
                zi.compress_type = zipfile.ZIP_DEFLATED
                z.writestr(zi, self.parts[n])


def next_rid(rels):
    used = {r['Id'] for r in rels}
    k = 1
    while f'rId{k}' in used:
        k += 1
    return f'rId{k}'


def check_template(pkgs):
    """양식 부분이 네 조각에서 바이트 동일한지 확인 — 다르면 레이아웃 공유가 틀리므로 멈춘다."""
    names = sorted({n for pk in pkgs.values() for n in pk if n.startswith(TEMPLATE_PREFIXES)})
    bad = []
    for n in names:
        hs = {k: sha(pk[n]) if n in pk else None for k, pk in pkgs.items()}
        if len(set(hs.values())) != 1:
            bad.append((n, hs))
    if bad:
        raise SystemExit('양식 부분이 조각마다 다르다: ' + repr(bad[:5]))
    return len(names)


def slide_parts(pkg):
    """presentation.xml 의 순서대로 슬라이드 partname."""
    pres = etree.fromstring(pkg['ppt/presentation.xml'])
    rels = {r['Id']: r for r in parse_rels(pkg, 'ppt/presentation.xml')}
    return [resolve('ppt/presentation.xml', rels[e.get(f'{R}id')]['Target']) for e in pres.find('p:sldIdLst', NS)]


def slide_ids(pkg):
    pres = etree.fromstring(pkg['ppt/presentation.xml'])
    return [int(e.get('id')) for e in pres.find('p:sldIdLst', NS)]


def notes_part(pkg, slide_part):
    for r in parse_rels(pkg, slide_part):
        if r['Type'] == RT_NOTES:
            return resolve(slide_part, r['Target'])
    return None


# ───────────────────────────── 도형 도우미 ─────────────────────────────
SHAPE_TAGS = ('sp', 'grpSp', 'pic', 'graphicFrame', 'cxnSp', 'AlternateContent', 'contentPart')


def sptree(tree):
    return tree.find('p:cSld/p:spTree', NS)


def tops(tree):
    return [e for e in sptree(tree) if etree.QName(e).localname in SHAPE_TAGS]


def first_shape(el):
    """AlternateContent 이면 Choice 쪽 도형."""
    if etree.QName(el).localname == 'AlternateContent':
        return el.find('mc:Choice', NS)[0]
    return el


def cnvpr(el):
    return first_shape(el).find('./*/p:cNvPr', NS)


def name_of(el):
    c = cnvpr(el)
    return c.get('name') if c is not None else None


def find_top(tree, name):
    hits = [e for e in tops(tree) if name_of(e) == name]
    if len(hits) != 1:
        raise SystemExit(f'최상위 도형 "{name}" 이 {len(hits)} 개')
    return hits[0]


def find_desc(el, name):
    """그룹 안까지 이름으로 찾기 (하나여야 한다)."""
    hits = [c.getparent().getparent() for c in el.iter(f'{P}cNvPr') if c.get('name') == name]
    if len(hits) != 1:
        raise SystemExit(f'도형 "{name}" 이 {len(hits)} 개')
    return hits[0]


def remove_tops(tree, names):
    for n in names:
        e = find_top(tree, n)
        e.getparent().remove(e)


def xfrms(el):
    """el 의 위치를 정하는 xfrm 들 (AlternateContent = Choice · Fallback 둘 다)."""
    ln = etree.QName(el).localname
    if ln == 'AlternateContent':
        out = []
        for branch in el:
            for child in branch:
                out += xfrms(child)
        return out
    if ln == 'graphicFrame':
        return [el.find('p:xfrm', NS)]
    if ln == 'grpSp':
        return [el.find('p:grpSpPr/a:xfrm', NS)]
    x = el.find('./p:spPr/a:xfrm', NS)
    return [x] if x is not None else []


def box(el):
    x = xfrms(el)[0]
    off, ext = x.find('a:off', NS), x.find('a:ext', NS)
    return int(off.get('x')), int(off.get('y')), int(ext.get('cx')), int(ext.get('cy'))


def union_box(els):
    bs = [box(e) for e in els]
    x0, y0 = min(b[0] for b in bs), min(b[1] for b in bs)
    x1, y1 = max(b[0] + b[2] for b in bs), max(b[1] + b[3] for b in bs)
    return x0, y0, x1 - x0, y1 - y0


def place(el, origin, dest, s):
    """원점 origin 을 dest 로 옮기며 s 배 (모든 xfrm · 표 격자 · 행 높이).  글자 크기는 scale_text 가 따로."""
    ox, oy = origin
    nx, ny = dest
    for x in xfrms(el):
        off, ext = x.find('a:off', NS), x.find('a:ext', NS)
        off.set('x', str(round(nx + (int(off.get('x')) - ox) * s)))
        off.set('y', str(round(ny + (int(off.get('y')) - oy) * s)))
        ext.set('cx', str(round(int(ext.get('cx')) * s)))
        ext.set('cy', str(round(int(ext.get('cy')) * s)))
    if etree.QName(el).localname == 'graphicFrame' and s != 1:
        for g in el.iter(f'{A}gridCol'):
            g.set('w', str(round(int(g.get('w')) * s)))
        for tr in el.iter(f'{A}tr'):
            tr.set('h', str(round(int(tr.get('h')) * s)))


def place_block(els, dest, s, font_s=1.0, cap_pt=None, keep_names=()):
    """여러 도형을 한 덩어리로 옮기고 키운다 (배치 = 덩어리 왼쪽 위 기준)."""
    x0, y0, _, _ = union_box(els)
    for e in els:
        place(e, (x0, y0), dest, s)
        scale_text(e, font_s, cap_pt, keep_names)


# 비례 확대 때 글자는 상자보다 2 % 덜 키운다 — PowerPoint 줄바꿈이 원본과 같게 (꽉 찬 줄이 넘어가지 않게)
FONT_MARGIN = 0.98


def _round_sz(v):
    """키운 글자 크기는 0.5 pt 단위로 내림 (9.7 → 9.5 pt 처럼 — 상자보다 커지지 않게 · 보기 좋은 크기)."""
    return max(100, int((v + 1e-6) // 50) * 50)


def scale_text(el, fs, cap_pt=None, keep_names=()):
    """글자 크기 · 글 상자 여백 · 들여쓰기 · 문단 간격 · 표 칸 여백을 fs 배.  cap_pt = 키울 때의 상한 (원래 더 크면 그대로).
    keep_names = 이 이름의 도형 안 글자는 그대로 둔다."""
    if fs == 1.0:
        return

    def kept(e):
        """keep_names 도형 안의 요소인가 (조상 도형의 이름으로 — lxml 프록시 id() 는 안정적이지 않다)."""
        for anc in e.iterancestors():
            ln = etree.QName(anc).localname
            if ln in ('sp', 'pic', 'graphicFrame', 'cxnSp'):
                c = anc.find('./*/p:cNvPr', NS)
                return c is not None and c.get('name') in keep_names
        return False

    def new_sz(sz):
        v = sz * fs
        if cap_pt is not None and fs > 1:
            v = max(sz, min(v, cap_pt * 100))
        return _round_sz(v)

    for e in el.iter(f'{A}rPr', f'{A}defRPr', f'{A}endParaRPr'):
        if e.get('sz') and not kept(e):
            e.set('sz', str(new_sz(int(e.get('sz')))))
    for bp in el.iter(f'{A}bodyPr'):
        if kept(bp):
            continue
        for k in ('lIns', 'tIns', 'rIns', 'bIns'):
            if bp.get(k):
                bp.set(k, str(round(int(bp.get(k)) * fs)))
    for pp in el.iter(f'{A}pPr'):
        if kept(pp):
            continue
        for k in ('marL', 'indent'):
            if pp.get(k):
                pp.set(k, str(round(int(pp.get(k)) * fs)))
    for sp in el.iter(f'{A}spcPts'):
        if not kept(sp):
            sp.set('val', str(round(int(sp.get('val')) * fs)))
    for tc in el.iter(f'{A}tcPr'):
        for k in ('marL', 'marR', 'marT', 'marB'):
            if tc.get(k):
                tc.set(k, str(round(int(tc.get(k)) * fs)))


def ungroup(tree, name):
    """회전 · 배율 없는 그룹을 풀어 자식들을 최상위로 (같은 자리 · 같은 겹침 순서)."""
    g = find_top(tree, name)
    x = g.find('p:grpSpPr/a:xfrm', NS)
    off, ext, cho, che = (x.find(f'a:{k}', NS) for k in ('off', 'ext', 'chOff', 'chExt'))
    if (ext.get('cx'), ext.get('cy')) != (che.get('cx'), che.get('cy')) or x.get('rot') not in (None, '0'):
        raise SystemExit(f'그룹 {name} 은 배율 · 회전이 있어 풀 수 없다')
    dx, dy = int(off.get('x')) - int(cho.get('x')), int(off.get('y')) - int(cho.get('y'))
    parent, idx = g.getparent(), g.getparent().index(g)
    kids = [c for c in g if etree.QName(c).localname in SHAPE_TAGS]
    for c in kids:
        for xx in xfrms(c):
            o = xx.find('a:off', NS)
            o.set('x', str(int(o.get('x')) + dx))
            o.set('y', str(int(o.get('y')) + dy))
    parent.remove(g)
    for i, c in enumerate(kids):
        parent.insert(idx + i, c)
    return kids


def set_box(el, x=None, y=None, w=None, h=None):
    xx = xfrms(el)[0]
    off, ext = xx.find('a:off', NS), xx.find('a:ext', NS)
    for k, v, node in (('x', x, off), ('y', y, off), ('cx', w, ext), ('cy', h, ext)):
        if v is not None:
            node.set(k, str(int(round(v))))


def max_shape_id(tree):
    return max(int(c.get('id')) for c in tree.iter(f'{P}cNvPr'))


def renumber_ids(el, start):
    k = start
    for c in el.iter(f'{P}cNvPr'):
        c.set('id', str(k))
        k += 1
    return k


# ───────────────────────────── 글 도우미 ─────────────────────────────
def para_text(p):
    return ''.join(t.text or '' for t in p.iter(f'{A}t'))


def runs_of(p):
    return p.findall('a:r', NS)


def replace_in_para(p, old, new):
    """문단 안 old → new (런 경계를 넘어도 된다).  첫 런에 새 글을 넣고 나머지 런에서 지운다.  바꾼 횟수."""
    n = 0
    while True:
        rs = runs_of(p)
        texts = [r.find('a:t', NS).text or '' for r in rs]
        full = ''.join(texts)
        i = full.find(old)
        if i < 0:
            return n
        pos, first = 0, True
        end = i + len(old)
        for r, t in zip(rs, texts):
            a, b = pos, pos + len(t)
            pos = b
            if b <= i or a >= end:
                continue
            lo, hi = max(i, a) - a, min(end, b) - a
            te = r.find('a:t', NS)
            te.text = t[:lo] + (new if first else '') + t[hi:]
            first = False
        n += 1
        if new.find(old) >= 0:  # 무한 반복 방지
            return n


def replace_text(el, old, new, expect=None):
    n = sum(replace_in_para(p, old, new) for p in el.iter(f'{A}p'))
    if expect is not None and n != expect:
        raise SystemExit(f'"{old}" → "{new}" 바꾼 횟수 {n} (기대 {expect})')
    return n


def make_run(tmpl_rpr, text, **attrs):
    """tmpl_rpr (a:rPr) 를 복제해 런 하나.  attrs: i='1', baseline='-25000', latin='Aptos' 등."""
    r = etree.Element(f'{A}r')
    rpr = copy.deepcopy(tmpl_rpr)
    rpr.tag = f'{A}rPr'
    for k in ('err', 'dirty', 'smtClean'):
        rpr.attrib.pop(k, None)
    latin = attrs.pop('latin', None)
    for k, v in attrs.items():
        if v is None:
            rpr.attrib.pop(k, None)
        else:
            rpr.set(k, v)
    if latin:
        la = rpr.find('a:latin', NS)
        if la is None:
            la = etree.SubElement(rpr, f'{A}latin')
        for k in list(la.attrib):
            del la.attrib[k]
        la.set('typeface', latin)
    r.append(rpr)
    etree.SubElement(r, f'{A}t').text = text
    return r


TOKEN_RE = re.compile(r'(\{[^{}]*\})')


def rich_runs(spec, plain_rpr, sym_rpr, latin_words=True):
    """간단 표기 → 런 목록.  {σ|ion} = 기울임 σ + 아래첨자 ion, {τ|^2} = 위첨자, {φ|SE}, {σ|0}.
    라틴 낱말 (영문자로 시작) 은 latin=Aptos, 나머지 (한글 · 숫자 · 기호) 는 plain_rpr 의 글꼴 그대로."""
    out = []
    for tok in TOKEN_RE.split(spec):
        if not tok:
            continue
        if tok.startswith('{'):
            base, sub = tok[1:-1].split('|')
            out.append(make_run(sym_rpr, base, i='1', baseline=None))
            if sub.startswith('^'):
                out.append(make_run(sym_rpr, sub[1:], i='0', baseline='30000'))
            else:
                out.append(make_run(sym_rpr, sub, i='1' if sub.isalpha() else '0', baseline='-25000'))
            continue
        # 라틴 낱말 / 나머지로 쪼갠다
        for piece in re.split(r'([A-Za-z][A-Za-z\-–]*(?:\s+[A-Za-z][A-Za-z\-–]*)*)', tok):
            if not piece:
                continue
            if latin_words and re.match(r'[A-Za-z]', piece):
                out.append(make_run(plain_rpr, piece, latin='Aptos'))
            else:
                out.append(make_run(plain_rpr, piece))
    return out


def set_para_runs(p, runs):
    for r in p.findall('a:r', NS) + p.findall('a:br', NS) + p.findall('a:fld', NS):
        p.remove(r)
    end = p.find('a:endParaRPr', NS)
    for r in runs:
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)


# 글 폭 추정 (배치 계산용 — 맑은 고딕 한글 1 em, 라틴 · 숫자는 평균 폭).  최종 확인은 렌더로 한다.
def text_width_em(s):
    w = 0.0
    for ch in s:
        o = ord(ch)
        if 0xAC00 <= o <= 0xD7A3 or 0x3130 <= o <= 0x318F or 0x4E00 <= o <= 0x9FFF:
            w += 1.0
        elif ch == ' ':
            w += 0.30
        elif ch.isdigit():
            w += 0.55
        elif ch.isupper():
            w += 0.64
        elif ch.islower():
            w += 0.50
        elif ch in '·•':
            w += 0.40
        elif ch in '—→–':
            w += 1.0 if ch in '—→' else 0.55
        else:
            w += 0.40
    return w


def estimate_lines(text, width_emu, size_pt, indent_emu=0):
    """단어 단위 줄바꿈 추정 — 줄 수."""
    avail = (width_emu - indent_emu) / 12700.0  # pt
    words = text.split(' ')
    lines, cur = 1, 0.0
    for wd in words:
        ww = text_width_em(wd) * size_pt
        sp = 0.30 * size_pt if cur > 0 else 0
        if cur + sp + ww <= avail or cur == 0:
            cur += sp + ww
        else:
            lines += 1
            cur = ww
    return lines


def textbox_height(el, width=None):
    """spAutoFit 글 상자의 예상 높이 (EMU) — 줄 높이 = 1.2 × 글자 크기 (원본 PDF 실측) + 문단 뒤 간격 + 여백."""
    sp = first_shape(el)
    bp = sp.find('.//a:bodyPr', NS)
    w = width if width is not None else box(el)[2]
    lins = int(bp.get('lIns', 91440)); rins = int(bp.get('rIns', 91440))
    tins = int(bp.get('tIns', 45720)); bins = int(bp.get('bIns', 45720))
    h = tins + bins
    for p in sp.iter(f'{A}p'):
        szs = [int(r.get('sz')) for r in p.iter(f'{A}rPr', f'{A}endParaRPr') if r.get('sz')]
        size = (max(szs) if szs else 1800) / 100.0
        ppr = p.find('a:pPr', NS)
        marl = int(ppr.get('marL', 0)) if ppr is not None else 0
        n = estimate_lines(para_text(p), w - lins - rins - marl, size) if para_text(p) else 1
        h += n * size * 1.2 * 12700
        sa = p.find('a:pPr/a:spcAft/a:spcPts', NS)
        if sa is not None:
            h += int(sa.get('val')) / 100.0 * 12700
    return int(h)


# ───────────────────────────── 슬라이드 작업대 ─────────────────────────────
class SlideWork:
    """원본 슬라이드 하나를 복제해 고치는 작업대 — 트리 · 출력 관계 목록 · 메모."""

    def __init__(self, deck, pkg, part):
        self.deck, self.pkg, self.part = deck, pkg, part
        self.tree = etree.fromstring(pkg[part])
        self.rels = deck.convert_rels(pkg, part, parse_rels(pkg, part))
        np_ = notes_part(pkg, part)
        self.notes = etree.fromstring(pkg[np_]) if np_ else None

    def import_el(self, src_pkg, src_part, el, index=None):
        """다른 슬라이드의 도형을 복제해 넣는다 — r:embed 등은 이 슬라이드의 새 관계로 바꾼다."""
        e = copy.deepcopy(el)
        src_rels = {r['Id']: r for r in parse_rels(src_pkg, src_part)}
        mapping = {}
        for node in e.iter():
            for k, v in list(node.attrib.items()):
                if not k.startswith(R) or v not in src_rels:
                    continue
                if v not in mapping:
                    conv = self.deck.convert_rels(src_pkg, src_part, [src_rels[v]])[0]
                    same = [r for r in self.rels if r['Type'] == conv['Type'] and r['Target'] == conv['Target']]
                    if same:
                        mapping[v] = same[0]['Id']
                    else:
                        conv['Id'] = next_rid(self.rels)
                        self.rels.append(conv)
                        mapping[v] = conv['Id']
                node.set(k, mapping[v])
        renumber_ids(e, max_shape_id(self.tree) + 1)
        st = sptree(self.tree)
        if index is None:
            st.append(e)
        else:
            st.insert(index, e)
        return e

    def prune_rels(self):
        """지운 도형만 쓰던 관계를 뺀다 (레이아웃 · 메모 관계는 남긴다)."""
        used = set()
        for node in self.tree.iter():
            for k, v in node.attrib.items():
                if k.startswith(R):
                    used.add(v)
        self.rels = [r for r in self.rels if r['Id'] in used or r['Type'] in (RT_LAYOUT, RT_NOTES)]

    def add_to(self, sld_id, notes=True):
        self.prune_rels()
        return self.deck.add_slide(self.tree, self.rels, sld_id, self.notes if notes else None)


# 온폭 패널 틀 (3–5 장과 같은 것) = part1.5 (2 장) 의 '그룹 2' — 탭 글 · 제목 글만 바꿔 쓴다
FRAME_NAME, FRAME_TAB, FRAME_TITLE = '그룹 2', '직사각형 28', 'TextBox 29'
# 내용 영역 (2 · 3 장 글머리 상자와 같은 좌우 · 패널 흰 바탕 안쪽)
CX0, CX1 = 591633, 9435005
CW = CX1 - CX0
CY0 = 1736326            # 글머리 첫 줄 (2–6 장 공통)
CY1 = 6330000            # 내용 아래 한계 (패널 흰 바탕 아래 끝 6435741 안쪽)


def copy_paras(dst_shape, src_shape):
    """dst 글 상자의 문단을 src 의 문단 복제로 바꾼다 (bodyPr · lstStyle 는 dst 것 유지)."""
    dtx = dst_shape.find('p:txBody', NS)
    for p in dtx.findall('a:p', NS):
        dtx.remove(p)
    for p in src_shape.find('p:txBody', NS).findall('a:p', NS):
        dtx.append(copy.deepcopy(p))


def put_frame(sw, frame_src, tab_src, title_src, index):
    """온폭 틀을 index 자리에 넣고 탭 · 제목 문단을 원래 패널 것으로 바꾼다."""
    fpkg, fpart, fel = frame_src
    fr = sw.import_el(fpkg, fpart, fel, index=index)
    copy_paras(find_desc(fr, FRAME_TAB), tab_src)
    copy_paras(find_desc(fr, FRAME_TITLE), title_src)
    return fr


def fit_textbox(el, x, y, w):
    """글 상자를 (x, y, 폭 w) 로 옮기고 높이를 추정값으로 맞춘다 (spAutoFit — PowerPoint 는 편집 때 다시 맞춘다)."""
    h = textbox_height(el, w)
    set_box(el, x=x, y=y, w=w, h=h)
    return h


# ───────────────────────────── 값 (TSV) ─────────────────────────────
COMPS = ('0:10', '3:7', '5:5', '7:3', '10:0')


def read_tsv(path):
    with open(path, encoding='utf-8') as fh:
        rows = list(csv.DictReader(fh, delimiter='\t'))
    by = {r['PC_SC']: r for r in rows}
    if tuple(by) != COMPS:
        raise SystemExit(f'TSV 조성 순서가 다르다: {tuple(by)}')
    return by


def fmt(v, nd):
    """반올림 = 사사오입 (0.1175 → 0.118 — 이진 부동소수의 0.117 을 피한다)."""
    q = Decimal(1).scaleb(-nd)
    return str(Decimal(str(v)).quantize(q, rounding=ROUND_HALF_UP))


def derived(tsv):
    """표 · 글머리 · 그래프에 쓰는 값 (모두 TSV 에서).  호출자 문안과 대조해 다르면 멈춘다."""
    g = lambda c, k: float(tsv[c][k])
    ion_h = {c: g(c, 'sigma_ion_hertz_mScm') for c in COMPS}
    ion_p = {c: g(c, 'sigma_ion_physics_mScm') for c in COMPS}
    e_h = {c: g(c, 'sigma_e_hertz_mScm') for c in COMPS}
    e_p = {c: g(c, 'sigma_e_physics_mScm') for c in COMPS}
    t2 = {c: g(c, 'tau2_hertz') for c in COMPS}
    t2_p = {c: g(c, 'tau2_physics') for c in COMPS}
    amam = {c: int(tsv[c]['AM_AM_contacts']) for c in COMPS}
    d = dict(ion_h=ion_h, ion_p=ion_p, e_h=e_h, e_p=e_p, t2=t2, t2_p=t2_p, amam=amam)
    d['t2_p_min'] = min(COMPS, key=lambda c: t2_p[c])
    d['ion_max'] = max(COMPS, key=lambda c: ion_h[c])
    d['t2_min'] = min(COMPS, key=lambda c: t2[c])
    e_list = [e_h[c] for c in COMPS]
    d['e_monotone'] = all(a > b for a, b in zip(e_list, e_list[1:]))
    cf = [CF_OVER_FULL_HERTZ[c] for c in COMPS]
    d['cf_range'] = (min(cf), max(cf))
    # 호출자 (1저자 확인용) 문안과 같은 숫자인지
    assert (fmt(ion_h['0:10'], 3), fmt(ion_h['7:3'], 3), d['ion_max']) == ('0.118', '0.179', '7:3'), d
    assert (tsv['0:10']['sigma_e_hertz_mScm'], tsv['10:0']['sigma_e_hertz_mScm'], d['e_monotone']) == ('4.77', '1.30', True)
    assert (fmt(t2['0:10'], 2), fmt(t2['7:3'], 2), d['t2_min']) == ('7.24', '5.03', '7:3'), d
    assert 3.85 <= d['cf_range'][0] and d['cf_range'][1] <= 4.25
    assert (amam['0:10'], amam['10:0']) == (8970, 733)
    assert fmt(ion_p['7:3'], 3) == fmt(ion_p['10:0'], 3) == '0.106'
    return d


# ───────────────────────────── 장별 손보기 ─────────────────────────────
def first_rpr(el, text_contains=None, italic=None):
    for r in el.iter(f'{A}r'):
        t = r.find('a:t', NS).text or ''
        rpr = r.find('a:rPr', NS)
        if text_contains is not None and text_contains not in t:
            continue
        if italic is not None and (rpr.get('i') == '1') != italic:
            continue
        return rpr
    raise SystemExit(f'런 서식 못 찾음 ({text_contains}, {italic})')


def join_break(el, which=0):
    """글 상자의 which 번째 줄바꿈 (a:br) 을 빈칸 하나로 — 좁은 패널에 맞춰 넣은 줄바꿈을 온폭에서 푼다."""
    brs = list(el.iter(f'{A}br'))
    br = brs[which]
    prev, nxt = br.getprevious(), br.getnext()
    before = (prev.findtext('a:t', namespaces=NS) or '') if prev is not None else ''
    after = (nxt.findtext('a:t', namespaces=NS) or '') if nxt is not None else ''
    if not before.endswith(' ') and not after.startswith(' '):   # 앞 · 뒤에 빈칸이 없을 때만 하나 넣는다
        r = etree.Element(f'{A}r')
        rpr = br.find('a:rPr', NS)
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        etree.SubElement(r, f'{A}t').text = ' '
        br.addprevious(r)
    br.getparent().remove(br)


def cell_lines(tc, width, size_pt):
    """표 칸의 예상 줄 수 (문단 · 줄바꿈 · 단어 줄바꿈)."""
    n = 0
    for p in tc.findall('.//a:p', NS):
        parts = [[]]
        for ch in p:
            ln = etree.QName(ch).localname
            if ln == 'br':
                parts.append([])
            elif ln == 'r':
                parts[-1].append(ch.find('a:t', NS).text or '')
        for seg in parts:
            n += estimate_lines(''.join(seg), width, size_pt) if ''.join(seg).strip() else 1
    return max(n, 1)


def edit_1a(sw, frame_src):
    """1a = 'DEM 원리 · 시뮬레이션 조건' — 왼쪽 패널 내용을 온폭으로."""
    t = sw.tree
    left = find_top(t, '그룹 2')
    tab, title = find_desc(left, '직사각형 28'), find_desc(left, 'TextBox 29')
    idx = sptree(t).index(left)
    tab, title = copy.deepcopy(tab), copy.deepcopy(title)
    remove_tops(t, ['그룹 2', '그룹 30', 'TextBox 77', '그룹 103', '그룹 94', '그룹 101'])
    put_frame(sw, frame_src, tab, title, idx)
    bul, tbl, foot = find_top(t, 'TextBox 8'), find_top(t, '표 40'), find_top(t, 'TextBox 76')
    fig = [find_top(t, n) for n in ('그룹 39', 'Insertion — 생략 표시', 'Compression — 가압 표시')]
    # 글머리 = 2 · 3 장과 같은 자리 · 폭 (9 pt 그대로)
    h_bul = fit_textbox(bul, CX0, CY0, CW)
    # 가운데 줄: 그림 (왼쪽, s_f 배 — 그림만 키우고 그림 밑 글은 9 pt 그대로) | 표 (오른쪽, 8 → 9 pt = 5 · 6 장 표와 같은 크기)
    s_f, s_t, gap = 1.35, 9 / 8, 300000
    fx, fy, fw, fh = union_box(fig)
    fig_w, fig_h = fw * s_f, fh * s_f
    tw_new = round(CW - fig_w - gap)
    grid = tbl.findall('.//a:tblGrid/a:gridCol', NS)
    c1 = 1213000                                         # '고체전해질 유효 영률' 9 pt 한 줄
    grid[0].set('w', str(c1))
    grid[1].set('w', str(tw_new - c1))
    scale_text(tbl, s_t)
    rows = tbl.findall('.//a:tr', NS)
    nat = []
    for r in rows:
        tcs = r.findall('a:tc', NS)
        lines = max(cell_lines(tc, w - 2 * 51435, 9) for tc, w in zip(tcs, (c1, tw_new - c1)))
        nat.append(lines * 9 * 1.2 * 12700 + 60000)
    k = fig_h / sum(nat)                                  # 표 높이 = 그림 높이 (위 · 아래 끝 맞춤)
    for r, b in zip(rows, nat):
        r.set('h', str(round(b * k)))
    # 각주: 표의 '*' 설명 — 온폭이라 첫 줄바꿈을 풀고 (괄호 문장 앞 줄바꿈은 둔다) 줄 바로 아래
    join_break(foot, 0)
    h_foot = textbox_height(foot, CW)
    block = fig_h + 110000 + h_foot
    slack = CY1 - (CY0 + h_bul) - block
    y_row = CY0 + h_bul + max(150000, slack * 0.45)        # 남는 높이는 위 · 아래로 고르게 (위를 조금 덜)
    place_block(fig, (CX0, y_row), s_f)
    set_box(tbl, x=CX1 - tw_new, y=y_row, w=tw_new, h=fig_h)
    set_box(foot, x=CX0, y=y_row + fig_h + 110000, w=CW, h=h_foot)
    return dict(s_fig=s_f, table_font=9, table_w=tw_new, fig_h=fig_h, table_row_stretch=k)


def edit_1b(sw, frame_src):
    """1b = '지배 방정식 · 계산 과정' — 오른쪽 패널 내용을 온폭으로 (수식은 OMML 그대로 · 크기만)."""
    t = sw.tree
    left, right = find_top(t, '그룹 2'), find_top(t, '그룹 30')
    tab, title = copy.deepcopy(find_desc(right, '직사각형 32')), copy.deepcopy(find_desc(right, 'TextBox 33'))
    idx = sptree(t).index(left)
    remove_tops(t, ['그룹 2', '그룹 30', 'TextBox 8', '그룹 39', '표 40', 'Insertion — 생략 표시',
                    'Compression — 가압 표시', 'TextBox 76'])
    put_frame(sw, frame_src, tab, title, idx)
    intro, eqs = find_top(t, 'TextBox 77'), find_top(t, '그룹 103')
    force, sym = find_top(t, '그룹 94'), find_top(t, '그룹 101')
    h_in = fit_textbox(intro, CX0, CY0, CW)
    top = CY0 + h_in + 150000
    avail = CY1 - top
    # 왼쪽: 수식 묶음 (Δt 흐름 · ①–⑤ 이름 · OMML 수식) — 그대로 s_e 배 (글자 · 수식 같이)
    ex, ey, ew, eh = box(eqs)
    s_e = min(1.3, avail / eh)
    eq_h = eh * s_e
    top += (avail - eq_h) * 0.35                        # 남는 높이를 위 · 아래로 (아래를 조금 더)
    # 오른쪽: 힘 그림 (위) + 기호 상자 (아래) — 위 · 아래 끝을 수식 묶음과 맞춘다
    gap, vgap = 300000, 150000
    wr = CW - ew * s_e - gap
    _, _, gw, gh = box(force)
    _, _, hw, hh = box(sym)
    s_h = min(1.25, wr / hw)
    s_g = min((eq_h - hh * s_h - vgap) / gh, wr / gw)
    place_block([eqs], (CX0, top), s_e, font_s=s_e * FONT_MARGIN)
    col_x = CX1 - wr
    # 힘 그림: 그림 위 글자 (하중 · F · M · i · j) 는 그림과 같은 배율 (11 pt 상한) · 아래 설명은 9 pt 상한
    place_block([force], (col_x + (wr - gw * s_g) / 2, top), s_g, font_s=s_g * FONT_MARGIN, cap_pt=11,
                keep_names=('TextBox 92',))
    scale_text(find_desc(force, 'TextBox 92'), min(s_g, 9 / 7))
    # 기호 상자: 상자는 s_h 배 · 글자는 8 → 9 pt (본문 크기)
    place_block([sym], (col_x + (wr - hw * s_h) / 2, top + eq_h - hh * s_h), s_h, font_s=9 / 8)
    return dict(s_eq=s_e, s_force=s_g, s_sym=s_h)


def edit_6a(sw, frame_src):
    """6a = '접촉 네트워크 개념' — 왼쪽 패널 내용을 온폭으로."""
    t = sw.tree
    tab, title = copy.deepcopy(find_top(t, '직사각형 28')), copy.deepcopy(find_top(t, 'TextBox 29'))
    idx = sptree(t).index(find_top(t, '그림 26'))
    remove_tops(t, ['그림 26', '사각형: 둥근 모서리 13', '직사각형 28', 'TextBox 29', '그룹 30', 'TextBox 80',
                    'Table 413', '예정 표지 10월 계산 예정'] + CHART_STUB)
    put_frame(sw, frame_src, tab, title, idx)
    bul = find_top(t, 'TextBox 36')
    h_bul = fit_textbox(bul, CX0, CY0, CW)
    top = CY0 + h_bul + 150000
    avail = CY1 - top
    rblock = [find_top(t, n) for n in R_BLOCK]   # 범례를 풀기 전에 (범례에도 'Freeform 10' 이 있다)
    ungroup(t, '그룹 504')
    legend = ungroup(t, '그룹 466')               # 범례 묶음 — 첫 자식 (점선 'Connector 379') 은 그림 오른쪽 경계선
    edge = [e for e in legend if name_of(e) == 'Connector 379']
    legend = [e for e in legend if name_of(e) != 'Connector 379']
    diag = [find_top(t, '그룹 467')] + edge       # 접촉망 그림 + 오른쪽 점선 (같이 움직여야 한다)
    dx, dy, dw, dh = union_box(diag)
    rx, ry, rw, rh = union_box(rblock)
    lx, ly, lw, lh = union_box(legend)
    # 왼쪽: 접촉망 그림 s_d 배 (막대 글자 9 → 10 pt) · 오른쪽: 범례 (위) + '접촉 하나의 저항' 상자 (아래)
    # — 범례 · 상자는 글자까지 같은 배율 (상자 안 글자 위치가 그림에 맞춰져 있다)
    s_d = min(1.5, avail / dh)
    d_h = dh * s_d
    place_block(diag, (CX0, top + (avail - d_h) / 2), s_d, font_s=min(s_d, 10 / 9))
    col_x = CX0 + dw * s_d + 400000
    wr = CX1 - col_x
    s_r = min(1.15, wr / rw)
    vgap = 220000
    s_l = max(1.0, min(1.1, (avail - rh * s_r - vgap) / lh))
    s_r = min(s_r, (avail - lh * s_l - vgap) / rh)
    if s_r < 1.0:
        raise SystemExit(f'6a 오른쪽 칸이 모자란다 (s_r={s_r:.2f})')
    used = lh * s_l + vgap + rh * s_r
    y0 = top + (avail - used) / 2
    place_block(legend, (col_x + (wr - lw * s_l) / 2, y0), s_l, font_s=s_l * FONT_MARGIN)
    place_block(rblock, (col_x + (wr - rw * s_r) / 2, y0 + lh * s_l + vgap), s_r, font_s=s_r * FONT_MARGIN)
    return dict(s_diag=s_d, s_legend=s_l, s_rblock=s_r)


# ───────────────────────────── 네이티브 그래프 (4 장 Origin 그래프 양식) ─────────────────────────────
# 4 장 그래프를 슬라이드에 붙인 실제 크기 (원본 PDF 실측): 눈금 글자 8.1 pt · 축 이름 9.8 pt · 값 6.3 pt · Arial ·
# 파랑 0000FF (기하 / Hertz) · 주황 FF8000 (Tabor 보정 / Physics) · 7:3 빨강 FF0000 · 왼쪽 · 아래 축만 · 바깥 눈금.
BLUE, ORANGE, RED, BLACK = '0000FF', 'FF8000', 'FF0000', '000000'
TICK_PT, AXIS_PT, LABEL_PT = 8.1, 9.8, 6.3


def make_chart(spec, cx, cy):
    """python-pptx 로 선 그래프 한 개를 만들어 (graphicFrame XML, chart XML, 내장 xlsx) 를 돌려준다.
    spec = dict(sym=('σ', 'ion'), unit='mS/cm', h=[…], p=[…], nd=3, ymax, major, numfmt)."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_MARK, XL_TICK_LABEL_POSITION
    from pptx.util import Pt, Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = 9906000, 6858000
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData(number_format='0.' + '0' * spec['nd'])
    cd.categories = list(COMPS)
    cd.add_series('Hertz', spec['h'])
    cd.add_series('Physics', spec['p'])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, 0, 0, Emu(cx), Emu(cy), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.name = 'Arial'
    ch.font.size = Pt(TICK_PT)
    ch.font.color.rgb = RGBColor.from_string(BLACK)
    for ser, col, pos in zip(ch.plots[0].series, (BLUE, ORANGE), (XL_LABEL_POSITION.ABOVE, XL_LABEL_POSITION.BELOW)):
        ser.smooth = False
        ser.format.line.color.rgb = RGBColor.from_string(col)
        ser.format.line.width = Pt(1.0)
        ser.marker.style = XL_MARKER_STYLE.SQUARE
        ser.marker.size = 5
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = RGBColor.from_string(col)
        ser.marker.format.line.color.rgb = RGBColor.from_string(col)
        pt = ser.points[COMPS.index('7:3')]
        pt.marker.style = XL_MARKER_STYLE.SQUARE
        pt.marker.size = 6
        pt.marker.format.fill.solid()
        pt.marker.format.fill.fore_color.rgb = RGBColor.from_string(RED)
        pt.marker.format.line.color.rgb = RGBColor.from_string(RED)
        dl = ser.data_labels
        dl.show_value = True
        dl.position = pos
        dl.number_format = '0.' + '0' * spec['nd']
        dl.number_format_is_linked = False
        dl.font.size = Pt(LABEL_PT)
        dl.font.name = 'Arial'
        dl.font.color.rgb = RGBColor.from_string(BLACK)
    va, ca = ch.value_axis, ch.category_axis
    va.minimum_scale, va.maximum_scale = 0, spec['ymax']
    va.major_unit, va.minor_unit = spec['major'], spec['major'] / 2
    va.has_major_gridlines = False
    va.has_minor_gridlines = False
    va.major_tick_mark = XL_TICK_MARK.OUTSIDE
    va.minor_tick_mark = XL_TICK_MARK.OUTSIDE
    va.tick_labels.number_format = spec['numfmt']
    va.tick_labels.number_format_is_linked = False
    va.tick_labels.font.size = Pt(TICK_PT)
    ca.major_tick_mark = XL_TICK_MARK.OUTSIDE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(TICK_PT)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    for ax in (va, ca):
        ax.format.line.color.rgb = RGBColor.from_string(BLACK)
        ax.format.line.width = Pt(0.75)
    # 축 이름 — 기호는 기울임 + 아래첨자 (덱 본문 σion 과 같은 표기)
    va.has_title = True
    tf = va.axis_title.text_frame
    p = tf.paragraphs[0]
    base, sub = spec['sym']
    for text, kw in ((base, dict(i=True)), (sub, dict(i=True, sub=True)), (f' ({spec["unit"]})', {})):
        r = p.add_run()
        r.text = text
        r.font.size = Pt(AXIS_PT)
        r.font.name = 'Arial'
        r.font.italic = kw.get('i', False)
        r.font.bold = False
        r.font.color.rgb = RGBColor.from_string(BLACK)
        if kw.get('sub'):
            r.font._rPr.set('baseline', '-25000')
    ca.has_title = True
    p = ca.axis_title.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = 'PC:SC (wt%)'
    r.font.bold = False
    r.font.size = Pt(AXIS_PT)
    r.font.name = 'Arial'
    r.font.color.rgb = RGBColor.from_string(BLACK)
    # 그래프 영역 · 그림 영역 = 채움 · 테두리 없음 (4 장 그래프처럼 프레임 없음)
    cs = ch._chartSpace
    C = f'{{{NS["c"]}}}'
    for title in cs.iter(f'{C}title'):     # 축 이름 기본값 = 굵게 → 4 장처럼 보통 굵기
        for dr in title.iter(f'{A}defRPr'):
            dr.set('b', '0')
            dr.set('sz', str(int(AXIS_PT * 100)))
    vt_body = cs.find(f'{C}chart/{C}plotArea/{C}valAx/{C}title/{C}tx/{C}rich/{A}bodyPr')
    vt_body.set('rot', '-5400000')         # 세로축 이름 = 90° 돌림 (PowerPoint 기본값을 명시)
    vt_body.set('vert', 'horz')

    def no_fill_no_line(parent, before_tags):
        sppr = parent.find(f'{C}spPr')
        if sppr is None:
            sppr = etree.SubElement(parent, f'{C}spPr')
            for tag in before_tags:  # 스키마 순서: spPr 은 이 태그들 앞
                hit = parent.find(f'{C}{tag}')
                if hit is not None:
                    hit.addprevious(sppr)
                    break
        for e in list(sppr):
            sppr.remove(e)
        etree.SubElement(sppr, f'{A}noFill')
        ln = etree.SubElement(sppr, f'{A}ln')
        etree.SubElement(ln, f'{A}noFill')

    no_fill_no_line(cs, ('txPr', 'externalData', 'printSettings', 'userShapes', 'extLst'))
    plot_area = cs.find(f'{C}chart/{C}plotArea')
    no_fill_no_line(plot_area, ('extLst',))
    # 그림 영역 테두리 = 축선과 같은 0.75 pt 검정 — 5 장 Origin 그래프처럼 위 · 오른쪽 선까지 닫힌 틀 (원칙 §2 프레임 통일)
    pa_ln = plot_area.find(f'{C}spPr/{A}ln')
    for e in list(pa_ln):
        pa_ln.remove(e)
    pa_ln.set('w', '9525')
    etree.SubElement(etree.SubElement(pa_ln, f'{A}solidFill'), f'{A}srgbClr').set('val', BLACK)
    # 둥근 모서리 끔 (PowerPoint 기본값이 둥근 테두리)
    rc = cs.find(f'{C}roundedCorners')
    if rc is None:
        rc = etree.Element(f'{C}roundedCorners')
        cs.find(f'{C}chart').addprevious(rc)
    rc.set('val', '0')
    buf = io.BytesIO()
    prs.save(buf)
    pk = load_package(buf)
    spart = 'ppt/slides/slide1.xml'
    stree = etree.fromstring(pk[spart])
    gfe = [e for e in tops(stree) if etree.QName(e).localname == 'graphicFrame'][0]
    srels = {r['Id']: r for r in parse_rels(pk, spart)}
    cid = gfe.find('.//c:chart', NS).get(f'{R}id')
    cpart = resolve(spart, srels[cid]['Target'])
    crels = parse_rels(pk, cpart)
    emb = [resolve(cpart, r['Target']) for r in crels if not r.get('TargetMode')]
    if len(emb) != 1:
        raise SystemExit(f'그래프 내장 통합문서가 {len(emb)} 개')
    return copy.deepcopy(gfe), pk[cpart], fixed_xlsx(pk[emb[0]]), [r for r in crels]


def fixed_xlsx(xlsx):
    """내장 통합문서의 작성 · 수정 시각을 고정한다 — 같은 입력이면 같은 출력 바이트 (재현용)."""
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for it in src.infolist():
            data = src.read(it.filename)
            if it.filename == 'docProps/core.xml':
                data = re.sub(rb'(<dcterms:(?:created|modified)[^>]*>)[^<]*', rb'\g<1>2026-10-06T00:00:00Z', data)
            zi = zipfile.ZipInfo(it.filename, date_time=(2026, 10, 6, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)
    return out.getvalue()


def add_chart(sw, spec, x, y, cx, cy, name):
    """make_chart 결과를 이 슬라이드 · 덱에 붙인다 (chart 부분 + 내장 xlsx + 관계)."""
    gfe, cxml, xlsx, crels = make_chart(spec, cx, cy)
    deck = sw.deck
    cpart = deck._new_name('ppt/charts', 'chart', 'xml')
    xpart = deck._new_name('ppt/embeddings', 'Microsoft_Excel_Worksheet', 'xlsx')
    deck._put(cpart, cxml, 'application/vnd.openxmlformats-officedocument.drawingml.chart+xml')
    deck._put(xpart, xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    r0 = crels[0]
    deck.parts[rels_name(cpart)] = rels_xml([dict(Id=r0['Id'], Type=r0['Type'], Target=rel_target(cpart, xpart))])
    rid = next_rid(sw.rels)
    sw.rels.append(dict(Id=rid, Type=RT + 'chart', Target=rel_target('ppt/slides/slide0.xml', cpart)))
    gfe.find('.//c:chart', NS).set(f'{R}id', rid)
    c = gfe.find('.//p:cNvPr', NS)
    c.set('name', name)
    set_box(gfe, x=x, y=y, w=cx, h=cy)
    renumber_ids(gfe, max_shape_id(sw.tree) + 1)
    sptree(sw.tree).append(gfe)
    return gfe


def edit_6b(sw, frame_src, d, legend_src, foot_src):
    """6b = '네트워크 해석 결과 (PC:SC 5 조성)' — 오른쪽 패널: 글머리 · 표 (결과 칸 채움) · 네이티브 그래프 둘 · 각주."""
    t = sw.tree
    grp = find_top(t, '그룹 30')
    tab, title = copy.deepcopy(find_desc(grp, '직사각형 32')), copy.deepcopy(find_desc(grp, 'TextBox 33'))
    idx = sptree(t).index(find_top(t, '그림 26'))
    stub = [n for n in CHART_STUB if n not in CHART_FRAMES]
    remove_tops(t, ['그림 26', '사각형: 둥근 모서리 13', '직사각형 28', 'TextBox 29', '그룹 30', 'TextBox 36',
                    '그룹 504', '예정 표지 10월 계산 예정'] + list(R_BLOCK) + stub)
    # 제목: '네트워크 해석 항목' → '네트워크 해석 결과 (PC:SC 5 조성)' (PC:SC 는 4 장 제목처럼 Aptos)
    p = title.find('.//a:p', NS)
    trpr = first_rpr(title)
    set_para_runs(p, [make_run(trpr, '네트워크 해석 결과 ('), make_run(trpr, 'PC:SC', latin='Aptos'),
                      make_run(trpr, ' 5 조성)')])
    put_frame(sw, frame_src, tab, title, idx)

    # ── 글머리: 원래 두 줄 (9 pt 로 — 덱 글머리 크기) + 데이터가 말하는 것 두 줄 ──
    bul = find_top(t, 'TextBox 80')
    for e in bul.iter(f'{A}rPr', f'{A}endParaRPr'):
        if e.get('sz') == '1000':
            e.set('sz', '900')
    replace_text(bul, '4장의', '5장의', expect=1)
    paras = bul.findall('.//a:p', NS)
    plain = first_rpr(bul, '접촉')
    sym = copy.deepcopy(plain)               # 기호 런 = 덱 본문 σion 표기 (Aptos 기울임 · 아래첨자)
    sym.find('a:latin', NS).set('typeface', 'Aptos')
    cs = sym.find('a:cs', NS)
    if cs is not None:
        cs.set('typeface', 'Cambria Math')
    ih, ip = d['ion_h'], d['ion_p']
    new_bullets = [
        f'Hertz 기준 {{σ|ion}}은 5 조성 중 7:3에서 가장 높다 ({fmt(ih["7:3"], 3)} mS/cm · 0:10의 '
        f'{fmt(ih["7:3"] / ih["0:10"], 1)} 배) — 5장의 Porosity 최소 조성과 같다',
        f'{{σ|e}}는 PC 비율이 클수록 낮아지고 (Hertz {d["e_h"]["0:10"]:.2f} → {d["e_h"]["10:0"]:.2f} mS/cm), '
        f'AM–AM 접촉 수도 {d["amam"]["0:10"]} → {d["amam"]["10:0"]} 개로 함께 줄어든다',
    ]
    for spec in new_bullets:
        np_ = copy.deepcopy(paras[-1])
        set_para_runs(np_, rich_runs(spec, plain, sym))
        paras[-1].addnext(np_)
        paras.append(np_)
    h_bul = fit_textbox(bul, CX0, CY0, CW)

    # ── 표: 결과 칸 채움 · 온폭 · 한 줄 행 ──
    tbl = find_top(t, 'Table 413')
    rows = tbl.findall('.//a:tr', NS)
    cell = lambda r, c: rows[r].findall('a:tc', NS)[c]
    head_rpr = first_rpr(cell(0, 2))
    set_para_runs(cell(0, 2).find('.//a:p', NS), [make_run(head_rpr, '결과 ('), make_run(head_rpr, 'Hertz', latin='Aptos'),
                                                   make_run(head_rpr, ')')])
    body_rpr = first_rpr(cell(1, 0))
    sym_rpr = first_rpr(cell(1, 2), italic=True)
    content = {                             # '약 4 배' = CF_OVER_FULL_HERTZ 3.9–4.2 (derived() 가 범위 확인)
        (1, 2): f'{{σ|ion}} {fmt(ih["0:10"], 3)} → {fmt(ih["7:3"], 3)} mS/cm (7:3 최대)',
        (2, 2): f'{{σ|e}} {d["e_h"]["0:10"]:.2f} → {d["e_h"]["10:0"]:.2f} mS/cm (PC 비율이 클수록 낮아짐)',
        (3, 1): '벌크 + 입계 + 접촉 (constriction) — 접촉마다 직렬',
        (3, 2): '접촉 저항을 빼면 {σ|ion} 약 4 배 (모델 내부 비 · Hertz)',
        (4, 0): '수송 tortuosity',          # 1저자 10-05 밤: τ² 라 쓰지 않는다 (값 = φ_SE·σ₀/σ_ion · 제곱근 아님)
        (4, 1): '{φ|SE} · {σ|0} / {σ|ion} ({σ|0} = 고체전해질 펠릿 3.0 mS/cm)',
        (4, 2): f'{fmt(d["t2"]["0:10"], 2)} (0:10) → {fmt(d["t2"]["7:3"], 2)} (7:3 최소)',
    }
    for (r, c), spec in content.items():
        tc = cell(r, c)
        ps = tc.findall('.//a:p', NS)
        for extra in ps[1:]:
            extra.getparent().remove(extra)
        set_para_runs(ps[0], rich_runs(spec, body_rpr, sym_rpr))
    widths = (1600000, 3700000, CW - 1600000 - 3700000)
    for g, w in zip(tbl.findall('.//a:tblGrid/a:gridCol', NS), widths):
        g.set('w', str(w))
    rows[0].set('h', '201168')
    for r in rows[1:]:
        r.set('h', '240000')
    y_tbl = CY0 + h_bul + 110000
    t_h = 201168 + 240000 * (len(rows) - 1)
    set_box(tbl, x=CX0, y=y_tbl, w=CW, h=t_h)

    # ── 각주 (4 장 각주 양식 복제: 9 pt · 회색) ──
    foot = sw.import_el(*foot_src)
    fp = foot.findall('.//a:p', NS)
    for extra in fp[1:]:
        extra.getparent().remove(extra)
    frpr = first_rpr(foot, 'DEM')
    frpr = copy.deepcopy(frpr)
    frpr.find('a:latin', NS).set('typeface', '맑은 고딕')
    foot_spec = ('Hertz = 기하 접촉 면적 · Physics = Tabor 보정 (5장과 같은 정의) · 값 = 웹앱 접촉망 계산 '
                 '(모델 값 · 실험 앵커 아님) · 조성당 전극 1 개')
    set_para_runs(fp[0], rich_runs(foot_spec, frpr, frpr))
    h_foot = textbox_height(foot, CW)
    set_box(foot, x=CX0, y=CY1 - h_foot + 60000, w=CW, h=h_foot)
    foot_top = CY1 - h_foot + 60000

    # ── 범례 줄 (4 장 범례 묶음 복제 · 가로 배치) — 두 그래프 공통, 표와 그래프 상자 사이 가운데 ──
    legend_top = y_tbl + t_h + 90000
    LEG_H = 230832
    lg = sw.import_el(*legend_src)
    texts = [e for e in lg.iter(f'{P}sp')]
    lines = [e for e in lg.iter(f'{P}cxnSp')]
    for sp_, label in zip(texts, ('Hertz (기하 접촉 면적)', 'Physics (Tabor 보정)')):
        p_ = sp_.find('.//a:p', NS)
        rpr = first_rpr(sp_)
        set_para_runs(p_, rich_runs(label, rpr, rpr, latin_words=False))
    gx = lg.find('p:grpSpPr/a:xfrm', NS)
    cho = gx.find('a:chOff', NS)
    ox, oy = int(cho.get('x')), int(cho.get('y'))
    seg, col2 = 186217, 1700000
    for k, (ln_, sp_) in enumerate(zip(lines, texts)):
        lx = ox + k * col2
        set_box(ln_, x=lx, y=oy + LEG_H // 2)
        set_box(sp_, x=lx + seg + 40000, y=oy, w=col2 - seg - 80000, h=LEG_H)
        ln_.find('p:spPr/a:ln', NS).set('w', '12700')      # 그래프 선 굵기 (1 pt) 와 같게
    total_w = 2 * col2
    gx.find('a:chExt', NS).set('cx', str(total_w))
    gx.find('a:chExt', NS).set('cy', str(LEG_H))
    set_box(lg, x=CX0 + (CW - total_w) / 2, y=legend_top, w=total_w, h=LEG_H)

    # ── 그래프 둘: 6 장 원래 틀 (흰 상자 · 제목) 을 키워 그 안에 네이티브 선 그래프 ──
    gap = 300000
    top = legend_top + LEG_H + 50000
    bw = (CW - gap) / 2
    bh = foot_top - 60000 - top
    specs = (
        dict(frame='Rectangle 414', title='TextBox 471', name='그래프 σion (Hertz · Physics)', sym=('σ', 'ion'),
             unit='mS/cm', h=[ih[c] for c in COMPS], p=[ip[c] for c in COMPS], nd=3, ymax=0.22, major=0.05,
             numfmt='[=0]0;0.00'),
        dict(frame='Rectangle 431', title='TextBox 488', name='그래프 σe (Hertz · Physics)', sym=('σ', 'e'),
             unit='mS/cm', h=[d['e_h'][c] for c in COMPS], p=[d['e_p'][c] for c in COMPS], nd=2, ymax=6, major=2,
             numfmt='0'),
    )
    for i, spec in enumerate(specs):
        x = CX0 + i * (bw + gap)
        fr, ti = find_top(t, spec['frame']), find_top(t, spec['title'])
        set_box(fr, x=x, y=top, w=bw, h=bh)
        th_ = box(ti)[3]
        set_box(ti, x=x + 45720, y=top + 36576, w=bw - 91440)
        cy_top = top + 36576 + th_ + 20000
        add_chart(sw, spec, round(x + 60000), round(cy_top), round(bw - 120000), round(top + bh - 40000 - cy_top),
                  spec['name'])
    return dict(table_rows=len(rows))


# 6 장 오른쪽 아래 '빈 그래프' 자리 (손으로 그린 축 · 눈금 · 축 이름) — 6b 에서 네이티브 그래프로 바꾼다
CHART_STUB = (['Rectangle 414', 'TextBox 471', 'Rectangle 431', 'TextBox 488'] +
              [f'Connector {n}' for n in (416, 418, 420, 422, 424, 426, 427, 433, 435, 437, 439, 441, 443, 444)] +
              [f'TextBox {n}' for n in (473, 475, 477, 479, 481, 484, 485, 490, 492, 494, 496, 498, 501, 502)])
CHART_FRAMES = ('Rectangle 414', 'TextBox 471', 'Rectangle 431', 'TextBox 488')
# 6 장 왼쪽 아래 '접촉 하나의 저항 = 직렬 연결' 상자
R_BLOCK = (['Rectangle 396', 'TextBox 106', 'Oval 398', 'Oval 399', 'Oval 401', 'Oval 402',
            'Freeform 11', 'Freeform 13', 'Freeform 14', 'Freeform 15', 'Oval 403', 'Oval 404'] +
           [f'TextBox {n}' for n in range(118, 125)] + ['Freeform 10', 'Connector 400'])


# ───────────────────────────── 메모 ─────────────────────────────
def notes_body(notes):
    for sp in notes.iter(f'{P}sp'):
        ph = sp.find('.//p:nvPr/p:ph', NS)
        if ph is not None and ph.get('type') == 'body':
            return sp
    raise SystemExit('메모 본문 자리 없음')


def set_notes_text(notes, paragraphs):
    """메모 본문을 문단 목록으로 바꾼다 — 원래 첫 런을 복제 (메모 런에는 rPr 이 없다)."""
    body = notes_body(notes)
    tx = body.find('p:txBody', NS)
    ps = tx.findall('a:p', NS)
    run0 = copy.deepcopy(ps[0].find('a:r', NS))
    for p in ps[1:]:
        tx.remove(p)
    tmpl = ps[0]
    for i, text in enumerate(paragraphs):
        p = tmpl if i == 0 else copy.deepcopy(tmpl)
        r = copy.deepcopy(run0)
        r.find('a:t', NS).text = text
        set_para_runs(p, [r])
        if i:
            tx.append(p)


NOTES_6A = [
    '접촉망 해석 = 입자 = 노드 · 접촉 = 저항 · Kirchhoff 로 푼다.  접촉 저항 = 벌크 + 접촉 (협착, R_c = 1/(2σa)) + 입계 '
    '— 접촉마다 직렬.',
    '이 해석에서 나오는 수송 지표 = 유효 전도도 σ 와 수송 tortuosity = φ_SE·σ₀/σ_ion (다음 장 · 제곱근 아님 = 문헌의 tortuosity '
    'factor) — 3 장의 Tortuosity '
    '(기하학적 · SE 접촉망 최단 경로) 와 다른 양이다.',
]


def notes_6b(d):
    """6b 메모 — 값의 출처 · 정의 · 한정어 (화면에 다 못 싣는 것)."""
    ip, ep = d['ion_p'], d['e_p']
    cf = ' · '.join(f'{CF_OVER_FULL_HERTZ[c]:.1f}' for c in COMPS)
    return [
        '값 = 웹앱 케이스 페이지의 접촉망 솔버 값 (FULL · 업로드 때 09-22 · 09-25 계산) — 1저자 전사 10-05, '
        'docs/data/ps45_r45_network_webapp_20261005/network_5comp.tsv.  Hertz = LIGGGHTS 기하 겹침 원판 면적 · '
        'Physics = Tabor · 부피 · 기하 상한 보정 (5 장 Coverage 와 같은 정의).',
        f'σ_contact-free / σ_full (이온 · Hertz) = {cf} (0:10 → 10:0) — 같은 망에서 접촉 저항 항만 뺀 모델 내부 비 '
        '(실험 대비 오차가 아니다).',
        '수송 tortuosity = φ_SE·σ₀/σ_ion (σ₀ = 3.0 mS/cm 펠릿값 · 제곱근 아님 = 문헌의 tortuosity factor) — 3 장의 Tortuosity '
        '(기하학적) 와 다른 양.',
        f'"7:3 최대 · 최소" 는 Hertz 기준이다 — Physics 로는 σ_ion 이 7:3 ({fmt(ip["7:3"], 4)}) · 10:0 '
        f'({fmt(ip["10:0"], 4)}) 같은 수준, σ_e 는 7:3 ({ep["7:3"]:.2f}) 이 10:0 ({ep["10:0"]:.2f}) 보다 낮고, '
        f'수송 tortuosity 최소는 {d["t2_p_min"]} ({fmt(d["t2_p"][d["t2_p_min"]], 2)}).',
        '조성당 전극 1 개 (시드 반복 없음) — 7:3 과 10:0 의 차이가 전극 간 산포보다 큰지는 미확인.  σ 는 접촉망 모델 값 '
        '(실험 앵커 아님).  그래프 = PowerPoint 차트 (오른쪽 클릭 → 데이터 편집).',
    ]


# ───────────────────────────── 조립 ─────────────────────────────
def build_edited(deck, pkgs, srcs, a):
    tsv = read_tsv(a.tsv)
    d = derived(tsv)
    src = [(pkgs[k], part, sid) for k, part, sid in srcs]
    p15 = slide_parts(pkgs['part15'])[0]
    frame_src = (pkgs['part15'], p15, find_top(etree.fromstring(pkgs['part15'][p15]), FRAME_NAME))
    p2s = slide_parts(pkgs['part2'])
    s4 = etree.fromstring(pkgs['part2'][p2s[0]])
    legend_src = (pkgs['part2'], p2s[0], find_desc(find_top(s4, '그룹 43'), '그룹 25'))
    foot_src = (pkgs['part2'], p2s[0], find_top(s4, 'TextBox 42'))
    new_id = max(sid for *_, sid in srcs) + 1
    info = {}

    # 1 장 → 1a · 1b
    pk, part, sid = src[0]
    w = SlideWork(deck, pk, part)
    info['1a'] = edit_1a(w, frame_src)
    w.add_to(sid)
    w = SlideWork(deck, pk, part)
    info['1b'] = edit_1b(w, frame_src)
    w.add_to(new_id)
    new_id += 1
    # 2 장 — 장 번호 참조 (2 · 4 · 5 → 3 · 5 · 6)
    pk, part, sid = src[1]
    w = SlideWork(deck, pk, part)
    replace_text(w.tree, '(2 · 4 · 5 장 공통)', '(3 · 5 · 6 장 공통)', expect=1)
    w.add_to(sid)
    # 3 장 — 천 단위 콤마 (보고자료 원칙 §6) · "1장과 같다" 는 1a 그대로라 그대로
    pk, part, sid = src[2]
    w = SlideWork(deck, pk, part)
    for old in ('4,588', '3,209', '2,293', '1,376'):
        replace_text(w.tree, old, old.replace(',', ''), expect=1)
    w.add_to(sid)
    # 4 장 — "2장의 네 지표" → 3장
    pk, part, sid = src[3]
    w = SlideWork(deck, pk, part)
    replace_text(w.tree, '2장의 네 지표', '3장의 네 지표', expect=1)
    w.add_to(sid)
    # 5 장 — 메모 "4 장과 같은 계산" → 5 장
    pk, part, sid = src[4]
    w = SlideWork(deck, pk, part)
    replace_text(w.notes, '4 장과 같은 계산', '5 장과 같은 계산', expect=1)
    w.add_to(sid)
    # 6 장 → 6a · 6b
    pk, part, sid = src[5]
    w = SlideWork(deck, pk, part)
    info['6a'] = edit_6a(w, frame_src)
    set_notes_text(w.notes, NOTES_6A)
    w.add_to(sid)
    w = SlideWork(deck, pk, part)
    info['6b'] = edit_6b(w, frame_src, d, legend_src, foot_src)
    set_notes_text(w.notes, notes_6b(d))
    w.add_to(new_id)
    for k, v in info.items():
        print(k, {kk: (round(vv, 3) if isinstance(vv, float) else vv) for kk, vv in v.items()})
    return info


# ───────────────────────────── 실행 ─────────────────────────────
PART_ORDER = ('part1', 'part15', 'part175', 'part2')


def collect_sources(pkgs):
    """원본 순서 1–6 장: [(조각 이름, 슬라이드 partname, sldId)]."""
    out = []
    for key in PART_ORDER:
        pk = pkgs[key]
        out += [(key, part, sid) for part, sid in zip(slide_parts(pk), slide_ids(pk))]
    if len(out) != 6:
        raise SystemExit(f'원본 장 수가 6 이 아니다: {len(out)}')
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--part1', required=True)
    ap.add_argument('--part15', required=True)
    ap.add_argument('--part175', required=True, help='pptx 또는 pptx 하나를 담은 zip')
    ap.add_argument('--part2', required=True)
    ap.add_argument('--tsv', default='docs/data/ps45_r45_network_webapp_20261005/network_5comp.tsv')
    ap.add_argument('--out', required=True)
    ap.add_argument('--merge-only', action='store_true', help='손보기 없이 6 장만 합친다 (대조용)')
    a = ap.parse_args(argv)
    pkgs = {k: load_package(getattr(a, k)) for k in PART_ORDER}
    n_tmpl = check_template(pkgs)
    srcs = collect_sources(pkgs)
    deck = Deck(pkgs['part1'])
    if a.merge_only:
        for key, part, sid in srcs:
            pk = pkgs[key]
            tree = etree.fromstring(pk[part])
            rels = deck.convert_rels(pk, part, parse_rels(pk, part))
            npart = notes_part(pk, part)
            deck.add_slide(tree, rels, sid, etree.fromstring(pk[npart]) if npart else None)
    else:
        build_edited(deck, pkgs, srcs, a)
    deck.finish(pkgs['part2']['docProps/app.xml'])
    deck.save(a.out)
    if getattr(deck, 'dropped', None):
        print('쓰이지 않아 뺀 부분:', ', '.join(deck.dropped))
    print(f'양식 부분 {n_tmpl} 개 동일 확인 · 장 {len(deck.slides)} · 부분 {len(deck.parts)} → {a.out}')


if __name__ == '__main__':
    main()
