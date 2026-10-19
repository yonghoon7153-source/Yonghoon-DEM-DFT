#!/usr/bin/env python3
"""Weekly-report slide (lab deck style of the 2026-10-12 report) — Nd Q3: Bader charges and bond-midpoint ELF.

    python3 tools/figures/weekly_nd_q3_slide_2026_10_19.py [--out docs/Weekly_2026_10_19_Nd_Q3_charge_ELF.pptx]

Numbers are read from the recovered raw files, not typed in:
    db/raw/nd_q3_bader_2026_10_19/nd_bader_summary.json   (Bader net charge = ZVAL - N, per site class)
    db/raw/nd_q3_elf_2026_10_19/nd_elf_bonds_mic.csv      (ELF at bond midpoints, minimum-image fixed re-read)
The host Li charge (+0.883, LPSCl1.6, same method) comes from the runlog text
(db/properties/nd_q3_bader_elf_runlog_2026_10_19.json · Bader_결과 · 사전등록_비교_그대로).

What this slide does not say: which Nd site is preferred (fixed-cell + EOS still running), decomposition
stability, or anything about transport. The S in PS4 site charges are left out on purpose (P/S partition
artifact in this Bader run); the chart shows only the site classes without that artifact.
"""
import argparse
import csv
import json
import pathlib

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt, Inches, Emu

REPO = pathlib.Path(__file__).resolve().parents[2]
BADER = REPO / "db/raw/nd_q3_bader_2026_10_19/nd_bader_summary.json"
ELF = REPO / "db/raw/nd_q3_elf_2026_10_19/nd_elf_bonds_mic.csv"
HOST_LI = 0.883

KO = "맑은 고딕"
NAVY, TITLE, RED, BUL, BLUE, GRAY = "001F5F", "1F497D", "A50021", "404040", "0070C0", "7F7F7F"
C_AWAY, C_NEAR = "8EA9C1", "C00000"


def rgb(h):
    return RGBColor.from_string(h)


def num(x, nd=2):
    """text number with a true minus sign (U+2212) — chart labels keep Excel's own format."""
    return f"{x:.{nd}f}".replace("-", "\u2212")


def style_run(run, size, color="000000", bold=False, latin=KO, ea=KO, sub=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = latin
    rpr = run._r.get_or_add_rPr()
    for old in rpr.findall(qn("a:ea")):
        rpr.remove(old)
    etree.SubElement(rpr, qn("a:ea"), typeface=ea)
    if sub:
        rpr.set("baseline", "-25000")


def add_runs(p, parts, size, color="000000", bold=False):
    """parts: list of (text, opts) — opts keys: color, bold, sub."""
    for text, o in parts:
        r = p.add_run()
        r.text = text
        style_run(r, size, o.get("color", color), o.get("bold", bold), sub=o.get("sub", False))


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    return tb, tf


def bullet(p, indent_in=0.3, space_before=6):
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", str(Emu(Inches(indent_in))))
    ppr.set("indent", str(-Emu(Inches(indent_in))))
    spc = etree.SubElement(ppr, qn("a:spcBef"))
    etree.SubElement(spc, qn("a:spcPts"), val=str(int(space_before * 100)))
    bc = etree.SubElement(ppr, qn("a:buClr"))
    etree.SubElement(bc, qn("a:srgbClr"), val=BUL)
    etree.SubElement(ppr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(ppr, qn("a:buChar"), char="•")


def slide_number(slide, x, y, w, h):
    tb, tf = textbox(slide, x, y, w, h, MSO_ANCHOR.BOTTOM)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"), lang="en-US", sz="1200")
    sf = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val="888888")
    etree.SubElement(rpr, qn("a:latin"), typeface="Calibri")
    t = etree.SubElement(fld, qn("a:t"))
    t.text = "‹#›"


def style_chart(ch, ytitle):
    ch.font.size = Pt(9)
    ch.font.name = "Arial"
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(9)
    va = ch.value_axis
    va.has_title = True
    va.axis_title.text_frame.text = ytitle
    r = va.axis_title.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(9)
    r.font.bold = False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("BFBFBF")
    va.major_gridlines.format.line.dash_style = MSO_LINE.DASH
    va.format.line.color.rgb = rgb("000000")
    ch.category_axis.format.line.color.rgb = rgb("000000")
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW


def load():
    b = json.loads(BADER.read_text(encoding="utf-8"))
    cat, sp = b["sites"]["per_category"], b["per_species"]
    rows = {r["bond"]: r for r in csv.DictReader(ELF.open(encoding="utf-8"))}
    return cat, sp, rows


def build(out):
    cat, sp, elf = load()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(Pt(780)), Emu(Pt(540))
    s = prs.slides.add_slide(prs.slide_layouts[6])

    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(Pt(15)), Emu(Pt(60)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(NAVY)
    bar.line.fill.background()

    _, tf = textbox(s, 0.66, 0.24, 7.0, 0.55)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "2. "; style_run(r, 24, TITLE, True, latin="Arial")
    r = p.add_run(); r.text = "DFT 관련"; style_run(r, 24, TITLE, True)

    _, tf = textbox(s, 1.05, 1.2, 9.4, 0.4)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "■  "; style_run(r, 12, RED, latin="Arial")
    add_runs(p, [("Nd 주변 전자 분포 — Bader 전하 · 결합 중점 ELF (Nd@Li 모델)", {})], 16, bold=True)

    sF, sFn = cat["S_free"]["net_mean"], cat["S_free_nearNd"]["net_mean"]
    cl, cln = cat["Cl"]["net_mean"], cat["Cl_nearNd"]["net_mean"]
    nd, li = sp["Nd"]["mean"], sp["Li"]["mean"]
    nd_x = [float(elf[k]["ELF_midpoint"]) for k in ("Nd-S", "Cl-Nd", "Nd-O")]
    ps = float(elf["P-S"]["ELF_midpoint"])
    _, tf = textbox(s, 1.3, 1.62, 9.2, 1.75)
    lines = [
        [("Nd", {}), ("2", {"sub": True}), ("O", {}), ("3", {"sub": True}),
         ("@LPSCl1.6 (Nd@Li · 120 원자) · QE PBE · Nd 4f-in-core — 같은 셀에서 Nd 옆 / 먼 음이온 비교", {})],
        [("Nd 옆 음이온이 전자를 덜 가짐 (자유 S ", {}), (num(sFn), {"color": BLUE}), (f" vs {num(sF)} e · Cl ", {}),
         (num(cln), {"color": BLUE}), (f" vs {num(cl)} e) → 전자가 음이온 쪽으로 더 가지 않음", {})],
        [("Nd ", {}), (f"+{nd:.2f}", {"color": BLUE}), (f" (형식 +3 의 {100 * nd / 3:.0f} %) · Li +{li:.2f} (호스트와 같음)"
                                                     " → Nd 는 전자를 다 내주지 않고 일부를 음이온 사이에 남김", {})],
        [("Nd–X 결합 중점 ELF ", {}), (f"{min(nd_x):.2f}–{max(nd_x):.2f}", {"color": BLUE}),
         (f" (P–S {ps:.2f}) → 공유 전자쌍은 없음 · 주로 이온 결합, Li–X 보다 공유성이 조금 더", {})],
    ]
    for i, parts in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        bullet(p, 0.3, 0 if i == 0 else 5)
        add_runs(p, parts, 12)

    # left chart — Bader charge, site classes without the P/S partition artifact
    cd = CategoryChartData()
    cd.categories = ["Free S", "Cl", "O (P–O)"]
    cd.add_series("Away from Nd", [sF, cl, cat["O_P"]["net_mean"]], number_format="0.00")
    cd.add_series("Next to Nd", [sFn, cln, cat["O_P_nearNd"]["net_mean"]], number_format="0.00")
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.85), Inches(3.4), Inches(4.65), Inches(2.95), cd).chart
    style_chart(ch, "Bader charge (e)")
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = -2.0, 0.0, 0.5
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.0", False
    pl = ch.plots[0]
    pl.gap_width, pl.overlap = 70, -10
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = "0.00", False
    dl.font.size = Pt(8)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for ser, col in zip(pl.series, (C_AWAY, C_NEAR)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(col)

    # right chart — ELF at bond midpoints (Nd bonds in red)
    order = [("P-S", "P–S"), ("O-P", "O–P"), ("Li-S", "Li–S"), ("Cl-Li", "Li–Cl"), ("Li-O", "Li–O"),
             ("Nd-S", "Nd–S"), ("Cl-Nd", "Nd–Cl"), ("Nd-O", "Nd–O")]
    cd = CategoryChartData()
    cd.categories = [lab for _, lab in order]
    v = [float(elf[k]["ELF_midpoint"]) for k, _ in order]
    isnd = ["Nd" in k for k, _ in order]
    cd.add_series("Host bonds", [x if not n else 0 for x, n in zip(v, isnd)], number_format="0.00")
    cd.add_series("Nd bonds", [x if n else 0 for x, n in zip(v, isnd)], number_format="0.00")
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.6), Inches(3.4), Inches(4.65), Inches(2.95), cd).chart
    style_chart(ch, "ELF at bond midpoint")
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0.0, 1.0, 0.5
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.0", False
    pl = ch.plots[0]
    pl.gap_width, pl.overlap = 60, 100
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = "0.00;;;", False
    dl.font.size = Pt(8)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for ser, col in zip(pl.series, ("4F81BD", C_NEAR)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(col)

    for x, text in ((0.85, "< Bader 전하 — Nd 옆 / 먼 음이온 >"), (5.6, "< 결합 중점 ELF (0.5 = 균일 전자기체) >")):
        _, tf = textbox(s, x, 6.38, 4.65, 0.3)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, [(text, {})], 12)

    _, tf = textbox(s, 0.85, 6.74, 9.4, 0.24)
    p = tf.paragraphs[0]
    n_near = cat["S_free_nearNd"]["n"]
    add_runs(p, [(f"※ Nd 2 개 · Nd 옆 음이온 각 {n_near} 개 · 구조 1 개 · 4f-in-core PP 로 낸 값 — "
                  "Nd 자리 선호 (EOS 진행 중) 와 분해 안정성은 따로", {})], 9, color=GRAY)

    _, tf = textbox(s, 0.24, 7.0, 3.5, 0.3, MSO_ANCHOR.BOTTOM)
    r = tf.paragraphs[0].add_run(); r.text = "Battery Materials Laboratory"; style_run(r, 12, GRAY, latin="Calibri", ea="Calibri")
    slide_number(s, 9.9, 7.0, 0.6, 0.3)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(REPO / "docs/Weekly_2026_10_19_Nd_Q3_charge_ELF.pptx"))
    a = ap.parse_args()
    print("->", build(pathlib.Path(a.out)))


if __name__ == "__main__":
    main()
