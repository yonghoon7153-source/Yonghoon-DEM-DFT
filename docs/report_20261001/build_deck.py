#!/usr/bin/env python3
"""이종기술 3세부 2차년도 진도점검 — 안용훈 담당 (DEM 미세구조 · 건식 믹싱) 슬라이드 초안 생성기 (2026-09-29).

    python3 docs/report_20261001/build_deck.py --template <3세부 템플릿.pptx> --out docs/report_20261001/report_20261001_draft.pptx

- 템플릿의 슬라이드 8 (2-1 · 2차년도 연구개발 추진 내용 · 자유 양식) 을 복제해 머리 (번호 · 제목 · 과제명) 와 서식을 그대로 쓴다.
  템플릿의 빈 추진 내용 슬라이드 8 · 9 는 빼고, 나머지 (1–7 · 10–11) 는 손대지 않는다.
- 수식 = PowerPoint 네이티브 수식 (OMML · a14:m) — PowerPoint 에서 수식을 눌러 바로 고칠 수 있다.
- 차트 = 네이티브 차트 (마우스 오른쪽 → 데이터 편집).  표 = 네이티브 표.
- 값의 출처: `docs/data/ps45_r45_union_20260929/r45_union_summary.tsv` (사양 5 조성 union) ·
  `docs/data/lhs_handover_20260928.csv` (130 설계점 — 경향은 이 스크립트가 계산해 `lhs_trends.json` 에 같이 쓴다) ·
  덱 `dem_scripts/ps_sweep_6mah_20260914/in.ps_7_3_r45.liggghts` (물성) · 믹서 `scripts/make_mixer_deck.py` · `scripts/measure_mixing_index.py` (식).
- ⬜ 칸 = 아직 값이 없는 자리 (전도도 · 굴곡도 = 웹앱 값 · 믹싱 판정 = 09-30 완주 뒤).
"""
import argparse
import copy
import csv
import json
import os
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt, Emu

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
MNS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
FONT = '맑은 고딕'
BLUE, ORANGE, GRAY, DARK = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x9A, 0xA0, 0xA8), RGBColor(0x26, 0x26, 0x26)

# ───────────────────────────── OMML (PowerPoint 수식) ─────────────────────────────

def _esc(t):
    return t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


class Eq:
    """작은 OMML 조립기 — 각 함수는 OMML 조각 문자열을 돌려준다."""

    def __init__(self, sz=1300):
        self.sz = sz

    def r(self, t, sty=None):
        mrpr = f'<m:rPr><m:sty m:val="{sty}"/></m:rPr>' if sty else ''
        return (f'<m:r>{mrpr}<a:rPr lang="en-US" altLang="ko-KR" sz="{self.sz}" dirty="0">'
                f'<a:latin typeface="Cambria Math" panose="02040503050406030204" pitchFamily="18" charset="0"/></a:rPr>'
                f'<m:t>{_esc(t)}</m:t></m:r>')

    def t(self, t):          # 똑바른 글자 (단위 · 이름)
        return self.r(t, 'p')

    def b(self, t):          # 굵은 똑바른 글자 (벡터)
        return self.r(t, 'b')

    @staticmethod
    def sub(e, s):
        return f'<m:sSub><m:e>{e}</m:e><m:sub>{s}</m:sub></m:sSub>'

    @staticmethod
    def sup(e, s):
        return f'<m:sSup><m:e>{e}</m:e><m:sup>{s}</m:sup></m:sSup>'

    @staticmethod
    def subsup(e, s, p):
        return f'<m:sSubSup><m:e>{e}</m:e><m:sub>{s}</m:sub><m:sup>{p}</m:sup></m:sSubSup>'

    @staticmethod
    def frac(n, d):
        return f'<m:f><m:num>{n}</m:num><m:den>{d}</m:den></m:f>'

    @staticmethod
    def nary(ch, lo, e, hi=None):
        hide = '' if hi else '<m:supHide m:val="1"/>'
        return (f'<m:nary><m:naryPr><m:chr m:val="{ch}"/><m:limLoc m:val="undOvr"/>{hide}</m:naryPr>'
                f'<m:sub>{lo}</m:sub><m:sup>{hi or ""}</m:sup><m:e>{e}</m:e></m:nary>')

    @staticmethod
    def rad(e):
        return f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{e}</m:e></m:rad>'

    @staticmethod
    def par(e, beg='(', end=')'):
        return f'<m:d><m:dPr><m:begChr m:val="{beg}"/><m:endChr m:val="{end}"/></m:dPr><m:e>{e}</m:e></m:d>'

    @staticmethod
    def cases(*rows):
        inner = ''.join(f'<m:e>{x}</m:e>' for x in rows)
        return (f'<m:d><m:dPr><m:begChr m:val="{{"/><m:endChr m:val=""/></m:dPr>'
                f'<m:e><m:eqArr>{inner}</m:eqArr></m:e></m:d>')


def add_equations(slide, x, y, w, h, eqs, sz=1300, gap_pt=4):
    """eqs = [(OMML 조각, 대체 텍스트), …] — 한 상자에 수식 문단 여러 개.  mc:AlternateContent (Choice = a14:m 수식 · Fallback = 평문)."""
    tree = slide.shapes._spTree
    nid = max(int(e.get('id')) for e in tree.iter('{%s}cNvPr' % NS['p'])) + 1

    def body(math):
        paras = []
        for omml, alt in eqs:
            if math:
                paras.append(f'<a:p><a:pPr><a:spcBef><a:spcPts val="{gap_pt * 100}"/></a:spcBef></a:pPr><a14:m>'
                             f'<m:oMathPara><m:oMathParaPr><m:jc m:val="left"/></m:oMathParaPr><m:oMath>{omml}</m:oMath></m:oMathPara>'
                             f'</a14:m><a:endParaRPr lang="en-US" altLang="ko-KR" sz="{sz}" dirty="0"/></a:p>')
            else:
                paras.append(f'<a:p><a:pPr><a:spcBef><a:spcPts val="{gap_pt * 100}"/></a:spcBef></a:pPr><a:r><a:rPr lang="en-US" altLang="ko-KR" sz="{sz}" dirty="0">'
                             f'<a:latin typeface="Cambria Math"/></a:rPr><a:t>{_esc(alt)}</a:t></a:r></a:p>')
        return ''.join(paras)

    def sp(math):
        return (f'<p:sp><p:nvSpPr><p:cNvPr id="{nid}" name="수식 {nid}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
                f'<p:spPr><a:xfrm><a:off x="{int(x)}" y="{int(y)}"/><a:ext cx="{int(w)}" cy="{int(h)}"/></a:xfrm>'
                f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
                f'<p:txBody><a:bodyPr wrap="square" lIns="72000" tIns="36000" rIns="36000" bIns="36000" rtlCol="0"><a:noAutofit/></a:bodyPr>'
                f'<a:lstStyle/>{body(math)}</p:txBody></p:sp>')

    xml = (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           f'xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}" xmlns:m="{MNS}">'
           f'<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">{sp(True)}</mc:Choice>'
           f'<mc:Fallback>{sp(False)}</mc:Fallback></mc:AlternateContent>')
    tree.append(etree.fromstring(xml))


# ───────────────────────────── 슬라이드 복제 · 정리 ─────────────────────────────

def dup_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith(('/slideLayout', '/notesSlide')):
            continue
        rid_map[rid] = (new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                        else new.part.relate_to(rel.target_part, rel.reltype))
    src_csld = src._element.find('{%s}cSld' % NS['p'])
    bg = src_csld.find('{%s}bg' % NS['p'])
    if bg is not None:
        new._element.find('{%s}cSld' % NS['p']).insert(0, copy.deepcopy(bg))
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(('}nvGrpSpPr', '}grpSpPr')):
            continue
        new.shapes._spTree.append(copy.deepcopy(el))
    for el in new.shapes._spTree.iter():
        for k in list(el.attrib):
            if k.startswith('{%s}' % NS['r']) and el.attrib[k] in rid_map:
                el.attrib[k] = rid_map[el.attrib[k]]
    return new


def shape(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    raise KeyError(name)


def drop(slide, *names):
    for n in names:
        s = shape(slide, n)
        s._element.getparent().remove(s._element)


def set_runs_text(tf, text):
    """첫 문단 · 첫 run 의 서식을 유지하고 글자만 바꾼다."""
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0]
    r0.text = text
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def set_bullets(box, items, size=12):
    """직사각형 14 (연파랑 글머리표 상자) — 첫 문단을 틀로 글머리표 문단을 만든다.  item = 문자열 또는 [(글자, 굵게?), …]."""
    tf = box.text_frame
    proto = copy.deepcopy(tf.paragraphs[0]._p)
    body = tf._txBody
    for p in list(tf.paragraphs):
        body.remove(p._p)
    for it in items:
        p = copy.deepcopy(proto)
        runs = p.findall('{%s}r' % NS['a'])
        rpr = copy.deepcopy(runs[0].find('{%s}rPr' % NS['a']))
        for r in runs:
            p.remove(r)
        for extra in p.findall('{%s}endParaRPr' % NS['a']):
            p.remove(extra)
        segs = [(it, False)] if isinstance(it, str) else it
        for txt, bold in segs:
            r = etree.SubElement(p, '{%s}r' % NS['a'])
            rp = copy.deepcopy(rpr)
            rp.set('sz', str(size * 100))
            rp.set('b', '1' if bold else '0')
            r.append(rp)
            t = etree.SubElement(r, '{%s}t' % NS['a'])
            t.text = txt
        body.append(p)


def text(slide, x, y, w, h, s, size=11, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, Cm(0.1))
    lines = s if isinstance(s, list) else [s]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color
    if fill is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill
    return tb


def label(slide, x, y, w, s):
    return text(slide, x, y, w, Cm(0.7), f'[ {s} ]', size=12, bold=True, align=PP_ALIGN.CENTER)


def placeholder(slide, x, y, w, h, s):
    from pptx.enum.shapes import MSO_SHAPE
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    b.fill.solid()
    b.fill.fore_color.rgb = RGBColor(0xF3, 0xF4, 0xF6)
    b.line.color.rgb = GRAY
    b.line.dash_style = 4  # dash
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = s if isinstance(s, list) else [s]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = ln
        r.font.size, r.font.name = Pt(10.5), FONT
        r.font.color.rgb = RGBColor(0x6B, 0x70, 0x78)
    return b


def frame(slide, x, y, w, h):
    from pptx.enum.shapes import MSO_SHAPE
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    b.adjustments[0] = 0.04
    b.fill.background()
    b.line.color.rgb = RGBColor(0xC9, 0xD3, 0xE6)
    b.line.width = Pt(1)
    b.shadow.inherit = False
    return b


def table(slide, x, y, w, rows, col_w=None, size=10, row_h=0.62, head_fill=RGBColor(0x2F, 0x4B, 0x7C)):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, x, y, w, Cm(row_h * nr))
    tb = gt.table
    if col_w:
        tot = sum(col_w)
        for j, cw in enumerate(col_w):
            tb.columns[j].width = int(w * cw / tot)
    for i, row in enumerate(rows):
        tb.rows[i].height = Cm(row_h)
        for j, v in enumerate(row):
            c = tb.cell(i, j)
            c.margin_left = c.margin_right = Cm(0.1)
            c.margin_top = c.margin_bottom = Cm(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(v)
            r.font.size, r.font.name = Pt(size), FONT
            r.font.bold = (i == 0)
            if i == 0:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                c.fill.solid()
                c.fill.fore_color.rgb = head_fill
            else:
                c.fill.solid()
                c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else RGBColor(0xF2, 0xF5, 0xFB)
                if str(v).startswith('⬜'):
                    r.font.color.rgb = ORANGE
    return tb


def chart(slide, x, y, w, h, cats, series, ytitle, ymin=None, ymax=None, fmt='0.0', colors=None, point_colors=None,
          legend=True, label_pos=XL_LABEL_POSITION.OUTSIDE_END, xtitle=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, cd)
    ch = gf.chart
    ch.has_title = False
    ch.font.name, ch.font.size = FONT, Pt(10)
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
        ch.legend.font.size = Pt(10)
    va = ch.value_axis
    if ymin is not None:
        va.minimum_scale = ymin
    if ymax is not None:
        va.maximum_scale = ymax
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE6, 0xEB)
    va.tick_labels.font.size = Pt(10)
    va.has_title = True
    va.axis_title.text_frame.text = ytitle
    va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    va.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(10)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if xtitle:
        ca.has_title = True
        ca.axis_title.text_frame.text = xtitle
        ca.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
        ca.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.font.size = Pt(9)
    dl.position = label_pos
    colors = colors or [BLUE, GRAY, ORANGE]
    for k, s in enumerate(ch.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = colors[k % len(colors)]
    if point_colors:
        for idx, col in point_colors.items():
            pt = ch.series[0].points[idx]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = col
    return ch


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ───────────────────────────── 데이터 ─────────────────────────────

def load_r45():
    rows = list(csv.DictReader(open(ROOT / 'docs/data/ps45_r45_union_20260929/r45_union_summary.tsv', encoding='utf-8'), delimiter='\t'))
    return {r['P_S']: r for r in rows}


def lhs_trends():
    rows = list(csv.DictReader(open(ROOT / 'docs/data/lhs_handover_20260928.csv', encoding='utf-8')))
    f = lambda r, k: float(r[k]) if r.get(k) not in (None, '') else np.nan
    U = np.array([f(r, 'porosity_union_exact_pct') for r in rows])
    se = np.array([f(r, 'se_of_solid_vol') for r in rows])
    ps = np.array([f(r, 'ps_frac') for r in rows])
    dse = np.array([f(r, 'd_se_um') for r in rows])
    dp = np.array([f(r, 'd_am_p_um') for r in rows])
    ok = np.isfinite(U) & np.isfinite(se) & np.isfinite(ps) & np.isfinite(dse)
    out = {'n': int(ok.sum()), 'source': 'docs/data/lhs_handover_20260928.csv · porosity_union_exact_pct'}
    edges = [(0.0, 0.2, '< 20'), (0.2, 0.3, '20–30'), (0.3, 0.4, '30–40'), (0.4, 0.5, '40–50'), (0.5, 0.71, '≥ 50')]
    out['se_bins'] = [dict(label=l, n=int((ok & (se >= a) & (se < b)).sum()), median=round(float(np.median(U[ok & (se >= a) & (se < b)])), 2))
                      for a, b, l in edges]
    pe = [(4.9, 7, '5–7'), (7, 9, '7–9'), (9, 11, '9–11'), (11, 13, '11–13'), (13, 15.1, '13–15')]
    okp = ok & np.isfinite(dp)
    out['dp_bins'] = [dict(label=l, n=int((okp & (dp >= a) & (dp < b)).sum()), median=round(float(np.median(U[okp & (dp >= a) & (dp < b)])), 2))
                      for a, b, l in pe]
    X = np.column_stack([np.ones_like(se), se, se ** 2, dse])
    b, *_ = np.linalg.lstsq(X[ok], np.log(U[ok]), rcond=None)
    res = np.log(U) - X @ b
    out['control_model'] = 'ln(union %) ~ 1 + SE/고체 + (SE/고체)² + d_SE  (최소제곱)'
    out['control_r2'] = round(float(1 - np.var(res[ok]) / np.var(np.log(U[ok]))), 3)
    out['ps_resid'] = []
    for v in np.round(np.arange(0, 1.01, 0.1), 1):
        k = ok & (np.abs(ps - v) < 0.01)
        out['ps_resid'].append(dict(p_frac=float(v), label=f'{int(round(v * 10))}:{10 - int(round(v * 10))}', n=int(k.sum()),
                                    median_pct=round(float(np.median(res[k]) * 100), 1)))
    return out


def crop(src, box, dst):
    im = Image.open(src)
    W, H = im.size
    im.crop((int(box[0] * W), int(box[1] * H), int(box[2] * W), int(box[3] * H))).save(dst)
    return dst


# ───────────────────────────── 슬라이드 ─────────────────────────────

def content_slide(prs, proto, num, headline, bullets, bsize=12):
    s = dup_slide(prs, proto)
    drop(s, '직사각형 17', '직사각형 12', '그룹 2', '그룹 30')
    set_runs_text(shape(s, 'TextBox 3').text_frame, num)
    lab = shape(s, '직사각형 10')
    set_runs_text(lab.text_frame, headline)
    lab.width = Cm(25.8)                                   # 긴 제목이 가운데로 자라며 왼쪽이 잘리던 것 (렌더 확인 09-29)
    lab.text_frame._txBody.find('{%s}bodyPr' % NS['a']).set('wrap', 'square')
    box = shape(s, '직사각형 14')
    set_bullets(box, bullets, size=bsize)
    return s


def build(template, out):
    prs = Presentation(template)
    proto = prs.slides[7]
    r45 = load_r45()
    lt = lhs_trends()
    (HERE / 'lhs_trends.json').write_text(json.dumps(lt, ensure_ascii=False, indent=1), encoding='utf-8')
    E = Eq(1300)
    L, W = Cm(1.0), Cm(25.8)
    Y0 = Cm(6.35)
    import tempfile
    tmp = Path(tempfile.mkdtemp(prefix='deck_img_'))          # 잘라낸 그림은 리포 밖 임시 폴더
    made = []

    # 2-1 목표 · 추진 내용
    s = content_slide(prs, proto, '2-1', '3세부 미세구조 · 공정 시뮬레이션 — 연구 목표 및 추진 내용', [
        [('목표  ', True), ('건식 황화물 복합양극의 미세구조 (조성 · 입도) 와 믹싱 공정을 시뮬레이션으로 예측해 이온 · 전자 경로가 좋은 설계 조건 제시', False)],
        'DEM 압축 모델로 전해질 함량 · 바이모달 비 (대:소 활물질) · 입경에 따른 공극률 · 굴곡도 · 전도도 정량화',
        '건식 믹싱 DEM 으로 활물질 · 전해질 응집을 구현하고 활물질 표면 점착 (코팅 모사) 변화에 따른 혼합 균일성 평가'])
    steps = [('① 전극 미세구조 모델', ['DEM 분말 투입 · 300 MPa 압축', '대 · 소 활물질 + 고체전해질']),
             ('② 설계인자별 구조 · 수송', ['공극률 · 굴곡도 · 이온/전자 전도도', '130 + 64 설계점 · 사양 5 조성']),
             ('③ 건식 믹싱 균일성', ['점착 (vdW) 응집 구현', '표면 점착 저감 (코팅 모사) 효과'])]
    from pptx.enum.shapes import MSO_SHAPE
    bw, gap = Cm(7.6), Cm(1.5)
    for k, (tt, sub) in enumerate(steps):
        x = L + k * (bw + gap)
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Y0 + Cm(0.3), bw, Cm(1.1))
        b.fill.solid()
        b.fill.fore_color.rgb = RGBColor(0x2F, 0x4B, 0x7C)
        b.line.fill.background()
        tf = b.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = tt
        r.font.size, r.font.bold, r.font.name = Pt(13), True, FONT
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        text(s, x, Y0 + Cm(1.5), bw, Cm(1.3), sub, size=11, align=PP_ALIGN.CENTER)
        if k < 2:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Cm(0.35), Y0 + Cm(0.55), Cm(0.8), Cm(0.6))
            a.fill.solid()
            a.fill.fore_color.rgb = GRAY
            a.line.fill.background()
    label(s, L, Y0 + Cm(3.3), W, '2차년도 추진 현황 (2026-09-29)')
    table(s, L, Y0 + Cm(4.1), W, [
        ['항목', '내용', '상태'],
        ['DEM 압축 모델', '대 · 소 활물질 + LPSCl · 300 MPa · 50 × 50 µm 주기 RVE', '완료'],
        ['설계점 데이터셋', '라틴 하이퍼큐브 130 설계점 + 전해질 과량 64 설계점 (공극률 · 두께 · 접촉 위상)', '인계 준비'],
        ['사양 5 조성 (6 mAh/cm²)', '대입자 9 µm · 대:소 = 10:0 · 7:3 · 5:5 · 3:7 · 0:10', '공극률 완료 · 전도도 정리 중'],
        ['건식 믹싱 캠페인', '층상 장입 → 8 회전 · 점착 사다리 6 조건 + 기준 (13 런)', '런 진행 중 (09-30 완주 예정)'],
        ['코팅 조건 확장', '고점착 조건 · 강성 민감도 사전 등록 · 외부 검토', '설계 검토 중'],
    ], col_w=[4.2, 14.5, 7.1], size=10.5, row_h=0.95)
    notes(s, '참고 스토리라인: 목표 → 입력 → 설계인자 → 기준 구조 → 영향 분석 → 데이터셋 → 요약.  상태 열은 2026-09-29 기준.')
    made.append(s)

    # 2-2 DEM 모델
    s = content_slide(prs, proto, '2-2', 'DEM 기반 복합양극 압축 모델', [
        '활물질 (NCM 대 · 소) 과 고체전해질 (LPSCl) 을 구형 입자로 두고 접촉력으로 운동방정식을 풀어 투입 → 침강 → 300 MPa 압축 → 이완을 재현 (LIGGGHTS)',
        '접촉 = 탄성 접촉 + 하중/제하 이력 (잔류 겹침) 으로 소성 압밀을 대리 · 마찰 · 구름 저항 포함',
        '전해질은 재배열 · 입계 미끄러짐을 담은 유효 탄성계수 1.35 GPa 사용 (실제 24 GPa · 순수 전해질 압축 겹침으로 보정)'], bsize=11)
    frame(s, L, Y0, Cm(14.2), Cm(12.2))
    label(s, L, Y0 + Cm(0.1), Cm(14.2), '지배 방정식')
    Fn, Ft, Mr = E.sub(E.b('F'), E.r('n,ij')), E.sub(E.b('F'), E.r('t,ij')), E.sub(E.b('M'), E.r('r,ij'))
    eqs = [
        (E.sub(E.r('m'), E.r('i')) + E.frac(E.r('d') + E.sub(E.b('v'), E.r('i')), E.r('dt')) + E.r('=') +
         E.nary('∑', E.r('j'), E.par(Fn + E.r('+') + Ft)) + E.r('+') + E.sub(E.r('m'), E.r('i')) + E.b('g'),
         'm_i dv_i/dt = Σ_j (F_n,ij + F_t,ij) + m_i g'),
        (E.sub(E.r('I'), E.r('i')) + E.frac(E.r('d') + E.sub(E.b('ω'), E.r('i')), E.r('dt')) + E.r('=') +
         E.nary('∑', E.r('j'), E.par(E.sub(E.r('R'), E.r('i')) + E.sub(E.b('n'), E.r('ij')) + E.r('×') + Ft + E.r('+') + Mr)),
         'I_i dω_i/dt = Σ_j (R_i n_ij × F_t,ij + M_r,ij)'),
        (E.sub(E.r('F'), E.r('n')) + E.r('=') + E.cases(
            E.sub(E.r('k'), E.r('1')) + E.r('δ') + E.r(',     ') + E.t('loading'),
            E.sub(E.r('k'), E.r('2')) + E.par(E.r('δ') + E.r('−') + E.sub(E.r('δ'), E.r('0'))) + E.r(',   ') + E.t('unloading')),
         'F_n = k1·δ (하중) · k2·(δ − δ0) (제하), k2 > k1 → 잔류 겹침 δ0'),
        (E.frac(E.r('1'), E.sup(E.r('E'), E.r('*'))) + E.r('=') + E.frac(E.r('1') + E.r('−') + E.subsup(E.r('ν'), E.r('i'), E.r('2')), E.sub(E.r('E'), E.r('i'))) +
         E.r('+') + E.frac(E.r('1') + E.r('−') + E.subsup(E.r('ν'), E.r('j'), E.r('2')), E.sub(E.r('E'), E.r('j'))) + E.r(',   ') +
         E.frac(E.r('1'), E.sup(E.r('R'), E.r('*'))) + E.r('=') + E.frac(E.r('1'), E.sub(E.r('R'), E.r('i'))) + E.r('+') + E.frac(E.r('1'), E.sub(E.r('R'), E.r('j'))),
         '1/E* = (1−νi²)/Ei + (1−νj²)/Ej ,  1/R* = 1/Ri + 1/Rj'),
        (E.par(E.sub(E.b('F'), E.r('t')), '|', '|') + E.r('≤') + E.r('μ') + E.par(E.sub(E.b('F'), E.r('n')), '|', '|') + E.r(',   ') +
         E.par(E.sub(E.b('M'), E.r('r')), '|', '|') + E.r('≤') + E.sub(E.r('μ'), E.r('r')) + E.sup(E.r('R'), E.r('*')) + E.par(E.sub(E.b('F'), E.r('n')), '|', '|'),
         '|F_t| ≤ μ|F_n| ,  |M_r| ≤ μ_r R*|F_n|'),
        (E.r('P') + E.r('=') + E.frac(E.sub(E.r('F'), E.t('plate')), E.sub(E.r('A'), E.t('box'))) + E.r('=') + E.t('300 MPa'),
         'P = F_plate / A_box = 300 MPa'),
    ]
    add_equations(s, L + Cm(0.2), Y0 + Cm(0.8), Cm(13.8), Cm(11.2), eqs, sz=1300, gap_pt=6)
    X2 = L + Cm(14.6)
    label(s, X2, Y0 + Cm(0.1), Cm(11.2), '입력 물성 (덱 값)')
    table(s, X2, Y0 + Cm(0.9), Cm(11.2), [
        ['상', '지름 (µm)', 'E (GPa)', 'ν', 'ρ (g/cm³)'],
        ['대입자 AM (다결정 NCM)', '9', '140', '0.25', '4.8'],
        ['소입자 AM (단결정 NCM)', '4', '140', '0.25', '4.8'],
        ['고체전해질 LPSCl', '1', '1.35 *', '0.30', '2.0'],
    ], col_w=[4.6, 1.7, 1.6, 1.1, 1.8], size=10, row_h=0.8)
    text(s, X2, Y0 + Cm(4.2), Cm(11.2), Cm(1.2), ['* 유효값 (실제 bulk 24 GPa) — 강체 구가 못 담는 재배열 · 입계 미끄러짐 · 미세 파쇄를',
                                                   '   탄성계수 하나에 담은 규약 (순수 전해질 압축 겹침으로 보정)'], size=9, color=RGBColor(0x55, 0x5B, 0x66))
    label(s, X2, Y0 + Cm(5.6), Cm(11.2), '압축 과정')
    table(s, X2, Y0 + Cm(6.4), Cm(11.2), [
        ['단계', '내용'],
        ['투입 · 침강', '질량비대로 무작위 투입 → 중력 침강'],
        ['압축', '플래튼 하강 → 300 MPa 도달'],
        ['이완', '플래튼 고정 → 잔류 응력 이완'],
        ['경계', 'x · y 주기 · 바닥 벽 · 50 × 50 µm'],
    ], col_w=[3.0, 8.2], size=10, row_h=0.72)
    text(s, X2, Y0 + Cm(10.2), Cm(11.2), Cm(1.6), ['접촉 모델: LIGGGHTS hooke/hysteresis + tangential history',
                                                    '+ rolling friction (CDT)  ·  덱 in.ps_*_r45.liggghts'], size=9, color=RGBColor(0x55, 0x5B, 0x66))
    notes(s, '수식은 PowerPoint 수식 개체 (더블클릭해 편집).  F_n 이력식은 LIGGGHTS hooke/hysteresis 의 하중/제하 강성 구조를 적은 것 — 정확한 강성 매개 (coefficientPlasticityDepth 등) 는 덱 참조.  '
             '물성 = 덱 값 (영률 · 압력은 덱 안에서 ×0.001 스케일, 길이는 ×1000 스케일).  1.35 GPa 연화 = CLAUDE.md E_SE calibration 절.')
    made.append(s)

    # 2-3 지표 정의
    s = content_slide(prs, proto, '2-3', '미세구조 · 수송 성능 지표 정의', [
        '공극률은 겹친 입자를 한 번만 센 기하 공극률 (union) 로 보고 — 구 부피 단순 합은 겹침을 두 번 세어 공극을 작게 셈',
        '굴곡도 = 전해질 접촉망의 최단 경로 길이 / 전극 두께',
        '이온 · 전자 전도도 = 입자 접촉망에 접촉 협착 저항 (Holm) 을 두고 Kirchhoff 방정식을 풀어 산출'])
    label(s, L, Y0 + Cm(0.1), Cm(13.6), '지표')
    table(s, L, Y0 + Cm(0.9), Cm(13.6), [
        ['지표', '기호', '단위', '뜻'],
        ['공극률 (union)', 'ε', '%', '입자에 속하지 않는 부피 비율'],
        ['두께', 'L', 'µm', '압축 · 이완 뒤 플래튼 높이'],
        ['굴곡도', 'τ', '–', '전해질 경로 길이 / 두께'],
        ['이온 전도도', 'σ_ion', 'mS/cm', '전해질 접촉망 유효 전도도'],
        ['전자 전도도', 'σ_e', 'mS/cm', '활물질 접촉망 유효 전도도'],
        ['전해질 배위수', 'CN_SE', '–', '전해질 1 개당 전해질 접촉 수'],
        ['피복률', 'θ', '%', '활물질 표면 중 전해질 접촉 비율'],
    ], col_w=[3.6, 1.8, 1.7, 6.5], size=10, row_h=0.78)
    X2 = L + Cm(14.2)
    frame(s, X2, Y0, Cm(11.6), Cm(12.2))
    label(s, X2, Y0 + Cm(0.1), Cm(11.6), '계산식')
    eqs = [
        (E.r('ε') + E.r('=') + E.r('1') + E.r('−') + E.frac(E.r('V') + E.par(E.par(E.nary('⋃', E.r('i'), E.sub(E.r('B'), E.r('i')))) + E.r('∩') + E.r('Ω')), E.r('V') + E.par(E.r('Ω'))),
         'ε = 1 − V(∪ B_i ∩ Ω) / V(Ω)'),
        (E.r('τ') + E.r('=') + E.frac(E.sub(E.r('L'), E.t('path')), E.sub(E.r('L'), E.r('z'))), 'τ = L_path / L_z'),
        (E.nary('∑', E.r('j'), E.frac(E.sub(E.r('φ'), E.r('i')) + E.r('−') + E.sub(E.r('φ'), E.r('j')), E.sub(E.r('R'), E.r('ij')))) + E.r('=') + E.r('0'),
         'Σ_j (φ_i − φ_j) / R_ij = 0'),
        (E.sub(E.r('R'), E.r('ij')) + E.r('=') + E.frac(E.r('1'), E.r('2') + E.r('σ') + E.sub(E.r('a'), E.r('ij'))), 'R_ij = 1 / (2σ a_ij)   (Holm 협착)'),
        (E.sub(E.r('σ'), E.t('eff')) + E.r('=') + E.frac(E.r('I') + E.sub(E.r('L'), E.r('z')), E.r('A') + E.r('Δ') + E.r('V')), 'σ_eff = I L_z / (A ΔV)'),
    ]
    add_equations(s, X2 + Cm(0.2), Y0 + Cm(0.8), Cm(11.2), Cm(8.6), eqs, sz=1300, gap_pt=8)
    text(s, X2 + Cm(0.3), Y0 + Cm(9.6), Cm(11.0), Cm(2.4), ['B_i: 입자 i 의 구 · Ω: 전극 상자 (x · y 주기)', 'a_ij: 접촉 반지름 · φ: 전위 · ΔV: 인가 전위차',
                                                             'union 은 무작위 점 4×10⁶ 개로 계산 (오차 ±0.02 %p)'], size=9.5, color=RGBColor(0x55, 0x5B, 0x66))
    notes(s, '굴곡도 = 웹앱 calc_tortuosity (거리 가중 최단 경로 / z 거리).  전도도 = 웹앱 network_conductivity (Holm 1967 협착 · Kirchhoff).  union = scripts/lhs_union_webapp.py.')
    made.append(s)

    # 2-4 설계인자
    s = content_slide(prs, proto, '2-4', '설계인자 정의 및 설계점 데이터셋', [
        '전해질 함량 (활물질 wt%) · 바이모달 비 (대:소) · 대 · 소 활물질 입경 · 전해질 입경을 핵심 설계인자로 선정',
        '이종기술 사양 (6 mAh/cm² · 대입자 9 µm) 을 기준으로 바이모달 비 5 조성 압축 계산',
        '라틴 하이퍼큐브로 130 설계점 (+ 전해질 과량 64 설계점) 을 만들어 설계공간 전반의 구조 데이터 확보'])
    label(s, L, Y0 + Cm(0.1), Cm(15.4), '설계인자')
    table(s, L, Y0 + Cm(0.9), Cm(15.4), [
        ['설계인자', '기준 (사양)', '탐색 범위 (130 설계점)', '단위'],
        ['활물질 함량 (나머지 = 전해질)', '81.6', '70 – 95', 'wt%'],
        ['바이모달 비 (대 : 소)', '7 : 3', '0:10 – 10:0', '질량비'],
        ['대입자 입경 D_P', '9', '5 – 15', 'µm'],
        ['소입자 입경 D_S', '4', '1 – 5', 'µm'],
        ['전해질 입경 D_SE', '1', '1 – 2', 'µm'],
        ['압력', '300', '300', 'MPa'],
        ['면용량', '6', '2', 'mAh/cm²'],
    ], col_w=[6.0, 2.8, 4.4, 2.2], size=10.5, row_h=0.8)
    X2 = L + Cm(16.0)
    label(s, X2, Y0 + Cm(0.1), Cm(9.8), '데이터셋')
    table(s, X2, Y0 + Cm(0.9), Cm(9.8), [
        ['묶음', '건수', '공극률 union (%)'],
        ['130 설계점', '130', '13.6  (6.1 – 28.8)'],
        ['전해질 과량 설계점', '64', '6.6  (5.4 – 9.4)'],
        ['사양 5 조성', '5', '17.7  (16.5 – 20.6)'],
    ], col_w=[4.2, 1.6, 4.0], size=10.5, row_h=0.9)
    text(s, X2, Y0 + Cm(4.9), Cm(9.8), Cm(3.5), ['공극률 = 중앙 (최소 – 최대)',
                                                  '설계점마다 두께 · 공극률 · 전해질/활물질 부피분율 · 피복률 · 접촉 위상 (배위수) 을 표로 정리해 ML 담당에게 인계 준비'],
         size=10, color=RGBColor(0x55, 0x5B, 0x66))
    notes(s, '130 = docs/data/lhs_design_20260818.csv 설계 (2 mAh/cm²) · 64 = lhsx (전해질 과량) · 공극률 = docs/data/lhs_union_20260927/README.md 표.  사양 5 조성 = r4.5 union (r45_union_summary.tsv).')
    made.append(s)

    # 2-5 기준 구조 (7:3)
    b73 = r45['7:3']
    s = content_slide(prs, proto, '2-5', '기준 전극 (대:소 = 7:3) 의 3차원 압축 구조', [
        f'6 mAh/cm² · 300 MPa 압축 뒤 두께 {float(b73["thickness_um"]):.1f} µm · 공극률 (union) {float(b73["porosity_union_exact_pct"]):.1f} %',
        '대입자 사이를 소입자와 전해질이 채우는 바이모달 충전 구조',
        '웹앱 3D 뷰어로 상별 분포 · 전해질 접촉망 · 공극 위치 확인'])
    label(s, L, Y0 + Cm(0.1), Cm(8.4), '3D 구조 (압축 뒤)')
    placeholder(s, L, Y0 + Cm(0.9), Cm(8.4), Cm(11.2), ['[그림 자리]', '웹앱 3D 뷰어', 'input_6mAh_real_4_D9 (7:3)', '상별 색: 대 · 소 활물질 · 전해질'])
    label(s, L + Cm(8.7), Y0 + Cm(0.1), Cm(8.4), '단면 · 전해질 접촉망')
    placeholder(s, L + Cm(8.7), Y0 + Cm(0.9), Cm(8.4), Cm(11.2), ['[그림 자리]', '2D 단면 (xz) 또는', '전해질 퍼콜레이션 망'])
    X3 = L + Cm(17.4)
    label(s, X3, Y0 + Cm(0.1), Cm(8.4), '주요 값')
    table(s, X3, Y0 + Cm(0.9), Cm(8.4), [
        ['항목', '값'],
        ['두께', f'{float(b73["thickness_um"]):.1f} µm'],
        ['공극률 (union)', f'{float(b73["porosity_union_exact_pct"]):.2f} %'],
        ['공극률 (구 부피 합)', f'{float(b73["porosity_sphere_web_pct"]):.2f} %'],
        ['대 / 소 활물질 수', '280 / 1,376'],
        ['전해질 입자 수', '약 15.9 만'],
        ['σ_ion (mS/cm)', '⬜ 웹앱 값'],
        ['σ_e (mS/cm)', '⬜ 웹앱 값'],
        ['굴곡도 τ', '⬜ 웹앱 값'],
    ], col_w=[4.4, 4.0], size=10, row_h=0.78)
    notes(s, '웹앱 케이스 260925_000559_082983 (input_6mAh_real_4_D9) · 덤프 atom_3220000.  구 부피 합은 겹침을 두 번 세는 웹앱 기본값 (union 과 병기).')
    made.append(s)

    # 2-6 바이모달 비
    order = ['0:10', '3:7', '5:5', '7:3', '10:0']
    uni = [float(r45[k]['porosity_union_exact_pct']) for k in order]
    sph = [float(r45[k]['porosity_sphere_web_pct']) for k in order]
    thk = [float(r45[k]['thickness_um']) for k in order]
    imin = int(np.argmin(uni))
    s = content_slide(prs, proto, '2-6', '바이모달 비 (대:소) 에 따른 공극률 — 7:3 에서 가장 치밀', [
        f'사양 5 조성에서 공극률 (union) 은 대:소 = 7:3 에서 {uni[3]:.1f} % 로 가장 낮고 소입자만 (0:10) {uni[0]:.1f} % 로 가장 높음 · 두께도 7:3 이 최소 ({thk[3]:.1f} µm)',
        '소입자가 대입자 사이 틈을 채우는 바이모달 충전 효과 — 대입자만 (10:0) 보다도 낮음',
        f'130 설계점에서 전해질 함량 · 입경을 보정해도 대입자 몫 0.7 (7:3) 부근에서 공극률이 가장 낮게 남음 (보정 R² {lt["control_r2"]:.2f})'], bsize=11)
    label(s, L, Y0 + Cm(0.1), Cm(12.6), '사양 5 조성 (6 mAh/cm² · 대입자 9 µm · 300 MPa)')
    chart(s, L, Y0 + Cm(0.8), Cm(12.6), Cm(8.4), order, [('공극률 union', uni), ('구 부피 합 (참고)', sph)], '공극률 (%)',
          ymin=10, ymax=22, xtitle='대 : 소 (질량비)', colors=[BLUE, GRAY], point_colors={imin: ORANGE})
    table(s, L, Y0 + Cm(9.4), Cm(12.6), [['대 : 소'] + order, ['두께 (µm)'] + [f'{v:.1f}' for v in thk]], size=10, row_h=0.7)
    X2 = L + Cm(13.2)
    label(s, X2, Y0 + Cm(0.1), Cm(12.6), '130 설계점 — 전해질 함량 · 입경 보정 뒤 공극률 편차')
    pr = lt['ps_resid']
    cats = [d['label'] for d in pr]
    vals = [d['median_pct'] for d in pr]
    ch = chart(s, X2, Y0 + Cm(0.8), Cm(12.6), Cm(8.4), cats, [('공극률 편차 (%)', vals)], '보정 뒤 공극률 편차 (%)', ymin=-20, ymax=15,
               fmt='+0.0;-0.0;0', xtitle='대 : 소 (질량비)', colors=[BLUE], point_colors={i: ORANGE for i, d in enumerate(pr) if d['label'] == '7:3'}, legend=False)
    text(s, X2, Y0 + Cm(9.3), Cm(12.6), Cm(2.6), [
        f'편차 = ln(공극률) 을 전해질 부피분율 (2차) · 전해질 입경으로 회귀한 잔차의 조성별 중앙값 (조성당 {min(d["n"] for d in pr)}–{max(d["n"] for d in pr)} 점 · 2 mAh/cm²)',
        '음수 = 같은 전해질 함량에서 더 치밀',
        '⚠ 사양 5 조성은 조성당 침대 1 개 — 시드 산포 미측정'], size=9, color=RGBColor(0x55, 0x5B, 0x66))
    notes(s, f'왼쪽 = docs/data/ps45_r45_union_20260929/r45_union_summary.tsv (union 정확 · MC ±0.02 %p).  오른쪽 = lhs_trends.json ps_resid ({lt["control_model"]}, R² {lt["control_r2"]}).  '
             '입경 (대 · 소) 을 보정에 더해도 7:3 편차 −15 % 수준 유지 (2026-09-29 점검).  탐색적 분석 — 사전 등록 판정 아님.')
    made.append(s)

    # 2-7 전해질 함량 · 입경
    s = content_slide(prs, proto, '2-7', '전해질 함량 · 입경의 영향 (130 설계점)', [
        f'공극률을 가장 크게 좌우하는 인자는 전해질 함량 — 고체 중 전해질 부피 20 % 미만 {lt["se_bins"][0]["median"]:.1f} % → 50 % 이상 {lt["se_bins"][-1]["median"]:.1f} % (중앙값)',
        '전해질이 대 · 소 활물질 틈을 채우므로 함량이 늘수록 공극이 줄고 두께가 얇아짐',
        '대입자 입경 (5 – 15 µm) 의 단독 효과는 작음 — 입경 구간별 중앙값 12 – 15 %'])
    label(s, L, Y0 + Cm(0.1), Cm(12.6), '고체 중 전해질 부피분율')
    chart(s, L, Y0 + Cm(0.8), Cm(12.6), Cm(9.4), [d['label'] + ' %' for d in lt['se_bins']], [('공극률 union 중앙값', [d['median'] for d in lt['se_bins']])],
          '공극률 union (%)', ymin=0, ymax=30, xtitle='고체 중 전해질 부피 (%)', colors=[BLUE], legend=False)
    text(s, L, Y0 + Cm(10.3), Cm(12.6), Cm(1.2), '구간별 점 수: ' + ' · '.join(f'{d["label"]} {d["n"]}' for d in lt['se_bins']), size=9, color=RGBColor(0x55, 0x5B, 0x66))
    X2 = L + Cm(13.2)
    label(s, X2, Y0 + Cm(0.1), Cm(12.6), '대입자 입경 D_P')
    chart(s, X2, Y0 + Cm(0.8), Cm(12.6), Cm(9.4), [d['label'] for d in lt['dp_bins']], [('공극률 union 중앙값', [d['median'] for d in lt['dp_bins']])],
          '공극률 union (%)', ymin=0, ymax=30, xtitle='대입자 지름 (µm)', colors=[GRAY], legend=False)
    text(s, X2, Y0 + Cm(10.3), Cm(12.6), Cm(1.2), '구간별 점 수: ' + ' · '.join(f'{d["label"]} {d["n"]}' for d in lt['dp_bins']) + ' (소입자만 15 점 제외)',
         size=9, color=RGBColor(0x55, 0x5B, 0x66))
    notes(s, '라틴 하이퍼큐브는 인자를 동시에 바꾸므로 구간 중앙값은 다른 인자가 섞인 주변 경향이다.  전해질 함량 효과가 가장 크다는 점은 보정 회귀 (R² 0.86) 로도 같다.')
    made.append(s)

    # 2-8 전도도 · 굴곡도
    s = content_slide(prs, proto, '2-8', '바이모달 비에 따른 전도도 · 굴곡도 (사양 5 조성)', [
        '5 조성의 이온 · 전자 전도도와 굴곡도를 웹앱 접촉망 해석으로 산출 — ⬜ 값 정리 중',
        '⬜ (값 입력 뒤 문장 확정) 7:3 의 σ_ion · σ_e · τ 를 단일 입도 (10:0 · 0:10) 와 비교',
        '공극률이 가장 낮은 7:3 이 이온 경로 연결성 (배위수 · 퍼콜레이션) 에서도 유리한지 확인'])
    label(s, L, Y0 + Cm(0.1), W, '사양 5 조성 — 구조 · 수송 지표')
    rows = [['대 : 소', '공극률 union (%)', '두께 (µm)', 'σ_ion (mS/cm)', 'σ_e (mS/cm)', '굴곡도 τ', 'CN_SE', '퍼콜레이션 (%)']]
    for k in order:
        rows.append([k, f'{float(r45[k]["porosity_union_exact_pct"]):.2f}', f'{float(r45[k]["thickness_um"]):.1f}', '⬜', '⬜', '⬜', '⬜', '⬜'])
    table(s, L, Y0 + Cm(0.9), W, rows, size=10.5, row_h=0.85)
    placeholder(s, L, Y0 + Cm(6.4), Cm(12.6), Cm(5.6), ['[차트 자리]', 'σ_ion · σ_e vs 대:소', '(값 입력 뒤 생성)'])
    placeholder(s, L + Cm(13.2), Y0 + Cm(6.4), Cm(12.6), Cm(5.6), ['[차트 자리]', '굴곡도 τ · CN_SE vs 대:소', '(값 입력 뒤 생성)'])
    notes(s, '값 = 웹앱 results/<case_id>/full_metrics.json (5 케이스: 6e4ff6 · 082983 · 123978 · bd85f9 · 0bee25).  ⬜ 칸은 값이 오면 채운다.')
    made.append(s)

    # 2-9 믹싱 모델
    s = content_slide(prs, proto, '2-9', '건식 믹싱 DEM — 활물질 · 전해질 응집 구현', [
        '회전 드럼에 대 · 소 활물질과 전해질 입자를 넣고 DEM 으로 혼합 과정을 계산 (100,000 입자 · 75 rpm)',
        '입자 간 van der Waals 부착을 SJKR 접촉 점착력으로 넣고 크기는 Bond 수 (점착력 / 입자 무게) 로 설정',
        '점착이 크면 활물질 응집체가 드럼과 함께 회전하고 작으면 자유 유동 — 응집 거동 재현'], bsize=11)
    frame(s, L, Y0, Cm(12.4), Cm(12.2))
    label(s, L, Y0 + Cm(0.1), Cm(12.4), '지배 방정식')
    eqs = [
        (E.sub(E.r('F'), E.r('n')) + E.r('=') + E.frac(E.r('4'), E.r('3')) + E.sup(E.r('E'), E.r('*')) + E.rad(E.sup(E.r('R'), E.r('*'))) + E.sup(E.r('δ'), E.r('3/2')) +
         E.r('−') + E.sub(E.r('γ'), E.r('n')) + E.sub(E.r('v'), E.r('n')) + E.r('−') + E.sub(E.r('F'), E.t('coh')),
         'F_n = (4/3) E* √R* δ^(3/2) − γ_n v_n − F_coh'),
        (E.sub(E.r('F'), E.t('coh')) + E.r('=') + E.sub(E.r('k'), E.r('c')) + E.r('·') + E.r('2π') + E.sup(E.r('R'), E.r('*')) + E.r('δ'),
         'F_coh = k_c · 2π R* δ   (SJKR, k_c = 점착 에너지 밀도)'),
        (E.t('Bo') + E.r('=') + E.frac(E.sub(E.r('F'), E.t('coh')), E.r('m') + E.r('g')), 'Bo = F_coh / (m g)'),
        (E.r('M') + E.par(E.r('t')) + E.r('=') + E.frac(E.subsup(E.r('S'), E.r('0'), E.r('2')) + E.r('−') + E.sup(E.r('S'), E.r('2')) + E.par(E.r('t')),
                                                        E.subsup(E.r('S'), E.r('0'), E.r('2')) + E.r('−') + E.subsup(E.r('S'), E.r('R'), E.r('2'))),
         'M(t) = (S0² − S²(t)) / (S0² − S_R²)'),
    ]
    add_equations(s, L + Cm(0.2), Y0 + Cm(0.8), Cm(12.0), Cm(7.6), eqs, sz=1300, gap_pt=8)
    text(s, L + Cm(0.3), Y0 + Cm(8.6), Cm(11.8), Cm(3.4), ['k_c: 점착 에너지 밀도 (J/m³) · Bo: Bond 수',
                                                           'S²: 셀별 활물질 부피분율의 분산',
                                                           'S0²: 층상 장입 직후 · S_R²: 무작위 균일 장입 (기준 런)',
                                                           'M = 0 미혼합 · 1 무작위 완전 혼합 (Lacey)'], size=9.5, color=RGBColor(0x55, 0x5B, 0x66))
    X2 = L + Cm(12.8)
    ab = crop(ROOT / 'docs/figures/mixer_summary_20260920.png', (0.0, 0.0, 0.585, 0.87), tmp / 'mix_ab.png')
    label(s, X2, Y0 + Cm(0.1), Cm(13.0), '응집 구현 예시 (드럼 단면)')
    s.shapes.add_picture(str(ab), X2 + Cm(1.25), Y0 + Cm(0.8), width=Cm(10.5))
    text(s, X2, Y0 + Cm(7.0), Cm(13.0), Cm(0.9), '⚠ 이전 세대 모델 (09-20) 예시 — 섬유 첨가제 포함 옛 덱 · 현 캠페인은 3 상 · 실제 입경 비 · 벽 부착으로 재설계', size=9, color=ORANGE)
    label(s, X2, Y0 + Cm(8.0), Cm(13.0), '계산 조건 (현 캠페인)')
    table(s, X2, Y0 + Cm(8.8), Cm(13.0), [
        ['항목', '값'],
        ['입자 (3 상)', '대 9 · 소 4 · 전해질 1 µm (조대화 ×151)'],
        ['배치 · 입자 수', '2 g · 100,000 개'],
        ['드럼', '반지름 13.1 mm · 75 rpm · 8 회전'],
    ], col_w=[3.4, 9.6], size=10, row_h=0.8)
    notes(s, '식 = scripts/make_mixer_deck.py (Hertz + SJKR · Bo = F_coh / W) · scripts/measure_mixing_index.py (Lacey).  그림 (a)(b) = docs/figures/mixer_summary_20260920.png 의 왼쪽 두 패널 — '
             '09-21 재설계 이전 덱 (세대 표지 · devlog v08).  (c) 막대는 싣지 않았다 (옛 세대 수치 · 철회된 앵커 라벨).  현 캠페인 = docs/reviews/mixer_layered_prereg_20260921.md.')
    made.append(s)

    # 2-10 코팅 모사 · 균일성
    s = content_slide(prs, proto, '2-10', '활물질 표면 점착 저감 (코팅 모사) 에 따른 혼합 균일성', [
        '층상 장입 (아래 활물질 · 위 전해질) 에서 8 회전 뒤 Lacey 혼합지수 M 으로 균일성 비교',
        '표면 점착을 전해질 수준으로 낮춘 조건 (LC · 코팅 모사) ↔ 점착이 큰 조건 (LA) 을 같은 시드끼리 비교 — 판정선 사전 등록 (차이 ≥ 0.10 · 시드 오차의 2 배 이상)',
        '13 런 진행 중 (09-29 기준 88 – 98 %) → 09-30 완주 뒤 판정 · 결과 반영'], bsize=11)
    dsg = crop(ROOT / 'docs/figures/mixer_campaign_design_20260928.png', (0.0, 0.0, 0.43, 1.0), tmp / 'mix_design.png')
    label(s, L, Y0 + Cm(0.1), Cm(10.2), '캠페인 설계')
    s.shapes.add_picture(str(dsg), L, Y0 + Cm(0.8), width=Cm(10.2))
    X2 = L + Cm(10.6)
    label(s, X2, Y0 + Cm(0.1), Cm(7.0), '점착 조건 (AM–AM Bo)')
    table(s, X2, Y0 + Cm(0.8), Cm(7.0), [
        ['조건', 'Bo_code', '시드'],
        ['LC (코팅 모사)', '0.00109', '3'],
        ['LB1', '0.1', '1'],
        ['LB2', '0.5', '1'],
        ['LB3', '1.5', '1'],
        ['LA (기준)', '3.0', '3'],
        ['L0 (음성 대조)', '0 (전 쌍)', '1'],
        ['E0 (무작위 기준)', '0 · 무회전', '3'],
    ], col_w=[3.3, 2.1, 1.6], size=10, row_h=0.7)
    X3 = X2 + Cm(7.4)
    label(s, X3, Y0 + Cm(0.1), Cm(7.8), '결과 (판정 뒤 입력)')
    table(s, X3, Y0 + Cm(0.8), Cm(7.8), [
        ['', 'M (8 회전 뒤)', '시드 SD'],
        ['LC', '⬜', '⬜'],
        ['LA', '⬜', '⬜'],
        ['Δ = LC − LA', '⬜', '⬜'],
        ['판정', '⬜', ''],
    ], col_w=[3.0, 2.8, 2.0], size=10, row_h=0.8)
    text(s, X2, Y0 + Cm(6.6), Cm(15.2), Cm(5.4), [
        '• 판정 칸 16×16×4 (한 변 1.64 mm = 대입자 1.2 개) · 같은 시드 3 쌍의 차이 평균',
        '• h1 = Δ ≥ 0.10 이고 Δ ≥ 2·SE — 그 사이는 "애매" 로 그대로 보고',
        '• 코팅 모사 = 모델에서 활물질 표면 점착 에너지 밀도를 낮춘 조건 (AM–AM · AM–벽 부착이 함께 변함)',
        '  → 실제 코팅 공정의 재현이 아니라 "표면 물성 변화" 효과의 모델 평가',
        '• 고점착 확장 (Bo 38) · 강성 민감도는 외부 검토 뒤 착수'], size=10)
    notes(s, '판정선 = docs/reviews/mixer_layered_prereg_20260921.md §2 · 결론 문안 §5 (Codex §9 문안: "사전 정의된 두 CED 행렬 · AM–AM 및 AM–벽 부착이 함께 다름 · 실제 코팅 재현으로 해석하지 않음").  '
             '⬜ 칸 = 09-30 완주 뒤 판독기 planned.M_final.')
    made.append(s)

    # 2-11 요약
    s = content_slide(prs, proto, '2-11', '요약 및 향후 계획', [
        'DEM 압축 모델로 전해질 함량 · 바이모달 비 · 입경에 따른 전극 미세구조를 정량화 (130 + 64 설계점 · 사양 5 조성)',
        '공극률은 전해질 함량이 가장 크게 좌우하고, 같은 함량에서는 바이모달 비 7:3 부근이 가장 치밀 (사양 5 조성 16.5 %)',
        '건식 믹싱 DEM 에 점착 응집을 구현했고, 표면 점착 저감 (코팅 모사) 의 균일성 효과를 사전 등록 판정으로 평가 중'])
    label(s, L, Y0 + Cm(0.1), W, '향후 계획')
    table(s, L, Y0 + Cm(0.9), W, [
        ['항목', '내용', '시기'],
        ['사양 5 조성 수송', '이온 · 전자 전도도 · 굴곡도 정리 → 7:3 수송 우위 확인', '10월 초'],
        ['믹싱 판정', '층상 캠페인 판정 (LC ↔ LA) → 결과 반영', '09-30 완주 뒤'],
        ['코팅 조건 확장', '고점착 · 강성 민감도 확인 런 (외부 검토 뒤)', '10월'],
        ['데이터셋 인계', '130 + 64 설계점 구조 지표 → 구조 예측 ML', '10월'],
        ['MPM 연계', '전해질 소성 변형 (형상) · 활물질 하중 분담', '진행 중'],
    ], col_w=[4.6, 16.2, 5.0], size=10.5, row_h=0.95)
    notes(s, '요약 문장의 7:3 = 사양 5 조성 union 최소 + 130 설계점 보정 편차 (탐색적).  수송 우위 문장은 2-8 값이 채워진 뒤에만 쓴다.')
    made.append(s)

    # 순서: 1–7 · 새 슬라이드 · 10–11 (템플릿 빈 추진 내용 8 · 9 는 뺀다)
    sld = prs.slides._sldIdLst
    ids = list(sld)
    orig, new = ids[:11], ids[11:]
    for el in ids:
        sld.remove(el)
    for el in orig[:7] + new + orig[9:11]:
        sld.append(el)
    for el in orig[7:9]:
        prs.part.drop_rel(el.get('{%s}id' % NS['r']))
    prs.save(out)
    return len(made)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--template', required=True)
    ap.add_argument('--out', default=str(HERE / 'report_20261001_draft.pptx'))
    a = ap.parse_args()
    n = build(a.template, a.out)
    print(f'{n} 장 → {a.out}')


if __name__ == '__main__':
    main()
