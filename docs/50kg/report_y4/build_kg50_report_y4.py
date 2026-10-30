#!/usr/bin/env python3
"""50kg/batch 4차년도 DEM 보고 7쪽 — 초안 1 생성기 (1쪽 완성 · 2–7쪽 틀).

    python3 -P docs/50kg/report_y4/build_kg50_report_y4.py                  # 1쪽 그림 → 덱 → 검사 (렌더 없음)
    python3 -P docs/50kg/report_y4/build_kg50_report_y4.py --render         # + LibreOffice PDF · 쪽 PNG · 렌더 검사
    python3 -P docs/50kg/report_y4/build_kg50_report_y4.py --determinism    # 같은 입력으로 두 번 만들어 바이트 대조
    python3 -P docs/50kg/report_y4/build_kg50_report_y4.py --template <계획 덱.pptx>   # 양식 상수 대조 (읽기만 · 리포 밖 파일)
    python3 -P docs/50kg/report_y4/build_kg50_report_y4.py --check-only <deck.pptx> [--pdf <deck.pdf>]

- 양식 = 1저자 50kg 4차년도 계획 덱 (리포 밖 · 사용자 업로드 원본) 의 쪽 크기 (10 × 7.5 in) · 테마 (색 · 글꼴 체계) · 내용 장 머리
  (그라데이션 사각 + 장 번호 · 과제명 · 쪽 제목 · 회색 띠) · 요약 띠 (회색 둥근 틀 + 파란 알약 제목 + ① ② 항목) · 점선 패널 둘 (괄호 제목) ·
  【 】 캡션 — 이 생성기 안의 상수로 다시 그린다 (양식의 그림 · 장 내용 · 내장 글꼴은 가져오지 않는다).  `--template` 은 상수를 양식과 대조만 한다.
- 글꼴 = 기존 생성기 (`docs/report_20261021/build_deck_v2.py`) 규약: 머리 · 탭 · 표 · 라벨 = 한글 KoPubWorld돋움체 Medium + 영어 Aptos ·
  설명 문장 (- 줄 · 캡션 · 회색 결론 줄) = 맑은 고딕 · 그림 영어 = Aptos.  ⚠ 보고 원칙 (`docs/report_making_principles.md` §4 · §8) 의
  "한글 = 맑은 고딕" 과 다르다 = ⬜ 1저자 결정 (README).
- 1쪽 숫자 = 전부 빌드 때 덱 · manifest · 끝 상태 덤프 · 물성 기록에서 읽는다 (`kg50_y4_sources.py`) — 검사 C8 이 쪽 글의 숫자를 하나씩 출처와 대조.
- 2–7쪽 = 틀: 계획 3판 (`docs/50kg/kg50_report_plan_fix_plan_20261030.md` §2) 제목 · 주 그림 · 보조 칸 (상태 표지) · 결론의 꼴 (회색 줄) ·
  §1 한계 문구.  가짜 곡선 없음.
- 결정적 출력: python-pptx 저장 뒤 zip 을 고정 시각 · 같은 순서로 다시 쓴다 (`--determinism` 이 확인).
"""
import argparse
import datetime as _dt
import hashlib
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from html import unescape
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE / 'figs'))
import kg50_y4_sources as src  # noqa: E402

OUT_DECK = HERE / 'kg50_report_y4_draft1.pptx'
OUT_CHECKS = HERE / 'checks_draft1.json'
RENDER_DIR = HERE / 'render'
EMU = 914400
NSA = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NSP = 'http://schemas.openxmlformats.org/presentationml/2006/main'
N_PAGES = 7

# ───────────────────────────── 양식 상수 (계획 덱에서 잰 값 · EMU · `--template` 이 대조) ─────────────────────────────
SLIDE_W, SLIDE_H = 9144000, 6858000                     # 10 × 7.5 in (4:3)
#: 레이아웃 '빈 화면' 의 머리 무리 '그룹 9' (벡터 도형만)
HDR_SQ = (0, 1, 929898, 740176)                          # 그라데이션 사각 '직사각형 6' (0094E9 9 % → 78CCC8 100 % · 45°)
HDR_BAND = (0, 740178, 9144000, 93715)                   # 회색 띠 '직사각형 7' (tx1 α 5 %)
HDR_LINE = (0, 740178, 9144000, 0)                       # 선 '직선 연결선[R] 8' (bg1 lumMod 85 % · 0.75 pt)
HDR_DECO = ((7725814, -1867, 1418185, 742045, 676140), (8170428, -1867, 973572, 742045, 231527))   # '자유형 4' · '자유형 5' (0094E9 α 5 % · 180° 회전)
#: 내용 장 (계획 덱 7 · 8 장) 의 머리 글 · 요약 띠 · 패널
HDR_NUM = (283650, 138022, 362600, 461665)               # 'TextBox 128' 장 번호 24 pt 굵게 흰색
HDR_PROJ = (1009838, 85527, 3499676, 230832)             # 'TextBox 10' 과제명 9 pt 회색 (white lumMod 50 %)
HDR_TITLE = (1031154, 332236, 307777)                    # 'TextBox 5' 쪽 제목 14 pt 굵게 (x, y, h · 폭 = 글 길이)
BAND_OUT = (258763, 897206, 8641044, 1178322, 0.11335)   # 'AutoShape 2' bg1 lumMod 95 %
BAND_IN = (360476, 994354, 8437618, 990371, 0.08127)     # 'AutoShape 2' 흰색 · 선 1 pt bg1 lumMod 75 %
BAND_TAB = (1820056, 880287, 5518458, 320182, 0.5)       # '모서리가 둥근 직사각형 132' 흰 알약 · 선 1 pt 0070C0 · 12 pt 굵게 0070C0
ITEM_NUM = ((559520, 1330248), (4768657, 1330248), 191630)        # '타원 147 · 148' 10 pt 0070C0
ITEM_TITLE = ((759693, 1326715), (5002368, 1330734), 261610)      # '직사각형 149 · 151' 11 pt
ITEM_DESC = ((782113, 1590931), (5024787, 1594950), 261610)       # 'TextBox 150 · 152' 11 pt "- …"
PANEL = ((258763, 2204863), (4640670, 2204863), 4235043, 4470257, 0.02377)   # '모서리가 둥근 직사각형 72' 점선 (bg1 lumMod 50 %)
PANEL_TAB_Y, PANEL_TAB_H = 2097005, 276999               # '직사각형 144 · 139' 12 pt "( … )" 흰 바탕
CONTENT_TOP = (PANEL_TAB_Y + PANEL_TAB_H) / 914400 + 0.03   # 패널 안 첫 글 = 탭 아래
SLDNUM = (7010401, 6491968, 2133599, 365125)             # 쪽 번호 자리 (12 pt · 오른쪽 · tx1 tint 75 %)
THEME_NAME = '한컴오피스'
THEME_COLORS = dict(dk2='3A3C84', lt2='FAF3DB', accent1='6182D6', accent2='FF843A', accent3='B2B2B2', accent4='FFD700',
                    accent5='289B6E', accent6='9D5CBB', hlink='0000FF', folHlink='800080')
THEME_FONT = '함초롬돋움'
#: 이 초안의 조정 (양식 패널 바닥 7.30 in → 6.97 in · 회색 결론 줄 자리)
PANEL_BOTTOM = 6.97
NOTE_BOX = (0.283, 7.03, 7.30, 0.25)

BLUE, GRAY_TXT, DARK, MID, LINE_GRAY, DASH_GRAY, FAINT = '0070C0', '7F7F7F', '262626', '595959', 'BFBFBF', '7F7F7F', 'F2F2F2'
KO, EN, DESC = 'KoPubWorld돋움체 Medium', 'Aptos', '맑은 고딕'
FONT_SETS = {'ko': (EN, KO, EN), 'desc': (DESC, DESC, DESC)}
PROJECT = '복합전극용 50kg/batch급 고성능 고체전해질 제조공정기술 및 장비개발'   # 양식 내용 장의 머리 과제명 (모든 장 같음)
SECTION_NO = '4'                                          # 계획 덱 4 장 (4차년도) 과 같은 장 번호 자리 — ⬜ 1저자 확인 (README)
SERIES_FMT = '4차년도 DEM 해석 — SE 입도 {sizes} µm 비교 (입도별 시드 {seeds})'

# 상태 표지 (자료 없는 칸) — 가짜 곡선 대신 이 글만
ST_NEW = '[새 지표 정의 · 시험 뒤]'
ST_PLAN = '[분석 계획 고정 뒤 · 계산 전]'
ST_FD = '[f(d) 계산 뒤]'
ST_ALL = '[네 입도 지표 산출 뒤]'
ST_NEXT = '[다음 보고 (3D 전기화학)]'
ST_P2AUX = '[D3 · D2 로그 반입 뒤 · D1 · D0.7 자료 전]'
STATUS_VOCAB = (ST_NEW, ST_PLAN, ST_FD, ST_ALL, ST_NEXT, ST_P2AUX)
#: 금지 낱말 (계획 §1 · 보고 원칙 §5) — 쪽 글 + notes 어디에도.  Stage-E 는 허용 문장 안에서만
FORBIDDEN = ('SE 이용률', '이용률', '활성 AM', '최적', '후막 반응', '로렌츠', 'Lorenz', '가장', '입증', 'LHS', '보간')
STAGE_E_OK = 'Stage-E σ_{ion} 은 쓰지 않음'

HDR_INS = (0.03, 0.03, 0.03, 0.03)   # 머리 글 상자 여백 (양식 상자 크기 그대로 · 추정이 보수적이라 여백만 줄임)
LH = 1.25                 # 줄 높이 추정 = 글자 크기 × 1.25 × 줄 간격 (KoPub 양식 상자 실측 1.215 em · LibreOffice 렌더 실측 1.20 em)
LH_DESC = 1.35            # 맑은 고딕 (설명 문장) — PowerPoint 한 줄 ≈ 1.33 em 추정 (글꼴 ascent + descent) → 보수적으로 1.35
FIT_SLACK = 1.0           # 추정 높이 ≤ 상자 안 높이 × FIT_SLACK
_LAT_FONT_PATH = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'   # 영문 폭 추정 (Aptos 보다 넓다 = 보수적)
_lat_font = None
_char_w = {}


def I(x):
    return int(round(x * EMU))


def inch(e):
    return e / EMU


# ───────────────────────────── 글 폭 · 줄 수 추정 (덱 배치와 검사가 같은 함수) ─────────────────────────────

#: 동아시아 모호 폭 (EA Ambiguous) 기호 — 1 em 로 잡는다 (LibreOffice 는 한글 문맥에서 아시아 글꼴로 그리고, PowerPoint 는 ea 글꼴로 그릴 수 있다 · 둘 중 넓은 쪽)
_AMBIG = set('·×—–µ²³¹°±→←↔≤≥≈√∼§※…‥″′')
BOLD_K = 1.12             # 굵은 영문 폭 / 보통 (DejaVu Sans Bold ÷ Regular · 보정 실측 1.11)
AUTOSPACE = 0.25          # 한글 ↔ 영문 · 숫자 경계마다 (공백 건너 · LibreOffice 아시아 자동 간격 · 보정 실측 0.23 em)


def _cjk(ch):
    cp = ord(ch)
    return (0x1100 <= cp <= 0x11FF or 0x2E80 <= cp <= 0x9FFF or 0xAC00 <= cp <= 0xD7A3 or 0xF900 <= cp <= 0xFAFF
            or 0xFF00 <= cp <= 0xFFEF or 0x2460 <= cp <= 0x24FF or 0x25A0 <= cp <= 0x25FF or 0x2190 <= cp <= 0x22FF
            or 0x3000 <= cp <= 0x303F or 0x3130 <= cp <= 0x318F)


SUPSUB_EM = 0.6           # 위 · 아래첨자 숫자 (U+2070–209F · 10⁵ 의 ⁵) = 영문 문자 · 0.6 em (DejaVu 0.4 em 보다 넓게)


def _wide(ch):
    return ch in _AMBIG or _cjk(ch)


def _latin_em(ch):
    global _lat_font
    if ch not in _char_w:
        if _lat_font is None:
            _lat_font = ImageFont.truetype(_LAT_FONT_PATH, 100)
        _char_w[ch] = _lat_font.getlength(ch) / 100.0
    return _char_w[ch]


def seq_widths(chars):
    """chars = [(글자, pt, 굵게)] → 글자마다 폭 (pt).  한글 · 한자 · 모호 폭 기호 = 1 em · 그 밖 = DejaVu Sans 폭 (굵게 × 1.12) ·
    한글 ↔ 영문 · 숫자 경계 + 0.25 em.  보정 (10-30 · LibreOffice 렌더 WenQuanYi + DejaVu): 한글 · 영문만 = 같음 · 섞인 글 = 렌더보다 넓거나 같음."""
    out, prev = [], None
    for c, pt, b in chars:
        if _cjk(c):
            w, scr = pt, 'A'
        elif c in _AMBIG:
            w, scr = pt, None
        else:
            em = SUPSUB_EM if 0x2070 <= ord(c) <= 0x209F else _latin_em(c)
            w = em * pt * (BOLD_K if b else 1.0)
            scr = 'L' if c.isalnum() else None
        if scr and prev and scr != prev:
            w += AUTOSPACE * pt
        if scr:
            prev = scr
        out.append(w)
    return out


def char_w_pt(ch, pt, bold=False):
    return seq_widths([(ch, pt, bold)])[0]


def wrap_lines(segs, width_pt):
    """segs = [(글, pt, 위아래첨자 여부, 굵게)] → 줄 수 (공백에서만 나눈다 = 보수적 · 한 낱말이 줄보다 길면 그 낱말을 여러 줄로)."""
    chars = [(c, pt * (0.65 if sub else 1.0), bold) for t, pt, sub, bold in segs for c in t]
    if not chars:
        return 1
    ws = seq_widths(chars)
    words, cur = [], []
    for (c, _, _), w in zip(chars, ws):
        if c == ' ':
            words.append(cur)
            cur = []
            words.append([(c, w)])
        elif c == '\n':
            words.append(cur)
            cur = []
            words.append(None)
        else:
            cur.append((c, w))
    words.append(cur)
    lines, x = 1, 0.0
    for wd in words:
        if wd is None:
            lines += 1
            x = 0.0
            continue
        if not wd:
            continue
        ww = sum(w for _, w in wd)
        if wd[0][0] == ' ':
            x += ww
            continue
        if x + ww <= width_pt + 1e-6:
            x += ww
        elif x == 0.0:
            n = max(1, math.ceil(ww / width_pt))
            lines += n - 1
            x = ww - (n - 1) * width_pt
        else:
            lines += 1
            if ww > width_pt:
                n = math.ceil(ww / width_pt)
                lines += n - 1
                x = ww - (n - 1) * width_pt
            else:
                x = ww
    return lines


def para_height_pt(segs, width_pt, spacing, lh=LH):
    size = max((s[1] for s in segs), default=10)
    n = wrap_lines(segs, width_pt)
    return n, n * size * lh * spacing


def lh_for(faces):
    return LH_DESC if DESC in faces else LH


# ───────────────────────────── 글 꾸밈 (아래 · 위첨자 표기) ─────────────────────────────

def parse_markup(text):
    """'D_{AM}' → 아래첨자 · 'x^{2}' → 위첨자.  → [(글, 'sub'|'sup'|None)]"""
    out, pos = [], 0
    for m in re.finditer(r'([_^])\{([^}]*)\}', text):
        if m.start() > pos:
            out.append((text[pos:m.start()], None))
        out.append((m.group(2), 'sub' if m.group(1) == '_' else 'sup'))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], None))
    return out


def plain(text):
    return ''.join(t for t, _ in parse_markup(text))


# ───────────────────────────── 도형 도우미 ─────────────────────────────

def _strip_style(sp):
    st = sp._element.find(qn('p:style'))
    if st is not None:
        sp._element.remove(st)


def _rgb(hexs):
    return RGBColor.from_string(hexs)


def _set_fill(sp, color=None):
    if color is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(color)


def _set_line(sp, color=None, w_pt=0.75, dash=None):
    ln = sp.line
    if color is None:
        ln.fill.background()
        return
    ln.color.rgb = _rgb(color)
    ln.width = Pt(w_pt)
    if dash:
        lnel = ln._get_or_add_ln()
        for e in lnel.findall(qn('a:prstDash')):
            lnel.remove(e)
        pd = etree.SubElement(lnel, qn('a:prstDash'))
        pd.set('val', dash)
        # prstDash 는 solidFill 뒤 · round 앞 (스키마 순서)
        lnel.remove(pd)
        sf = lnel.find(qn('a:solidFill'))
        lnel.insert(list(lnel).index(sf) + 1 if sf is not None else 0, pd)


def _set_typefaces(rpr, kind):
    latin, ea, cs = FONT_SETS[kind]
    for tag in ('latin', 'ea', 'cs'):
        for e in rpr.findall(qn('a:' + tag)):
            rpr.remove(e)
    after = [rpr.find(qn('a:' + t)) for t in ('sym', 'hlinkClick', 'hlinkMouseOver', 'rtl', 'extLst')]
    after = [a for a in after if a is not None]
    els = []
    for tag, face in (('latin', latin), ('ea', ea), ('cs', cs)):
        e = etree.Element(qn('a:' + tag))
        e.set('typeface', face)
        els.append(e)
    if after:
        i0 = list(rpr).index(after[0])
        for k, e in enumerate(els):
            rpr.insert(i0 + k, e)
    else:
        rpr.extend(els)


def add_runs(p, text, size, color=DARK, bold=False, kind='ko'):
    for t, mode in parse_markup(text):
        r = p.add_run()
        r.text = t
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = _rgb(color)
        rpr = r._r.get_or_add_rPr()
        rpr.set('lang', 'ko-KR')
        rpr.set('altLang', 'en-US')
        if mode:
            rpr.set('baseline', '-25000' if mode == 'sub' else '30000')
        _set_typefaces(rpr, kind)


def setup_tf(sp, *, wrap=True, anchor=MSO_ANCHOR.TOP, ins=(0.04, 0.02, 0.04, 0.02)):
    tf = sp.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Emu(I(v)) for v in ins)
    return tf


def fill_paras(tf, paras, size, color=DARK, bold=False, kind='ko', align=PP_ALIGN.LEFT, spacing=1.0):
    """paras = [글 | [(글, 색 | None, 굵게 | None)]]."""
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        segs = [(para, None, None)] if isinstance(para, str) else para
        for t, c, b in segs:
            for j, part in enumerate(t.split('\n')):
                if j:
                    p.add_line_break()
                    brpr = p._p.findall(qn('a:br'))[-1].get_or_add_rPr()
                    brpr.set('sz', str(int(size * 100)))
                    _set_typefaces(brpr, kind)
                if part:
                    add_runs(p, part, size, color=c or color, bold=bold if b is None else b, kind=kind)
        epr = p._p.get_or_add_endParaRPr()
        epr.set('lang', 'ko-KR')
        epr.set('sz', str(int(size * 100)))
        _set_typefaces(epr, kind)


def textbox(slide, name, box, paras, size, *, color=DARK, bold=False, kind='ko', align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            spacing=1.0, wrap=True, ins=(0.04, 0.02, 0.04, 0.02), fill=None):
    x, y, w, h = box
    sp = slide.shapes.add_textbox(Emu(I(x)), Emu(I(y)), Emu(I(w)), Emu(I(h)))
    sp.name = name
    tf = setup_tf(sp, wrap=wrap, anchor=anchor, ins=ins)
    fill_paras(tf, paras, size, color=color, bold=bold, kind=kind, align=align, spacing=spacing)
    if fill:
        _set_fill(sp, fill)
    return sp


def autoshape(slide, name, kind_shape, box_emu, *, fill=None, line=None, line_w=0.75, dash=None, adj=None):
    x, y, w, h = box_emu
    sp = slide.shapes.add_shape(kind_shape, Emu(x), Emu(y), Emu(w), Emu(h))
    sp.name = name
    _strip_style(sp)
    _set_fill(sp, fill)
    _set_line(sp, line, line_w, dash)
    if adj is not None:
        sp.adjustments[0] = adj
    return sp


def shape_text(sp, paras, size, *, color=DARK, bold=False, kind='ko', align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
               spacing=1.0, ins=(0.04, 0.02, 0.04, 0.02), wrap=True):
    tf = setup_tf(sp, wrap=wrap, anchor=anchor, ins=ins)
    fill_paras(tf, paras, size, color=color, bold=bold, kind=kind, align=align, spacing=spacing)
    return sp


def est_box_height(paras, size, width_in, spacing=1.0, ins=(0.04, 0.02, 0.04, 0.02), bold=False, kind='ko'):
    """글 상자에 필요한 높이 (in) — 덱 배치용 (검사는 저장된 XML 에서 같은 함수로 다시 잰다)."""
    wpt = (width_in - ins[0] - ins[2]) * 72
    lh = lh_for(FONT_SETS[kind])
    tot = 0.0
    for para in paras:
        t = para if isinstance(para, str) else ''.join(s[0] for s in para)
        segs = [(tt, size, mode is not None, bold) for tt, mode in parse_markup(t)]
        tot += para_height_pt(segs, wpt, spacing, lh)[1]
    return tot / 72 + ins[1] + ins[3]


# ───────────────────────────── 덱 바탕 (테마 · 레이아웃) ─────────────────────────────

def _deco_xml(i, x, y, w, h, top_w):
    return f'''<p:sp xmlns:p="{NSP}" xmlns:a="{NSA}"><p:nvSpPr><p:cNvPr id="{10 + i}" name="hdr/deco{i}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>
<p:spPr><a:xfrm rot="10800000"><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>
<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="l" t="t" r="r" b="b"/><a:pathLst><a:path w="{w}" h="{h}">
<a:moveTo><a:pt x="{w}" y="{h}"/></a:moveTo><a:lnTo><a:pt x="0" y="{h}"/></a:lnTo><a:lnTo><a:pt x="0" y="0"/></a:lnTo><a:lnTo><a:pt x="{top_w}" y="0"/></a:lnTo><a:close/>
</a:path></a:pathLst></a:custGeom><a:solidFill><a:srgbClr val="0094E9"><a:alpha val="5000"/></a:srgbClr></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>'''


def header_group_xml():
    sx, sy, sw, sh = HDR_SQ
    bx, by, bw, bh = HDR_BAND
    lx, ly, lw, _ = HDR_LINE
    decos = ''.join(_deco_xml(k + 1, *d) for k, d in enumerate(HDR_DECO))
    return f'''<p:grpSp xmlns:p="{NSP}" xmlns:a="{NSA}"><p:nvGrpSpPr><p:cNvPr id="9" name="hdr"/><p:cNvGrpSpPr/><p:nvPr userDrawn="1"/></p:nvGrpSpPr>
<p:grpSpPr><a:xfrm><a:off x="0" y="-1867"/><a:ext cx="9144000" cy="835760"/><a:chOff x="0" y="-1867"/><a:chExt cx="9144000" cy="835760"/></a:xfrm></p:grpSpPr>
{decos}
<p:sp><p:nvSpPr><p:cNvPr id="13" name="hdr/sq"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{sx}" y="{sy}"/><a:ext cx="{sw}" cy="{sh}"/></a:xfrm>
<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:gradFill flip="none" rotWithShape="1"><a:gsLst><a:gs pos="9000"><a:srgbClr val="0094E9"/></a:gs><a:gs pos="100000"><a:srgbClr val="78CCC8"/></a:gs></a:gsLst>
<a:lin ang="2700000" scaled="1"/><a:tileRect/></a:gradFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>
<p:sp><p:nvSpPr><p:cNvPr id="14" name="hdr/band"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{bx}" y="{by}"/><a:ext cx="{bw}" cy="{bh}"/></a:xfrm>
<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:schemeClr val="tx1"><a:alpha val="5000"/></a:schemeClr></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>
<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="15" name="hdr/line"/><p:cNvCxnSpPr/><p:nvPr userDrawn="1"/></p:nvCxnSpPr><p:spPr><a:xfrm><a:off x="{lx}" y="{ly}"/><a:ext cx="{lw}" cy="0"/></a:xfrm>
<a:prstGeom prst="line"><a:avLst/></a:prstGeom><a:ln w="9525"><a:solidFill><a:schemeClr val="bg1"><a:lumMod val="85000"/></a:schemeClr></a:solidFill></a:ln></p:spPr></p:cxnSp>
</p:grpSp>'''


def apply_theme(prs):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    root.set('name', THEME_NAME)
    te = root.find(qn('a:themeElements'))
    cs = te.find(qn('a:clrScheme'))
    cs.set('name', THEME_NAME)
    for tag, val in THEME_COLORS.items():
        el = cs.find(qn('a:' + tag))
        for c in list(el):
            el.remove(c)
        etree.SubElement(el, qn('a:srgbClr')).set('val', val)
    for tag, sys_val, last in (('dk1', 'windowText', '000000'), ('lt1', 'window', 'FFFFFF')):
        el = cs.find(qn('a:' + tag))
        for c in list(el):
            el.remove(c)
        s = etree.SubElement(el, qn('a:sysClr'))
        s.set('val', sys_val)
        s.set('lastClr', last)
    fs = te.find(qn('a:fontScheme'))
    fs.set('name', THEME_NAME)
    for which in ('majorFont', 'minorFont'):
        f = fs.find(qn('a:' + which))
        for c in list(f):
            f.remove(c)
        for tag, face in (('latin', THEME_FONT), ('ea', ''), ('cs', 'Times New Roman')):
            etree.SubElement(f, qn('a:' + tag)).set('typeface', face)
        hang = etree.SubElement(f, qn('a:font'))
        hang.set('script', 'Hang')
        hang.set('typeface', THEME_FONT)
    part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def new_presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    apply_theme(prs)
    keep = [l for l in prs.slide_layouts if l.name == 'Blank'][0]
    for l in list(prs.slide_layouts):
        if l is not keep:
            prs.slide_layouts.remove(l)
    keep._element.find(qn('p:cSld')).set('name', '빈 화면')
    tree = keep.shapes._spTree
    for sp in list(tree.iterchildren(qn('p:sp'))):
        tree.remove(sp)
    tree.append(etree.fromstring(header_group_xml()))
    cp = prs.core_properties
    cp.title = '50kg/batch 4차년도 DEM 해석 7쪽 — 초안 1'
    cp.subject = 'SE 입도 0.7 · 1 · 2 · 3 µm 비교 (1쪽 완성 · 2–7쪽 틀)'
    cp.author = ''
    cp.last_modified_by = ''
    cp.keywords = ''
    cp.comments = 'docs/50kg/report_y4/build_kg50_report_y4.py'
    cp.revision = 1
    cp.created = cp.modified = _dt.datetime(2026, 10, 30, 0, 0, 0)
    cp.last_printed = _dt.datetime(2026, 10, 30, 0, 0, 0)
    return prs, keep


# ───────────────────────────── 장 틀 (머리 · 요약 띠 · 패널 · 결론 줄 · 쪽 번호) ─────────────────────────────

def header(slide, title):
    textbox(slide, 'hdr/sq/num', tuple(inch(v) for v in HDR_NUM), [SECTION_NO], 24, color='FFFFFF', bold=True, align=PP_ALIGN.CENTER,
            wrap=False, ins=HDR_INS, anchor=MSO_ANCHOR.MIDDLE)
    px, py, _, ph = HDR_PROJ
    textbox(slide, 'hdr/proj', (inch(px), inch(py), est_text_width_in(PROJECT, 9) + 0.12, inch(ph)), [PROJECT], 9, color=GRAY_TXT, wrap=False,
            ins=HDR_INS, anchor=MSO_ANCHOR.MIDDLE)
    x, y, h = HDR_TITLE
    w = est_text_width_in(title, 14, bold=True) + 0.12
    textbox(slide, 'hdr/title', (inch(x), inch(y), w, inch(h)), [title], 14, color=DARK, bold=True, wrap=False, ins=HDR_INS,
            anchor=MSO_ANCHOR.MIDDLE)


def est_text_width_in(text, pt, bold=False):
    return sum(seq_widths([(c, pt * (0.65 if mode else 1.0), bold) for t, mode in parse_markup(text) for c in t])) / 72


def band(slide, series, items):
    autoshape(slide, 'band', MSO_SHAPE.ROUNDED_RECTANGLE, BAND_OUT[:4], fill=FAINT, adj=BAND_OUT[4])
    autoshape(slide, 'band/in', MSO_SHAPE.ROUNDED_RECTANGLE, BAND_IN[:4], fill='FFFFFF', line=LINE_GRAY, line_w=1.0, adj=BAND_IN[4])
    tw = I(est_text_width_in(series, 12, bold=True) + 0.30)               # 탭 폭 = 글 + 여백 (양식 6.035 in 은 그 덱 글 길이) · 띠 가운데
    tab = autoshape(slide, 'band/in/tab', MSO_SHAPE.ROUNDED_RECTANGLE, (BAND_OUT[0] + (BAND_OUT[2] - tw) // 2, BAND_TAB[1], tw, BAND_TAB[3]),
                    fill='FFFFFF', line=BLUE, line_w=1.0, adj=BAND_TAB[4])
    shape_text(tab, [series], 12, color=BLUE, bold=True, anchor=MSO_ANCHOR.MIDDLE, wrap=False, ins=(0.1, 0.0, 0.1, 0.0))
    in_right = inch(BAND_IN[0] + BAND_IN[2])
    for k, (title, desc) in enumerate(items):
        (nx, ny), d = ITEM_NUM[k], ITEM_NUM[2]
        c = autoshape(slide, f'band/in/i{k + 1}/num', MSO_SHAPE.OVAL, (nx, ny, d, d), fill='FFFFFF', line=BLUE, line_w=1.0)
        shape_text(c, [str(k + 1)], 10, color=BLUE, anchor=MSO_ANCHOR.MIDDLE, wrap=False, ins=(0, 0, 0, 0))
        tx, ty = ITEM_TITLE[k]
        dx, dy = ITEM_DESC[k]
        w_title = (in_right - 0.12 - inch(tx)) if k else (inch(ITEM_NUM[1][0]) - 0.15 - inch(tx))
        w_desc = (in_right - 0.12 - inch(dx)) if k else (inch(ITEM_NUM[1][0]) - 0.15 - inch(dx))
        textbox(slide, f'band/in/i{k + 1}/title', (inch(tx), inch(ty), w_title, inch(ITEM_TITLE[2])), [title], 11, wrap=False,
                ins=(0.05, 0.02, 0.05, 0.02), anchor=MSO_ANCHOR.MIDDLE)
        textbox(slide, f'band/in/i{k + 1}/desc', (inch(dx), inch(dy), w_desc, inch(ITEM_DESC[2])), [desc], 11, kind='desc', wrap=True,
                ins=(0.05, 0.02, 0.05, 0.02), anchor=MSO_ANCHOR.MIDDLE)


def panel(slide, side, tab_text):
    (x, y), w, h, adj = PANEL[0 if side == 'L' else 1], PANEL[2], PANEL[3], PANEL[4]
    h = I(PANEL_BOTTOM) - y
    autoshape(slide, side, MSO_SHAPE.ROUNDED_RECTANGLE, (x, y, w, h), fill=None, line=DASH_GRAY, line_w=0.75, dash='sysDash', adj=adj)
    tw = est_text_width_in(tab_text, 12) + 0.24
    tx = inch(x) + (inch(w) - tw) / 2
    t = textbox(slide, f'{side}/tab', (tx, inch(PANEL_TAB_Y), tw, inch(PANEL_TAB_H)), [tab_text], 12, color='000000', align=PP_ALIGN.CENTER,
                wrap=False, fill='FFFFFF', anchor=MSO_ANCHOR.MIDDLE, ins=(0.05, 0.0, 0.05, 0.0))
    return (inch(x), inch(y), inch(w), inch(h)), t


def note_line(slide, text):
    textbox(slide, 'note', NOTE_BOX, [text], 9, color=GRAY_TXT, kind='desc', wrap=True, anchor=MSO_ANCHOR.MIDDLE, ins=(0.05, 0.02, 0.05, 0.02))


def page_number(slide):
    x, y, w, h = (inch(v) for v in SLDNUM)
    sp = textbox(slide, 'pgnum', (x, y, w, h), [''], 12, color='404040', kind='desc', align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
                 wrap=False, ins=(0.1, 0.05, 0.1, 0.05))
    p = sp.text_frame.paragraphs[0]._p
    for r in p.findall(qn('a:r')):
        p.remove(r)
    fld = etree.Element(qn('a:fld'))
    fld.set('id', '{6B3E2F1A-5C4D-4E8F-9A1B-2C3D4E5F6A7B}')
    fld.set('type', 'slidenum')
    rpr = etree.SubElement(fld, qn('a:rPr'))
    rpr.set('lang', 'ko-KR')
    rpr.set('sz', '1200')
    sf = etree.SubElement(rpr, qn('a:solidFill'))
    etree.SubElement(sf, qn('a:srgbClr')).set('val', '404040')
    _set_typefaces(rpr, 'desc')
    etree.SubElement(fld, qn('a:t')).text = '‹#›'
    p.insert(list(p).index(p.find(qn('a:endParaRPr'))), fld)


def stack_text(slide, name, x, y, w, paras, size, *, kind='ko', spacing=1.0, color=DARK, bold=False, ins=(0.04, 0.02, 0.04, 0.02),
               pad=0.02):
    """글 상자를 추정 높이 (+ 여유) 로 놓고 다음 y 를 돌려준다."""
    h = est_box_height(paras, size, w, spacing, ins, bold=bold, kind=kind) + pad
    textbox(slide, name, (x, y, w, h), paras, size, kind=kind, spacing=spacing, color=color, bold=bold, ins=ins)
    return y + h


def frame_box(slide, name, box, status):
    sp = autoshape(slide, name, MSO_SHAPE.RECTANGLE, tuple(I(v) for v in box), fill=None, line=DASH_GRAY, line_w=0.75, dash='dash')
    shape_text(sp, [status], 11, color=GRAY_TXT, anchor=MSO_ANCHOR.MIDDLE, ins=(0.1, 0.05, 0.1, 0.05))
    return sp


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ───────────────────────────── 표 ─────────────────────────────

def _cell_border(cell, color=LINE_GRAY, w_pt=0.5):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('lnL', 'lnR', 'lnT', 'lnB'):
        for el in tcPr.findall(qn('a:' + tag)):
            tcPr.remove(el)
    for k, tag in enumerate(('lnL', 'lnR', 'lnT', 'lnB')):
        ln = etree.Element(qn('a:' + tag))
        ln.set('w', str(int(w_pt * 12700)))
        sf = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(sf, qn('a:srgbClr')).set('val', color)
        tcPr.insert(k, ln)


TBL_PT, TBL_INS = 8, (0.03, 0.005, 0.03, 0.005)


def table(slide, name, x, y, col_w, rows, *, head=True, merges=(), row_min=0.17):
    """rows = [[글 …]] · merges = [((r0, c0), (r1, c1))] · 높이 = 칸마다 추정 줄 수 (덱 배치) — 표 끝 y 를 돌려준다."""
    nr, nc = len(rows), len(rows[0])
    heights = []
    for r, row in enumerate(rows):
        hmax = row_min
        for c, txt in enumerate(row):
            if txt is None:
                continue
            span_w = col_w[c]
            for (r0, c0), (r1, c1) in merges:
                if (r0, c0) == (r, c):
                    span_w = sum(col_w[c0:c1 + 1])
            need_h = est_box_height([txt], TBL_PT, span_w, 1.0, TBL_INS)
            for (r0, c0), (r1, c1) in merges:
                if (r0, c0) == (r, c):
                    need_h /= (r1 - r0 + 1)
            hmax = max(hmax, need_h)
        heights.append(round(hmax + 0.005, 3))
    gf = slide.shapes.add_table(nr, nc, Emu(I(x)), Emu(I(y)), Emu(I(sum(col_w))), Emu(I(sum(heights))))
    gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for k in ('firstRow', 'bandRow', 'firstCol', 'lastRow', 'lastCol', 'bandCol'):
        tblPr.set(k, '0')
    for el in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(el)
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Emu(I(cw))
    for i, hh in enumerate(heights):
        tbl.rows[i].height = Emu(I(hh))
    for (r0, c0), (r1, c1) in merges:
        tbl.cell(r0, c0).merge(tbl.cell(r1, c1))
    for r, row in enumerate(rows):
        for c, txt in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = (Emu(I(v)) for v in TBL_INS)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = (head and r == 0) or c == 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(FAINT if is_head else 'FFFFFF')
            _cell_border(cell)
            if txt is None or cell.is_spanned:
                continue
            tf = cell.text_frame
            tf.word_wrap = True
            fill_paras(tf, [txt], TBL_PT, color=DARK, bold=(head and r == 0), kind='ko',
                       align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    return y + sum(heights)


# ───────────────────────────── 1쪽 (완성) ─────────────────────────────

class Nums:
    """1쪽 숫자 출처 장부 — 쪽 글에 쓰는 숫자는 여기서 만든 문자열이다 (검사 C8 이 쪽 글의 숫자를 이 장부와 대조)."""

    def __init__(self):
        self.src = {}

    def __call__(self, text, source):
        for tok in re.findall(r'\d+(?:\.\d+)?', text):
            self.src.setdefault(tok, set()).add(source)
        return text


def page1(prs, beds, common, plan, meta, nums):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    row = plan[0]
    header(s, nums(row['title'], 'plan §2 row 1 title (4차년도)'))
    d50s = [nums(src.fmt_g(b['deck_vals']['d50_se_um']), f"deck {b['case']} r_SE × 2") for b in beds]
    labels = [nums(src.bed_label(b), f"deck {b['case']} r_SE × 2") for b in beds]
    seeds = nums(str(common['seeds_per_size']), 'beds per SE size (sources.BEDS)')
    series = SERIES_FMT.format(sizes=' · '.join(d50s), seeds=seeds)
    nums(series, 'series line = plan year 4 + deck D50 + beds per size')
    ratios = [nums(src.fmt_g(common['d_am_um'] / b['deck_vals']['d50_se_um'], 2), f"deck r_AM_P / r_SE ({b['case']})") for b in beds]
    band(s, series, [
        ('해석 설계 — SE 입도만 다른 네 침대', '- 덱 실행 줄 차이 = SE 반지름 · 시드 · SE 수 · 이름'),
        (f"비교 범위 — 단분산 대리 입도 {' · '.join(d50s)} µm",
         f"- 4차년도 목표 {d50s[0]} µm 포함 · D_{{AM}}/D_{{SE}} {ratios[0]} – {ratios[-1]}"),
    ])
    # ── 왼쪽 패널: 조건 차이표 + 크기비 축 ──
    (lx, ly, lw, lh), _ = panel(s, 'L', '(해석 조건 · 크기비)')
    x0, w0 = lx + 0.10, lw - 0.20
    y = CONTENT_TOP
    y = stack_text(s, 'L/s1', x0, y, w0, ['• 조건 차이표 — 공통 조건 · 침대별 차이'], 10, spacing=1.0)
    c = common
    rA = [
        ['활물질 (AM)', f"NCM D{nums(src.fmt_g(c['d_am_um']), 'deck r_AM_P × 2')} 단일 (r {nums(src.fmt_g(c['r_am_um']), 'deck r_AM_P')} µm) · "
                       f"{nums(src.fmt_g(c['rho_am_gcc']), 'deck pts1 density')} g/cm³"],
        ['고체전해질 (SE)', f"LPSCl · {nums(src.fmt_g(c['rho_se_gcc']), 'deck pts3 density')} g/cm³ · 단분산"],
        ['SE 영률 — 물성', f"{nums(src.fmt_g(round(c['anchors']['dft_evrh_GPa'])), 'round(DFT E_VRH) · anchors doc')} GPa "
                          f"(DFT E_{{VRH}} {nums(src.fmt_g(c['anchors']['dft_evrh_GPa'], 2), 'anchors doc DFT E_VRH')} · 문헌)"],
        ['SE 영률 — DEM 유효 입력', f"{nums(src.fmt_g(c['E_se_GPa'], 2), 'deck youngsModulus type 3 / 0.001')} GPa (선행 이종기술 DEM 과 같은 규약)"],
        ['NCM 영률 (DEM 입력)', f"{nums(src.fmt_g(c['E_am_GPa']), 'deck youngsModulus type 1 / 0.001')} GPa"],
        ['조성 · 면용량', f"AM:SE {nums(src.fmt_g(c['w_am'] * 100), 'deck pdd_mix pts1')}:{nums(src.fmt_g(c['w_se'] * 100), 'deck pdd_mix pts3')} wt% · "
                        f"{nums(src.fmt_g(c['loading_target_mah_cm2']), 'manifest loading_target_mah_cm2')} mAh/cm² "
                        f"({nums(src.fmt_g(c['cap_mah_g']), 'manifest cap_mah_g')} mAh/g)"],
        ['가압 · 해석 영역', f"{nums(src.fmt_g(c['target_MPa']), 'deck target_press / 0.001')} MPa · "
                           f"{nums(src.fmt_g(c['box_x_um']), 'deck reg_box x')} × {nums(src.fmt_g(c['box_y_um']), 'deck reg_box y')} µm (주기 x · y)"],
    ]
    y = table(s, 'L/tA', x0 + 0.02, y + 0.02, [1.56, w0 - 0.04 - 1.56], rA, head=False) + 0.06
    end_rows = []
    for b in beds:
        es = b['endstate']
        if es:
            end_rows.append((nums(str(es['n_am']), f"end dump {b['case']} type 1"), nums(src.fmt_count(es['n_se']), f"end dump {b['case']} type 3")))
        else:
            end_rows.append((b['status'].replace(' ', '\n', 1), None))
    rB = [
        ['침대'] + labels,
        ['SE D50 (µm)'] + d50s,
        ['D_{AM}/D_{SE}'] + ratios,
        ['SE 수 (덱 계획)'] + [nums(src.fmt_count(b['deck_vals']['n_se_plan']), f"deck SE 검사 줄 {b['case']}") for b in beds],
        ['AM 수 (덱 계획)', f"{nums(str(c['n_am_plan']), 'deck AM 검사 줄 (네 덱 같음)')} (네 침대 같음)"] + [None] * (len(beds) - 1),
        ['AM 수 (끝 상태)'] + [e[0] for e in end_rows],
        ['SE 수 (끝 상태)'] + [e[1] for e in end_rows],
    ]
    merges = [((4, 1), (4, len(beds)))] + [((5, k + 1), (6, k + 1)) for k, e in enumerate(end_rows) if e[1] is None]
    cw = (w0 - 0.04 - 1.10) / len(beds)
    y = table(s, 'L/tB', x0 + 0.02, y, [1.10] + [cw] * len(beds), rB, head=True, merges=merges) + 0.06
    y = stack_text(s, 'L/s2', x0, y, w0, ['• 크기비 축 — D_{AM}/D_{SE} · 이상화 삼각 목'], 10)
    rm = meta['size_ratio']
    fw, fh = rm['size_in']
    fx = lx + (lw - fw) / 2
    s.shapes.add_picture(str(HERE / 'figs' / 'p1_size_ratio.png'), Emu(I(fx)), Emu(I(y + 0.02)), Emu(I(fw)), Emu(I(fh))).name = 'L/fig2'
    y += 0.02 + fh + 0.03
    throat = nums(src.fmt_g(common['throat_um'], 3), 'computed (2/√3 − 1) · D_AM from deck r_AM_P')
    nums('2/√3 − 1', 'formula constant (idealized planar throat)')
    cap2 = [f"【 점선 = 세 AM (D{src.fmt_g(c['d_am_um'])}) 이 맞닿은 평면 목의 내접원 d = (2/√3 − 1)·D_{{AM}} = {throat} µm · "
            f"이상화 기준선 (실제 침대의 문턱 아님) · 왼쪽 = 그 목 모식도 】"]
    stack_text(s, 'L/cap2', fx, y, fw, cap2, 7, kind='desc', color=MID)
    # ── 오른쪽 패널: 같은 배율 단면 넷 ──
    (rx, ry, rw, rh), _ = panel(s, 'R', '(같은 배율 단면)')
    x1, w1 = rx + 0.10, rw - 0.20
    y = CONTENT_TOP
    sec = meta['sections']
    y0 = nums(src.fmt_g(sec['y0_um']), 'slice y = box_y / 2 (deck reg_box)')
    y = stack_text(s, 'R/s1', x1, y, w1, [f'• 끝 상태 x–z 단면 — 같은 배율 (y = {y0} µm 평면)'], 10)
    y = stack_text(s, 'R/d1', x1 + 0.06, y, w1 - 0.06,
                   ['- DEM (이산요소법) 끝 상태 = 판 고정 이완 끝 (가압 상태 · 언로딩 없음)',
                    '- 벽 · 판과 겹친 입자 부분은 벽 · 판 선에서 자름'], 8, kind='desc', spacing=1.2)
    fw, fh = sec['size_in']
    fx = rx + (rw - fw) / 2
    fy = y + 0.03
    s.shapes.add_picture(str(HERE / 'figs' / 'p1_sections.png'), Emu(I(fx)), Emu(I(fy)), Emu(I(fw)), Emu(I(fh))).name = 'R/fig1'
    for b in beds:
        if b['endstate'] is not None:
            continue
        lab = src.bed_label(b)
        ax, ay, aw, ah = sec['axes'][lab]['box_in']
        bw, bh = aw - 0.10, 0.42
        st = b['status']
        two = [st[:st.index(' ')], st[st.index(' ') + 1:]]          # '[D1' / '해석 진행 중]' — 칸 폭 안 두 줄
        textbox(s, f'R/fig1/ph-{lab}', (fx + ax + (aw - bw) / 2, fy + ay + (ah - bh) / 2, bw, bh), [' '.join(two).replace(' ', '\n', 1)], 8,
                color=GRAY_TXT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill='FFFFFF', ins=(0.02, 0.02, 0.02, 0.02))
    y = fy + fh + 0.04
    # 범례 (그림 밖 · 편집 가능한 글 — 보고 원칙 §2)
    leg = [('SE (LPSCl)', 'se'), ('AM (NCM)', 'am'), ('기공', 'pore'), ('판 · 바닥 벽', 'wall')]
    lx_cur = fx + 0.10
    for text, key in leg:
        if key == 'wall':
            sw = autoshape(s, f'R/leg/{key}', MSO_SHAPE.RECTANGLE, (I(lx_cur), I(y + 0.06), I(0.16), I(0.035)), fill='000000')
        else:
            colr = {'se': sec['colors']['SE'][1:], 'am': sec['colors']['AM'][1:], 'pore': 'FFFFFF'}[key]
            sw = autoshape(s, f'R/leg/{key}', MSO_SHAPE.RECTANGLE, (I(lx_cur), I(y + 0.02), I(0.11), I(0.11)), fill=colr,
                           line='7F7F7F' if key == 'pore' else None, line_w=0.5)
        tw = est_text_width_in(text, 8) + 0.10
        textbox(s, f'R/leg/{key}-t', (inch(sw.left + sw.width) + 0.03, y, tw, 0.16), [text], 8, wrap=False, ins=(0.02, 0.0, 0.02, 0.0),
                anchor=MSO_ANCHOR.MIDDLE)
        lx_cur = inch(sw.left + sw.width) + 0.03 + tw + 0.16
    y += 0.20
    steps = []
    for b in beds:
        es = b['endstate']
        if es:
            steps.append(f"{src.bed_label(b)} step {nums(str(es['step']), 'end dump TIMESTEP ' + b['case'])} · 판 "
                         f"{nums(src.fmt_g(round(es['plate_z_um'], 1), 1), 'end mesh z ' + b['case'])} µm")
    cap1 = [f"【 끝 상태 덤프 ({' · '.join(steps)}) · x 주기 · 칸마다 같은 배율 · 빈 칸 = 자료 전 】"]
    stack_text(s, 'R/cap1', fx, y, fw, cap1, 7, kind='desc', color=MID)
    note_line(s, f"결론: 비교 범위 = SE D50 {' · '.join(d50s)} µm 단분산 대리 입도 (4차년도 목표 {d50s[0]} µm 포함) · 입도별 시드 {seeds}")
    page_number(s)
    srcs = [f"{src.bed_label(b)}: {b['deck']} (sha256 {b['sha256'][:12]}…)" for b in beds]
    ends = [f"{src.bed_label(b)}: {b['endstate']['tgz']} (sha256 {b['endstate']['tgz_sha256'][:12]}…) · step {b['endstate']['step']}"
            for b in beds if b['endstate']]
    notes(s, '1쪽 출처 (빌드 때 읽음 · kg50_y4_sources.py):\n덱 — ' + ' / '.join(srcs) + '\n끝 상태 — ' + ' / '.join(ends) +
          f"\n물성 — {src.ANCHORS} (DFT E_VRH {c['anchors']['dft_evrh_GPa']} GPa · Bazzoun 2026 DEM 문헌 입력 {c['anchors']['bazzoun_GPa']} GPa) · "
          "SE 영률 두 줄 = 계획 §1 KG50B-06 (모델 계열 = 선행 보고와 같음 · 입계 기전 = 실측 보고 인용만)\n"
          "NCM 140 GPa = DEM 입력 (출처 서술 없음 · CL-90)\n"
          "덱 넷 실행 줄 = 케이스 이름 · SE 반지름 · 시드 · 계획 SE 수 (검사 구간) 밖에서 같음 (생성기 대조 · normalize_deck)\n"
          "감쇠 계수 = 덱 넷 같음 (질량당 감쇠는 SE 크기에 따라 다름 · Q5 대조 전 — v2 덱 README §5 한계)\n"
          f"계획 = {src.PLAN} §2 1쪽 · 단면 = y = {sec['y0_um']} µm 평면 · 그림 = figs/make_figs_p1.py")
    return s


# ───────────────────────────── 2–7쪽 (틀) ─────────────────────────────

FRAMES = {
    2: dict(items=[('국소 점유 — AM 표면 거리별 SE 점유', '- 네 입도 같은 µm 거리 구간 · 정확 합집합'),
                   ('전체 압밀 — 기공도–판 압력', '- 가압 가지 · 고정 높이 이완 가지 따로')],
            L=dict(tab='(AM 표면 거리별 SE 점유)', s='• AM 표면에서 거리 구간별 SE 점유 분율 (같은 µm 구간 · 정확 합집합)',
                   d=['- 껍질 분율의 분모 · 상 겹침 처리 = 계산 전 정의 고정', '- 기하 반례 시험 먼저 (빈 SE · 단일 구 · 중복 구 · 경계 절단 · 회전)'],
                   status=ST_NEW),
            R=dict(tab='(기공도–판 압력)', s='• 기공도 (ε_{sphere} · 전극 부피 기준) – 판 압력',
                   d=['- 가압 가지 (판 이동) · 고정 높이 이완 가지 따로 · 언로딩 없음 (가압 상태)'], status=ST_P2AUX),
            limits='KG50B-01 (국소 점유 = 정의 그대로의 이름) · KG50B-05 (껍질 분율 분모 · 상 겹침 = 계산 전 정의 고정 + 기하 반례 시험 먼저)'),
    3: dict(items=[('피복 지도 — AM 구면 기하 캡 합집합', '- 등면적 투영 · 대표 AM = 사전 규칙'),
                   ('접촉 수 · 크기 — AM 마다 따로', '- 0 접촉 AM 은 따로 수 · 위치')],
            L=dict(tab='(기하 캡 합집합 피복)', s='• AM 구면의 기하 캡 합집합 지도 (등면적 투영) + 전체 AM 피복 분포',
                   d=['- 대표 AM 선택 · 구면 사상 · 캡 겹침 = 계산 전 사전 규칙', '- 기하 반례 시험 먼저 (단일 구 · 중복 구 · 원판 ↔ 캡)'], status=ST_NEW),
            R=dict(tab='(AM 별 접촉 수 · 면적)', s='• AM 마다 (SE 접촉 수, 평균 기하 면적)',
                   d=['- SE 접촉 0 개 AM = 따로 수 · 위치 표시'], status=ST_NEW),
            limits='KG50B-05 (대표 AM 선택 · 구면 사상 · 겹침 = 사전 규칙 · 기하 반례 시험 먼저)'),
    4: dict(items=[('두 tortuosity — 기하 ↔ 수송 짝', '- 같은 침대 · 다른 양 · 비로 묶지 않음'),
                   ('전류 · 소산 — 상위 접촉 · 협착 몫', '- 높이 단면 32 곳 평균 · 경로 보장 아님')],
            L=dict(tab='(기하 · 수송 tortuosity 짝)', s='• 기하 tortuosity (최단 경로 / 두께) ↔ 수송 tortuosity (φ_{SE}·σ₀/σ_{eff})',
                   d=['- 수송 tortuosity = 제곱근 아님 · f 와 한 수송 묶음 (독립 성공 둘로 세지 않음)', '- φ = 전극 부피 기준 (ε_{sphere} 주)'],
                   status=ST_PLAN),
            R=dict(tab='(전류 강조 3D · 소산 몫)', s='• 높이 단면 32 곳 평균으로 단자 전류 80 % 를 나르는 상위 접촉 (경로 보장 아님)',
                   d=['- 협착 ↔ bulk 소산 몫 = 같은 해의 Σ I²R_{c} / Σ I²R · 단자 전력과 합 대조'], status=ST_PLAN),
            limits='KG50B-01 (선형 1 V 이온 망 전류를 전기화학 반응 지표로 읽지 않음 · 전류 그림 이름 = 웹앱 정의 그대로) · KG50B-03 (f · 수송 tortuosity = 한 묶음 · φ 분모 명시)'),
    5: dict(items=[('구조 전달계수 — 공통 σ₀ 의 f(d)', '- 원 솔버 σ · Hertz 주 · Physics 민감도'),
                   ('손익분기선 — 조건부 시나리오', '- 1 µm 쪽이 높은 조건: f(1)/f(3) > {r13}')],
            L=dict(tab='(공통 σ₀ 의 f(d))', s='• SE 입도별 f = σ_{eff}/σ₀ (원 솔버 σ · Hertz 주 · Physics 민감도)',
                   d=['- 입도 비교 = 원 솔버 σ · f 만 · ' + STAGE_E_OK, '- 절대값 인증 아님 · 입도별 시드 1'], status=ST_PLAN),
            R=dict(tab='(손익분기선 · 조건부)', s='• f(1)/f(3) > {s3}/{s1} = {r13} 이면 1 µm 쪽 σ_{eff} 가 높은 시나리오',
                   d=['- σ_{SE} 실측 (3차년도): 1 µm {s1} · 3 µm {s3} mS/cm',
                      '- 0.7 µm 선: f(0.7)/f(3) > {s3}/{s07} = {r07} (D50 {d07_sample} µm 시료 · 대응 미확인 표지)',
                      '- 2 µm 실측 = 결측 · 펠릿 σ 대입 곱 = 조건부 시나리오 (표로만)'], status=ST_FD),
            limits='KG50B-02 (주 = 공통 σ₀ 의 f(d) · 보조 = 손익분기선 조건부 · 입도 우열 결론 문장 없음) · KG50B-07 (Stage-E σ_ion 은 D0.7 에만 가정 감점 → 보고에 안 씀)'),
    6: dict(items=[('응력 분포 — AM · SE 입자 응력', '- Love–Weber · 같은 도구 · 같은 규약끼리'),
                   ('z 구간 — SE 분율 · 전위 강하 · 소산 몫', '- 순 이온 전류 = z 에 평평 → 안 그림')],
            L=dict(tab='(AM · SE 응력 분포)', s='• AM · SE 입자 응력 분포 (Love–Weber)',
                   d=['- 같은 도구 · 같은 규약끼리만 비교 · 프레임 미인증 한정', '- 가압 상태 (판 고정 이완 끝) · 언로딩 없음'], status=ST_PLAN),
            R=dict(tab='(z 구간 SE 분율 · 전위 · 소산)', s='• 두께 (z) 구간별 SE 분율 · 전위 강하 · 소산 몫',
                   d=['- 순 이온 전류는 보존식으로 z 에 평평 → 그리지 않음'], status=ST_NEW),
            note='가압 상태의 비균일성 (반응률 · 균열 · 용량 해석 아님)',
            limits='KG50B-01 (p6 = z 별 SE 분율 · 전위 강하 · 소산 · 응력 · 반응률 · 균열 · 용량 해석 아님)'),
    7: dict(items=[('지표 × 입도 표 — 상태 넷', '- 일치 · 반대/비단조 · 보류 · 계산/모형 한계'),
                   ('다음 검증 — 3D 전기화학', '- 국부 전류밀도 · Li 농도 (계획 덱 8 장)')],
            L=dict(tab='(지표 × 입도 표)', s='• 지표 × 입도 ({sizes} µm) 표 — 칸마다 상태 넷 중 하나',
                   d=['- 상태 넷 = 방향 일치 · 반대 / 비단조 · 판정 보류 · 계산 / 모형 한계',
                      '- 비교 순서 {order} · 주 대비 {first} ↔ {last} · 고정 유한 영역 · 입도별 시드 1'], status=ST_ALL),
            R=dict(tab='(다음 검증)', s='• 다음 = 3D 전기화학 (국부 전류밀도 · Li 농도) 에서 볼 것',
                   d=['- DEM 구조 기반 · 이번 보고에 결과 없음'], status=ST_NEXT),
            limits='KG50B-04 (고정 유한 영역 · 입도별 시드 1 · 시드 산포를 침대 간 오차로 쓰지 않음)'),
}


def plan_numbers():
    """5쪽 손익분기선 숫자 — 계획 · DEM 계획 문서에서 읽고 산술 대조."""
    t_plan = src.rp(src.PLAN).read_text(encoding='utf-8')
    t_dem = src.rp('docs/50kg/dem_plan_20261026.md').read_text(encoding='utf-8')
    m = re.search(r'1 µm σ_ion ([\d.]+) · 3 µm ([\d.]+) mS/cm · 개발 SE \(D50 ([\d.]+) µm\) 이온전도도 ([\d.]+) mS/cm', t_dem)
    src.need(m, 'dem_plan §0: 실측 σ 줄을 찾지 못했다')
    s1, s3, d07, s07 = m.group(1), m.group(2), m.group(3), m.group(4)
    r13, r07 = float(s3) / float(s1), float(s3) / float(s07)
    src.need(f'{s3}/{s1} = {r13:.3f}' in t_plan, f'plan §2: "{s3}/{s1} = {r13:.3f}" 가 없다')
    src.need(f'0.7 = {r07:.3f}' in t_plan, f'plan §2: "0.7 = {r07:.3f}" 가 없다')
    return dict(s1=s1, s3=s3, s07=s07, d07_sample=d07, r13=f'{r13:.3f}', r07=f'{r07:.3f}')


def fill_keys(text, d):
    """{키} 를 d 의 값으로 — 위 · 아래첨자 표기 (_{…} · ^{…}) 와 d 에 없는 키는 그대로."""
    return re.sub(r'(?<![_^])\{(\w+)\}', lambda m: str(d[m.group(1)]) if m.group(1) in d else m.group(0), text)


def frame_page(prs, k, plan, beds, common, pn, extra=None):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    row = plan[k - 1]
    spec = FRAMES[k]
    d50s = [src.fmt_g(b['deck_vals']['d50_se_um']) for b in beds]
    fmtd = dict(pn, sizes=' · '.join(d50s), order=' → '.join(d50s), first=d50s[0], last=d50s[-1])
    header(s, row['title'])
    series = SERIES_FMT.format(sizes=' · '.join(d50s), seeds=common['seeds_per_size'])
    band(s, series, [(fill_keys(t, fmtd), fill_keys(d, fmtd)) for t, d in spec['items']])
    placed = {}
    for side in ('L', 'R'):
        ps = spec[side]
        (px, py, pw, ph), _ = panel(s, side, ps['tab'])
        x0, w0 = px + 0.10, pw - 0.20
        y = CONTENT_TOP
        y = stack_text(s, f'{side}/s1', x0, y, w0, [fill_keys(ps['s'], fmtd)], 10)
        y = stack_text(s, f'{side}/d1', x0 + 0.06, y, w0 - 0.06, [fill_keys(d, fmtd) for d in ps['d']], 8, kind='desc', spacing=1.2)
        placed[side] = (x0, w0, py, ph, y, ps)
    ytop = max(v[4] for v in placed.values()) + 0.06          # 두 칸 = 같은 위치 · 같은 크기 (보고 원칙 §3)
    for side, (x0, w0, py, ph, _, ps) in placed.items():
        box = (x0 + 0.04, ytop, w0 - 0.08, py + ph - 0.10 - ytop)
        if extra and side in extra:
            extra[side](s, box)
        else:
            frame_box(s, f'{side}/fig', box, ps['status'])
    note_line(s, f"쓸 수 있는 결론의 꼴: {spec.get('note', row['conclusion'])}")     # 낱말 = 계획 그대로 (C3 · 괄호만 읽기 쉽게)
    page_number(s)
    drafted = ''
    if extra and 'R' in extra and hasattr(extra['R'], 'meta'):
        mm = extra['R'].meta
        drafted = ('\n보조 그림 = 초안 (figs/make_fig_p2_aux.py) — 로그 (리포 밖 · 메인 세션 반입 예정): ' +
                   ' / '.join(f"{lab}: {d['log_name']} sha256 {d['log_sha256']} ({d['log_bytes']} B) · 판 시작 {d['plate_z0_um']:.3f} µm · "
                              f"압축 step {d['s0']}–{d['s_end']} ({d['n_runs']} run) · 예측 끝 판 {d['z_end_pred_um']:.4f} ↔ 끝 mesh {d['z_end_mesh_um']:.4f} µm"
                              for lab, d in mm['beds'].items()))
    notes(s, f"틀 (자료 받기 전) — 계획 {src.PLAN} §2 {k}쪽 · 주 그림: {row['main']} · 보조: {row['aux']} · 결론의 꼴: {row['conclusion']}{drafted}\n"
             f"한계: {spec['limits']}\n순서 (계획 §4): 3판 + 분석 계획 → 새 지표 도구 (기하 반례 시험 먼저) → 자료 받기 (D3 · D2 두 프레임) → D3 · D2 그림 → D1 → D0.7")
    return s


# ───────────────────────────── 저장 (결정적) ─────────────────────────────

def save_deterministic(prs, path):
    prs.save(str(path))
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zi.create_system = 3
            z.writestr(zi, data)


P2_META = HERE / 'figs' / 'p2_aux_meta.json'
P2_PNG = HERE / 'figs' / 'p2_aux_compaction.png'


def p2_aux_placer(beds):
    """2쪽 보조 칸 = 초안 그림 (figs/make_fig_p2_aux.py · 로그 = 리포 밖) — meta · png 가 있을 때만.  범례 · 빈 침대 상태 · 캡션 = 그림 밖 글."""
    if not (P2_META.exists() and P2_PNG.exists()):
        return None
    meta = json.loads(P2_META.read_text(encoding='utf-8'))
    src.need(hashlib.sha256(P2_PNG.read_bytes()).hexdigest() == meta['files']['p2_aux_compaction.png'], 'p2 aux png ≠ meta sha256')
    have = [src.bed_label(b) for b in beds if src.bed_label(b) in meta['beds']]

    def place(slide, box):
        x, y, w, h = box
        fw, fh = meta['size_in']
        fx = x + (w - fw) / 2
        slide.shapes.add_picture(str(P2_PNG), Emu(I(fx)), Emu(I(y)), Emu(I(fw)), Emu(I(fh))).name = 'R/fig'
        ly = y + fh + 0.04
        cur = fx + 0.10
        for b in beds:                                         # 범례 = 입도 순서 (0.7 → 3) · 자료 없는 침대 = 상태 글
            lab = src.bed_label(b)
            if lab in have:
                col = meta['colors'][lab].lstrip('#')
                autoshape(slide, f'R/leg2/{lab}', MSO_SHAPE.RECTANGLE, (I(cur), I(ly + 0.065), I(0.20), I(0.03)), fill=col)
                text = f"SE {src.fmt_g(b['deck_vals']['d50_se_um'])} µm"
                tw = est_text_width_in(text, 8) + 0.08
                textbox(slide, f'R/leg2/{lab}-t', (cur + 0.23, ly, tw, 0.16), [text], 8, wrap=False, ins=(0.02, 0.0, 0.02, 0.0),
                        anchor=MSO_ANCHOR.MIDDLE)
                cur += 0.23 + tw + 0.12
            else:
                tw = est_text_width_in(b['status'], 8) + 0.08
                textbox(slide, f'R/leg2/{lab}-st', (cur, ly, tw, 0.16), [b['status']], 8, color=GRAY_TXT, wrap=False,
                        ins=(0.02, 0.0, 0.02, 0.0), anchor=MSO_ANCHOR.MIDDLE)
                cur += tw + 0.10
        m = meta['beds']
        fixed = ' · '.join(f"{lab} {m[lab]['porosity_fixed_height_pct']:.2f} %" for lab in have)
        cap = [f"【 초안 · ε = 1 − ΣV_{{구}} / (50 × 50 µm² × 판 높이) · 판 높이 = PLATE HEIGHT − 판 속도 × dt × 압축 step (끝 mesh 와 대조) · "
               f"(a) 판 접촉 뒤 · (b) 판 고정 ε {fixed} 】"]
        stack_text(slide, 'R/cap2', fx, ly + 0.22, fw, cap, 7, kind='desc', color=MID)

    place.meta = meta
    return place


def build(out=OUT_DECK, make_figs=True, extra_p2=None):
    if make_figs:
        import make_figs_p1
        meta = make_figs_p1.build_all(HERE / 'figs')
    else:
        meta = json.loads((HERE / 'figs' / 'p1_figs_meta.json').read_text(encoding='utf-8'))
    beds, common = src.load_all()
    plan = src.plan_rows()
    pn = plan_numbers()
    prs, _ = new_presentation()
    nums = Nums()
    page1(prs, beds, common, plan, meta, nums)
    if extra_p2 is None:
        pl = p2_aux_placer(beds)
        extra_p2 = {'R': pl} if pl else None
    for k in range(2, N_PAGES + 1):
        frame_page(prs, k, plan, beds, common, pn, extra=(extra_p2 if k == 2 else None))
    save_deterministic(prs, out)
    return dict(beds=beds, common=common, plan=plan, meta=meta, nums=nums, pn=pn)


# ───────────────────────────── 검사 (저장된 덱을 다시 읽는다) ─────────────────────────────

class Checks:
    def __init__(self):
        self.items = []

    def add(self, name, ok, detail=''):
        self.items.append(dict(check=name, ok=bool(ok), detail='' if ok else str(detail)[:600]))

    @property
    def failed(self):
        return [c for c in self.items if not c['ok']]


def _xml_box(sp):
    return sp.left, sp.top, sp.width, sp.height


def _shape_paras(sp):
    """저장된 XML 에서 문단마다 [(글, pt, 첨자)] · 줄 간격 — 추정 검사용."""
    out = []
    el = sp._tc if hasattr(sp, '_tc') else sp._element          # 표 칸 (a:tc) 또는 도형
    txBody = el.find('.//' + qn('p:txBody'))
    if txBody is None:
        txBody = el.find('.//' + qn('a:txBody'))
    if txBody is None:
        return out
    for p in txBody.findall(qn('a:p')):
        ppr = p.find(qn('a:pPr'))
        spc = 1.0
        if ppr is not None:
            sp_el = ppr.find(qn('a:lnSpc') + '/' + qn('a:spcPct'))
            if sp_el is not None:
                spc = int(sp_el.get('val')) / 100000
        segs, faces = [], set()
        for r in p.iterchildren(qn('a:r'), qn('a:fld'), qn('a:br')):      # 문서 순서 (줄바꿈 자리 그대로)
            rpr = r.find(qn('a:rPr'))
            if rpr is not None and rpr.find(qn('a:ea')) is not None:
                faces.add(rpr.find(qn('a:ea')).get('typeface'))
            sz = int(rpr.get('sz')) / 100 if rpr is not None and rpr.get('sz') else (segs[-1][1] if segs else 18)
            if r.tag == qn('a:br'):
                segs.append(('\n', sz, False, False))
                continue
            sub = rpr is not None and rpr.get('baseline') not in (None, '0')
            bold = rpr is not None and rpr.get('b') == '1'
            t = r.find(qn('a:t'))
            segs.append((t.text or '', sz, sub, bold))
        out.append((segs, spc, lh_for(faces)))
    return out


def _insets(sp):
    bp = sp._element.find('.//' + qn('a:bodyPr'))
    d = dict(lIns=91440, tIns=45720, rIns=91440, bIns=45720)
    if bp is not None:
        for k in d:
            if bp.get(k) is not None:
                d[k] = int(bp.get(k))
    return d, (bp.get('wrap') if bp is not None else None)


def text_fit(sp):
    """(줄 수 목록, 추정 높이 in, 안 높이 in, 한 줄 폭 초과 여부) — 저장된 상자 기준."""
    d, wrap = _insets(sp)
    w_pt = (sp.width - d['lIns'] - d['rIns']) / EMU * 72
    inner_h = (sp.height - d['tIns'] - d['bIns']) / EMU
    lines, tot, too_wide = [], 0.0, False
    for segs, spc, lh in _shape_paras(sp):
        if not ''.join(sg[0] for sg in segs).strip():
            size = max((sg[1] for sg in segs), default=10)
            lines.append(1)
            tot += size * lh * spc / 72
            continue
        if wrap == 'none':
            chars = [(c, pt * (0.65 if sub else 1.0), b) for t, pt, sub, b in segs for c in t]
            lines_c, cur = [], []
            for ch in chars:
                if ch[0] == '\n':
                    lines_c.append(cur)
                    cur = []
                else:
                    cur.append(ch)
            lines_c.append(cur)
            wpt = max(sum(seq_widths(lc)) for lc in lines_c)
            too_wide = too_wide or wpt > w_pt + 0.5
            n = 1 + sum(sg[0].count('\n') for sg in segs)
        else:
            n = wrap_lines(segs, w_pt)
        size = max((sg[1] for sg in segs), default=10)
        lines.append(n)
        tot += n * size * lh * spc / 72
    return lines, tot, inner_h, too_wide


ONE_LINE = ('hdr/', '/tab', '/title', '/num', 'pgnum', 'note', '/leg/')


def _parent(name, names):
    parts = name.split('/')
    for k in range(len(parts) - 1, 0, -1):
        cand = '/'.join(parts[:k])
        if cand in names:
            return cand
    return None


def _overlap(a, b, tol=I(0.005)):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    dx = min(ax + aw, bx + bw) - max(ax, bx)
    dy = min(ay + ah, by + bh) - max(ay, by)
    return dx > tol and dy > tol


def _inside(child, parent, tol=I(0.01)):
    cx, cy, cw, ch = child
    px, py, pw, ph = parent
    return cx >= px - tol and cy >= py - tol and cx + cw <= px + pw + tol and cy + ch <= py + ph + tol


def table_est_height(gf):
    tbl = gf.table
    heights = []
    for r_i, row in enumerate(tbl.rows):
        hmax = row.height
        for c_i, cell in enumerate(row.cells):
            if cell.is_spanned:
                continue
            span_w = sum(tbl.columns[j].width for j in range(c_i, c_i + cell.span_width))
            ml, mr, mt, mb = cell.margin_left, cell.margin_right, cell.margin_top, cell.margin_bottom
            w_pt = (span_w - ml - mr) / EMU * 72
            tot = 0.0
            for segs, spc, lh in _shape_paras(cell):
                if not ''.join(sg[0] for sg in segs).strip():
                    continue
                n = wrap_lines(segs, w_pt)
                tot += n * max(sg[1] for sg in segs) * lh * spc / 72
            need_h = I(tot) + mt + mb
            span_h = sum(tbl.rows[j].height for j in range(r_i, r_i + cell.span_height))
            if need_h > span_h:
                hmax = max(hmax, row.height + (need_h - span_h))
        heights.append(hmax)
    return sum(heights)


def slide_boxes(slide):
    """이름 → 상자 (EMU) · 표 = 추정 높이 · 무리 (그룹) 는 없음 (이 덱은 슬라이드에 무리를 쓰지 않는다)."""
    boxes = {}
    for sp in slide.shapes:
        x, y, w, h = _xml_box(sp)
        if getattr(sp, 'has_table', False) and sp.has_table:
            h = max(h, table_est_height(sp))
        boxes[sp.name] = (x, y, w, h)
    return boxes


def all_text(slide, with_tables=True):
    parts = []
    for sp in slide.shapes:
        if sp.has_text_frame:
            parts.append(sp.text_frame.text)
        if with_tables and getattr(sp, 'has_table', False) and sp.has_table:
            parts.extend(c.text for r in sp.table.rows for c in r.cells)
    return '\n'.join(parts)


def _norm(s):
    return re.sub(r'\s+', ' ', s).strip()


def check_deck(path, build_info=None, template=None):
    C = Checks()
    prs = Presentation(str(path))
    slides = list(prs.slides)
    C.add('C1 slide count = 7', len(slides) == N_PAGES, len(slides))
    C.add('C1 slide size = template (10 × 7.5 in)', (prs.slide_width, prs.slide_height) == (SLIDE_W, SLIDE_H), (prs.slide_width, prs.slide_height))
    plan = src.plan_rows()
    # C2 제목 = 계획 3판 §2 · C3 회색 결론 줄 ⊇ 계획 결론의 꼴 (낱말 순서대로)
    for i, s in enumerate(slides, 1):
        t = [sp for sp in s.shapes if sp.name == 'hdr/title']
        C.add(f'C2 p{i} title = plan §2 row {i}', bool(t) and _norm(t[0].text_frame.text) == _norm(plan[i - 1]['title']),
              t[0].text_frame.text if t else 'no title')
        n = [sp for sp in s.shapes if sp.name == 'note']
        words = re.findall(r'[\w.]+', plan[i - 1]['conclusion'])
        got = n[0].text_frame.text if n else ''
        pos, ok = 0, bool(n)
        for wd in words:
            j = got.find(wd, pos)
            if j < 0:
                ok = False
                break
            pos = j + len(wd)
        C.add(f'C3 p{i} grey note carries the plan conclusion words in order', ok, f'{words} ⟂ {got}')
    # C4 금지 낱말 · 인용 금지 등록부
    texts = [all_text(s) for s in slides]
    notes_t = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else '' for s in slides]
    hits = [(i, f) for i, t in enumerate(texts, 1) for f in FORBIDDEN if f in t]
    hits += [(i, f'notes:{f}') for i, t in enumerate(notes_t, 1) for f in FORBIDDEN if f in t]
    C.add('C4 forbidden wording absent (plan §1 · principles §5)', not hits, hits)
    ok_e = plain(STAGE_E_OK)
    bad_stage = [(i, m.start()) for i, t in enumerate(texts, 1) for m in re.finditer('Stage-E', t) if t[m.start():m.start() + len(ok_e)] != ok_e]
    C.add(f'C4 "Stage-E" only inside "{ok_e}" (KG50B-07 limit)', not bad_stage, bad_stage)
    ban = json.loads(src.rp('docs/reviews/claims.json').read_text(encoding='utf-8')).get('quotation_ban', [])
    pats = [q.get('pattern') if isinstance(q, dict) else q for q in ban]
    allt = '\n'.join(texts + notes_t)
    C.add('C4 quotation_ban patterns absent (claims.json)', not [p for p in pats if p and p in allt], [p for p in pats if p and p in allt])
    # C5 글 상자 맞춤 · 한 줄 상자
    fit_bad, one_bad, wide_bad = [], [], []
    for i, s in enumerate(slides, 1):
        for sp in s.shapes:
            if not sp.has_text_frame or not sp.text_frame.text.strip():
                continue
            lines, need_h, inner_h, too_wide = text_fit(sp)
            if need_h > inner_h * FIT_SLACK + 1e-6:
                fit_bad.append((i, sp.name, round(need_h, 3), round(inner_h, 3)))
            if any(k in sp.name for k in ONE_LINE) and (max(lines) > 1 + sum(sp.text_frame.text.count(c) for c in ('\n', '\v'))):
                one_bad.append((i, sp.name, lines))
            if too_wide:
                wide_bad.append((i, sp.name))
    C.add('C5 every text box: estimated text height ≤ box (conservative widths: Hangul 1 em · Latin DejaVu)', not fit_bad, fit_bad)
    C.add('C5 one-line boxes (header · tabs · titles · numbers · legend · note) stay one line', not one_bad, one_bad)
    C.add('C5 no-wrap boxes: text width ≤ box', not wide_bad, wide_bad)
    # C6 · C7 상자 겹침 · 포함 (이름 = 계층 · 탭 = 부모 윗선에 걸침 · 머리 띠 안 침범)
    ov_bad, in_bad, hdr_bad = [], [], []
    hdr_bottom = HDR_BAND[1] + HDR_BAND[3]
    for i, s in enumerate(slides, 1):
        boxes = slide_boxes(s)
        names = list(boxes)
        for nm, b in boxes.items():
            par = _parent(nm, boxes)
            if not nm.startswith('hdr/') and b[1] < hdr_bottom:
                hdr_bad.append((i, nm, round(inch(b[1]), 3)))
            if par is None:
                if not _inside(b, (0, 0, SLIDE_W, SLIDE_H)):
                    in_bad.append((i, nm, 'outside slide'))
                continue
            pb = boxes[par]
            if nm.endswith('/tab'):
                cy = b[1] + b[3] / 2
                ok = abs(cy - pb[1]) <= I(0.09) and b[0] >= pb[0] and b[0] + b[2] <= pb[0] + pb[2]
                if not ok:
                    in_bad.append((i, nm, f'tab not centred on {par} top edge'))
            elif not _inside(b, pb):
                in_bad.append((i, nm, f'not inside {par}'))
        for a_i in range(len(names)):
            for b_i in range(a_i + 1, len(names)):
                a, b = names[a_i], names[b_i]
                if a.startswith(b + '/') or b.startswith(a + '/'):
                    continue
                if _overlap(boxes[a], boxes[b]):
                    ov_bad.append((i, a, b))
    C.add('C6 no overlapping shapes except parent ⊃ child (names = hierarchy)', not ov_bad, ov_bad)
    C.add('C6 children inside their parent · tabs straddle the parent top edge · roots inside the slide', not in_bad, in_bad)
    C.add('C6 nothing (but header text) inside the template header band (y < 0.911 in)', not hdr_bad, hdr_bad)
    # C9 · C10 1쪽 그림 · 상태 글
    meta = json.loads((HERE / 'figs' / 'p1_figs_meta.json').read_text(encoding='utf-8'))
    s1 = slides[0]
    pics = {sp.name: sp for sp in s1.shapes if sp.shape_type == 13}
    sc = {}
    for nm, key, stem in (('R/fig1', 'sections', 'p1_sections'), ('L/fig2', 'size_ratio', 'p1_size_ratio')):
        sp = pics.get(nm)
        if sp is None:
            C.add(f'C9 p1 picture {nm} present', False, 'missing')
            continue
        w_in, h_in = meta[key]['size_in']
        sx, sy = inch(sp.width) / w_in, inch(sp.height) / h_in
        sc[nm] = (sx, sy)
        C.add(f'C9 p1 {nm} placed at 100 % (figure size {w_in} × {h_in} in)', abs(sx - 1) < 0.005 and abs(sy - 1) < 0.005, (sx, sy))
        C.add(f'C9 p1 {nm} bytes = figs/{stem}.png = meta sha256',
              sp.image.blob == (HERE / 'figs' / f'{stem}.png').read_bytes()
              and hashlib.sha256(sp.image.blob).hexdigest() == meta[key]['files'][f'{stem}.png'], 'differs')
    ax = meta['sections']['axes']
    ws = {round(v['box_in'][2], 4) for v in ax.values()}
    hs = {round(v['box_in'][3], 4) for v in ax.values()}
    C.add('C9 p1 four section panels: same size (same µm per inch)', len(ws) == 1 and len(hs) == 1, (ws, hs))
    beds_meta = {k: v for k, v in ax.items()}
    ph_bad = []
    for lab, v in beds_meta.items():
        lab_shapes = [sp for sp in s1.shapes if sp.name == f'R/fig1/ph-{lab}']
        if v['data']:
            if lab_shapes:
                ph_bad.append((lab, 'label on a data panel'))
            continue
        if len(lab_shapes) != 1:
            ph_bad.append((lab, f'{len(lab_shapes)} labels'))
            continue
        sp = lab_shapes[0]
        fig = pics['R/fig1']
        axb = (fig.left + I(v['box_in'][0]), fig.top + I(v['box_in'][1]), I(v['box_in'][2]), I(v['box_in'][3]))
        if _norm(sp.text_frame.text) != _norm(v['status']) or not _inside(_xml_box(sp), axb, tol=0):
            ph_bad.append((lab, sp.text_frame.text, 'not the status or not inside its panel'))
    C.add('C10 p1 placeholder panels: label = registered status, inside its empty panel; no label on data panels', not ph_bad, ph_bad)
    # C11 틀 쪽: 칸 둘 · 상태 표지
    fr_bad = []
    for i, s in enumerate(slides[1:], 2):
        figs_ = [sp for sp in s.shapes if sp.name in ('L/fig', 'R/fig')]
        pics_ = [sp for sp in figs_ if sp.shape_type == 13]
        stat = [_norm(sp.text_frame.text) for sp in figs_ if sp.shape_type != 13]
        pic_ok = all(i == 2 and sp.name == 'R/fig' and sp.image.blob == P2_PNG.read_bytes() for sp in pics_)
        if len(figs_) != 2 or any(t not in STATUS_VOCAB for t in stat) or not pic_ok:
            fr_bad.append((i, stat, len(pics_)))
    C.add('C11 pages 2–7: main + aux boxes, each a registered status label (or a drafted figure)', not fr_bad, fr_bad)
    # C12 글꼴 · 크기
    f_bad = []
    allowed = {EN, KO, DESC}
    for i, s in enumerate(slides, 1):
        for el in s._element.iter(qn('a:rPr')):
            faces = {el.find(qn('a:' + t)).get('typeface') if el.find(qn('a:' + t)) is not None else None for t in ('latin', 'ea')}
            if not faces <= allowed or el.get('sz') is None or int(el.get('sz')) < 700:
                f_bad.append((i, faces, el.get('sz')))
    C.add('C12 every run: explicit fonts from {Aptos, KoPubWorld돋움체 Medium, 맑은 고딕} · size ≥ 7 pt', not f_bad, f_bad[:10])
    # C8 1쪽 숫자 = 출처 장부
    if build_info is not None:
        allowed_tok = build_info['nums'].src
        STRUCT = {SECTION_NO: 'section number (same slot as plan deck chapter 4)'}
        t1 = '\n'.join(sp.text_frame.text for sp in s1.shapes if sp.has_text_frame and sp.name != 'pgnum')
        t1 += '\n' + '\n'.join(c.text for sp in s1.shapes if getattr(sp, 'has_table', False) and sp.has_table for r in sp.table.rows for c in r.cells)
        t1n = re.sub(r'10(?=[⁰¹²³⁴⁵⁶⁷⁸⁹])', ' ', t1)       # "× 10⁵" 의 10 = 지수 표기
        toks = re.findall(r'\d+(?:\.\d+)?', t1n)
        unk = sorted({t for t in toks if t not in allowed_tok and t not in STRUCT})
        C.add('C8 every number on page 1 comes from the build-time source ledger (deck · manifest · dump · anchors · plan)', not unk, unk)
        build_info['num_ledger'] = {k: sorted(v) for k, v in sorted(allowed_tok.items())}
    if template:
        for nm, ok, det in template_checks(template):
            C.add(nm, ok, det)
    return C


# ───────────────────────────── 양식 대조 (--template · 읽기만) ─────────────────────────────

def template_checks(path):
    """계획 덱 (리포 밖) 을 python-pptx 로 읽어 이 생성기의 양식 상수와 대조한다 — 아무것도 복사하지 않는다."""
    out = []
    t = Presentation(str(path))
    out.append(('T1 template slide size = 10 × 7.5 in', (t.slide_width, t.slide_height) == (SLIDE_W, SLIDE_H), (t.slide_width, t.slide_height)))
    lay = [l for m in t.slide_masters for l in m.slide_layouts if l.name == '빈 화면']
    grp = [sp for sp in lay[0].shapes if sp.shape_type == 6] if lay else []
    if grp:
        kids = {sp.name: (sp.left, sp.top, sp.width, sp.height) for sp in grp[0].shapes}
        want = {'직사각형 6': HDR_SQ, '직사각형 7': HDR_BAND, '자유형 4': HDR_DECO[0][:4], '자유형 5': HDR_DECO[1][:4]}
        diff = {k: (kids.get(k), v) for k, v in want.items() if kids.get(k) != tuple(v)}
        out.append(('T2 layout "빈 화면" header group geometry = constants', not diff, diff))
    else:
        out.append(('T2 layout "빈 화면" header group geometry = constants', False, 'no group'))
    s7 = t.slides[6]
    by_id = {sp.shape_id: sp for sp in s7.shapes}
    want7 = {129: HDR_NUM, 11: HDR_PROJ, 141: BAND_OUT[:4], 142: BAND_IN[:4], 143: BAND_TAB[:4],
             144: (*PANEL[0], PANEL[2], PANEL[3]), 139: (*PANEL[1], PANEL[2], PANEL[3]),
             148: (*ITEM_NUM[0], ITEM_NUM[2], ITEM_NUM[2]), 149: (*ITEM_NUM[1], ITEM_NUM[2], ITEM_NUM[2]),
             150: (*ITEM_TITLE[0], by_id[150].width if 150 in by_id else 0, ITEM_TITLE[2]),
             151: (*ITEM_DESC[0], by_id[151].width if 151 in by_id else 0, ITEM_DESC[2]), 3: SLDNUM}
    diff = {k: ((by_id[k].left, by_id[k].top, by_id[k].width, by_id[k].height) if k in by_id else None, v)
            for k, v in want7.items() if k not in by_id or (by_id[k].left, by_id[k].top, by_id[k].width, by_id[k].height) != tuple(v)}
    out.append(('T3 plan-deck slide 7 header · band · panel geometry = constants', not diff, diff))
    title = by_id.get(6)
    out.append(('T3 plan-deck slide 7 page title box (x, y, h) = constants', title is not None and (title.left, title.top, title.height) == HDR_TITLE,
                (title.left, title.top, title.height) if title else None))
    faces = set()
    for sp in s7.shapes:
        for el in sp._element.iter(qn('a:rPr')):
            f = el.find(qn('a:ea'))
            if f is not None:
                faces.add(f.get('typeface'))
    out.append(('T4 plan-deck slide 7 Korean typeface includes KoPubWorld돋움체 Medium', KO in faces, faces))
    th = t.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(th.blob)
    cs = root.find(qn('a:themeElements') + '/' + qn('a:clrScheme'))
    got = {tag: cs.find(qn('a:' + tag)).find(qn('a:srgbClr')).get('val') for tag in THEME_COLORS}
    out.append(('T5 theme colours = constants', got == THEME_COLORS, got))
    mf = root.find(qn('a:themeElements') + '/' + qn('a:fontScheme') + '/' + qn('a:minorFont') + '/' + qn('a:latin')).get('typeface')
    out.append(('T5 theme minor latin font = 함초롬돋움', mf == THEME_FONT, mf))
    return out


# ───────────────────────────── 렌더 (LibreOffice → PDF → PNG) · 렌더 검사 ─────────────────────────────

def render(deck, outdir=RENDER_DIR, pdf_out=None, dpi=150):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    pdf_out = Path(pdf_out or deck.with_suffix('.pdf'))
    with tempfile.TemporaryDirectory() as td:
        prof = Path(td) / 'lo_profile'
        subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation=file://{prof}', '--convert-to', 'pdf',
                        '--outdir', td, str(deck)], check=True, capture_output=True, timeout=600)
        shutil.copyfile(Path(td) / (deck.stem + '.pdf'), pdf_out)
    for old in outdir.glob(deck.stem + '-*.png'):
        old.unlink()
    subprocess.run(['pdftoppm', '-r', str(dpi), '-png', str(pdf_out), str(outdir / deck.stem)], check=True, timeout=600)
    return pdf_out, sorted(outdir.glob(deck.stem + '-*.png'))


def _region_text(pg, box, sx, sy, pad=2.0):
    """상자 (EMU) 안에 중심이 있는 렌더 낱말을 줄 · 왼쪽 순서로 공백 없이 잇는다 (sx · sy = pt per in)."""
    x, y, w, h = box
    x0, y0, x1, y1 = inch(x) * sx - pad, inch(y) * sy - pad, inch(x + w) * sx + pad, inch(y + h) * sy + pad
    ws = [wd for wd in pg['words'] if x0 <= (wd['x0'] + wd['x1']) / 2 <= x1 and y0 <= (wd['y0'] + wd['y1']) / 2 <= y1]
    ws.sort(key=lambda wd: (wd['y0'] + wd['y1']) / 2)
    lines, cur = [], []
    for wd in ws:                     # 줄 묶기 = 세로 중심이 낱말 높이 반 안 (글꼴마다 y0 가 달라 반올림 묶기는 쓰지 않는다)
        if cur and abs((wd['y0'] + wd['y1']) / 2 - (cur[-1]['y0'] + cur[-1]['y1']) / 2) > 0.5 * (wd['y1'] - wd['y0']):
            lines.append(cur)
            cur = []
        cur.append(wd)
    if cur:
        lines.append(cur)
    return ''.join(wd['t'] for ln in lines for wd in sorted(ln, key=lambda wd: wd['x0']))


def render_checks(pdf, deck):
    C = Checks()
    with tempfile.TemporaryDirectory() as td:
        html = Path(td) / 'bbox.html'
        subprocess.run(['pdftotext', '-bbox', str(pdf), str(html)], check=True)
        srcx = html.read_text(encoding='utf-8')
    pages = []
    for pm in re.finditer(r'<page width="([\d.]+)" height="([\d.]+)">(.*?)</page>', srcx, flags=re.S):
        w, h, body = float(pm.group(1)), float(pm.group(2)), pm.group(3)
        words = [dict(x0=float(a), y0=float(b), x1=float(c), y1=float(e), t=unescape(t))
                 for a, b, c, e, t in re.findall(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>', body)]
        pages.append(dict(w=w, h=h, words=words))
    C.add('R1 rendered pages = 7', len(pages) == N_PAGES, len(pages))
    over, outside = [], []
    for i, pg in enumerate(pages, 1):
        ws = pg['words']
        for a in ws:
            if a['x0'] < -0.5 or a['y0'] < -0.5 or a['x1'] > pg['w'] + 0.5 or a['y1'] > pg['h'] + 0.5:
                outside.append((i, a['t']))
        for k in range(len(ws)):
            for b in ws[k + 1:]:
                dx = min(ws[k]['x1'], b['x1']) - max(ws[k]['x0'], b['x0'])
                dy = min(ws[k]['y1'], b['y1']) - max(ws[k]['y0'], b['y0'])
                if dx > 1.0 and dy > 1.0:
                    over.append((i, ws[k]['t'], b['t'], round(dx, 1), round(dy, 1)))
    C.add('R2 no two rendered words overlap by > 1 pt in both directions', not over, over[:20])
    C.add('R3 every rendered word inside the page', not outside, outside[:20])
    # R4 상자마다 고유 낱말이 그 상자 안에 그려졌는가 (넘침 · 밀림 검출)
    prs = Presentation(str(deck))
    miss, outbox = [], []
    for i, (s, pg) in enumerate(zip(prs.slides, pages), 1):
        sx, sy = pg['w'] / inch(SLIDE_W) / 72 * 72, pg['h'] / inch(SLIDE_H) / 72 * 72      # pt per in
        shape_tokens = {}
        for sp in s.shapes:
            txt = sp.text_frame.text if sp.has_text_frame else ''
            if getattr(sp, 'has_table', False) and sp.has_table:
                txt = ' '.join(c.text for r in sp.table.rows for c in r.cells)
            shape_tokens[sp.name] = [t for t in re.split(r'\s+', txt) if len(t) >= 2]
        counts = {}
        for toks in shape_tokens.values():
            for t in set(toks):
                counts[t] = counts.get(t, 0) + 1
        rendered = {}
        for wd in pg['words']:
            rendered.setdefault(wd['t'], []).append(wd)
        boxes = slide_boxes(s)
        for nm, toks in shape_tokens.items():
            x, y, w, h = boxes[nm]
            bx0, by0, bx1, by1 = inch(x) * sx - 2, inch(y) * sy - 2, inch(x + w) * sx + 2, inch(y + h) * sy + 2
            for t in set(toks):
                if counts[t] != 1 or len(rendered.get(t, [])) != 1:
                    continue
                wd = rendered[t][0]
                if not (wd['x0'] >= bx0 and wd['x1'] <= bx1 and wd['y0'] >= by0 and wd['y1'] <= by1):
                    outbox.append((i, nm, t, round(wd['x0']), round(wd['y0'])))
        # 글꼴이 바뀌는 자리 (숫자 ↔ 한글) 에서 pdftotext 가 낱말을 가르므로 공백을 뺀 글로 · 그 상자 자리의 낱말만 대조한다
        for sp in s.shapes:
            if sp.name == 'hdr/title' and re.sub(r'\s+', '', sp.text_frame.text) not in _region_text(pg, boxes[sp.name], sx, sy):
                miss.append((i, sp.text_frame.text))
    C.add('R4 unique words of each text box render inside that box (± 2 pt)', not outbox, outbox[:20])
    C.add('R5 page titles rendered in the title box (whitespace-insensitive)', not miss, miss)
    st_miss, on_fig = [], []
    for i, (s, pg) in enumerate(zip(prs.slides, pages), 1):
        sx, sy = pg['w'] / inch(SLIDE_W), pg['h'] / inch(SLIDE_H)
        boxes = slide_boxes(s)
        for sp in s.shapes:
            t = sp.text_frame.text if sp.has_text_frame else ''
            if _norm(t) in STATUS_VOCAB or sp.name.startswith('R/fig1/ph-'):
                if re.sub(r'\s+', '', t) not in _region_text(pg, boxes[sp.name], sx, sy):
                    st_miss.append((i, _norm(t)))
        # R6 그림 위 글 = 그 그림의 자식 (상태 글) 뿐 — 쪽 글이 그림을 덮지 않는다
        boxes = slide_boxes(s)
        for sp in s.shapes:
            if sp.shape_type != 13:
                continue
            x, y, w, h = boxes[sp.name]
            kids = [boxes[n] for n in boxes if n.startswith(sp.name + '/')]
            for wd in pg['words']:
                cx, cy = (wd['x0'] + wd['x1']) / 2 / sx * EMU, (wd['y0'] + wd['y1']) / 2 / sy * EMU
                if x <= cx <= x + w and y <= cy <= y + h:
                    if not any(kx <= cx <= kx + kw and ky <= cy <= ky + kh for kx, ky, kw, kh in kids):
                        on_fig.append((i, sp.name, wd['t']))
    C.add('R5 status labels rendered (whitespace-insensitive)', not st_miss, st_miss)
    C.add('R6 no slide text rendered over a figure except that figure\'s own status labels', not on_fig, on_fig[:20])
    return C


# ───────────────────────────── 실행 ─────────────────────────────

def report(checks, label):
    n, f = len(checks.items), len(checks.failed)
    print(f'{label}: {n} checks · failed {f}')
    for c in checks.failed:
        print('  ✗', c['check'], '—', c['detail'][:300])
    return f


def _shape(slide, name):
    hits = [sp for sp in slide.shapes if sp.name == name]
    assert len(hits) == 1, (name, len(hits))
    return hits[0]


def _set_first_run(sp, text):
    runs = [r for p in sp.text_frame.paragraphs for r in p.runs]
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''


def _mut_number(prs):
    gf = _shape(prs.slides[0], 'L/tB')
    cell = gf.table.cell(3, gf.table.columns.__len__() - 1)
    r = cell.text_frame.paragraphs[0].runs[0]
    r.text = str(int(r.text) + 1)


def _ban_pattern():
    """인용 금지 등록부 (claims.json quotation_ban) 의 첫 패턴 — 이 파일에 금지값을 적지 않으려고 실행 때 읽는다."""
    ban = json.loads(src.rp('docs/reviews/claims.json').read_text(encoding='utf-8')).get('quotation_ban', [])
    pats = [q.get('pattern') if isinstance(q, dict) else q for q in ban]
    return next(p for p in pats if p and len(p) >= 4)


def _mut_font(prs):
    rpr = next(prs.slides[1]._element.iter(qn('a:rPr')))
    rpr.find(qn('a:ea')).set('typeface', 'Arial')


MUTATIONS = [
    # (이름, 바꾸기, 걸려야 할 검사 이름 머리)
    ('shape overlap (p1 L/s2 moved into table B)', lambda p: setattr(_shape(p.slides[0], 'L/s2'), 'top', _shape(p.slides[0], 'L/tB').top + I(0.1)),
     'C6 no overlapping'),
    ('text overflow (p1 band item desc made long)', lambda p: _set_first_run(_shape(p.slides[0], 'band/in/i1/desc'),
                                                                            '- 덱 실행 줄 차이 = SE 반지름 · 시드 · SE 수 · 이름 · 감쇠 · 이웃 · 판 속도 · 목표 압력 · 덤프 간격'),
     'C5 every text box'),
    ('forbidden wording (p3 note "최적 입도")', lambda p: _set_first_run(_shape(p.slides[2], 'note'), '쓸 수 있는 결론의 꼴: 최적 입도 3 µm'), 'C4 forbidden'),
    ('retracted value (p1 note carries a quotation_ban pattern)', lambda p: _set_first_run(_shape(p.slides[0], 'note'), '결론: ' + _ban_pattern()),
     'C4 quotation_ban'),
    ('Stage-E outside the limit phrase (p5)', lambda p: _set_first_run(_shape(p.slides[4], 'L/d1'), '- Stage-E σ 로 비교'), 'C4 "Stage-E"'),
    ('page-1 number not from a source (table B 6184 → 6185)', _mut_number, 'C8 every number'),
    ('figure not at 100 % (p1 R/fig1 width + 2 %)', lambda p: setattr(_shape(p.slides[0], 'R/fig1'), 'width', int(_shape(p.slides[0], 'R/fig1').width * 1.02)),
     'C9 p1 R/fig1 placed at 100 %'),
    ('placeholder label off its panel (p1 D1 label moved 1 in left)', lambda p: setattr(_shape(p.slides[0], 'R/fig1/ph-D1'), 'left',
                                                                                         _shape(p.slides[0], 'R/fig1/ph-D1').left - I(1.0)),
     'C10 p1 placeholder'),
    ('title not the plan title (p3)', lambda p: _set_first_run(_shape(p.slides[2], 'hdr/title'), '접촉 피복'), 'C2 p3'),
    ('note without the plan conclusion (p5)', lambda p: _set_first_run(_shape(p.slides[4], 'note'), '쓸 수 있는 결론의 꼴: 없음'), 'C3 p5'),
    ('font outside the set (p2 first run ea = Arial)', _mut_font, 'C12 every run'),
    ('shape inside the header band (p4 L/s1 at y 0.5 in)', lambda p: setattr(_shape(p.slides[3], 'L/s1'), 'top', I(0.5)), 'C6 nothing'),
    ('status label not registered (p6 R/fig)', lambda p: _set_first_run(_shape(p.slides[5], 'R/fig'), '[완료]'), 'C11 pages 2–7'),
    ('root shape off the slide (p7 note at x 9.5 in)', lambda p: setattr(_shape(p.slides[6], 'note'), 'left', I(9.5)), 'C6 children inside'),
]


def selftest(deck, info):
    """검사기 시험 (변이) — 만든 덱을 하나씩 망가뜨려 그 검사가 실패하는지 본다 (검사기가 맞나 · 같은 덱 원본은 0 실패여야 한다)."""
    out = []
    base = check_deck(deck, build_info=info)
    out.append(dict(mutation='(none — the built deck)', expect='0 failures', ok=not base.failed, failed=[c['check'] for c in base.failed]))
    for name, fn, expect in MUTATIONS:
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / 'm.pptx'
            prs = Presentation(str(deck))
            fn(prs)
            prs.save(str(p))
            C = check_deck(p, build_info=info)
            failed = [c['check'] for c in C.failed]
            out.append(dict(mutation=name, expect=expect, ok=any(f.startswith(expect) for f in failed), failed=failed))
    bad = [o for o in out if not o['ok']]
    print(f'selftest (checker mutations): {len(out)} cases · not caught {len(bad)}')
    for o in bad:
        print('  ✗', o['mutation'], '— expected', o['expect'], '· got', o['failed'])
    return out, len(bad)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default=str(OUT_DECK))
    ap.add_argument('--render', action='store_true', help='LibreOffice → PDF · 쪽 PNG (render/) · 렌더 검사')
    ap.add_argument('--no-figs', action='store_true', help='그림을 다시 만들지 않는다 (figs/ 의 것을 쓴다)')
    ap.add_argument('--template', help='계획 덱 (리포 밖) — 양식 상수 대조만')
    ap.add_argument('--check-only', help='기존 덱 검사만')
    ap.add_argument('--pdf', help='--check-only 와 같이: 그 덱의 PDF 렌더 검사')
    ap.add_argument('--determinism', action='store_true', help='임시 폴더에 두 번 만들어 바이트 대조')
    ap.add_argument('--selftest', action='store_true', help='검사기 변이 시험 (덱을 하나씩 망가뜨려 검사가 걸리는지)')
    ap.add_argument('--json', default=str(OUT_CHECKS))
    a = ap.parse_args()
    if a.determinism:
        with tempfile.TemporaryDirectory() as td:
            p1, p2 = Path(td) / 'a.pptx', Path(td) / 'b.pptx'
            build(p1, make_figs=True)
            build(p2, make_figs=True)
            h1, h2 = (hashlib.sha256(p.read_bytes()).hexdigest() for p in (p1, p2))
            print('determinism:', 'same bytes' if h1 == h2 else 'DIFFERENT', h1[:16], h2[:16])
            return 0 if h1 == h2 else 1
    if a.check_only:
        C = check_deck(Path(a.check_only), template=a.template)
        nf = report(C, 'deck checks')
        if a.pdf:
            nf += report(render_checks(Path(a.pdf), Path(a.check_only)), 'render checks')
        return 1 if nf else 0
    out = Path(a.out)
    info = build(out, make_figs=not a.no_figs)
    C = check_deck(out, build_info=info, template=a.template)
    nf = report(C, 'deck checks')
    rec = dict(deck=out.name, deck_sha256=hashlib.sha256(out.read_bytes()).hexdigest(), checks=C.items,
               page1_number_ledger=info.get('num_ledger'), figs_meta='figs/p1_figs_meta.json',
               decks={src.bed_label(b): dict(path=b['deck'], sha256=b['sha256']) for b in info['beds']})
    if a.selftest:
        st, nbad = selftest(out, info)
        rec['selftest'] = st
        nf += nbad
    if a.render:
        pdf, pngs = render(out)
        R = render_checks(pdf, out)
        nf += report(R, 'render checks')
        rec['render'] = dict(pdf=pdf.name, pngs=[p.name for p in pngs], checks=R.items)
    rec['failed'] = nf
    Path(a.json).write_text(json.dumps(rec, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print('written', out, '·', a.json)
    return 1 if nf else 0


if __name__ == '__main__':
    sys.exit(main())
