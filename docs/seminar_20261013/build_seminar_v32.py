#!/usr/bin/env python3
"""연구세미나 덱 생성기 v3.2 (2026-10-14 · 1저자 결정 — 18 장 · 194 / LHS / ML / 데이터셋 내용 전부 뺌 · 지배방정식 1 장 · Problem 4 1 장).

v3.2 = v3.1 덱 (`build_seminar_v3.py` 를 **그대로** 불러 만든 Codex 종결본) 위의 차분이다 — 같은 코드 경로라 손대지 않은 장의 글은 v3.1 과 같다:
  · 순서 = question → model → what works → what does not → in progress → plan (18 장 · 아래 V32_ORDER)
  · 새 장 둘 = 3 장 DEM 지배방정식 (1저자 10-09 슬라이드의 식 OMML · 그림 · 관계를 그대로 복사 · 한글 라벨 → 영어) · 11 장 Problem 4 (판 반력 "300 MPa" 의 속)
  · 바뀐 장 = 5 (p06 실험 앵커 행) · 7 (p03 문헌) · 9 (p08 Lee R/Q · σ_VGCF* 3.41) · 10 (p09 문헌 notes) · 12–14 (믹서 문헌) · 15 (계획) · 16 (요약 한 줄) · 17 (참고문헌)
  · v3.1 p10 (194 · 데이터셋) 삭제 · 참고문헌 = 인용된 것만 · 번호 다시 (CITE_MAP) · 내용 장마다 다리 줄 하나 (바닥글 띠 · 같은 자리)

    python3 build_seminar_v32.py --template <10-13 덱.pptx> --eq-deck <10-09 지배방정식.pptx> --repo <리포> --figs figs_v3 --lit <크롭 폴더> --out <out.pptx>
    python3 build_seminar_v32.py --selftest-only <deck.pptx> --repo <리포> --figs figs_v3 [--v31 <v3.1 deck>] [--eq-deck <…>] [--json <기록.json>]   # 검사만
    python3 build_seminar_v32.py --dump <deck.pptx> --out <slide_text_v32.md>
"""
import argparse
import copy
import hashlib
import io
import json
import re
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build_seminar_v3 as v3  # noqa: E402  (v3.1 생성기 · 도우미 · v3.1 덱 자체)

NS = dict(v3.NS)
NS.update({'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
           'a14': 'http://schemas.microsoft.com/office/drawing/2010/main',
           'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
A, P_, R_ = NS['a'], NS['p'], NS['r']
RED, BLUE, NAVY, GRAY, DARK = v3.RED, v3.BLUE, v3.NAVY, v3.GRAY, v3.DARK
EMU = 914400

# ───────────────────────────── v3.2 구조 (검사 · 생성 공용) ─────────────────────────────
N_SLIDES = 18
#: v3.2 장 ← v3.1 장 번호 (문자열 = 새 장)
V32_ORDER = [1, 2, 'EQ', 5, 6, 7, 3, 4, 8, 9, 'P4', 11, 12, 13, 14, 15, 16, 17]
#: 손대지 않은 장 (v3.2 → v3.1) — 글 · notes · 표 · 그림은 v3.1 과 같다 (CITE_MAP 번호 바꿈만) · 자리 · 다리 줄만 다를 수 있다
UNTOUCHED = {1: 1, 2: 2, 4: 5, 6: 7, 8: 4, 18: 17}
#: 바뀐 장 (v3.2 → v3.1 원본) · 새 장 = 3 · 11
CHANGED = {5: 6, 7: 3, 9: 8, 10: 9, 12: 11, 13: 12, 14: 13, 15: 14, 16: 15, 17: 16}
NEW_SLIDES = (3, 11)
#: v3.1 → v3.2 인용 번호 (1 · 8 · 9 그대로 · 2–6 · 10 · 11 · 13 · 14 = v3.1 어디에도 인용 없음 → 뺌)
CITE_MAP = {7: 2, 15: 3, 12: 10, 16: 11}
EXPECT = {1: ('', 'Research Seminar'), 2: ('Introduction', 'A coupled DEM–MPM–transport workflow'),
          3: ('Model', 'DEM governing equations: particle motion and contact forces'),
          4: ('DEM → MPM', 'Fixed-shape DEM limits and the material point method (MPM)'),
          5: ('MPM', 'Effective material inputs and their status'),
          6: ('Transport', 'Voxelizing the compacted microstructure into a transport mesh'),
          7: ('DEM', 'Composition-dependent porosity in five simulated beds'),
          8: ('DEM', 'From the contact network to Li⁺ transport'),
          9: ('Limits', 'Problems 1 and 2: grid dependence and absolute σ(e⁻)'),
          10: ('Limits', 'Problem 3: Load sharing during compaction'),
          11: ('Limits', 'Problem 4: what the 300 MPa plate reaction contains'),
          12: ('Mixing', 'Layered-bed mixing: unresolved mixing-index contrast'),
          13: ('Mixing', 'Exploratory cohesion settings produce different mixing indices'),
          14: ('Mixing', 'Surface energy and coating: SE–SE adhesion is the open term'),
          15: ('Plan', 'Tests for unresolved model limits'),
          16: ('Summary', 'Current findings and limitations'),
          17: ('References', None), 18: ('', 'Thank you')}
#: 다리 줄 (내용 장 2–16 · 한 문장 · 새 숫자 없음 · 장 밖 주장 없음 · 같은 자리)
BRIDGES = {2: 'Question: how do packing and compaction shape transport paths?',
           3: 'Next: what a fixed-sphere contact law cannot capture',
           4: 'Next: which inputs are properties and which are effective',
           5: 'Next: from the compacted structure to a transport mesh',
           6: 'Model chain set — next: results from the simulated beds',
           7: 'Next: the same five beds as a Li⁺ contact network',
           8: 'Next: what does not work — grid and absolute σ(e⁻)',
           9: 'Next: how load is shared between NCM and SE contacts',
           10: 'Next: what the nominal plate pressure contains',
           11: 'Next: in progress — dry mixing before pressing',
           12: 'Next: exploratory cohesion settings',
           13: 'Next: why SE–SE adhesion is the open term',
           14: 'Next: the plan for these open items',
           15: 'Next: summary of findings and limitations',
           16: 'Questions and discussion'}
BRIDGE_BOX = (3.70, 7.02, 5.60, 0.28)          # 바닥글 띠 (로고 · "Battery Materials Lab." 오른쪽 · 쪽 번호 왼쪽) · 12 pt 기울임 · 오른쪽 맞춤
BRIDGE_PT = 12
CONTENT_BOTTOM = 6.97                          # 내용 도형 바닥 상한 (바닥글 로고 윗선 6.969 in)
#: 단어 상한 (본문 · 해석 줄 · 표 밖 글 — 출처 줄 · 크롭 출처 · 꼬리표 · 절 표지 · 다리 줄 · 무리 (그림 안 라벨) 제외)
WIDE = (5, 6, 9, 11, 12, 13, 14, 15)
NOT_COUNTED = set(v3.NOT_COUNTED) | {'Bridge'}
BLANK_SLIDES = (4, 6, 12, 13)                  # 빈 자리 (발표자 그림) — 출처 줄에 'figure: presenter'
#: 194 · LHS · ML · 데이터셋 금지 (장 글 + notes) — 예외 하나: 14 장 notes 의 "about 194×" (접착 산술 · 데이터셋 아님 · v3.1 Codex 종결 문장)
DATASET_BAN = [r'194(?!×)', r'(?i)latin', r'(?i)hypercube', r'(?i)machine[\s-]*learning', r'(?i)data\s*set', r'\bLHS\b',
               r'\bML\b', r'(?i)20-bed', r'(?i)\bv1\.[23]\b', r'(?i)\bOOF\b', r'(?i)launch approval']
HANGUL = re.compile('[가-힣㄰-㆏]')
#: v3.1 검사 목록의 장 번호를 v3.2 로 (v3.1 p10 = 삭제 · 그 필수 문구는 9 장으로 옮김)
REQUIRED = {5: ['j. electrochem. soc. 168, 030538'], 6: ['vgcf e = 10 gpa'], 10: ['experimental porosity'],
            9: ['order-of-magnitude agreement', 'σ_vgcf* = 3.41', 'effective value', 'r = 2.07', 'q = 0.84',
                'registered model-sensitivity prediction', 'not within the registered band'],
            12: ['unresolved', 'overlap condition was not met'], 13: ['bond number', 'cohesion-energy densit', 'no verdict'],
            14: ['assumed geometric-mean', 'per unit area', 'literature:'], 16: ['unresolved diagnostic mixing-index contrast'],
            17: ['[9] m. so et al., j. electrochem. soc. 168, 030538 (2021)'],
            3: ['translational and rotational equations of motion', 'rolling-resistance torque', 'global damping'],
            7: ['literature', '[13]', 'not composite beds'],
            11: ['global viscous damping', '57.9 %', '165 mpa', 'model-internal', '750 n'],
            15: ['unloading', 'generation 2', 'insertion', 'platen removal', '1.64 g cm⁻³', 'ε_union', 'heckel', 'history test passed']}
REQUIRED_TABLE = {5: ['lab anchor (planned)', 'eis + dc polarization', '0.1 / 0.2 / 0.3 g', 'polycrystalline', 'single-crystal',
                      'small and large particles']}
NOTES_REQUIRED = {6: ['100 gpa'], 12: ['does not establish compliance', 'hersey 1975 [4]', 'egermann & orr 1983 [5]'],
                  14: ['0.309', 'not a necessary or sufficient condition', 'ordered mixing', 'interactive', 'σ < σ_r', 'kim 2012 [7]'],
                  3: ['docs/dem_governing_equations_report_20261009.md', 'docs/dem_governing_equations_background_20261011.md',
                      'demp-01', '[2]', '[1]', 'non-saturated branch'],
                  5: ['planned lab anchor', 'no results'],
                  7: ['faka 2026 [13]', 'thompson 2026 [12]', 'neither is a validation'],
                  9: ['not "within band"', 'not an experimental hit', 'eight-origin', 'on hold', 'not a fibre property',
                      'docs/reviews/lee_abs_sigma_e_prereg_20260923.md', '3.41'],
                  10: ['thompson 2026 [12]', 'faka 2026 [13]'],
                  11: ['demp-01', 'mv-01', 'mv-02', 'mv-08', 'only after', 'docs/data/ps73_damping_branch_20261007/readme.md'],
                  13: ['kissel 2026 [6]'],
                  15: ['ulkr-01', 'dem_unloading_kit_fix_plan_20261014', 'h0–h5', 'mv-03', 'plan only']}
SLIDE_FORBIDDEN = {13: ['reference cohesion'], 5: ['against independent experiments: pending'],
                   11: ['static 300'], 15: ['194', '20-bed', 'launch approval'], 3: ['하중', '모델']}
#: 참고문헌 (번호 · 화면 글 · 정본 카드 slug · 카드에서 찾을 서지 낱말) — 인용 안 된 v3.1 항목 (2–6 · 10 · 11 · 13 · 14) 은 뺐다
REFS = [
    (1, 'S. Luding, Granul. Matter 10, 235–246 (2008)', 'luding2008_cohesive_frictional_contact_models', ['Luding', '10', '235', '2008']),
    (2, 'LIGGGHTS-PUBLIC 3.8.0 (commit 3d5c00f, 2024), github.com/CFDEMproject/LIGGGHTS-PUBLIC', None, ['3d5c00f']),
    (3, 'A. Sakuda, A. Hayashi, M. Tatsumisago, Sci. Rep. 3, 2261 (2013)', 'sakuda2013_sulfide_mechanical_property', ['Sakuda', '2261', '2013']),
    (4, 'J.A. Hersey, Powder Technol. 11, 41–44 (1975)', 'hersey1975_ordered_mixing_concept', ['Hersey', '11', '41–44', '1975']),
    (5, 'H. Egermann, N.A. Orr, Powder Technol. 36, 117–118 (1983)', 'egermann1983_ordered_vs_interactive_mixture_nomenclature',
     ['Egermann', 'Orr', '36', '117–118', '1983']),
    (6, 'M. Kissel et al., Nat. Commun. 17, 3215 (2026)', 'kissel2026_mechanofusion_derived_cathode_composite_microstructures_scalable',
     ['Kissel', '17', '3215', '2026']),
    (7, 'Y. Kim, H. Lee, S. Kang, J. Mater. Chem. 22, 12874–12881 (2012)', 'kim2012_lno_lco_surface_energy_morphology_dft',
     ['Kim', 'Lee', 'Kang', '22', '12874–12881', '2012']),
    (8, 'A.M. Bazzoun et al., J. Power Sources 661, 238682 (2026)', 'bazzoun2026_dem_fem_rnm_ionic', ['Bazzoun', '661', '238682', '2026']),
    (9, 'M. So et al., J. Electrochem. Soc. 168, 030538 (2021)', 'so2021_dem_fabrication_degradation_ductile_particles', ['So', '168', '030538', '2021']),
    (10, 'D.J. Lee et al., Nat. Commun. 16, 4200 (2025)', 'lee2025_corolling_dryprocess_lpscl_ptfe', ['Lee', '16', '4200', '2025']),
    (11, 'S. Wang, M. Yan, Y. Li, C. Vinado, J. Yang, J. Power Sources 393, 75–82 (2018)', 'wang2018_lco_nmc_electronic_ionic_conductivity_vs_ni',
     ['Wang', '393', '75', '2018']),
    (12, 'O. Thompson et al., Adv. Funct. Mater. 36, e75195 (2026)', 'thompson2026_high_pressure_densification_lps_lpscl_xrd_tomography',
     ['Thompson', '36', 'e75195', '2026']),
    (13, 'V. Faka et al., Adv. Energy Mater. 16, e05186 (2026)', 'faka2026_particle_size_consolidation_strain_lpscl', ['Faka', '16', 'e05186', '2026']),
]
SOFTWARE_PIN = {2: ['docs/data/liggghts_add_pair_pin.json', 'docs/data/liggghts_3d5c00f_dem_eq_source_20261009']}
LIT_BRANCH = v3.LIT_BRANCH
#: 지배방정식 장 — 원 슬라이드에서 복사하는 식 (OMML) · 그림 · 라벨 (cNvPr id)
EQ_IDS = dict(frame=106, tab=109, title=110, every=405, strip=(406, 407, 408, 409, 410, 411, 412),
              labels=(413, 415, 417, 420, 422, 424, 425), eqs=(414, 416, 418, 419, 421, 423, 426, 427),
              seps=(428, 429, 430, 431), force=8)
EQ_NAMES = ('eq1', 'eq2', 'eq3_1', 'eq3_2', 'eq4_1', 'eq4_2', 'eq5_1', 'eq5_2')
#: 복사한 도형의 이름 (원 cNvPr id → 영어 이름 · 식 eq* · 라벨 eq_label_* · 띠 · 구분선은 원 이름 그대로)
EQ_RENAME = {106: 'Panel frame', 107: 'Panel frame picture', 108: 'Panel inner', 109: 'Panel tab', 110: 'Panel title', 405: 'Step every dt',
             8: 'Force diagram', 10: 'Force diagram picture', 11: 'Force label load', 12: 'Force label i', 13: 'Force label j',
             14: 'Force label Fn', 15: 'Force label Ft', 16: 'Force label Mr', 17: 'Force label mg', 18: 'Force caption'}
#: 원 슬라이드 그림 이름 ↔ 새 장 그림 이름 (바이트 대조)
EQ_PIC_PAIRS = (('Picture 21', 'Force diagram picture'), ('그림 26', 'Panel frame picture'))
TRANS = {'DEM 모델': 'DEM model', '지배 방정식·계산 과정': 'Governing equations · calculation steps',
         '① 병진 운동': '① Translation', '② 회전 운동': '② Rotation', '③ 법선 접촉력': '③ Normal contact force',
         '④ 접선 마찰 · 구름 저항': '④ Tangential friction · rolling resistance', '⑤ 접촉 강성': '⑤ Contact stiffness', '하중': 'Load'}
FORCE_CAPTION = ('Forces on particle i', '(same model for particle–particle', 'and particle–wall contacts)')
EQ_LABELS_EN = ['① Translation', '② Rotation', '③ Normal contact force', '④ Tangential friction · rolling resistance', '⑤ Contact stiffness',
                'Load', 'Forces on particle i', 'DEM model', 'Governing equations · calculation steps', 'Contact detection',
                'Contact forces ③–⑤', 'Time integration ①–②']
#: Problem 4 그림 · 값 (기록 문자열 그대로 · figs_v3/make_seminar_figs_v32.py)
P4_FIG = 'F_p4'
SUMMARY_P4_LINE = True                        # 요약 장에 Problem 4 한 줄 (지시 목록 밖 · 일관성 · 끄려면 False)


# ───────────────────────────── 공용 읽기 ─────────────────────────────

def iter_shapes(shapes):
    """무리 (group) 안까지 (python-pptx 가 아는 도형만 — mc:AlternateContent 는 따로)."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from iter_shapes(sh.shapes)


def head_of(slide):
    for sh in slide.shapes:
        if sh.name == 'Rectangle 4' and sh.has_text_frame:
            paras = [p.text for p in sh.text_frame.paragraphs if p.text.strip()]
            return paras
    return []


def section_of(slide):
    for sh in slide.shapes:
        if sh.name == 'TextBox 20' and sh.has_text_frame:
            return sh.text_frame.text.strip()
    return ''


def renum(text, cmap=CITE_MAP):
    def rep(mo):
        nums = [n.strip() for n in mo.group(1).split(',')]
        return '[' + ', '.join(str(cmap.get(int(n), int(n))) for n in nums) + ']'
    return re.sub(r'\[(\d+(?:\s*,\s*\d+)*)\]', rep, text)


def cites(text):
    out = set()
    for mo in re.finditer(r'\[(\d+(?:\s*,\s*\d+)*)\]', text):
        nums = [int(n) for n in mo.group(1).split(',')]
        if nums == [0, 1]:          # Lacey 범위 [0, 1] — 인용 아님
            continue
        out.update(nums)
    return out


def slide_signature(slide, mapping=None):
    """글 (무리 안 포함 · 다리 줄 제외) · 표 · 그림 바이트 · notes — 자리는 보지 않는다."""
    f = (lambda t: renum(t)) if mapping else (lambda t: t)
    texts, tables, pics = [], [], []
    for sh in iter_shapes(slide.shapes):
        if sh.name == 'Bridge':
            continue
        if sh.has_text_frame:
            texts.append((sh.name, tuple(f(p.text) for p in sh.text_frame.paragraphs)))
        if getattr(sh, 'has_table', False) and sh.has_table:
            tables.append((sh.name, tuple(tuple(f(c.text) for c in r.cells) for r in sh.table.rows)))
        if sh.shape_type == 13:
            pics.append((sh.name, hashlib.sha1(sh.image.blob).hexdigest()))
    notes = f(slide.notes_slide.notes_text_frame.text) if slide.has_notes_slide else ''
    return dict(texts=sorted(texts), tables=sorted(tables), pics=sorted(pics), notes=notes)


def all_slide_text(slide, with_groups=True):
    parts = []
    shapes = iter_shapes(slide.shapes) if with_groups else slide.shapes
    for sh in shapes:
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if getattr(sh, 'has_table', False) and sh.has_table:
            parts += [c.text for r in sh.table.rows for c in r.cells]
    return '\n'.join(parts)


def xml_text(slide):
    """슬라이드 XML 의 모든 a:t (mc:AlternateContent 안 · 무리 안 포함)."""
    return '\n'.join(t.text or '' for t in slide._element.iter('{%s}t' % A))


def shape_box(sh):
    return (sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU)


# ───────────────────────────── 검사 (고치기 전 v3.1 덱에서 실패 확인 → 생성 → 전부 통과) ─────────────────────────────

class Checks:
    def __init__(self):
        self.items = []

    def add(self, name, ok, detail=''):
        self.items.append((name, bool(ok), detail if not ok else ''))

    @property
    def failed(self):
        return [(n, d) for n, ok, d in self.items if not ok]


def _git_show(repo, rel):
    r = subprocess.run(['git', '-C', str(repo), 'show', f'{LIT_BRANCH}:{rel}'], capture_output=True)
    return r.stdout.decode('utf-8', 'replace') if r.returncode == 0 else None


def _c14n(el):
    return etree.tostring(el, method='c14n')


def _source_eq(eq_deck):
    """원 지배방정식 슬라이드 — 식 이름 → (a14:m C14N, 대체 그림 바이트) · 힘 그림 · 틀 그림 바이트."""
    src = Presentation(str(eq_deck))
    ss = src.slides[0]
    blobs = {rid: rel.target_part.blob for rid, rel in ss.part.rels.items() if rel.reltype.endswith('/image')}
    out = {}
    for ac in ss._element.iter('{%s}AlternateContent' % NS['mc']):
        name = ac.find('.//{%s}cNvPr' % P_).get('name')
        m = ac.find('.//{%s}m' % NS['a14'])
        blip = ac.find('{%s}Fallback//{%s}blip' % (NS['mc'], A))
        out[name] = (_c14n(m), blobs[blip.get('{%s}embed' % R_)])
    pics = {}
    for pic in ss._element.iter('{%s}pic' % P_):
        nm = pic.find('.//{%s}cNvPr' % P_).get('name')
        blip = pic.find('.//{%s}blip' % A)
        pics[nm] = blobs[blip.get('{%s}embed' % R_)]
    return out, pics


def selftest_v32(deck, repo, figs=None, v31=None, eq_deck=None):
    """v3.2 요구 + v3.1 검사 (장 번호 바꿈).  모든 실패를 모은다 (첫 실패에서 멈추지 않는다)."""
    repo = Path(repo)
    C = Checks()
    prs = Presentation(str(deck))
    slides = list(prs.slides)
    n = len(slides)
    C.add('A1 slide count = 18', n == N_SLIDES, f'{n}')
    texts = [all_slide_text(s) for s in slides]
    notes = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else '' for s in slides]
    low = [t.lower() for t in texts]
    lown = [t.lower() for t in notes]
    get = lambda i: (texts[i - 1], notes[i - 1]) if i <= n else ('', '')

    # A2 순서 · 제목 · 절 표지
    for i, (sec, head) in EXPECT.items():
        if i > n:
            C.add(f'A2 slide {i} exists', False, 'missing')
            continue
        s = slides[i - 1]
        if i in (1, 18):
            C.add(f'A2 slide {i} = {head}', head in texts[i - 1], texts[i - 1][:60])
            continue
        if i == 17:
            C.add('A2 slide 17 = References', section_of(s) == 'References', section_of(s))
            continue
        hp = head_of(s)
        C.add(f'A2 slide {i} section "{sec}"', section_of(s) == sec, section_of(s))
        C.add(f'A2 slide {i} title', bool(hp) and hp[0] == head, (hp[0] if hp else '—')[:70])

    # A3 194 · LHS · ML · 데이터셋 없음 (장 글 · 무리 안 · notes)
    for i in range(1, n + 1):
        t, nt = get(i)
        xt = xml_text(slides[i - 1])            # 장 XML 의 모든 a:t (무리 · 표 · 식 대체 포함)
        for pat in DATASET_BAN:
            for where, body in (('slide', xt), ('notes', nt)):
                for mo in re.finditer(pat, body):
                    ctx = body[max(0, mo.start() - 25): mo.end() + 25].replace('\n', ' ')
                    C.add(f'A3 no dataset content ({where} {i}: {pat})', False, ctx)
    C.add('A3 dataset sweep ran', True)
    # A3b 남은 "194" = 14 장 notes 의 접착 산술 하나뿐
    hits194 = [(i, where) for i in range(1, n + 1) for where, body in (('slide', xml_text(slides[i - 1])), ('notes', notes[i - 1]))
               for _ in re.finditer('194', body)]
    C.add('A3b only remaining "194" = adhesion ratio "about 194×" (slide 14 notes)',
          hits194 == [] or (hits194 == [(14, 'notes')] and 'about 194×' in notes[13]), str(hits194))

    # A4 다리 줄 — 내용 장 2–16 마다 하나 · 같은 자리 · 12 pt 기울임 · 한 줄 · 숫자 없음
    geoms = set()
    for i in range(2, 17):
        if i > n:
            C.add(f'A4 bridge on slide {i}', False, 'missing slide')
            continue
        br = [sh for sh in slides[i - 1].shapes if sh.name == 'Bridge']
        if len(br) != 1:
            C.add(f'A4 bridge on slide {i}', False, f'{len(br)} bridge shapes')
            continue
        sh = br[0]
        txt = sh.text_frame.text
        geoms.add(tuple(round(v, 3) for v in shape_box(sh)))
        sizes = {r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size}
        ital = all(r.font.italic for p in sh.text_frame.paragraphs for r in p.runs)
        C.add(f'A4 bridge slide {i} text = registered', txt == BRIDGES[i], txt)
        C.add(f'A4 bridge slide {i} 12 pt italic', sizes == {BRIDGE_PT} and ital, str(sizes))
        C.add(f'A4 bridge slide {i} one line', v3.est_lines(txt, BRIDGE_PT, BRIDGE_BOX[2] - 0.15) == 1, txt)
        C.add(f'A4 bridge slide {i} no digits', not re.search(r'[0-9]', txt), txt)
    C.add('A4 bridge placement identical on all content slides', len(geoms) == 1, str(geoms))
    for i in (1, 17, 18):
        if i <= n:
            C.add(f'A4 no bridge on slide {i}', not any(sh.name == 'Bridge' for sh in slides[i - 1].shapes))

    # A5 손대지 않은 장 = v3.1 (CITE_MAP 번호 바꿈만 · 자리 · 다리 줄 제외)
    if v31:
        old = list(Presentation(str(v31)).slides)
        for i, j in UNTOUCHED.items():
            if i > n:
                C.add(f'A5 untouched slide {i} (= v3.1 p{j:02d})', False, 'missing')
                continue
            a, b = slide_signature(slides[i - 1]), slide_signature(old[j - 1], mapping=True)
            for k in ('texts', 'tables', 'pics', 'notes'):
                C.add(f'A5 untouched slide {i} (= v3.1 p{j:02d}) {k}', a[k] == b[k], 'differs')
    else:
        C.add('A5 untouched-slide comparison needs --v31', False, 'no v3.1 deck given')

    # A6 참고문헌 = 정본 카드 slug · 인용된 것만 · 번호 1..13 연속
    refs_txt = texts[16] if n >= 17 else ''
    entries = [ln.strip() for ln in refs_txt.split('\n') if re.match(r'\s*\[\d+\]', ln)]
    nums = [int(re.match(r'\s*\[(\d+)\]', e).group(1)) for e in entries]
    C.add('A6 reference numbers 1..13 contiguous', nums == list(range(1, len(REFS) + 1)), str(nums))
    want = {f'[{k}] {t}' for k, t, _, _ in REFS}
    C.add('A6 reference lines = registered list', set(entries) == want, str(sorted(set(entries) ^ want))[:200])
    cited = set()
    for i in range(1, n + 1):
        if i == 17:
            continue
        cited |= cites(texts[i - 1]) | cites(notes[i - 1])
    C.add('A6 every reference cited (slides or notes, refs slide excluded)', cited >= set(nums) and bool(nums),
          f'uncited {sorted(set(nums) - cited)}')
    C.add('A6 every citation has a reference', cited <= set(nums), f'dangling {sorted(cited - set(nums))}')
    for k, t, slug, toks in REFS:
        if slug is None:
            ok = all((repo / p).exists() for p in SOFTWARE_PIN[k])
            C.add(f'A6 ref [{k}] software pin files exist', ok, str(SOFTWARE_PIN[k]))
            continue
        card = _git_show(repo, f'litdb/papers/{slug}.md')
        C.add(f'A6 ref [{k}] canonical card exists ({slug})', card is not None, slug)
        if card:
            miss = [x for x in toks if x not in card]
            C.add(f'A6 ref [{k}] bibliographic fields found in card', not miss, str(miss))

    # A7 장별 필수 · 금지 문구
    for i, rs in REQUIRED.items():
        if i <= n:
            miss = [r for r in rs if r not in low[i - 1]]
            C.add(f'A7 required wording slide {i}', not miss, str(miss))
    for i, rs in REQUIRED_TABLE.items():
        if i <= n:
            tab = ' '.join(c.text for sh in slides[i - 1].shapes if sh.has_table for r in sh.table.rows for c in r.cells).lower()
            miss = [r for r in rs if r not in tab]
            C.add(f'A7 required table wording slide {i}', not miss, str(miss))
    for i, rs in NOTES_REQUIRED.items():
        if i <= n:
            miss = [r for r in rs if r not in lown[i - 1]]
            C.add(f'A7 required notes slide {i}', not miss, str(miss))
    for i, fs in SLIDE_FORBIDDEN.items():
        if i <= n:
            bad = [f for f in fs if f in texts[i - 1] or f in xml_text(slides[i - 1])]
            C.add(f'A7 forbidden wording slide {i}', not bad, str(bad))
    sn11 = notes[10] if n >= 11 else ''
    bad_static = [s for s in re.split(r'(?<=[.;])\s', sn11) if 'static 300' in s and 'only after' not in s]
    C.add('A7 slide 11 notes: "static 300 MPa" only inside the conditional sentence', not bad_static, str(bad_static)[:200])
    slide_txt_all = '\n'.join(texts)
    bad = [f for f in v3.FORBIDDEN if f in slide_txt_all]
    C.add('A7 v3.1 flagged phrases absent', not bad, str(bad))
    bad_ci = [(i, f) for i, x in enumerate(low, 1) for f in v3.FORBIDDEN_CI if f in x]
    C.add('A7 v3.1 case-insensitive flagged phrases absent', not bad_ci, str(bad_ci))
    bad_n = [(i, f) for i, x in enumerate(notes, 1) for f in v3.NOTES_FORBIDDEN if f in x]
    C.add('A7 v3.1 notes-forbidden phrases absent', not bad_n, str(bad_n))
    hang = [i for i in range(1, n + 1) if HANGUL.search(xml_text(slides[i - 1]))]
    C.add('A7 no Korean text on slides (notes may quote Korean headings)', not hang, str(hang))

    # A8 새 장 3 — 식 OMML · 대체 그림 · 그림 = 원 슬라이드 바이트 그대로 · 라벨 영어
    if n >= 3:
        s3 = slides[2]
        acs = list(s3._element.iter('{%s}AlternateContent' % NS['mc']))
        names = [ac.find('.//{%s}cNvPr' % P_).get('name') for ac in acs]
        C.add('A8 slide 3 has the 8 equation objects', sorted(names) == sorted(EQ_NAMES), str(names))
        ok_ns = all(ac.find('{%s}Choice' % NS['mc']) is not None and ac.find('{%s}Choice' % NS['mc']).get('Requires') == 'a14'
                    and 'a14' in ac.find('{%s}Choice' % NS['mc']).nsmap for ac in acs) and bool(acs)
        C.add('A8 equation Choice keeps Requires="a14" with the a14 namespace in scope', ok_ns)
        x3 = xml_text(s3)
        miss = [t for t in EQ_LABELS_EN if t not in x3]
        C.add('A8 slide 3 English labels present', not miss, str(miss))
        if eq_deck:
            src_eq, src_pics = _source_eq(eq_deck)
            blobs = {rid: rel.target_part.blob for rid, rel in s3.part.rels.items() if rel.reltype.endswith('/image')}
            for ac in acs:
                nm = ac.find('.//{%s}cNvPr' % P_).get('name')
                m = ac.find('.//{%s}m' % NS['a14'])
                blip = ac.find('{%s}Fallback//{%s}blip' % (NS['mc'], A))
                same_m = nm in src_eq and _c14n(m) == src_eq[nm][0]
                same_f = nm in src_eq and blip is not None and blobs.get(blip.get('{%s}embed' % R_)) == src_eq[nm][1]
                C.add(f'A8 {nm} math XML byte-identical to the 10-09 slide (C14N)', same_m)
                C.add(f'A8 {nm} fallback image identical to the 10-09 slide', same_f)
            got = {}
            for pic in s3._element.iter('{%s}pic' % P_):
                nm = pic.find('.//{%s}cNvPr' % P_).get('name')
                bl = pic.find('.//{%s}blip' % A)
                got[nm] = blobs.get(bl.get('{%s}embed' % R_)) if bl is not None else None
            for src_nm, nm in EQ_PIC_PAIRS:
                C.add(f'A8 picture "{nm}" identical to the 10-09 slide ("{src_nm}")',
                      got.get(nm) is not None and got.get(nm) == src_pics.get(src_nm))
        else:
            C.add('A8 equation byte check needs --eq-deck', False, 'no source deck given')

    # A9 새 그림 = 기록값에서만 (Problem 4)
    if figs:
        figs = Path(figs)
        csvp = figs / f'{P4_FIG}.csv'
        C.add('A9 F_p4 files exist (png · svg · csv)', all((figs / f'{P4_FIG}.{e}').exists() for e in ('png', 'svg', 'csv')))
        if csvp.exists():
            sys.path.insert(0, str(figs))
            import importlib
            fm = importlib.import_module('make_seminar_figs_v32')
            readme = (repo / fm.README).read_text(encoding='utf-8')
            import csv as _csv
            with open(csvp, encoding='utf-8') as f:
                rows = list(_csv.reader(f))[1:]
            C.add('A9 F_p4.csv values = recorded strings of README', all(r[2] in readme for r in rows) and len(rows) == 11, str(rows)[:120])
            try:
                fm.selftest(repo)
                C.add('A9 make_seminar_figs_v32 --selftest', True)
            except AssertionError as e:
                C.add('A9 make_seminar_figs_v32 --selftest', False, str(e)[:160])
        if n >= 11:
            png = figs / f'{P4_FIG}.png'
            pics11 = [sh for sh in slides[10].shapes if sh.shape_type == 13]
            C.add('A9 slide 11 carries F_p4.png byte-identical', png.exists() and any(sh.image.blob == png.read_bytes() for sh in pics11))

    # A10 v3.1 검사 그대로 (장 번호 바꿈) — 철회값 · 내부 ID · 팔 코드 · 꺾쇠 · 빈 자리 · 출처 줄 · 겹침 · 단어 · 글씨 · 그림
    txt = '\n'.join(texts + notes + [xml_text(s) for s in slides])
    ban = json.loads((repo / 'docs/reviews/claims.json').read_text(encoding='utf-8')).get('quotation_ban', [])
    pats = [q.get('pattern') if isinstance(q, dict) else q for q in ban]
    hits = [p for p in pats if p and p in txt]
    C.add('A10 banned strings (claims.json quotation_ban) = 0', not hits, str(hits))
    leak = re.findall(v3.INTERNAL, '\n'.join(texts + [xml_text(s) for s in slides]))
    C.add('A10 internal identifiers on slides = 0', not leak, str(leak))
    codes = re.findall(v3.ARM_CODES, txt)
    C.add('A10 mixer arm codes (slides + notes) = 0', not codes, str(codes))
    angle = [i for i, t in enumerate(texts, 1) if re.search(r'(^|\n)< .* >($|\n)', t)]
    C.add('A10 angle-bracket captions = 0', not angle, str(angle))
    for i in BLANK_SLIDES:
        if i <= n:
            src = [sh.text_frame.text for sh in slides[i - 1].shapes if sh.name == 'TextBox 69']
            C.add(f'A10 blank-figure slide {i} source line says "figure: presenter"', bool(src) and 'figure: presenter' in src[0], str(src))
    long_src = [(i, len(sh.text_frame.text)) for i, s in enumerate(slides, 1) for sh in s.shapes if sh.name == 'TextBox 69' and len(sh.text_frame.text) > 110]
    C.add('A10 source lines ≤ 110 characters', not long_src, str(long_src))
    ov = overlap_check_v32(slides)
    C.add('A10 body text above the figure zone (overlap check incl. groups)', not ov, str(ov))
    low_b = bottom_check(slides)
    C.add(f'A10 content above the footer band (bottom ≤ {CONTENT_BOTTOM} in)', not low_b, str(low_b))
    stats = slide_stats_v32(prs)
    heavy = [(i, w) for i, (w, _, _) in enumerate(stats, 1) if 2 <= i <= 16 and w > (120 if i in WIDE else 100)]
    C.add('A10 words ≤ 100 / 120 (small print, tables, groups excluded)', not heavy, str(heavy))
    small = [(i, bp) for i, (_, bp, _) in enumerate(stats, 1) if bp and min(bp) < v3.BODY_PT]
    C.add(f'A10 body bullets ≥ {v3.BODY_PT} pt', not small, str(small))
    cap_small = [(i, c) for i, (_, _, c) in enumerate(stats, 1) if c is not None and c < v3.CAP_PT]
    C.add(f'A10 note lines ≥ {v3.CAP_PT} pt', not cap_small, str(cap_small))
    pics = sum(1 for s in slides for sh in iter_shapes(s.shapes) if sh.shape_type == 13)
    C.add('A10 pictures ≥ 14', pics >= 14, str(pics))
    if figs:
        svg = (Path(figs) / 'F_limits.svg').read_text(encoding='utf-8')
        labels = [k for k in ('0:10', '3:7', '5:5', '7:3', '10:0') if k in svg]
        C.add('A10 F_limits composition labels 5/5', len(labels) == 5, str(labels))
    for i in (12, 13, 14):
        if i <= n:
            C.add(f'A10 mixing slide {i} keeps "cohesi"', 'cohesi' in texts[i - 1])

    # A12 장마다 도형 id 고유 (식의 Choice · Fallback 은 한 도형이라 같은 id)
    dups = [(i, d) for i, s in enumerate(slides, 1) for d in [_dup_ids(s)] if d]
    C.add('A12 shape ids unique per slide', not dups, str(dups)[:200])

    # A11 바뀐 · 새 장 notes 의 리포 경로 = 실재
    for i in sorted(set(CHANGED) | set(NEW_SLIDES)):
        if i > n:
            continue
        paths = sorted(set(p.rstrip('.,;:)') for p in re.findall(r'(?:docs|scripts|webapp)/[\w./§·-]+', notes[i - 1])))
        missing = [p for p in paths if not (repo / p.split('§')[0].rstrip('/ ')).exists()]
        C.add(f'A11 slide {i} notes cite repo sources that exist', bool(paths) and not missing, f'paths {len(paths)} missing {missing}')
    return C


def _in_fallback(el):
    p = el.getparent()
    while p is not None:
        if etree.QName(p).localname == 'Fallback':
            return True
        p = p.getparent()
    return False


def _dup_ids(slide):
    ids = [c.get('id') for c in slide._element.iter('{%s}cNvPr' % P_) if not _in_fallback(c)]
    return sorted({i for i in ids if ids.count(i) > 1})


def _unique_ids(slide):
    """같은 장 안에서 겹친 도형 id 를 새로 매긴다 (처음 것은 그대로 · 식의 Fallback id = Choice id).  글 · 자리 · 관계는 그대로."""
    root = slide._element
    cnv = [c for c in root.iter('{%s}cNvPr' % P_)]
    top = max(int(c.get('id')) for c in cnv)
    seen = set()
    for c in cnv:
        if _in_fallback(c):
            continue
        i = int(c.get('id'))
        if i in seen:
            top += 1
            c.set('id', str(top))
            i = top
        seen.add(i)
    for ac in root.iter('{%s}AlternateContent' % NS['mc']):
        ch = ac.find('{%s}Choice//{%s}cNvPr' % (NS['mc'], P_))
        fb = ac.find('{%s}Fallback//{%s}cNvPr' % (NS['mc'], P_))
        if ch is not None and fb is not None:
            fb.set('id', ch.get('id'))


def overlap_check_v32(slides):
    """v3.1 겹침 검사 + 무리 (group) 도형도 그림 띠 윗선에 넣는다 (지배방정식 장)."""
    bad = []
    for i, s in enumerate(slides, 1):
        body = [sh for sh in s.shapes if sh.name == 'Rectangle 4']
        if not body or i == 1:
            continue
        paras = [p.text for p in body[0].text_frame.paragraphs if p.text.strip()]
        if len(paras) < 2:
            continue
        head, bullets = paras[0], paras[1:]
        est = v3.est_body_bottom(head, bullets)
        tops = [sh.top / EMU for sh in s.shapes
                if sh.shape_type in (13, 6) or sh.has_table or sh.name.startswith('Flow') or sh.name in v3.ZONE_NAMES
                or sh.name in ('Callout', 'Callout damping', 'Gen2 head', 'Gen2 note') or sh.name.startswith('Gen2 box')]
        if not tops:
            continue
        if len(head) > v3.CHARS_HEAD or v3.est_lines(head, v3.HEAD_PT, v3.HEAD_W) != 1 or est > min(tops) + 0.02:
            bad.append((i, round(est, 2), round(min(tops), 2), len(head)))
    return bad


def bottom_check(slides):
    bad = []
    for i, s in enumerate(slides, 1):
        if i in (1, 17, 18):
            continue
        for sh in s.shapes:
            if sh.name == 'Bridge' or sh.top is None:
                continue
            bot = (sh.top + sh.height) / EMU
            if bot > CONTENT_BOTTOM + 1e-6:
                bad.append((i, sh.name, round(bot, 3)))
    return bad


def slide_stats_v32(prs):
    """v3.1 slide_stats 와 같다 — 다만 다리 줄도 작은 글 (단어 수 제외)."""
    out = []
    for s in prs.slides:
        words, bullet_pt, cap_pt = 0, set(), []
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if not t or t in v3.FOOT or re.fullmatch(r'\d+', t) or sh.name in NOT_COUNTED:
                continue
            words += len(t.split())
            if sh.name == 'Rectangle 4':
                for p in sh.text_frame.paragraphs[1:]:
                    for r in p.runs:
                        if r.font.size:
                            bullet_pt.add(r.font.size.pt)
            if sh.name == 'Caption':
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            cap_pt.append(r.font.size.pt)
        out.append((words, bullet_pt, min(cap_pt) if cap_pt else None))
    return out


def report_checks(C, out_json=None, label=''):
    total, failed = len(C.items), C.failed
    print(f'{label} checks: {total} · failed {len(failed)}')
    for nme, d in failed[:400]:
        print(f'  ✗ {nme}  {d}')
    if out_json:
        Path(out_json).write_text(json.dumps(dict(label=label, total=total, failed=len(failed),
                                                  failures=[dict(check=nme, detail=d) for nme, d in failed]),
                                             ensure_ascii=False, indent=1), encoding='utf-8')
    return len(failed)


# ───────────────────────────── 글 전문 (무리 · 식 포함) ─────────────────────────────

def dump_text_v32(deck, out):
    prs = Presentation(str(deck))
    lines = ['# 연구세미나 덱 v3.2 (18 장 · 1저자 10-14 결정) — 장별 글 전문 (생성기 출력에서 추출 · 표 · 출처 줄 · 무리 안 라벨 · 식 이름 · 다리 줄 · notes 포함)', '']

    def walk(shapes, prefix=''):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(sh.shapes, prefix + sh.name + '/')
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                for p in sh.text_frame.paragraphs:
                    if p.text.strip():
                        lines.append(f'- [{prefix}{sh.name}] {p.text}')
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows:
                    lines.append(f'- [{prefix}{sh.name}] ' + ' | '.join(c.text for c in r.cells))
            if sh.shape_type == 13:
                lines.append(f'- [picture] {prefix}{sh.name}')

    for i, s in enumerate(prs.slides, 1):
        lines.append(f'## p{i:02d}')
        walk(s.shapes)
        for ac in s._element.iter('{%s}AlternateContent' % NS['mc']):
            nm = ac.find('.//{%s}cNvPr' % P_).get('name')
            lines.append(f'- [equation] {nm} (OMML copied from the 10-09 slide · fallback image kept)')
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            lines.append('- notes: ' + s.notes_slide.notes_text_frame.text)
        lines.append('')
    Path(out).write_text('\n'.join(lines), encoding='utf-8')


# ───────────────────────────── 생성 (v3.1 덱 위 차분) ─────────────────────────────

def _xfrm_of(el):
    """도형 요소의 a:xfrm 목록 (mc:AlternateContent = Choice · Fallback 둘)."""
    tag = etree.QName(el).localname
    if tag == 'AlternateContent':
        return [x for sp in el.iter('{%s}sp' % P_) for x in sp.findall('{%s}spPr/{%s}xfrm' % (P_, A))]
    if tag == 'grpSp':
        return el.findall('{%s}grpSpPr/{%s}xfrm' % (P_, A))
    return el.findall('{%s}spPr/{%s}xfrm' % (P_, A))


def _set_box(el, x=None, y=None, w=None, h=None):
    for xf in _xfrm_of(el):
        off, ext = xf.find('{%s}off' % A), xf.find('{%s}ext' % A)
        if x is not None:
            off.set('x', str(int(round(x * EMU))))
        if y is not None:
            off.set('y', str(int(round(y * EMU))))
        if w is not None:
            ext.set('cx', str(int(round(w * EMU))))
        if h is not None:
            ext.set('cy', str(int(round(h * EMU))))


def _get_box(el):
    xf = _xfrm_of(el)[0]
    off, ext = xf.find('{%s}off' % A), xf.find('{%s}ext' % A)
    return (int(off.get('x')) / EMU, int(off.get('y')) / EMU, int(ext.get('cx')) / EMU, int(ext.get('cy')) / EMU)


def _find_top(root, cid):
    """cNvPr id → 그 도형의 맨 바깥 요소 (mc:AlternateContent 이면 그것)."""
    for c in root.iter('{%s}cNvPr' % P_):
        if c.get('id') != str(cid):
            continue
        shp = c.getparent().getparent()          # cNvPr → nv*Pr → 도형 (sp · pic · grpSp · cxnSp · graphicFrame)
        par = shp.getparent()
        if par is not None and etree.QName(par).localname in ('Choice', 'Fallback'):
            return par.getparent()               # mc:AlternateContent
        return shp
    return None


def _translate(el):
    for t in el.iter('{%s}t' % A):
        if t.text in TRANS:
            t.text = TRANS[t.text]
            rpr = t.getparent().find('{%s}rPr' % A)
            if rpr is not None:
                rpr.set('lang', 'en-US')
                rpr.attrib.pop('err', None)


def _rebuild_caption(el, lines):
    """힘 그림 캡션 (한글 런 여럿 + 줄바꿈) → 영어 두 줄 (같은 rPr · 같은 pPr)."""
    p = el.find('.//{%s}p' % A)
    rpr = copy.deepcopy(p.find('{%s}r/{%s}rPr' % (A, A)))
    rpr.set('lang', 'en-US')
    rpr.attrib.pop('err', None)
    for ch in list(p):
        if etree.QName(ch).localname in ('r', 'br', 'endParaRPr'):
            p.remove(ch)
    for k, txt in enumerate(lines):
        if k:
            br = etree.SubElement(p, '{%s}br' % A)
            br.append(copy.deepcopy(rpr))
        r = etree.SubElement(p, '{%s}r' % A)
        r.append(copy.deepcopy(rpr))
        etree.SubElement(r, '{%s}t' % A).text = txt


class EqCopier:
    """원 슬라이드 (10-09) → 새 장: 요소를 깊은 복사하고 그림 관계 (r:embed) 를 새 장의 그림 부분으로 다시 잇는다 · cNvPr id 를 새로 매긴다."""

    def __init__(self, eq_deck, dst_slide, id_base=3000):
        self.src = Presentation(str(eq_deck))
        self.ss = self.src.slides[0]
        self.root = self.ss._element
        self.dst = dst_slide
        self.blobs = {rid: rel.target_part.blob for rid, rel in self.ss.part.rels.items() if rel.reltype.endswith('/image')}
        self.rid = {}
        self.ids = {}
        self.next_id = id_base

    def _new_rid(self, old):
        if old not in self.rid:
            _, rid = self.dst.part.get_or_add_image_part(io.BytesIO(self.blobs[old]))
            self.rid[old] = rid
        return self.rid[old]

    def take(self, cid):
        el = _find_top(self.root, cid)
        assert el is not None, cid
        el = copy.deepcopy(el)
        for e in el.iter():
            for k in list(e.attrib):
                if k in ('{%s}embed' % R_, '{%s}link' % R_) and e.attrib[k] in self.blobs:
                    e.attrib[k] = self._new_rid(e.attrib[k])
        for c in el.iter('{%s}cNvPr' % P_):
            old = c.get('id')
            if int(old) in EQ_RENAME:
                c.set('name', EQ_RENAME[int(old)])
            if old not in self.ids:
                self.ids[old] = str(self.next_id)
                self.next_id += 1
            c.set('id', self.ids[old])
        _translate(el)
        return el


def _make_group(name, cid, children, box):
    x, y, w, h = (int(round(v * EMU)) for v in box)
    g = etree.Element('{%s}grpSp' % P_, nsmap={'p': P_, 'a': A, 'r': R_})
    nv = etree.SubElement(g, '{%s}nvGrpSpPr' % P_)
    c = etree.SubElement(nv, '{%s}cNvPr' % P_)
    c.set('id', str(cid))
    c.set('name', name)
    etree.SubElement(nv, '{%s}cNvGrpSpPr' % P_)
    etree.SubElement(nv, '{%s}nvPr' % P_)
    gp = etree.SubElement(g, '{%s}grpSpPr' % P_)
    xf = etree.SubElement(gp, '{%s}xfrm' % A)
    etree.SubElement(xf, '{%s}off' % A, x=str(x), y=str(y))
    etree.SubElement(xf, '{%s}ext' % A, cx=str(w), cy=str(h))
    etree.SubElement(xf, '{%s}chOff' % A, x=str(x), y=str(y))
    etree.SubElement(xf, '{%s}chExt' % A, cx=str(w), cy=str(h))
    for ch in children:
        g.append(ch)
    return g


_DJV = {}


def _djv_lines(text, pt, width, bold=False):
    """줄 수 어림 — DejaVu Sans (LibreOffice 가 라벨 글꼴 맑은 고딕 대신 쓰는 글꼴 · fc-match) · 없으면 v3 추정."""
    from PIL import ImageFont
    path = '/usr/share/fonts/truetype/dejavu/DejaVuSans%s.ttf' % ('-Bold' if bold else '')
    if not Path(path).exists():
        return v3.est_lines(text, pt, width)
    key = (pt, bold)
    if key not in _DJV:
        _DJV[key] = ImageFont.truetype(path, int(round(pt * 4)))
    f, maxw = _DJV[key], width * 72 * 4
    lines, cur = 1, ''
    for w in text.split(' '):
        trial = (cur + ' ' + w).strip()
        if f.getlength(trial) <= maxw:
            cur = trial
        else:
            lines += 1
            cur = w
    return lines


def _label_height(el, width):
    """라벨 글 상자 높이 어림 (문단마다 줄 수 × 글씨 × 1.2 × 105 % 줄 간격)."""
    h = 0.0
    for p in el.iter('{%s}p' % A):
        runs = p.findall('{%s}r' % A)
        txt = ''.join((r.findtext('{%s}t' % A) or '') for r in runs)
        if not txt.strip():
            continue
        rps = [r.find('{%s}rPr' % A) for r in runs if r.find('{%s}rPr' % A) is not None]
        sz = max((int(rp.get('sz', '1000')) for rp in rps), default=1000) / 100
        bold = any(rp.get('b') == '1' for rp in rps)
        h += _djv_lines(txt, sz, width, bold) * sz * 1.05 * 1.2 / 72
    return h


def _fallback_fit(ac, blobs, sx):
    """식 하나 (mc:AlternateContent) — 원 슬라이드에서 PowerPoint 가 만든 대체 그림은 **가로로 sx 배 눌린 무리** 안 크기로 그려졌다.
    눌림을 푼 자리 (상자 폭 그대로) 에서 그림이 제 크기 · 비율로 보이도록 fillRect 를 다시 셈하고, 글자 높이가 상자보다 크면 상자 높이를 늘린다
    (가운데 유지).  OMML · 그림 바이트는 그대로 · Choice 와 Fallback 의 xfrm 은 같게.  → (새 높이, 위로 옮길 양 dy (in))."""
    fb = ac.find('{%s}Fallback' % NS['mc'])
    xf = fb.find('.//{%s}xfrm' % A)
    w = int(xf.find('{%s}ext' % A).get('cx')) / EMU
    h = int(xf.find('{%s}ext' % A).get('cy')) / EMU
    blip = fb.find('.//{%s}blip' % A)
    st = blip.getparent().find('{%s}stretch' % A)
    fr = st.find('{%s}fillRect' % A)
    if fr is None:
        fr = etree.SubElement(st, '{%s}fillRect' % A)
    l, t, r, b = (int(fr.get(k, '0')) / 100000 for k in ('l', 't', 'r', 'b'))
    from PIL import Image
    img = Image.open(io.BytesIO(blobs[blip.get('{%s}embed' % R_)])).convert('RGBA')
    W, H = img.size
    bx0, by0, bx1, by1 = img.getchannel('A').point(lambda v: 255 if v > 24 else 0).getbbox()
    img_l, img_r = l * w * sx, w * sx * (1 - r)                 # 원 슬라이드 위 그림 가장자리 (옛 상자 왼쪽 · 위 기준 in)
    img_t, img_b = t * h, h * (1 - b)
    iw, ih = img_r - img_l, img_b - img_t
    ink_x1 = img_l + bx1 / W * iw
    ink_y0, ink_y1 = img_t + by0 / H * ih, img_t + by1 / H * ih
    ink_w = (bx1 - bx0) / W * iw
    # 상자 폭 = 그림 오른쪽 가장자리 (fillRect 오른쪽 0 · LibreOffice 는 양수 inset 이 섞이면 fillRect 를 무시한다) · 높이 = 글자가 넘칠 때만 늘림 (가운데)
    w2 = img_r                                                   # 글자는 그림 안에 있다 (ink_x1 ≤ img_r)
    assert ink_x1 <= img_r + 1e-9
    if ink_y0 >= -0.005 and ink_y1 <= h + 0.005:
        h2, dy = h, 0.0
    else:                                                        # 옛 상자 ∪ 글자 (+0.01) — 그림 가장자리 안에서만 (fillRect ≤ 0 유지)
        top = min(0.0, max(img_t, ink_y0 - 0.01))
        bottom = max(h, min(img_b, ink_y1 + 0.01))
        h2, dy = bottom - top, top                               # dy = 새 상자 윗변 − 옛 상자 윗변
    l2, r2 = img_l / w2, 1 - img_r / w2
    t2, b2 = (img_t - dy) / h2, 1 - (img_b - dy) / h2
    assert max(l2, r2, t2, b2) <= 1e-9, ('fillRect must stay non-positive', l2, r2, t2, b2)
    for k, v in (('l', l2), ('t', t2), ('r', r2), ('b', b2)):
        fr.set(k, str(int(round(v * 100000))))
    for x in _xfrm_of(ac):
        x.find('{%s}ext' % A).set('cx', str(int(round(w2 * EMU))))
        x.find('{%s}ext' % A).set('cy', str(int(round(h2 * EMU))))
    return h2, dy, ink_w


def build_eq_slide(s, eq_deck, ft):
    """3 장: 오른쪽 = 식 판 (틀 그림 · 탭 · 판 제목 · 계산 순서 띠 · 라벨 ①–⑤ · 식 OMML 여덟 · 구분선) · 왼쪽 = 힘 그림 무리 (라벨 영어).
    식 상자 폭 = 원 3.461 in (PowerPoint 실제 식 폭 ≤ 3.34 in — 원 슬라이드는 무리를 가로 0.717 배로 눌러 넘쳤다) · 대체 그림 = _fallback_fit."""
    cp = EqCopier(eq_deck, s)
    g3 = _find_top(cp.root, 3)                                   # 원 슬라이드의 바깥 무리 (식 판) — 가로 눌림 배율
    gx = g3.find('{%s}grpSpPr/{%s}xfrm' % (P_, A))
    sx = int(gx.find('{%s}ext' % A).get('cx')) / int(gx.find('{%s}chExt' % A).get('cx'))
    PX0, PW = 3.00, 6.65
    PY0 = ft
    PH = 6.95 - PY0
    IN_T, IN_H, IN_L, IN_W = 0.0738, 0.9101, 0.0290, 0.9502     # 틀 무리 안 흰 판 (원 child 좌표의 분율)
    CX0, CX1 = PX0 + IN_L * PW + 0.08, PX0 + (IN_L + IN_W) * PW - 0.08
    EQW = 3.461
    EX0 = CX1 - EQW
    LW = EX0 - 0.06 - CX0
    kids = []
    frame = cp.take(EQ_IDS['frame'])
    _set_box(frame, PX0, PY0, PW, PH)                           # 무리 (그림 + 흰 판) — chOff/chExt 그대로 = 늘이기
    kids.append(frame)
    tab = cp.take(EQ_IDS['tab'])
    _set_box(tab, PX0 + 0.7873 * PW, PY0 - 0.003 * PH, 0.2056 * PW, 0.055 * PH)
    kids.append(tab)
    title = cp.take(EQ_IDS['title'])
    _set_box(title, PX0 + 0.20, PY0 + 0.04, 4.30, 0.30)
    kids.append(title)
    # 계산 순서 띠 (every Δt · 상자 넷 · 삼각형 셋) — 폭 = 판 안쪽에 맞춤 · 높이 0.30
    ys, hs = PY0 + IN_T * PH + 0.06, 0.30
    ev = cp.take(EQ_IDS['every'])
    _set_box(ev, CX0, ys, 0.50, hs)
    kids.append(ev)
    x = CX0 + 0.54
    scale = (CX1 - x - 3 * 0.076 - 6 * 0.03) / 5.104
    for cid in EQ_IDS['strip']:
        el = cp.take(cid)
        _, _, w0, h0 = _get_box(el)
        if w0 < 0.2:                                             # 삼각형
            _set_box(el, x, ys + hs / 2 - h0 / 2)
            x += w0 + 0.03
        else:
            w = w0 * scale
            _set_box(el, x, ys, w, hs)
            x += w + 0.03
        kids.append(el)
    # 줄 ①–⑤ — 식 상자 높이 · 대체 그림 자리 (눌림 풀기)
    lab = {cid: cp.take(cid) for cid in EQ_IDS['labels']}
    eqs = {cid: cp.take(cid) for cid in EQ_IDS['eqs']}
    seps = [cp.take(cid) for cid in EQ_IDS['seps']]
    blobs = {rid: rel.target_part.blob for rid, rel in s.part.rels.items() if rel.reltype.endswith('/image')}
    fit = {cid: _fallback_fit(eqs[cid], blobs, sx) for cid in EQ_IDS['eqs']}
    y = ys + hs + 0.08
    gap = 0.035

    def place_eqs(cids, y0, x0=EX0):
        yy = y0
        for cid in cids:
            h2, dy, _ = fit[cid]
            _set_box(eqs[cid], x0, yy)
            yy += h2
        return yy

    def row(label_cid, eq_cids):
        nonlocal y
        lh = _label_height(lab[label_cid], LW)
        _set_box(lab[label_cid], CX0, y, LW, lh)
        y = max(y + lh, place_eqs(eq_cids, y))

    def sep(k):
        nonlocal y
        _set_box(seps[k], CX0, y + gap, CX1 - CX0, 0)
        y += 2 * gap

    row(413, [414]); sep(0)
    row(415, [416]); sep(1)
    row(417, [418, 419]); sep(2)
    # ④ = 머리 한 줄 (전폭) + 힌트 둘 (왼쪽) · 식 둘 (오른쪽)
    lh = _label_height(lab[420], CX1 - CX0)
    _set_box(lab[420], CX0, y, CX1 - CX0, lh)
    y += lh + 0.01
    for hint, eq in ((422, 421), (424, 423)):
        h2 = fit[eq][0]
        hh = _label_height(lab[hint], LW)
        _set_box(lab[hint], CX0, y + max(0.0, (h2 - hh) / 2), LW, hh)
        place_eqs([eq], y)
        y += max(h2, hh)
    sep(3)
    row(425, [426, 427])
    inner_bottom = PY0 + (IN_T + IN_H) * PH
    assert y <= inner_bottom - 0.02, ('equation panel overflows the inner panel', round(y, 3), round(inner_bottom, 3))
    kids += [lab[c] for c in (413, 415, 417, 420, 422, 424, 425)] + [eqs[c] for c in EQ_IDS['eqs']] + seps
    panel = _make_group('Equation panel', cp.next_id, kids, (PX0, PY0 - 0.02, PW, PH + 0.02))
    cp.next_id += 1
    s.shapes._spTree.append(panel)
    # 힘 그림 무리 — 고르게 줄임 (chOff/chExt 그대로 · off/ext 만)
    force = cp.take(EQ_IDS['force'])
    for sp in force.iter('{%s}sp' % P_):
        nm = sp.find('.//{%s}cNvPr' % P_).get('name')
        if nm == 'Force caption':
            _rebuild_caption(sp, FORCE_CAPTION)
    _, _, w0, h0 = _get_box(force)
    fw = PX0 - 0.15 - 0.35
    fs = fw / w0
    fy = PY0 + 0.25
    _set_box(force, 0.35, fy, fw, h0 * fs)
    s.shapes._spTree.append(force)
    return dict(force_bottom=fy + h0 * fs, left_w=fw, panel_bottom=PY0 + PH, rows_end=y, inner_bottom=inner_bottom, sx=sx,
                label_w=LW, eq_x=EX0, fits={EQ_NAMES[k]: tuple(round(v, 3) for v in fit[c]) for k, c in enumerate(EQ_IDS['eqs'])})


def _prune_image_rels(slide):
    used = {v for el in slide._element.iter() for k, v in el.attrib.items() if k.startswith('{%s}' % R_)}
    for rid, rel in list(slide.part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in used:
            slide.part.drop_rel(rid)


def _renumber_citations(slide):
    for t in slide._element.iter('{%s}t' % A):
        if t.text and '[' in t.text:
            t.text = renum(t.text)
    if slide.has_notes_slide:
        for t in slide.notes_slide._element.iter('{%s}t' % A):
            if t.text and '[' in t.text:
                t.text = renum(t.text)


def build_v32(template, repo, figs, lit, eq_deck, out, v31_out=None):
    repo, figs, out = Path(repo), Path(figs), Path(out)
    lit = Path(lit) if lit else out.resolve().parent / 'lit'
    v31 = Path(v31_out) if v31_out else out.with_name(out.stem + '_base_v31.pptx')
    v3.build(template, repo, figs, lit, str(v31))                 # v3.1 = Codex 종결본 (같은 코드 경로)
    D31 = json.loads(Path(str(v31) + '.sources.json').read_text(encoding='utf-8'))
    prs = Presentation(str(v31))
    old = {i: s for i, s in enumerate(prs.slides, 1)}
    PROTO = copy.deepcopy(v3.shape(old[3], 'Caption')._element)   # 글 상자 원형 (v3.1 해석 줄과 같은 꼴)
    R = RED

    def box(slide, b, paras, name, size=v3.CAP_PT, color=GRAY, line=1.1, align=None, italic=None):
        return v3.add_box(slide, PROTO, b, [[(t, color)] if isinstance(t, str) else t for t in paras], name,
                          size=size, line=line, align=align, italic=italic)

    def caption(slide, b, paras, align=None):
        return box(slide, b, paras, 'Caption', align=align)

    def lit_tag(slide, b, text):
        return v3.add_box(slide, PROTO, b, [[('Literature: ', DARK, True), (text, DARK)]], 'Lit tag', size=12.5, line=1.1)

    def source(slide, text):
        assert len(text) <= 110, (len(text), text)
        sp = v3.shape(slide, 'TextBox 69')
        v3.set_text(sp, [(text, GRAY)], size=10, italic=True)
        v3.place(sp, (2.45, 0.93, 7.2, 0.27))

    def new_slide(section, head, bullets, src, ft_min=v3.FT_MIN):
        s = v3.dup_slide(prs, old[3])
        for sh in list(s.shapes):
            if sh.name not in ('TextBox 20', 'Rectangle 4', 'TextBox 69'):
                sh._element.getparent().remove(sh._element)
        _prune_image_rels(s)
        v3.set_runs_text(v3.shape(s, 'TextBox 20').text_frame, section)
        bx = v3.shape(s, 'Rectangle 4')
        v3.set_body(bx, head, bullets)
        ft = max(ft_min, round(v3.est_body_bottom(head, [v3.seg_text(b) for b in bullets]) + v3.FT_MARGIN, 2))
        v3.place(bx, (0.61, v3.BODY_TOP, 8.84, ft - v3.BODY_TOP - 0.05))
        source(s, src)
        return s, ft

    D = {}
    # ── 0 v3.1 장의 인용 번호 바꿈 (새 글을 넣기 전에 · 참고문헌 장은 다시 쓴다) ─────────
    for i, s in old.items():
        if i != 16:
            _renumber_citations(s)

    # ── 3 (새) DEM 지배방정식 ───────────────────────────────────────────
    s, ft = new_slide('Model', 'DEM governing equations: particle motion and contact forces', [],
                      'This work — implemented equations of LIGGGHTS-PUBLIC 3.8.0 [2]; linear hysteretic law [1]', ft_min=0)
    ft = round(ft + 0.02, 2)
    geo = build_eq_slide(s, eq_deck, ft)
    # 원 슬라이드 판 안의 한 줄 (한글) → 영어 · 판에 자리가 없어 왼쪽 열 (힘 그림 아래) 로
    cy1 = round(geo['force_bottom'] + 0.55, 2)
    box(s, (0.35, cy1, geo['left_w'], 1.00),
        [[('Per particle: ', DARK), ('translational and rotational equations of motion', NAVY, True), ('; per contact pair: contact force and '
          'rolling-resistance torque.', DARK)]], 'Callout', size=12, line=1.1, align=PP_ALIGN.LEFT)
    cy2 = round(cy1 + 1.08, 2)
    box(s, (0.35, cy2, geo['left_w'], 1.05),
        [[('Global damping −γv_i in ①: ', NAVY, True), ('γ = 0.5 during stabilization and compression, 1e-5 during settling and relaxation '
          '(input-deck units) — see Problem 4.', DARK)]], 'Callout damping', size=12, line=1.1, align=PP_ALIGN.LEFT)
    assert cy2 + 1.05 <= CONTENT_BOTTOM, ('left column too long', cy2)
    print('equation slide geometry:', json.dumps(geo))
    eq_slide = s
    D['EQ'] = ('docs/dem_governing_equations_report_20261009.md §1 · §3; docs/dem_governing_equations_background_20261011.md §2.3 · §3.1 · §5.2 · §5.3 · '
               '§6 · §7 · §8.2 · §9 · §11; equations ①–⑤, force diagram and process strip copied from the author 10-09 slide (OMML objects and pictures '
               'unchanged · Korean labels translated · project header removed); DEMP-01 (global damping term, open)')
    v3.notes(s, 'Sources: docs/dem_governing_equations_report_20261009.md §1 (literature form ↔ implemented form ↔ source lines of LIGGGHTS-PUBLIC 3.8.0, '
                'commit 3d5c00f) and §3 (deck values); docs/dem_governing_equations_background_20261011.md §2.3 (order of one time step: contact '
                'detection → contact forces → integration → update), §3.1 (translational and rotational equations, eqs 1–2), §5.2 (k1 from E*, R*, m* '
                'and V), §5.3 (hooke/hysteresis), §6 (tangential spring with Coulomb limit), §7 (rolling resistance, CDT), §8.2 (global viscous damping), '
                '§9 (compression procedure), §11 (limits). Equations ①–⑤, the force diagram and the process strip are copied from the author\'s 10-09 '
                'slide: equation objects and pictures unchanged, Korean labels translated, project header removed. Qualifiers: these are the implemented '
                'equations of the LIGGGHTS-PUBLIC 3.8.0 hooke/hysteresis model [2], read from the source code, not re-derived by a run; the linear '
                'hysteretic law (loading stiffness k1, unloading/reloading stiffness k2, adhesion branch −k_c δ) follows Luding 2008 [1]; ③ shows the '
                'non-saturated branch only (the saturated branch and the k2(δmax) interpolation are in the report); ④ shows the spring part (tangential '
                'damping in non-slip steps is omitted); k_t = k_n in this model; the rolling torque uses k_n δ, not the hysteretic normal force; '
                'k1 depends on mass and the characteristic velocity V (k1 ∝ m*^(1/5) V^(2/5)), so V acts as a stiffness input in quasi-static compaction; '
                'the SE modulus 1.35 GPa entering E* is a calibrated input. The global damping term −γv_i in ① takes a large share of the plate reaction '
                'during compression (ledger DEMP-01, open) — see the Problem 4 slide. Particle–particle and particle–wall contacts use the same model.')

    # ── 11 (새) Problem 4 — 판 반력 "300 MPa" 의 속 ─────────────────────
    s, ft = new_slide('Limits', 'Problem 4: what the 300 MPa plate reaction contains',
                      [[('The 300 MPa stop criterion reads the plate reaction during compression, which includes ', None),
                        ('global viscous damping', R), (': in one 7:3 bed the drag is ', None),
                        ('57.9 % of the target force', R), (' (750 N, deck units).', None)],
                       [('At fixed platen the same bed relaxes to 165 MPa; a ', None), ('servo branch holding the target force', R),
                        (' (damping 1e-5) instead compacts it further (111.05 → 105.84 µm) — ', None), ('model-internal', R),
                        (' until the open checks pass.', None)]],
                      'This work — DEM, one 7:3 bed: stop, relaxation and servo-hold branch')
    fh = v3.FZ_BOTTOM - ft - 0.50
    v3.place_pic(s, figs / f'{P4_FIG}.png', (0.45, ft, 5.45, fh), name='Fig plate reaction')
    rows = [('Before citing the hold state', 'Status'),
            ('Convergence flag vs. static equilibrium', 'open'),
            ('Plate re-created at servo start: contact history reset', 'open'),
            ('Atom-count guard vs. conservation (13 SE particles lost)', 'open')]
    v3.add_table(s, (6.10, ft + 0.15, 3.35, 2.10), rows, (2.65, 0.70), size=11, name='Table checks', row_h=(0.34, 0.55, 0.62, 0.62))
    caption(s, (0.45, ft + fh + 0.04, 9.0, 0.45),
            ['(a) one 7:3 bed: stop step (damping 0.5), relaxed at fixed platen (original run), servo hold at the target force (damping 1e-5)',
             '(b) ε_sphere = sphere-volume porosity, not the exact union; shaded = band registered before the run; hatched = checks still open'])
    p4_slide = s
    D['P4'] = ('findings.json DEMP-01 (open) · MV-01 · MV-02 · MV-08 (open); docs/data/ps73_damping_branch_20261007/README.md §1 (165.34 MPa), §4 (B1 · B2 '
               'registered before the run), §12-4 (stop 300.906 MPa · 111.052 µm · ε_sphere 14.953 % · drag 434.18 N = 57.89 %; servo-hold end 299.274 MPa · '
               '105.842 µm · ε_sphere 10.769 %); figure F_p4 redrawn from these recorded strings only; codex_review_model_validation_20261008 Q11 wording')
    v3.notes(s, 'Sources: docs/reviews/findings.json DEMP-01 (open: global viscous damping 0.5 during stabilization and compression, 1e-5 during '
                'relaxation; the stop criterion |F_plate|/A ≥ 0.30 deck MPa includes the damping drag; one 7:3 bed: drag −γΣv_z = 434.3 N = 57.9 % of '
                'the 750 N target (deck units: 0.30 deck MPa × plate area 0.0025 in deck units; docs/dem_governing_equations_background_20261011.md §9); '
                'after the plate stops the network pressure relaxes to 165 MPa), MV-01, MV-02 and MV-08 (open); '
                'docs/data/ps73_damping_branch_20261007/README.md §1 (original run, relaxation end at fixed platen 165.34 MPa), §4 (predictions B1 and B2 '
                'registered before the run), §12-4 (branch B on ibb, job 271368: stop step 3,125,000, 111.052 µm, ε_sphere 14.953 %, 300.906 MPa, '
                'drag 434.18 N = 57.89 %; servo hold at 750 N with damping 1e-5 to step 3,735,000, 122 chunks, CONVERGED flag, 105.842 µm, ε_sphere '
                '10.769 %, 299.274 MPa; atoms 160,420 → 160,407, 13 SE lost, AM unchanged; plane-load NCM–NCM share 52.1 → 58.1 %, judgement on hold). '
                'B1 equals the registered stop state; B2 (104.9 µm, band 103–107 µm; ε_sphere 10.0 %, band 8.3–11.7 %) contains the result. '
                'Figure F_p4 is redrawn from these recorded values only (docs/seminar_20261013/figs_v3/make_seminar_figs_v32.py). '
                'Qualifiers: one bed, one composition, one checkpoint; model-internal diagnostic values. (i) The CONVERGED flag does not certify static '
                'equilibrium (MV-01). (ii) The servo deck re-creates the plate and wall fixes, which also resets the normal plastic contact history, so B '
                'is a combined protocol (history reset + damping drop + force hold); the registered follow-up protocol includes a reset of the plate '
                'contact history (MV-02). (iii) The atom-count guard does not certify conservation; the 13 lost particles are SE (MV-08). '
                'Wording (docs/reviews/codex_review_model_validation_20261008.md Q11): relative to the existing dynamic stop state, the registered '
                'follow-up protocol changed thickness and sphere-volume porosity by −5.210 µm and −4.184 %p; only after separate equilibrium and '
                'conservation checks pass may the end point be read as the nominal static 300 MPa state; the comparison is limited to this bed and '
                'loading history. ε_sphere is the sphere-volume porosity (overlap counted twice), a different convention from the exact-union porosity '
                'shown earlier. The porosity and load-sharing slides use the original protocol (generation 1).')

    # ── 5 (p06) — 실험 앵커 행 (계획) ──────────────────────────────────
    s = old[6]
    tbl = v3.shape(s, 'Table inputs').table
    for j, txt in enumerate(('Lab anchor (planned)', 'EIS + DC polarization, no results yet: NCM pellets 0.1 / 0.2 / 0.3 g, polycrystalline and '
                                                     'single-crystal; LPSCl pellets, small and large particles')):
        cell = tbl.cell(5, j)
        cell.text_frame.paragraphs[0].runs[0].text = txt
    for k, rh in enumerate((0.34, 0.46, 0.44, 0.46, 0.46, 0.78)):
        tbl.rows[k].height = Inches(rh)
    gf = v3.shape(s, 'Table inputs')
    tb_bottom = (gf.top + sum(r.height for r in tbl.rows)) / EMU
    for sh in v3.shapes_named(s, 'Caption'):
        v3.place(sh, (0.45, round(tb_bottom + 0.06, 2), 4.9, 0.40))
    v3.notes(s, s.notes_slide.notes_text_frame.text + ' Planned lab anchor (replaces the earlier "validation pending" row; no results yet): '
                'docs/session_20260923_progress.md 10-14 entries — EIS plan: LPSCl pellet 0.1 g (10 mm diameter, 300 MPa) as the σ₀ reference; '
                'NCM polycrystalline and single-crystal pellets of 0.1 / 0.2 / 0.3 g (six pellets, EIS + DC polarization); composite and Li-In '
                'electron-blocking cells after the NCM pellets; the small-particle SE pellet also gives a geometric porosity (true density unverified) '
                'as the pure-SE compaction check; a large-particle LPSCl pellet tests the particle-size factor (listed by the author on 10-14 for this '
                'deck; the progress log of 10-14 evening still records it as a recommendation awaiting an answer).')

    # ── 7 (p03) — 문헌 꼬리표 (순수 SE · 다이 안) ────────────────────────
    s = old[3]
    caps = v3.shapes_named(s, 'Caption')
    cb = caps[0].top / EMU + caps[0].height / EMU
    lit_tag(s, (0.45, round(cb + 0.02, 2), 9.0, 0.30),
            'pure LPSCl in the die at 0.3 GPa, porosity ≈ 13 % (D50 40 µm) to ≈ 26 % (4 µm) [13] — not composite beds')
    v3.notes(s, s.notes_slide.notes_text_frame.text + ' Literature context (pure SE, not composite beds; different convention and state): '
                'Faka 2026 [13] — in-die (at pressure) relative density of pure Li6PS5Cl at 0.3 GPa read from their Fig. S2: porosity ≈ 12.9 % (D50 40 µm), '
                '16.0 % (20 µm), 22.5 % (11 µm), 26.3 % (4 µm); relative density uses the Rietveld crystal density; no compliance or elastic-compression '
                'correction is described. Thompson 2026 [12] — in situ CT (1.277 µm pixel) of pure LPSCl under quasi-isostatic pressure: intergranular '
                'pores mostly closed at 479 MPa; no pores resolved at 1155 MPa. Neither is a validation of the five DEM beds. '
                'Cards: litdb faka2026_particle_size_consolidation_strain_lpscl · thompson2026_high_pressure_densification_lps_lpscl_xrd_tomography '
                '(docs/seminar_20261013/slide_sources_v32.json).')

    # ── 9 (p08) — Lee R / Q 옮김 · σ_VGCF* 3.41 ─────────────────────────
    s = old[8]
    bx = v3.shape(s, 'Rectangle 4')
    head = 'Problems 1 and 2: grid dependence and absolute σ(e⁻)'
    bullets = [[('The electronic ratio decreases across the three tested voxel spacings (0.4, 0.3, 0.25 µm); the ionic ratio crosses 1. ', None),
                ('Grid convergence has not been established', R), ('.', None)],
               [('Voxel predictions ', None), ('exceed', R), (' the available DC-polarization estimates; measurement conditions remain ', None),
                ('unverified', R), ('. A registered re-run calibrates carbon on SBE alone (', None), ('σ_VGCF* = 3.41 S cm⁻¹', R),
                (', an effective value); the DBE prediction (order only) is pending.', None)]]
    v3.set_body(bx, head, bullets)
    ft9 = max(v3.FT_MIN, round(v3.est_body_bottom(head, [v3.seg_text(b) for b in bullets]) + v3.FT_MARGIN, 2))
    v3.place(bx, (0.61, v3.BODY_TOP, 8.84, ft9 - v3.BODY_TOP - 0.05))
    pic = v3.shape(s, 'Fig limits')
    assert abs(pic.top / EMU - ft9) < 0.02, ('figure zone moved', pic.top / EMU, ft9)
    cap = v3.shapes_named(s, 'Caption')[0]                          # v3.1 해석 줄 셋 그대로 ([16] → [11] 만) · 자리 · 줄 간격만 (그림 바로 아래 · 1.0)
    fig_bottom = (pic.top + pic.height) / EMU
    v3.place(cap, (0.45, round(fig_bottom + 0.01, 2), 9.0, 0.56))
    for p in cap.text_frame.paragraphs:
        p.line_spacing = 1.0
    # Lee 2025 비교 (v3.1 p10 상태표의 그 행 · R / Q) — 머리 없는 두 줄 표 (표 = 단어 수 제외 · v3.1 과 같은 규칙)
    gf = v3.add_table(s, (0.45, round(fig_bottom + 0.65, 2), 9.0, 0.40),
                      [('Lee 2025 [10]', 'R = 2.07, model / literature: order-of-magnitude agreement, not within the registered band'),
                       ('PTFE 1 vs 0.5 wt%', 'Q = 0.84, retained after PTFE: registered model-sensitivity prediction met, not an experimental hit')],
                      (1.45, 7.55), size=10.5, name='Table absolute', row_h=(0.20, 0.20))
    for j in range(2):                                              # 첫 줄도 본문 꼴 (add_table 의 머리 행 색을 되돌린다)
        c = gf.table.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for r in c.text_frame.paragraphs[0].runs:
            r.font.bold = (j == 0)
            r.font.color.rgb = RGBColor.from_string(NAVY if j == 0 else DARK)
    nt = s.notes_slide.notes_text_frame.text
    old_b = '(DC-polarization values 4.47 and 12.4 mS/cm, measurement conditions unverified; model 54.0 and 70.6; re-run registered, no results yet)'
    assert old_b in nt
    nt = nt.replace(old_b, '(DC-polarization values 4.47 and 12.4 mS/cm, measurement conditions unverified; model 54.0 and 70.6; re-run registered: '
                           'S1–S2 done on 10-14, σ_VGCF* = 3.41 S/cm interpolated in ln σ_e – ln σ_VGCF between 2.618 and 3.927 S/cm (local slope 0.545); '
                           'S3 (8 origins × SBE and DBE) running — acceptance (J1) and order (J2) verdicts pending; docs/session_20260923_progress.md 10-14)')
    v3.notes(s, nt + ' Qualifiers for σ_VGCF* (registration §4): a lumped effective value at vox 0.15 — fibre contact resistance and unresolved '
                'geometry in one carbon conductivity (grid not converged; not a fibre property); calibrated on one SBE point, DBE predicted (only the '
                'order is judged); experimental conditions unverified, so "the model reproduces the experiment" is not claimed (one calibration point '
                'plus one order check). Absolute-conductivity comparison (moved from the removed progress slide): docs/reviews/lee_abs_sigma_e_prereg_20260923.md '
                '§5 (R = 2.074 ± 0.004, Tier 1 "order-of-magnitude", not "within band"; Q = 0.840 = registered model-sensitivity prediction met, not an '
                'experimental hit; comparison with Lee 2025 [10]). Historical eight-origin study (moved from the removed progress slide): '
                'docs/reviews/sdcp_gain_prereg_v2_20260816.md — the eight-origin (voxel 0.15 µm) study: numerical ratios and the h1 inference are on hold '
                'in the claims ledger (CL-33/34/41), so no result is reported. "Origin" = voxel-grid origin offsets of the same bed, not independent '
                'replicate beds. Moved from the slide body to these notes (the (c) note line carries it): The NCM–SE comparison uses a common reference '
                'conductivity and a specified contact model.')

    # ── 10 (p09) — notes 문헌 (언로딩) ───────────────────────────────────
    s = old[9]
    v3.notes(s, s.notes_slide.notes_text_frame.text + ' Literature on unloading (pure LPSCl; context for the springback question above, not a validation): '
                'Thompson 2026 [12] — on decompression the lattice shift is reversible while peak broadening and the conductivity gain remain '
                '(quasi-isostatic cell; decompression only to ≈1.1–1.8 GPa, not to zero load; conductivity from a separate cell under similar pressure '
                'conditions). Faka 2026 [13] — distinguishes in-die (at pressure) from '
                'out-of-die Heckel analysis, the latter including elastic recovery; their own data are in-die only.')

    # ── 12–14 (p11–p13) — 믹서 문헌 ─────────────────────────────────────
    ORD = ('Mixing in which small particles adhere to carrier-particle surfaces was called ordered mixing by Hersey 1975 [4]; Egermann & Orr 1983 [5] '
           'proposed calling this type interactive mixing and reserved "ordered" for homogeneity beyond random (σ < σ_R) (wording of docs/hetero_thesis.md §2-5).')
    s = old[11]
    v3.notes(s, s.notes_slide.notes_text_frame.text + ' Literature on the index (context): Lacey\'s M is normalized to a random-mixture variance; '
                'Hersey 1975 [4] noted that homogeneity measures referenced to a random mixture are not satisfactory for ordered mixtures. ' + ORD +
                ' The M values here use the random reference.')
    s = old[12]
    v3.notes(s, s.notes_slide.notes_text_frame.text + ' Literature on mixer stress (context, different system and mixer): Kissel 2026 [6] — coarse-grained '
                'DEM of a high-intensity mixer (one particle species, no cohesion term; simulated net specific power agreed with the experiment at high '
                'rotor speed only) for a halide-SE coating on '
                'single-crystal NCM: coating coverage followed the stress intensity rather than the number of stress events. This is context for '
                'mixer-side process variables, not a comparison with the drum runs.')
    s = old[13]
    caps = v3.shapes_named(s, 'Caption')
    cb = caps[0].top / EMU + caps[0].height / EMU
    lit_tag(s, (0.45, round(cb + 0.02, 2), 9.0, 0.30),
            'fines adhering to carriers — ordered mixing [4], interactive mixture [5]; clean-slab DFT γ: LiNiO₂ < LiCoO₂ [7]')
    v3.notes(s, s.notes_slide.notes_text_frame.text + ' Literature (canonical cards): ' + ORD + ' NCM side of the surface-energy ratio: Kim 2012 [7] — '
                'clean-slab DFT (GGA+U, 0 K) surface energies of the end members LiNiO₂ and LiCoO₂ (not NCM) are lower for LiNiO₂ on all computed '
                'facets; these are bond-breaking energies of clean slabs, not the dispersion adhesion W of particle contacts, so the uncoated NCM / SE '
                'ratio stays open (docs/hetero_thesis.md §2-8).')

    # ── 15 (p14) — 계획: 194 줄 뺌 · 언로딩 · DEM 프로토콜 세대 2 ──────────
    s = old[14]
    for sh in v3.shapes_named(s, 'Conclusion'):
        sh._element.getparent().remove(sh._element)
    li, ri = v3.shape(s, 'Audit findings items'), v3.shape(s, 'Decisions recorded on Sep 12 items')
    v3.set_items(ri, ['AM load sharing: final hypothesis verdict pending',
                      'Unloading (springback vs. measured thickness): history test passed; kit re-review HOLD, fixes ratified',
                      'Mixing: an adhesion parameter for SE–SE contacts'], size=v3.BODY_PT)
    v3.place(li, (0.55, 2.66, 4.25, 2.25))
    v3.place(ri, (5.05, 2.66, 4.40, 2.25))
    gy = 5.00
    box(s, (0.55, gy, 8.90, 0.32), [[('DEM protocol generation 2', NAVY, True), (' — plan, no runs yet', GRAY)]], 'Gen2 head', size=14, line=1.0)
    steps = ['insertion', 'stabilization', 'compression', 'servo', 'relaxation', 'platen removal', 'unloading']
    bw, gap = 1.12, 0.17
    x0 = 0.55 + (8.90 - (len(steps) * bw + (len(steps) - 1) * gap)) / 2
    for k, st in enumerate(steps):
        x = x0 + k * (bw + gap)
        sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(gy + 0.40), Inches(bw), Inches(0.44))
        sp.name = f'Gen2 box {k + 1}'
        sp.adjustments[0] = 0.18
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        col = NAVY if st in ('servo', 'platen removal', 'unloading') else GRAY
        sp.line.color.rgb = RGBColor.from_string(col); sp.line.width = Pt(1.25)
        sp.shadow.inherit = False
        tf = sp.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
        r = p0.add_run(); r.text = st; r.font.size = Pt(12); r.font.name = 'Calibri'; r.font.bold = col == NAVY
        r.font.color.rgb = RGBColor.from_string(col)
        if k < len(steps) - 1:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.02), Inches(gy + 0.40 + 0.22 - 0.07), Inches(gap - 0.04), Inches(0.14))
            ar.name = f'Gen2 arrow {k + 1}'
            ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor.from_string(GRAY); ar.line.fill.background(); ar.shadow.inherit = False
    box(s, (0.55, gy + 0.94, 8.90, 0.62),
        [[('SE true density 1.64 g cm⁻³ (lab standard, unverified) · E_SE set so that pure-SE beds reach ε_union = 10–13 % after '
           '300 MPa and relaxation (ε_sphere kept internally) · Heckel analysis redone', DARK)]], 'Gen2 note', size=13, line=1.1)
    v3.notes(s, 'Sources: docs/reviews/contact_resistance_pipeline_draft_v3_1_20261007.md; docs/reviews/fam_platen_prereg_20260812.md §12; '
                'docs/reviews/dem_unloading_ps45_prereg_20261009.md §12–13; docs/hetero_thesis.md §4. Grid sensitivity: no new grid campaign is promised here. '
                'Unloading: the contact-history test on ibb passed (H0–H5 all PASS; docs/reviews/dem_unloading_history_test_20261014/README.md); the Codex '
                're-review of the unloading kit returned HOLD (12 findings ULKR-01 … ULKR-12; docs/reviews/codex_review_dem_unloading_kit_20261014.md); the '
                'fix plan docs/reviews/dem_unloading_kit_fix_plan_20261014.md was ratified on 10-14 (Q1–Q9 as recommended; docs/session_20260923_progress.md, 10-14 afternoon 2 entry); the main unloading run waits '
                'for the Codex re-check. DEM protocol generation 2 — plan only, no runs and no results (docs/session_20260923_progress.md, 10-14 evening '
                'entries): stages insertion → stabilization → compression → servo (holding the target force) → relaxation → platen removal → unloading, '
                'so the unloading logic is a prerequisite; SE true density 1.64 g/cm³ to match the lab value (true density unverified; canonical cards '
                'give 1.64 or 1.86–1.87); E_SE re-calibrated so that a pure-SE bed after 300 MPa and relaxation has a union porosity of 10–13 % (lab '
                'standard; the small-particle SE pellet geometric porosity is the check), with ε_sphere still recorded internally — choosing ε_union alone '
                'cannot resolve a mismatch with a mass-based porosity (docs/reviews/codex_review_model_validation_20261008.md MV-03); Heckel analysis '
                'redone, because a slope against the nominal stop pressure is an apparent coefficient of the old protocol (same review, Q2). Design points: '
                'compression damping, a static stop rule, no plate re-creation at the servo switch, conservation checks; generation-1 beds stay as '
                'labelled results. Next: plan document with canonical literature cards → author ratification → Codex.')

    # ── 16 (p15) — 요약: Problem 4 한 줄 (SUMMARY_P4_LINE) ───────────────
    s = old[15]
    sb = [[('DEM and contact networks: ', BLUE), ('composition trends from one bed per composition', None)]]
    if SUMMARY_P4_LINE:
        sb.append([('DEM loading: ', BLUE), ('the nominal 300 MPa plate reaction includes damping; a revised protocol is planned', None)])
    sb += [[('MPM: ', BLUE), ('plastic deformation, with compaction agreement dependent on initialization and loading protocol', None)],
           [('Voxel transport: ', BLUE), ('ideal-interface calculations with grid dependence and unresolved absolute conductivity', None)],
           [('Mixing: ', BLUE), ('unresolved diagnostic mixing-index contrast under an unmet overlap condition; exploratory cohesion results without a verdict', None)],
           [('Final model-validation verdicts remain pending', None)]]
    v3.set_body(v3.shape(s, 'Rectangle 4'), 'Current findings and limitations', sb, line=1.3, after=8)
    v3.place(v3.shape(s, 'Rectangle 4'), (0.61, 1.08, 8.84, 4.6))
    if SUMMARY_P4_LINE:
        v3.notes(s, s.notes_slide.notes_text_frame.text + ' The "DEM loading" line summarizes the Problem 4 slide (docs/reviews/findings.json DEMP-01, '
                    'open; docs/data/ps73_damping_branch_20261007/README.md §12-4) and the generation-2 plan (docs/session_20260923_progress.md 10-14).')

    # ── 17 (p16) — 참고문헌 (인용된 것만 · 번호 다시) ─────────────────────
    s = old[16]
    left = [('h', 'Contact mechanics and DEM')] + [f'[{k}] {t}' for k, t, _, _ in REFS[0:3]] + \
           [('h', 'Powder mixing and surface energy')] + [f'[{k}] {t}' for k, t, _, _ in REFS[3:7]]
    right = [('h', 'Solid-state battery electrodes and electrolytes')] + [f'[{k}] {t}' for k, t, _, _ in REFS[7:]]
    v3.set_items(v3.shape(s, 'Refs left'), left, size=12, line=1.1, before=3)
    v3.set_items(v3.shape(s, 'Refs right'), right, size=12, line=1.1, before=3)
    v3.place(v3.shape(s, 'Refs left'), (0.50, 1.00, 4.45, 5.85))
    v3.place(v3.shape(s, 'Refs right'), (5.05, 1.00, 4.45, 5.85))
    v3.notes(s, 'All literature entries ([1], [3]–[13]) are taken from the canonical literature cards (original PDFs, litdb canonical branch); '
                'no bibliographic field was filled from memory; [2] is the software repository at the pinned commit (docs/data/liggghts_add_pair_pin.json; '
                'version and commit date 2024-06-07 from docs/data/liggghts_3d5c00f_dem_eq_source_20261009/README.md and '
                'docs/dem_governing_equations_report_20261009.md). '
                'Card slugs: docs/seminar_20261013/slide_sources_v32.json. Entries not cited anywhere in v3.2 were removed and the list was renumbered '
                '(v3.1 → v3.2: [7] → [2], [15] → [3], [12] → [10], [16] → [11]; [1], [8], [9] unchanged). The constriction law (Holm/Maxwell) and the '
                'MLS-MPM numerical method have no canonical card yet and are cited as the models used, without a bibliographic entry. Literature crops on '
                'the DEM-limits, MPM-inputs and contact-network slides are reproduced from the papers of [1], [9] and [8] (canonical figure crops).')

    # ── v3.1 p10 (194 · 데이터셋) 삭제 · 순서 바꿈 ─────────────────────────
    new_ids = {'EQ': eq_slide, 'P4': p4_slide}
    sldIdLst = prs.slides._sldIdLst
    by_slide = {}
    for sld, sl in zip(list(sldIdLst), list(prs.slides)):
        by_slide[id(sl._element)] = sld
    order_elems = []
    for key in V32_ORDER:
        sl = new_ids[key] if isinstance(key, str) else old[key]
        order_elems.append(by_slide[id(sl._element)])
    drop = by_slide[id(old[10]._element)]
    prs.part.drop_rel(drop.get('{%s}id' % R_))
    sldIdLst.remove(drop)
    for sld in list(sldIdLst):
        sldIdLst.remove(sld)
    for sld in order_elems:
        sldIdLst.append(sld)

    # ── 다리 줄 (내용 장 2–16 · 바닥글 띠 · 같은 자리) ──────────────────────
    slides = list(prs.slides)
    for i, txt in BRIDGES.items():
        sl = slides[i - 1]
        v3.add_box(sl, PROTO, BRIDGE_BOX, [[(txt, '1F497D')]], 'Bridge', size=BRIDGE_PT, line=1.0, align=PP_ALIGN.RIGHT, italic=True)
    for sl in slides:                                              # v3.1 의 add_box 가 원형 id 를 그대로 복사해 생긴 겹침 (글 무관)
        _unique_ids(sl)

    prs.save(str(out))
    # 출처 (v3.2 장 번호)
    S = {}
    for i, key in enumerate(V32_ORDER, 1):
        S[str(i)] = D[key] if isinstance(key, str) else renum(D31[str(key)])
    S['5'] += ('; planned lab anchor row = docs/session_20260923_progress.md 10-14 (EIS + DC polarization: NCM 0.1 / 0.2 / 0.3 g PC · SC; LPSCl small and '
               'large particles)')
    S['7'] += '; literature tag = faka2026_particle_size_consolidation_strain_lpscl (Fig. S2 read values) · notes thompson2026_high_pressure_densification_lps_lpscl_xrd_tomography'
    S['9'] += ('; moved from v3.1 p10: lee_abs_sigma_e_prereg_20260923 §5 (R 2.074 · Q 0.840 · Lee 2025 [10]) · sdcp_gain_prereg_v2_20260816 (eight-origin, on hold); '
               'σ_VGCF* 3.41 S/cm = docs/session_20260923_progress.md 10-14 (S2) + sdcp_sigma_e_dcp_anchor_registration_20261012 §3–§4 qualifiers')
    S['10'] += '; notes literature = thompson2026 (decompression) · faka2026 (in-die vs out-of-die Heckel)'
    S['12'] += '; notes literature = hersey1975_ordered_mixing_concept · egermann1983_ordered_vs_interactive_mixture_nomenclature (hetero_thesis §2-5 wording)'
    S['13'] += '; notes literature = kissel2026_mechanofusion_derived_cathode_composite_microstructures_scalable'
    S['14'] += '; literature tag + notes = hersey1975 · egermann1983 · kim2012_lno_lco_surface_energy_morphology_dft'
    S['15'] = ('Rint v3.1; fam_platen_prereg §12; dem_unloading_ps45_prereg §13; hetero_thesis §4; unloading = dem_unloading_history_test_20261014 (H0–H5 PASS) · '
               'codex_review_dem_unloading_kit_20261014 (HOLD, ULKR-01..12) · dem_unloading_kit_fix_plan_20261014 (ratified); generation 2 = '
               'docs/session_20260923_progress.md 10-14 evening entries · codex_review_model_validation_20261008 MV-03 · Q2 (194-bed / 20-bed launch line removed)')
    S['16'] += '; + DEM loading line (Problem 4 · generation 2 plan)' if SUMMARY_P4_LINE else ''
    S['17'] = 'references = canonical cards only: ' + ' · '.join(f'[{k}] {slug or "software pin " + SOFTWARE_PIN[k][0]}' for k, _, slug, _ in REFS)
    Path(str(out) + '.sources.json').write_text(json.dumps(S, ensure_ascii=False, indent=1), encoding='utf-8')
    return S, v31


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--template')
    ap.add_argument('--eq-deck', help='1저자 10-09 DEM 지배방정식 슬라이드 (식 OMML · 그림 원본)')
    ap.add_argument('--repo')
    ap.add_argument('--figs')
    ap.add_argument('--lit')
    ap.add_argument('--out')
    ap.add_argument('--v31', help='v3.1 덱 (손대지 않은 장 대조용 · 생성 때는 자동)')
    ap.add_argument('--json', help='검사 결과 기록 (JSON)')
    ap.add_argument('--sources-out', help='장별 출처 JSON 사본 (리포의 slide_sources_v32.json)')
    ap.add_argument('--selftest-only', help='기존 덱 검사만')
    ap.add_argument('--dump', help='덱의 장별 글 전문을 --out (md) 로')
    a = ap.parse_args()
    if a.dump:
        dump_text_v32(a.dump, a.out)
        return 0
    if a.selftest_only:
        C = selftest_v32(a.selftest_only, a.repo, a.figs, a.v31, a.eq_deck)
        return 1 if report_checks(C, a.json, label=Path(a.selftest_only).name) else 0
    S, v31 = build_v32(a.template, a.repo, a.figs, a.lit, a.eq_deck, a.out, a.v31)
    if a.sources_out:
        Path(a.sources_out).write_text(json.dumps(S, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    C = selftest_v32(a.out, a.repo, a.figs, v31, a.eq_deck)
    nf = report_checks(C, a.json, label=Path(a.out).name)
    print('written', a.out)
    return 1 if nf else 0


if __name__ == '__main__':
    sys.exit(main())
