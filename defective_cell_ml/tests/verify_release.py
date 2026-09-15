"""Exercise the distributable package using artificial data in a new directory."""
from pathlib import Path
import argparse
import json
import subprocess
import sys
from datetime import datetime,timezone
import numpy as np
import pandas as pd
from openpyxl import load_workbook

PACKAGE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(PACKAGE))
from config import digest,write_json
from export_for_origin import versions


def main():
    parser=argparse.ArgumentParser(description='Synthetic release checks; requires a new output directory')
    parser.add_argument('--output',required=True)
    args=parser.parse_args()
    out=Path(args.output).resolve()
    if out.exists():
        parser.error('Output exists; choose a new directory')
    out.mkdir(parents=True)
    report={'started_utc':datetime.now(timezone.utc).isoformat(),'status':'running',
            'data_kind':'synthetic_only','versions':versions(),'commands':[],'checks':[],
            'code_sha256':{str(p.relative_to(PACKAGE)):digest(p) for p in sorted(PACKAGE.rglob('*.py'))}}

    def execute(name,arguments):
        command=[sys.executable,'-X','utf8',*map(str,arguments)]
        result=subprocess.run(command,cwd=PACKAGE,encoding='utf-8',errors='replace',capture_output=True,timeout=600)
        # Saved evidence stays portable and omits machine-specific filesystem roots.
        def portable(text):
            return text.replace(str(out),'{validation_output}').replace(str(PACKAGE),'{package}')
        (out/f'{name}.log').write_text(portable(result.stdout+'\n'+result.stderr),encoding='utf-8')
        report['commands'].append({'name':name,'arguments':[portable(str(v)) for v in arguments],
                                    'returncode':result.returncode,'log':f'{name}.log'})
        if result.returncode:
            raise RuntimeError(f'{name} failed; see {out/name}.log')
        print(f'PASS {name}',flush=True)

    def check(name):
        report['checks'].append({'name':name,'status':'passed'})
        print(f'PASS {name}',flush=True)

    try:
        execute('01_unit_tests',['-m','unittest','discover','-s','tests','-v'])
        execute('02_full_default_demo',['export_for_origin.py','--demo','--output',out/'demo'])
        raw=['--cycle-dir','examples/raw/cycle','--eis-dir','examples/raw/eis',
             '--selection-manifest','examples/selection_manifest.csv','--config','examples/config_group.json']
        execute('03_prepare_raw',['extract_capacity.py',*raw,'--output',out/'prepared.csv'])
        execute('04_train_raw_group',['export_for_origin.py',*raw,'--synthetic-data','--output',out/'raw_group'])
        execute('05_train_csv_group',['export_for_origin.py','--features',out/'prepared.csv',
                '--config','examples/config_group.json','--synthetic-data','--output',out/'csv_group'])
        left=pd.read_csv(out/'raw_group'/'oof_predictions.csv')
        right=pd.read_csv(out/'csv_group'/'oof_predictions.csv')
        pd.testing.assert_frame_equal(left,right,check_exact=False,rtol=1e-8,atol=1e-12)
        check('raw_and_prepared_csv_reproduce_same_OOF')
        for run_name in ['demo','raw_group','csv_group']:
            folder=out/run_name
            state=json.loads((folder/'run.json').read_text(encoding='utf-8'))
            assert state['status']=='completed' and state['data_kind']=='synthetic_demo'
            metrics=json.loads((folder/'metrics.json').read_text(encoding='utf-8'),
                               parse_constant=lambda value:(_ for _ in ()).throw(ValueError(value)))
            for row in metrics['classification']:
                assert sum(row[k] for k in ['TP','FN','FP','TN'])==row['n_samples']==12
            preds=pd.read_csv(folder/'oof_predictions.csv')
            assert preds['cell_number'].is_unique and len(preds)==12
            assert np.isfinite(preds[['pred_capacity_Ah','pred_cc_frac','score_logistic']]).all().all()
            wb=load_workbook(folder/'results_for_origin.xlsx',read_only=True,data_only=False)
            assert len(wb.sheetnames)==11 and wb['07_clf_cellwise'].max_row==13
            assert not any(cell.data_type=='f' for sheet in wb for row in sheet for cell in row)
            wb.close()
            hashes=json.loads((folder/'artifact_hashes.json').read_text(encoding='utf-8'))
            assert all(digest(folder/path)==expected for path,expected in hashes.items())
        check('all_three_runs_complete_12_cells_11_sheets_finite_metrics_and_valid_hashes')
        traces=json.loads((out/'raw_group'/'folds.json').read_text(encoding='utf-8'))
        for target,folds in traces.items():
            assert len(folds)==3
            for fold in folds:
                test=set(fold['test_cells']); groups=set(fold['test_groups'])
                assert not test & set(fold['training_cells'])
                assert np.isfinite(fold['best_inner_score'])
                for inner in fold['inner_folds']:
                    assert not groups & (set(inner['train_groups'])|set(inner['validation_groups']))
                    assert not set(inner['train_groups']) & set(inner['validation_groups'])
        check('all_three_targets_keep_outer_and_inner_groups_separate')
        execute('06_prepare_unlabeled_eis',['extract_capacity.py','--config','config.json',
                '--eis-dir','examples/raw/new_eis','--output',out/'new_features.csv'])
        execute('07_predict_from_saved_models',['predict.py','--run',out/'raw_group',
                '--features',out/'new_features.csv','--output',out/'new_predictions.csv'])
        prediction=pd.read_csv(out/'new_predictions.csv')
        assert prediction['cell_number'].tolist()==list(range(61,73))
        assert np.isfinite(prediction[['predicted_capacity','predicted_cc_fraction','logistic_score']]).all().all()
        assert not prediction[[c for c in prediction if c.endswith('_seen_in_training')]].any().any()
        check('saved_models_predict_12_unlabeled_new_identifiers')
        report['status']='passed'
    except Exception as exc:
        report.update(status='failed',error_type=type(exc).__name__,error=str(exc))
        raise
    finally:
        report['finished_utc']=datetime.now(timezone.utc).isoformat()
        write_json(out/'release_verification.json',report)
    print(f'All release checks passed: {out}',flush=True)


if __name__=='__main__':
    main()
