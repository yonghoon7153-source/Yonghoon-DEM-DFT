"""시험 먼저 기록 — v3.1 생성기 · 그림 스크립트의 새 검사를 Codex 가 받은 v3 덱 (sha256 baf2bb98…) 과
v3 커밋 (0c6f5c84e) 의 그림 스크립트에 돌려 실패 목록을 모은다 (assert 대신 수집).  리포 · 덱은 읽기만 한다.

python3 testfirst_v31.py <repo> <v3 덱.pptx> <v3.1 덱.pptx> <out.json>
"""
import importlib.util, json, subprocess, sys, tempfile
from pathlib import Path

repo, deck_v3, deck_v31, out = (Path(a) for a in sys.argv[1:5])
V3_COMMIT = '0c6f5c84e'   # 덱 v3 커밋 (그림 스크립트 = ZIP v6 판 · 96a0cc02b 와 같음)


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


gen = load(repo / 'docs/seminar_20261013/build_seminar_v3.py', 'gen31')
from pptx import Presentation  # noqa: E402


def deck_checks(deck):
    prs = Presentation(deck)
    texts = gen.slide_texts(prs)
    notes = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else '' for s in prs.slides]
    low = [x.lower() for x in texts]
    return {
        'FORBIDDEN_CI (slides)': [(i, f) for i, x in enumerate(low, 1) for f in gen.FORBIDDEN_CI if f in x],
        'NOTES_FORBIDDEN': [(i, f) for i, x in enumerate(notes, 1) for f in gen.NOTES_FORBIDDEN if f in x],
        'SLIDE_FORBIDDEN': [(i, f) for i, fs in gen.SLIDE_FORBIDDEN.items() for f in fs if f in texts[i - 1]],
        'REQUIRED (slides) missing': [(i, r) for i, rs in gen.REQUIRED.items() for r in rs if r not in low[i - 1]],
        'NOTES_REQUIRED missing': [(i, r) for i, rs in gen.NOTES_REQUIRED.items() for r in rs if r not in notes[i - 1].lower()],
    }


new = load(repo / 'docs/seminar_20261013/figs_v3/make_seminar_figs.py', 'figs31')
old_src = subprocess.run(['git', '-C', str(repo), 'show', V3_COMMIT + ':docs/seminar_20261013/figs_v3/make_seminar_figs.py'],
                         capture_output=True, text=True, check=True).stdout


def old_fig_checks():
    """v3 커밋의 그림 스크립트로 그림을 다시 그려 새 검사 (글자 겹침 · 라벨 · 값) 를 적용한다."""
    res = {}
    with tempfile.TemporaryDirectory() as d:
        p = Path(d, 'old_figs.py')
        p.write_text(old_src, encoding='utf-8')
        old = load(p, 'figs_v3_old')
        pairs = []
        orig = old._save

        def _save(fig, out_dir, name, *a, **k):
            pairs.extend((name, x, y) for x, y in new._text_overlaps(fig))
            return orig(fig, out_dir, name, *a, **k)
        old._save = _save
        o = Path(d, 'out'); o.mkdir()
        old.fig_poro(o); old.fig_trans(o); old.fig_limits(o); old.fig_load(o); old.fig_vgcf(o)
        old.fig_dataset(o, repo)
        old.fig_mix1(o); old.fig_mix2(o); old.fig_mix3(o); old.fig_mix3_scheme(o)
        svg = {n: (o / f'{n}.svg').read_text(encoding='utf-8') for n in ('F_mix1', 'F_mix2', 'F_mix3', 'F_mix3_scheme')}
        res['text overlaps (pairs)'] = pairs
        res['zero-cohesion wording in svg'] = [n for n in ('F_mix1', 'F_mix2', 'F_mix3') if 'removed' in svg[n] or 'non-cohesive' in svg[n]]
        res['low-cohesion comparator label missing'] = [n for n in ('F_mix1', 'F_mix3') if 'low NCM' not in svg[n]]
        res['F_mix2 Bo / CED labels'] = [k for k in ('38.4', '1152', '14.42', 'Bo', 'CED') if k not in svg['F_mix2']] + (['reference present'] if 'reference' in svg['F_mix2'] else [])
        res['scheme assumptions missing'] = [k for k in ('Assumed', 'Conditional', 'J m') if k not in svg['F_mix3_scheme']]
        res['scheme survival claim present'] = [k for k in ('favoured', 'holds on most') if k in svg['F_mix3_scheme']]
        rows = (o / 'F_mix2.csv').read_text(encoding='utf-8').splitlines()
        res['F_mix2.csv Bo 384 12x12x3 value'] = rows[1 + 1 * 3 + 1] if len(rows) > 5 else rows
    return res


report = {'deck_v3_sha256': subprocess.run(['sha256sum', str(deck_v3)], capture_output=True, text=True).stdout.split()[0],
          'deck_v3': deck_checks(deck_v3), 'deck_v31': deck_checks(deck_v31), 'figs_v3_commit': old_fig_checks()}
out.write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding='utf-8')
for k, v in report['deck_v3'].items():
    print('v3  ', k, len(v))
for k, v in report['deck_v31'].items():
    print('v3.1', k, len(v))
for k, v in report['figs_v3_commit'].items():
    print('figs v3', k, v if not isinstance(v, list) else len(v))
