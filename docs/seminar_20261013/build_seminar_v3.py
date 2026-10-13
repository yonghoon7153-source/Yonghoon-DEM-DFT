#!/usr/bin/env python3
"""연구세미나 덱 생성기 v3 "보는 판" (2026-10-13 · 1저자 지시 — 양식 = 10-13 연구세미나 덱 사본 · 그 덱 4 장의 문법).

    python3 build_seminar.py --template <10-13 덱.pptx> --repo <리포> --figs <그림 폴더> --lit <문헌 크롭 폴더> --out <out.pptx>
    python3 build_seminar.py --selftest-only <deck.pptx> --repo <리포> --figs <그림 폴더>     # 검사만 (고치기 전 덱 = 실패해야 한다)
    python3 build_seminar.py --dump <deck.pptx> --out <slide_text.md>                       # 장별 글 전문

장 문법 (10-13 덱 4 장): 절 표지 24 pt · ■ 제목 18 pt (한 줄 · ≤ 62 자) + 소불릿 ≤ 2 (15 pt · 핵심어 빨강) · 그림 띠 (본문 높이 추정값 아래 · 바닥 6.45 in) ·
우상단 출처 줄 10 pt 기울임 · 그림 아래 해석 조건 한 줄 (12 pt 회색 · 꺾쇠 캡션 없음) · 문헌 크롭 = 위 꼬리표 (12.5 pt) + 아래 출처 (9.5 pt) ·
**빈 자리 = 아무 도형도 없음** (출처 줄의 `figure: presenter` 가 무엇을 붙일지 적는다).
앞 15 장 → 10 장 (5+6 · 7+8 · 9+10 · 11+12 · 14+15 합침) + 믹서 3 장 (Plan 앞) · 참고문헌 16 그대로.  Codex 가 닫은 값 · 문장은 그대로 (자리만 옮김 · notes 에 전문).
믹서 용어 = 분야 표준 (팔 코드 없음): cohesive / non-cohesive NCM · Bond number · Lacey mixing index · drum revolutions · bin grid · registered "no segregation" category.
"""
import argparse
import copy
import json
import math
import re
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.util import Inches, Pt

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
A = NS['a']
DARK, GRAY, NAVY, RED, BLUE, LINE = '1F2937', '6B7280', '1F497D', 'FF5050', '0070C0', 'D1D5DB'
BODY_PT, CAP_PT = 15, 12
N_SLIDES = 17
BODY_TOP, FZ_BOTTOM, FT_MIN = 1.08, 6.45, 2.95
CHARS_HEAD = 62                                   # ■ 제목 글자 수 상한 (한 줄)
HEAD_PT, HEAD_W, BUL_W = 18, 8.35, 8.0            # 제목 18 pt · 제목 줄 폭 · 불릿 줄 폭 (상자 8.84 − 들여쓰기 − 여백) in
HEAD_LINE_H, BUL_LINE_H, BUL_GAP, BODY_PAD, FT_MARGIN = 0.46, 0.30, 0.07, 0.10, 0.12   # 렌더 실측 (DejaVu · 125 % 줄 간격)
#: 줄바꿈 추정 글꼴 — Carlito (Calibri 와 글자 폭 같음 · fonts-crosextra-carlito) 가 있으면 그것 · 없으면 DejaVu (더 넓다 → 보수적)
_WRAP_FONTS = ('/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
_FONTS = {}


def _font(pt):
    from PIL import ImageFont
    if pt not in _FONTS:
        path = next(p for p in _WRAP_FONTS if Path(p).exists())
        _FONTS[pt] = ImageFont.truetype(path, int(round(pt * 4)))
    return _FONTS[pt]


def est_lines(text, pt, width_in):
    """탐욕 줄바꿈으로 줄 수를 센다 (실제 글꼴 폭 · 4 px/pt)."""
    f = _font(pt)
    maxw = width_in * 72 * 4
    lines, cur = 1, ''
    for w in text.split(' '):
        trial = (cur + ' ' + w).strip()
        if f.getlength(trial) <= maxw:
            cur = trial
        else:
            lines += 1
            cur = w
    return lines


def est_body_bottom(head, bullet_texts):
    """본문 상자 바닥 추정 (in) — ■ 제목 줄 수 + 불릿 줄 수 × 줄 높이.  그림 띠는 이 값 + 여유 아래에 둔다 (selftest 가 같은 식으로 겹침을 검사)."""
    h = BODY_PAD + HEAD_LINE_H * est_lines(head, HEAD_PT, HEAD_W)
    for t in bullet_texts:
        h += BUL_LINE_H * est_lines(t, BODY_PT, BUL_W) + BUL_GAP
    return BODY_TOP + h


# ───────────────────────────── 양식 복제 도우미 (v2.1 그대로) ─────────────────────────────

def dup_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith(('/slideLayout', '/notesSlide')):
            continue
        rid_map[rid] = (new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                        else new.part.relate_to(rel.target_part, rel.reltype))
    src_csld = src._element.find('{%s}cSld' % NS['p'])
    bg = src_csld.find('{%s}bg' % NS['p'])
    if bg is not None:
        new._element.find('{%s}cSld' % NS['p']).insert(0, copy.deepcopy(bg))
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(('}nvGrpSpPr', '}grpSpPr')):
            continue
        new.shapes._spTree.append(copy.deepcopy(el))
    for el in new.shapes._spTree.iter():
        for k in list(el.attrib):
            if k.startswith('{%s}' % NS['r']) and el.attrib[k] in rid_map:
                el.attrib[k] = rid_map[el.attrib[k]]
    return new


def shape(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    raise KeyError(name)


def shapes_named(slide, name):
    return [s for s in slide.shapes if s.name == name]


def drop(slide, *names):
    for n in names:
        for s in shapes_named(slide, n):
            s._element.getparent().remove(s._element)


def set_runs_text(tf, text):
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0]
    r0.text = text
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def _run(p, text, proto_rpr, color=None, bold=None, size=None, italic=None):
    r = etree.SubElement(p, '{%s}r' % A)
    rp = copy.deepcopy(proto_rpr) if proto_rpr is not None else etree.Element('{%s}rPr' % A)
    if color:
        for sf in rp.findall('{%s}solidFill' % A):
            rp.remove(sf)
        sf = etree.Element('{%s}solidFill' % A)
        c = etree.SubElement(sf, '{%s}srgbClr' % A); c.set('val', color)
        rp.insert(0, sf)
    if bold is not None:
        rp.set('b', '1' if bold else '0')
    if italic is not None:
        rp.set('i', '1' if italic else '0')
    if size:
        rp.set('sz', str(int(size * 100)))
    r.append(rp)
    t = etree.SubElement(r, '{%s}t' % A); t.text = text
    return r


def _spacing(paragraph, line=None, before=None, after=None):
    if line is not None:
        paragraph.line_spacing = line
    if before is not None:
        paragraph.space_before = Pt(before)
    if after is not None:
        paragraph.space_after = Pt(after)


def set_body(box, head, bullets, head_size=None, bullet_size=BODY_PT, line=1.25, after=5):
    """양식 본문 상자: 문단 0 = ■ 제목 (18 pt 상속) · 문단 1.. = • 소불릿 (15 pt · 핵심어 빨강 FF5050 = 양식 색)."""
    tf = box.text_frame
    body = tf._txBody
    paras = tf.paragraphs
    p_head = copy.deepcopy(paras[0]._p)
    p_bul = copy.deepcopy(paras[1]._p) if len(paras) > 1 else copy.deepcopy(paras[0]._p)
    for p in list(paras):
        body.remove(p._p)
    def fill(p_proto, segs, size):
        p = copy.deepcopy(p_proto)
        runs = p.findall('{%s}r' % A)
        rpr = copy.deepcopy(runs[0].find('{%s}rPr' % A))
        for r in runs:
            p.remove(r)
        for e in p.findall('{%s}endParaRPr' % A):
            p.remove(e)
        for txt, col in segs:
            _run(p, txt, rpr, color=col, size=size)
        body.append(p)
    fill(p_head, [(head, None)], head_size)
    for b in bullets:
        segs = [(b, None)] if isinstance(b, str) else b
        fill(p_bul, segs, bullet_size)
    for p in tf.paragraphs[1:]:
        _spacing(p, line=line, after=after)


def set_items(box, items, size=None, line=1.12, before=4):
    tf = box.text_frame
    body = tf._txBody
    paras = tf.paragraphs
    protos = [copy.deepcopy(p._p) for p in paras]
    for p in list(paras):
        body.remove(p._p)
    for it in items:
        kind, segs = ('b', it) if isinstance(it, (str, list)) else it
        proto = protos[0] if kind == 'b' or len(protos) == 1 else protos[1]
        p = copy.deepcopy(proto)
        runs = p.findall('{%s}r' % A)
        rpr = copy.deepcopy(runs[0].find('{%s}rPr' % A))
        for r in runs:
            p.remove(r)
        for e in p.findall('{%s}endParaRPr' % A):
            p.remove(e)
        segs = [(segs, None)] if isinstance(segs, str) else segs
        for txt, col in segs:
            _run(p, txt, rpr, color=col, size=size)
        body.append(p)
    for p in tf.paragraphs:
        _spacing(p, line=line, before=before)


def set_text(box, segs, size=None, bold=None, italic=None):
    tf = box.text_frame
    p0 = tf.paragraphs[0]
    runs = p0.runs
    rpr = copy.deepcopy(runs[0]._r.find('{%s}rPr' % A))
    for r in runs:
        p0._p.remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    segs = [(segs, None)] if isinstance(segs, str) else segs
    for txt, col in segs:
        _run(p0._p, txt, rpr, color=col, size=size, bold=bold, italic=italic)


def place(sp, box):
    x, y, w, h = box
    sp.left, sp.top, sp.width, sp.height = Inches(x), Inches(y), Inches(w), Inches(h)


def place_pic(slide, path, box, name=None):
    """box = (x, y, w, h) in — 비율 유지 · 가운데.  파일이 없으면 실패 (빈 자리는 이 함수를 부르지 않는다)."""
    x, y, w, h = box
    path = Path(path)
    assert path.exists(), path
    im = Image.open(path)
    iw, ih = im.size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    pic = slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    if name:
        pic.name = name
    return pic


def add_box(slide, proto_el, box, text_paras, name, size=None, line=None, align=None, italic=None):
    el = copy.deepcopy(proto_el)
    slide.shapes._spTree.append(el)
    sp = slide.shapes[-1]
    sp.name = name
    place(sp, box)
    tf = sp.text_frame
    tf.word_wrap = True
    body = tf._txBody
    paras = tf.paragraphs
    proto = copy.deepcopy(paras[0]._p)
    for p in list(paras):
        body.remove(p._p)
    for segs in text_paras:
        p = copy.deepcopy(proto)
        runs = p.findall('{%s}r' % A)
        rpr = copy.deepcopy(runs[0].find('{%s}rPr' % A)) if runs else None
        for r in runs:
            p.remove(r)
        for e in p.findall('{%s}endParaRPr' % A):
            p.remove(e)
        for seg in segs:
            txt, col, b = (seg + (None, None))[:3] if isinstance(seg, tuple) else (seg, None, None)
            _run(p, txt, rpr, color=col, bold=b, size=size, italic=italic)
        body.append(p)
    for p in tf.paragraphs:
        if line is not None:
            _spacing(p, line=line)
        if align is not None:
            p.alignment = align
    return sp


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ───────────────────────────── v3 요소 ─────────────────────────────

def _tc_borders(cell, color=LINE, w_pt=0.75, sides=('B',)):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('lnL', 'lnR', 'lnT', 'lnB'):
        for el in tcPr.findall('{%s}%s' % (A, tag)):
            tcPr.remove(el)
    idx = 0
    for tag in ('lnL', 'lnR', 'lnT', 'lnB'):
        ln = etree.Element('{%s}%s' % (A, tag))
        if tag[-1] in sides:
            ln.set('w', str(int(w_pt * 12700)))
            sf = etree.SubElement(ln, '{%s}solidFill' % A)
            c = etree.SubElement(sf, '{%s}srgbClr' % A); c.set('val', color)
        else:
            etree.SubElement(ln, '{%s}noFill' % A)
        tcPr.insert(idx, ln); idx += 1


def add_table(slide, box, rows, col_w, size=11.5, name='Table', row_h=None, head_fill=NAVY):
    """표 도형 (양식 스타일 없음): 머리 행 = 남색 바탕 · 흰 글씨 · 본문 = 흰 바탕 · 아래 가는 선 · 첫 열 굵게."""
    x, y, w, h = box
    nrows, ncols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for k in ('firstRow', 'bandRow', 'firstCol'):
        tblPr.set(k, '0')
    for el in tblPr.findall('{%s}tableStyleId' % A):
        tblPr.remove(el)
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    if row_h:
        for i, rh in enumerate(row_h):
            tbl.rows[i].height = Inches(rh)
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.name = 'Calibri'
            if i == 0:
                r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(head_fill)
            else:
                r.font.bold = (j == 0)
                r.font.color.rgb = RGBColor.from_string(NAVY if j == 0 else DARK)
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _tc_borders(cell, sides=('B',))
    return gf


def add_flow(slide, y0=3.25):
    """2 장 흐름도 = PPT 도형 (둥근 상자 4 · 화살표 3 · 보정 표적 띠 3) — 글은 v2.1 F1 그림과 같다 (SEM-01 · SEMV2-04)."""
    boxes = [('DEM packing', 'soft-sphere DEM\n(contact law)', NAVY),
             ('Plastic compaction', 'MPM: J2 plasticity\nDEM scaffold', NAVY),
             ('Voxel transport', 'finite-volume solver\nionic · electronic · thermal', NAVY),
             ('Electrode response', 'reaction distribution\n(in development)', GRAY)]
    x0, w, h, gap = 0.75, 1.85, 1.2, 0.37
    for i, (head, sub, col) in enumerate(boxes):
        x = x0 + i * (w + gap)
        sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y0), Inches(w), Inches(h))
        sp.name = f'Flow box {i + 1}'
        sp.adjustments[0] = 0.14
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sp.line.color.rgb = RGBColor.from_string(col); sp.line.width = Pt(1.25)
        sp.shadow.inherit = False
        tf = sp.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.04)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
        r = p0.add_run(); r.text = head; r.font.size = Pt(14); r.font.bold = True; r.font.name = 'Calibri'; r.font.color.rgb = RGBColor.from_string(col)
        for line_ in sub.split('\n'):
            p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = line_; r.font.size = Pt(10.5); r.font.name = 'Calibri'; r.font.color.rgb = RGBColor.from_string(DARK)
        if i < len(boxes) - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w + 0.06), Inches(y0 + h / 2 - 0.13), Inches(gap - 0.12), Inches(0.26))
            ar.name = f'Flow arrow {i + 1}'
            ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor.from_string(GRAY); ar.line.fill.background(); ar.shadow.inherit = False
    targets = ['calibration target:\ncompaction data', 'calibration target:\npure-SE compaction', 'comparison target:\nconductivity data']
    for i, txt in enumerate(targets):
        x = x0 + i * (w + gap)
        cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + w / 2), Inches(y0 + h), Inches(x + w / 2), Inches(y0 + h + 0.4))
        cn.name = f'Flow tie {i + 1}'
        cn.line.color.rgb = RGBColor.from_string(GRAY); cn.line.width = Pt(0.75); cn.line.dash_style = MSO_LINE.ROUND_DOT
        tb = slide.shapes.add_textbox(Inches(x - 0.12), Inches(y0 + h + 0.42), Inches(w + 0.24), Inches(0.55))
        tb.name = f'Flow target {i + 1}'
        tf = tb.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.02); tf.margin_top = tf.margin_bottom = Inches(0.0)
        for k, line_ in enumerate(txt.split('\n')):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = line_; r.font.size = Pt(11); r.font.name = 'Calibri'; r.font.color.rgb = RGBColor.from_string(GRAY)


def seg_text(b):
    return b if isinstance(b, str) else ''.join(t for t, _ in b)


# ───────────────────────────── 문헌 크롭 (litdb 정본 브랜치 · 사본을 리포에 두지 않는다) ─────────────────────────────

LIT_BRANCH = 'origin/claude/friendly-meitner-lldvar'
LIT_FILES = ('bazzoun2026_dem_fem_rnm_ionic/fig_2.png',
             'luding2008_cohesive_frictional_contact_models/fig_1.png',
             'so2021_dem_fabrication_degradation_ductile_particles/fig_4.png')
SO_CROP = (0.0, 0.0, 0.475, 1.0)     # So 2021 Fig. 4 → (a) 패널 (상대 밀도 vs 압력 · Sakuda 실험점) — 폭 · 높이 분율 (l, t, r, b)


def fetch_lit(lit, repo):
    """`--lit` 폴더에 크롭이 없으면 litdb 정본 브랜치 `litdb/figures/<slug>/` 에서 받는다 (git show) · So 2021 은 (a) 패널만 잘라 쓴다."""
    import subprocess
    lit = Path(lit)
    lit.mkdir(parents=True, exist_ok=True)
    for rel in LIT_FILES:
        dst = lit / rel
        if not dst.exists():
            dst.parent.mkdir(parents=True, exist_ok=True)
            data = subprocess.run(['git', '-C', str(repo), 'show', f'{LIT_BRANCH}:litdb/figures/{rel}'], check=True, capture_output=True).stdout
            dst.write_bytes(data)
    crop = lit / 'so2021_fig4a_crop.png'
    if not crop.exists():
        im = Image.open(lit / LIT_FILES[2])
        w, h = im.size
        l, t, r, b = SO_CROP
        im.crop((int(w * l), int(h * t), int(w * r), int(h * b))).save(crop)
    return lit


# ───────────────────────────── 덱 내용 ─────────────────────────────

def build(template, repo, figs, lit, out):
    repo = Path(repo); figs = Path(figs)
    lit = fetch_lit(lit or Path(out).resolve().parent / 'lit', repo)
    prs = Presentation(template)
    P = {k: prs.slides[i - 1] for k, i in dict(title=1, pic=4, two=11, summary=15, thanks=16, refs=17).items()}
    CONC_PROTO = copy.deepcopy(shape(P['summary'], 'Conclusion')._element)      # 글 상자 원형 (Calibri)
    n_orig = len(prs.slides)
    D = {}

    def note_line(slide, box, text, color=GRAY, size=CAP_PT, align=None):
        """그림 아래 해석 조건 (꺾쇠 없음 · ≥ 12 pt · 이름 'Caption' = 검사 대상)."""
        paras = [text] if isinstance(text, str) else text
        return add_box(slide, CONC_PROTO, box, [[(t, color)] for t in paras], 'Caption', size=size, line=1.1, align=align)

    def lit_tag(slide, box, text):
        return add_box(slide, CONC_PROTO, box, [[('Literature: ', DARK, True), (text, DARK)]], 'Lit tag', size=12.5, line=1.1)

    def credit(slide, box, text):
        return add_box(slide, CONC_PROTO, box, [[(text, GRAY)]], 'Credit', size=9.5, line=1.1, italic=True)

    def source(slide, text):
        """우상단 출처 줄 (양식 TextBox 69 · 10 pt 기울임 · 오른쪽 맞춤 · ≤ 110 자) — 빈 자리는 `figure: presenter — …`."""
        assert len(text) <= 110, (len(text), text)
        sp = shape(slide, 'TextBox 69')
        set_text(sp, [(text, GRAY)], size=10, italic=True)
        place(sp, (2.45, 0.93, 7.2, 0.27))

    def body_slide(section, head, bullets, src=None):
        """본문 장 + 그림 띠 윗선 ft (본문 높이 추정값 + 여유 · 최소 2.95) 를 돌려준다."""
        assert len(head) <= CHARS_HEAD and est_lines(head, HEAD_PT, HEAD_W) == 1, (len(head), head)
        s = dup_slide(prs, P['pic'])
        drop(s, 'New picture')
        set_runs_text(shape(s, 'TextBox 20').text_frame, section)
        box = shape(s, 'Rectangle 4')
        set_body(box, head, bullets)
        ft = max(FT_MIN, round(est_body_bottom(head, [seg_text(b) for b in bullets]) + FT_MARGIN, 2))
        place(box, (0.61, BODY_TOP, 8.84, ft - BODY_TOP - 0.05))
        if src:
            source(s, src)
        else:
            drop(s, 'TextBox 69')
        return s, ft

    def two_column(section, head, bullet, left_title, left_items, right_title, right_items, bottom=None, bottom_box=(0.55, 5.75, 8.90, 0.65)):
        s = dup_slide(prs, P['two'])
        set_runs_text(shape(s, 'TextBox 20').text_frame, section)
        set_body(shape(s, 'Rectangle 4'), head, [bullet])
        place(shape(s, 'Rectangle 4'), (0.61, 1.08, 8.84, 1.05))
        lt, li = shape(s, 'Audit findings'), shape(s, 'Audit findings items')
        rt, ri = shape(s, 'Decisions recorded on Sep 12'), shape(s, 'Decisions recorded on Sep 12 items')
        set_text(lt, left_title, size=15); place(lt, (0.55, 2.20, 4.25, 0.36))
        set_text(rt, right_title, size=15); place(rt, (5.05, 2.20, 4.40, 0.36))
        set_items(li, left_items, size=BODY_PT); place(li, (0.55, 2.66, 4.25, 2.75))
        set_items(ri, right_items, size=BODY_PT); place(ri, (5.05, 2.66, 4.40, 2.75))
        concs = shapes_named(s, 'Conclusion')
        if bottom:
            set_text(concs[0], bottom, size=14); place(concs[0], bottom_box)
        else:
            concs[0]._element.getparent().remove(concs[0]._element)
        concs[1]._element.getparent().remove(concs[1]._element)
        return s

    R = RED
    # ── 1 표지 ────────────────────────────────────────────────────────────
    s = dup_slide(prs, P['title'])
    set_runs_text(shape(s, '직사각형 12').text_frame, '[date — to be set]')
    tb = shape(s, 'Rectangle 4').text_frame
    p1 = tb.paragraphs[1]
    r0 = p1.runs[0]
    r0.text = 'Particle-resolved multiscale simulation of sulfide all-solid-state composite cathodes:'
    rpr = copy.deepcopy(r0._r.find('{%s}rPr' % A))
    for r in p1.runs[1:]:
        r._r.getparent().remove(r._r)
    for br in p1._p.findall('{%s}br' % A):
        p1._p.remove(br)
    br = etree.SubElement(p1._p, '{%s}br' % A)
    if rpr is not None:
        br.append(copy.deepcopy(rpr))
    _run(p1._p, 'from powder packing to transport', rpr)
    au = shape(s, '직사각형 10')
    place(au, (0.68, 2.98, 8.40, 0.75))
    atf = au.text_frame
    atf.margin_top = atf.margin_bottom = 0
    for p in list(atf.paragraphs):
        if not p.text.strip():
            p._p.getparent().remove(p._p)
    for i, p in enumerate(atf.paragraphs):
        for r in p.runs:
            r.font.size = Pt(16 if i == 0 else 14)
        p.line_spacing = Pt(18); p.space_before = Pt(0); p.space_after = Pt(0)
    D[1] = 'title'
    notes(s, 'Title. Date to be set by the presenter.')

    # ── 2 Introduction — workflow (PPT 도형) ─────────────────────────────
    s, ft = body_slide('Introduction', 'A coupled DEM–MPM–transport workflow',
                       [[('Each stage has specific ', None), ('calibration targets', R), (' and unresolved validation tasks.', None)],
                        [('Agreement between models sharing inputs is a ', None), ('consistency check, not independent validation', R), ('.', None)]])
    add_flow(s, y0=ft + 0.35)
    note_line(s, (0.6, ft + 2.75, 8.8, 0.36), 'DEM packing → plastic compaction → voxel transport → electrode response; targets are what each stage is calibrated or compared against', align=PP_ALIGN.CENTER)
    D[2] = 'CLAUDE.md frame [1]–[5]; Codex 10-13 SEM-01 wording; flow chart = PPT shapes (same words as the v2.1 F1 figure)'
    notes(s, 'No numbers. Calibration targets per stage: DEM — compaction experiments; MPM (material point method) — pure-electrolyte compaction (lab reference sample, unpublished); '
             'voxel transport — comparison with conductivity measurements (not yet matched). Agreement between stages that share inputs is a consistency check (SEM-01).')

    # ── 3 DEM — porosity (3D four stages + porosity vs composition) ─────
    s, ft = body_slide('DEM', 'Composition-dependent porosity in five simulated beds',
                       [[('Soft-sphere DEM with a linear hysteretic contact law, pressed to 300 MPa; ', None), ('one bed per composition, no seed replicates', R), ('.', None)],
                        [('One 7:3 bed: porosity 31.8 → 25.1 → 16.7 → 16.5 % (final state relaxed at fixed platen); ', None), ('7:3 (large : small NCM) is the lowest of the five beds', R), ('.', None)]],
                       src='This work — DEM, five NCM–SE beds (one bed per composition)')
    fh = FZ_BOTTOM - ft - 0.55
    place_pic(s, repo / 'docs/data/ps73_compaction_curve_20261006/stages_20261012/aligned_ckpt/t1_geometry/ps73_geom3d_4panel_porosity_preview.png', (0.45, ft + 0.1, 5.7, fh - 0.1), name='Fig 3D stages')
    place_pic(s, figs / 'F_poro.png', (6.3, ft, 3.15, fh), name='Fig porosity')
    note_line(s, (0.45, ft + fh + 0.04, 9.0, 0.5), ['Left: four stages of one 7:3 bed, 0 · 47 · 296 · 165 MPa (last = relaxed at fixed platen); porosity = exact union of spheres.',
                                                    'Right: final porosity of the five beds (one bed each; the red bar is the minimum).'])
    D[3] = 'ps73 stages README §6 (porosity exact union 31.81 · 25.09 · 16.68 · 16.49); report_20261021 G1 (20.59 · 18.67 · 17.67 · 16.49 · 16.90) — figure F_poro redrawn from G1'
    notes(s, 'Sources: docs/data/ps73_compaction_curve_20261006/stages_20261012/README.md §6; docs/report_20261021/summary.json G1 (figure F_poro redrawn from the same values). '
             'Qualifiers: one bed per composition, no seed replicates; porosity = exact union of spheres; "lowest of five" is a statement about these five single beds, '
             'not an optimum ratio; the 165 MPa panel is the relaxed state at fixed platen, not an unloaded state.')

    # ── 4 DEM → transport (two-panel chart + Bazzoun 2026 network crop) ─
    s, ft = body_slide('DEM', 'From the contact network to Li⁺ transport',
                       [[('Contact-network approach [8]: bulk resistance + ', None), ('Holm/Maxwell constriction', R), (' using Hertz contact radii; chosen σ₀ = 1.5 mS cm⁻¹ (pellet-literature range).', None)],
                        [('Transport tortuosity T = φ_SE σ₀ / σ_eff: 7.2 → 5.0 (0:10 → 7:3). Contact constriction accounts for ', None), ('73–75 % of Joule dissipation', R), (' in this network model.', None)]],
                       src='This work (left) · Bazzoun et al., J. Power Sources 661, 238682 (2026), Fig. 2 (right) [8]')
    fh = FZ_BOTTOM - ft - 0.55
    place_pic(s, figs / 'F_trans.png', (0.45, ft, 4.95, fh), name='Fig transport')
    lit_tag(s, (5.6, ft, 3.85, 0.3), 'SE resistor network [8]')
    pic = place_pic(s, lit / 'bazzoun2026_dem_fem_rnm_ionic/fig_2.png', (5.6, ft + 0.36, 3.85, fh - 0.36 - 0.3), name='Lit Bazzoun')
    credit(s, (5.6, pic.top / 914400 + pic.height / 914400 + 0.03, 3.85, 0.26), 'Reproduced from Bazzoun et al., J. Power Sources 661, 238682 (2026), Fig. 2')
    note_line(s, (0.45, ft + fh + 0.04, 4.95, 0.5), ['(a) transport tortuosity (Hertz), (b) σ(Li⁺) at σ₀ = 1.5 mS cm⁻¹; one bed per composition;', 'T is the literature "tortuosity factor" (not its square root); red = extreme of the five'])
    D[4] = 'ps45_r45_network_g2 README (tau2 7.2433 · 5.9733 · 5.4495 · 5.0297 · 5.2864; σ_ion 0.0587 · 0.0730 · 0.0814 · 0.0893 · 0.0838 at σ0 = 1.5; constriction share 73–75 %); literature crop = litdb figures bazzoun2026_dem_fem_rnm_ionic/fig_2'
    notes(s, 'Sources: docs/data/ps45_r45_network_g2_20261007/README.md (values, σ0 = 1.5 section, decomposition section); figure F_trans redrawn from those values (y-axes start at 0). '
             'Qualifiers: transport tortuosity T = φ_SE σ0 / σ_eff (the literature "tortuosity factor"; not its square root); Hertz H0 = primary; '
             'the 73–75 % share = ΣI²R(constriction) / ΣI²R(total) of the same FULL network solution — a model-internal ratio, not evidence that contact necks explain 75 % of a real electrode. '
             'No seed replicates for these five beds; the 7:3 vs 10:0 difference is not resolved. [8] = Bazzoun 2026 (same network approach; the crop shows their resistor-network scheme). '
             'A 3D current-pathway view of the 7:3 bed is on the transport slide (presenter figure).')

    # ── 5 DEM limits + MPM (Luding 2008 crop + blank for the MPM 3D view) ─
    s, ft = body_slide('DEM → MPM', 'Fixed-shape DEM limits and the material point method (MPM)',
                       [[('No shape change and no yield cap: ', None), ('overlap is the only plasticity proxy', R), ('; the softened SE modulus (1.35 vs ≈ 22 GPa) is an ', None), ('effective parameter, not a property', R), (' [1, 7].', None)],
                        [('J2 plasticity on a background grid: ', None), ('shape change and void filling', R), (' that a fixed-shape sphere cannot represent; active material is kept rigid (DEM skeleton), only the electrolyte flows.', None)]],
                       src='Luding 2008, Fig. 1 [1] (left) · figure: presenter — MPM 3D view of the DEM-scaffold bed (right)')
    fh = FZ_BOTTOM - ft - 0.55
    lit_tag(s, (0.5, ft, 3.9, 0.3), 'hysteretic contact law [1]')
    pic = place_pic(s, lit / 'luding2008_cohesive_frictional_contact_models/fig_1.png', (0.5, ft + 0.36, 3.9, fh - 0.36 - 0.3), name='Lit Luding')
    credit(s, (0.5, pic.top / 914400 + pic.height / 914400 + 0.03, 3.9, 0.26), 'Reproduced from Luding, Granul. Matter 10, 235 (2008), Fig. 1')
    note_line(s, (0.5, ft + fh + 0.04, 3.9, 0.5), ['δ = overlap; piece-wise linear loading / unloading branches;', 'the implementation used here is the LIGGGHTS hysteretic law [7]'])
    D[5] = 'dem_governing_equations_background_20261011.md §5.1 · §10.3 · §11; DEMP-01 (57.9 %, notes); CLAUDE.md DEM↔MPM frame; literature crop = litdb figures luding2008_cohesive_frictional_contact_models/fig_1; right half = presenter figure'
    notes(s, 'Sources: docs/dem_governing_equations_background_20261011.md (§5.3; §10.3 literature counts; §11 limits); CLAUDE.md MPM sections (2026-06-07 … 06-17). '
             'In one 7:3 bed, global viscous damping carries 58 % of the plate reaction during compression — an open issue (damping share 57.9 % measured on one 7:3 bed during compression; ledger DEMP-01, open). '
             'Selected literature survey (22 studies; author classification, not a systematic review): softened SE modulus 0/22; yield or hardness cap 7/22; linear hysteretic law 1/22; global damping in dry compaction none. '
             'Young\'s modulus reference: our DFT E_VRH = 22.06 GPa (literature DEM input 22.1 GPa, Bazzoun 2026 [8]); the bulk modulus K / B0 = 26.23 GPa is a separate quantity; the glass-electrolyte value 24 GPa (Sakuda 2013 [15]) is a different material system. '
             'The DEM statement is about this implementation (fixed spherical shapes, hooke/hysteresis), not about DEM in general. '
             'MPM parameters: 2D E = 1.53 GPa, σy = 0.15 GPa; 3D E = 1.53 GPa, ν = 0.49, σy = 0.30 GPa (next slide). '
             'Presenter figure (right, blank): MPM 3D render of the DEM-scaffold bed (AM fixed, SE material points), e.g. python3 scripts/viz_mpm_morphology.py --se <SE material points .npy> --scaffold <AM scaffold .csv> --y 0.5 --out morphology.png '
             '(actual files required; MPM renders are not tracked in the repository). Do not present the MPM image as an independent validation of the DEM bed.')

    # ── 6 MPM inputs (table) + So 2021 crop ─────────────────────────────
    s, ft = body_slide('MPM', 'Effective material inputs and their status',
                       [[('Bulk modulus from DFT; ', None), ('shear softening and yield stress', R), (' from a pure-electrolyte calibration sweep.', None)],
                        [('Historical and current scaffold calculations use different electrolyte initializations and run configurations; ', None), ('they are not one validation series', R), ('.', None)]],
                       src='This work (left) · So et al., J. Power Sources 508, 230344 (2021), Fig. 4a (right) [9]')
    fh = FZ_BOTTOM - ft - 0.55
    rows = [('Input', 'Value and status'),
            ('LPSCl reference', 'E ≈ 22 GPa (our DFT E_VRH; literature DEM input 22.1 GPa [8])'),
            ('DEM effective', 'E = 1.35 GPa — reproduces compaction, not a property'),
            ('MPM', 'E = 1.53 GPa, ν = 0.49, σy = 0.30 GPa; K = 25.5 ≈ DFT 26.2 GPa'),
            ('Calibration target', 'pure-SE porosity 10 % at 300 MPa (lab reference sample, unpublished)'),
            ('Validation', 'against independent experiments: pending')]
    add_table(s, (0.45, ft + 0.05, 4.9, 2.75), rows, (1.45, 3.45), size=11.5, name='Table inputs', row_h=(0.36, 0.55, 0.42, 0.55, 0.55, 0.36))
    note_line(s, (0.45, ft + fh + 0.04, 4.9, 0.5), 'Tested alternatives (volumetric cap, real E) did not remove the need for effective parameters.', color=BLUE)
    lit_tag(s, (5.55, ft, 3.9, 0.3), 'DEM vs SE density–pressure data [9]')
    pic = place_pic(s, lit / 'so2021_fig4a_crop.png', (5.55, ft + 0.36, 3.9, fh - 0.36 - 0.3), name='Lit So')
    credit(s, (5.55, pic.top / 914400 + pic.height / 914400 + 0.03, 3.9, 0.26), 'Reproduced from So et al., J. Power Sources 508, 230344 (2021), Fig. 4a')
    D[6] = 'CLAUDE.md 3D MPM (ν 0.49 · K 25.5 ≈ B0 26.23); DFT E_VRH 22.06; Bazzoun 2026 E_SE 22.1; CL-94 (10 % = calibration target, source unverified); DPC cross-check; ' \
           'scaffold status (notes) = fam_platen_prereg §11 · §12-6-2; literature crop = litdb figures so2021_dem_fabrication_degradation_ductile_particles/fig_4 (panel a)'
    notes(s, 'Sources: CLAUDE.md "3D MPM compaction — 3-fix calibration" (K = E/(3(1−2ν)) = 25.5 GPa vs DFT B0 26.23); our DFT E_VRH 22.06 GPa; Bazzoun 2026 [8] uses E_SE = 22.1 GPa for Li6PS5Cl; '
             'the sulfide-glass value 24 GPa (75Li2S·25P2S5, Sakuda 2013 [15]) is a different material system and is not used on the slide. '
             'σy = 0.30 GPa was selected in the pure-SE calibration sweep (0.15 → 5.6 %, 0.20 → 6.7 %, 0.25 → 9.0 %, 0.30 → 10.0 %), inside the literature range. '
             'Claims ledger CL-94: the pure-SE 10 % value is a calibration target (lab standard sample, unpublished); it is not a literature anchor — our earlier literature attribution was withdrawn. '
             'Shear-only softening (ν = 0.49) keeps K = 25.5 GPa close to the DFT reference B0 = 26.23 GPa; softening the bulk modulus over-crushed the bed. '
             '"Tested alternatives" = DPC volumetric cap and real-E runs (docs/mpm_dpc_cap_crosscheck.md); the statement is limited to the tested models. '
             'Scaffold calculations (AM skeleton from the DEM bed) — status: Earlier: cell-filled SE, bidirectional servo — historical, not a DEM-seeded validation. '
             'Current registered runs: DEM-seeded SE, stress-controlled servo — f = 0 completed; f(a) execution failed. AM-load-sharing hypothesis: verdict pending. '
             'The current statuses refer to the registered servo runs (--protocol servo; primary endpoint = final thickness after servo convergence). '
             'AM load sharing remains a candidate explanation for over-compaction; the final hypothesis verdict is pending. '
             'The earlier thickness agreement is not shown: it comes from a different initialization and loading protocol and cannot serve as the validation value of the DEM-seeded runs (SEM-02). '
             'Literature crop: So et al. 2021 calibrate a contact-plasticity parameter on the relative density vs pressure of a sulfide SE (their Fig. 4a, experimental points from Sakuda et al.) — a literature example of calibrating an effective contact parameter on pure-SE compaction data.')

    # ── 7 Transport: voxel mesh + additives (blank 3D current flow + VGCF chart) ─
    s, ft = body_slide('Transport', 'Voxelizing the compacted microstructure into a transport mesh',
                       [[('Faces are joined by a harmonic mean with ', None), ('no interface-resistance term', R), (': an ideal-interface voxel model — an ', None), ('upper bound only within the same geometry, properties and boundaries', R), ('.', None)],
                        [('VGCF as diameter-preserving segments, PTFE as a centre-line insulator, polymer as spheres; a registered 104-arm study keeps the σ(e⁻) ', None), ('order across 1–4 wt% VGCF', R), (' on all tested grids (72/72).', None)]],
                       src='figure: presenter — 3D current-flow views, contact network and voxel (left) · this work, 104-arm study (right)')
    fh = FZ_BOTTOM - ft - 0.55
    place_pic(s, figs / 'F_vgcf.png', (5.3, ft, 4.15, fh), name='Fig VGCF')
    note_line(s, (5.3, ft + fh + 0.04, 4.15, 0.5), ['σ(e⁻) vs VGCF wt%: campaign means (96 primary arms), three voxel sizes;', 'no contact resistance'])
    D[7] = 'CLAUDE.md §5 two-solver table; CL-81 (ideal-interface / contact-free branch); CL-60 (centerline); phase_a_rep100 README (ORDER-ROBUST 72/72, replay 0.00000 %); report_20261021 G6 (figure F_vgcf redrawn); left half = presenter figure'
    notes(s, 'Sources: CLAUDE.md "σ 를 내는 솔버는 둘이다" table and CL-81. Finite-volume conservation on the phase grid gives electrical (σ, φ) and thermal (κ, T) transport on the same structure. '
             'The voxel solver has no interface-resistance term by default (the film mechanism exists but is off): this makes it an ideal-interface model and an upper bound only for the same geometry, phase properties and boundary conditions — not a physical upper bound; '
             'geometric necks still cost resistance inside the voxel grid. The 4.04× (contact-free / full network, n = 157) is a ratio inside the contact-network model. '
             'Additives: carbon fibres (VGCF) as diameter-preserving segments, PTFE as a centre-line insulator, a conductive polymer as true-diameter spheres. Pre-registered 104-arm study: the electronic-conductivity order across 1–4 wt% VGCF holds across the tested voxel sizes and grid-origin offsets (72/72) — '
             'docs/data/phase_a_rep100_20260930/README.md (104 arms, 72/72 adjacent-composition comparisons with no threshold violation; order only is the registered estimand; increments are technical report only; replay difference 0.00000 %); '
             'docs/report_20261021 G6 (voxel σe means; contact resistance not included; figure F_vgcf redrawn from G6). Stamp conventions: CLAUDE.md SR-01 (sphere stamp, centerline PTFE = CL-60). '
             'Presenter figures (left, blank): (i) contact-network current pathways of the 7:3 bed — scripts/network_current.py dump <7:3 result folder> → webapp DEM 3D viewer "current flow" (same bed, same H0/FULL ionic solution, 1 V, colour = j = |I|/A; '
             '"top 50 % of current" = edges carrying the top half of the current share, not the top half by count); (ii) STEP3 phase grid and the 1 V current-density field of the same voxel solution (A cm⁻²; state origin, voxel size, phase mapping, electrode rule, run hash).')

    # ── 8 Limits 1 + 2 (three-panel) ───────────────────────────────────
    s, ft = body_slide('Limits', 'Problems 1 and 2: grid dependence and absolute σ(e⁻)',
                       [[('The electronic ratio decreases across the three tested voxel spacings (0.4, 0.3, 0.25 µm); the ionic ratio crosses 1. ', None), ('Grid convergence has not been established', R), ('.', None)],
                        [('Voxel predictions ', None), ('exceed', R), (' the available DC-polarization estimates; measurement conditions remain ', None), ('unverified', R), ('. The NCM–SE comparison uses a common reference conductivity and a specified contact model.', None)]],
                       src='This work — voxel transport solver (a, b) · contact network, five NCM–SE beds (c)')
    fh = FZ_BOTTOM - ft - 0.55
    place_pic(s, figs / 'F_limits.png', (0.4, ft, 9.1, fh), name='Fig limits')
    note_line(s, (0.45, ft + fh + 0.04, 9.0, 0.5), ['(a) same bed pair and convention; no converged conductivity ratio is reported · (b) log scale; conditions not yet matched ·',
                                                    '(c) no carbon/binder; Hertz H0; common NCM reference 4.1 mS cm⁻¹ [16]; size factor off; one bed per composition'])
    D[8] = 'CL-24 (live: 1.4215 → 1.1621 → 1.0849; 1.0742 → 0.9926 → 0.9908); CL-41 hold; CL-72 retracted; sdcp_sigma_e_dcp_anchor_registration_20261012 (54.0 · 70.6 ↔ 4.47 · 12.4); ps45_sigmae_factor_off_20261013 (0.391 · 0.406 · 0.402 · 0.375 · 0.467); Wang 2018 [16]'
    notes(s, 'Sources: claims ledger CL-24 (public values at vox 0.4/0.3/0.25); CL-41 (vox ≤ 0.15 values on hold — not shown); CL-72 ("lower bound" retracted). '
             'The three public points admit an exact power-law fit (three points, three parameters), so no statement about power laws or extrapolation is made from them (SEM-03); '
             'the earlier manuscript headline at 0.4 µm is superseded. Finer grids reach the memory limit of one GPU (peak RSS 35.6 GB at vox 0.115; 0.10 µm ≈ 54 GB). '
             'Panel (b): docs/reviews/sdcp_sigma_e_dcp_anchor_registration_20261012.md (DC-polarization values 4.47 and 12.4 mS/cm, measurement conditions unverified; model 54.0 and 70.6; re-run registered, no results yet). '
             'Panel (c): docs/data/ps45_sigmae_factor_off_20261013/README.md (0.391 · 0.406 · 0.402 · 0.375 · 0.467 mS/cm; raw minimum 0.3747). Footnote (Codex §6-1): computed effective electronic conductivity of the '
             'DEM particle arrangement; Hertz H0 / Holm–Maxwell constriction; common reference conductivity 4.1 mS/cm for polycrystalline and single-crystal NCM (Wang 2018 [16], SPS pellet NMC811 at 20 °C, designated reference); '
             'no empirical size factor; absolute values conditional on this reference and contact model; single bed per composition, scatter not evaluated; carbon- and binder-free AM–SE composite.')

    # ── 9 Limits 3 — load sharing (stacked bar + SE von Mises 3D) ──────
    s, ft = body_slide('Limits', 'Problem 3: Load sharing during compaction',
                       [[('NCM–NCM contacts contribute ', None), ('48–60 % of vertical stress', R), (' after relaxation at fixed platen position; ', None), ('7:3 is the minimum', R), ('.', None)],
                        [('The porosity ranking agrees for the selected five compositions (exploratory comparison, n = 5).', None)]],
                       src='This work — DEM, five beds (left) · one 7:3 bed, four stages (right)')
    fh = FZ_BOTTOM - ft - 0.55
    place_pic(s, figs / 'F_load.png', (0.45, ft, 3.5, fh), name='Fig load share')
    place_pic(s, repo / 'docs/data/ps73_compaction_curve_20261006/stages_20261012/aligned_ckpt/t3_vm_se/ps73_se_vm3d_4panel_nowall_linear_0_150_preview.png', (4.05, ft, 5.4, fh), name='Fig SE von Mises')
    note_line(s, (0.45, ft + fh + 0.04, 9.0, 0.5), ['Left: contribution to σ_zz by contact type, 21 planes, load-weighted; relaxed at fixed platen.',
                                                    'Right: SE von Mises stress in one 7:3 bed, four stages (0 · 47 · 296 · 165 MPa), 0–150 MPa; wall SE excluded.'])
    D[9] = 'ps45_plane_load_20261007 ps45_v2/plane_load.json contribution_sigma_zz_pct (AM–AM 59.56 · 54.63 · 52.37 · 48.46 · 56.91 %; AM–SE 31.01 · 34.66 · 36.22 · 39.14 · 32.95; SE–SE 9.43 · 10.71 · 11.41 · 12.39 · 10.13 — figure F_load redrawn); RB-01/02 (interpretation on hold); ps73 stages §6 (SE σ_VM)'
    notes(s, 'Sources: docs/data/ps45_plane_load_20261007/ps45_v2/plane_load.json (21-plane load-weighted shares; figure F_load redrawn from contribution_sigma_zz_pct: NCM–NCM 59.56 · 54.63 · 52.37 · 48.46 · 56.91 % for 0:10 → 10:0; '
             'all NCM–NCM contacts, not large–large — the 0:10 bed has no large particles); state = relaxed at fixed platen (mean plane load 158–175 MPa, not the 300 MPa peak). '
             'Codex 2026-10-09 verdict: correlation observed; "springback" interpretation and stiffness inversion on hold (RB-01/RB-02); experimental thickness = selected 3-of-n average. '
             'SE stress figure: ps73 stages README §6 (wall SE excluded, linear 0–150 MPa). A 3D view of NCM particle stress / force chains of the same bed can be added by the presenter from the webapp.')

    # ── 10 Progress — registered studies (table) + 194-bed dataset ─────
    s, ft = body_slide('Progress', 'Registered studies and a 194-bed dataset for machine learning',
                       [[('The table summarizes the current status of each registered study.', None)],
                        [('130 + 64 Latin-hypercube beds (composition, particle sizes) → structure and transport descriptors, released as v1.2 (', None), ('conditional use as model descriptors', R), (').', None)]],
                       src='This work — registered studies (left) · dataset v1.2 (right)')
    fh = FZ_BOTTOM - ft - 0.6
    rows = [('Registered study', 'Status'),
            ('Historical eight-origin study', 'Conductivity estimates and inference remain on hold'),
            ('VGCF ordering study (104 arms)', 'Ordering holds across tested grids and offsets (72/72)'),
            ('Absolute-conductivity comparison [12]', 'R = 2.07 (Tier 1); Q = 0.84'),
            ('194-bed dataset v1.2', 'Released for conditional use as model descriptors'),
            ('Updated transport software', 'Seven reviews completed; production remains pending')]
    add_table(s, (0.45, ft, 4.6, 2.85), rows, (2.05, 2.55), size=11, name='Table studies', row_h=(0.34, 0.5, 0.5, 0.5, 0.5, 0.5))
    note_line(s, (0.45, ft + fh + 0.04, 4.6, 0.6), ['R: model / literature conductivity. Q: conductivity retained after PTFE treatment', '(registered model-sensitivity prediction).'])
    place_pic(s, figs / 'F_dataset.png', (5.15, ft, 4.35, fh), name='Fig dataset')
    note_line(s, (5.15, ft + fh + 0.04, 4.35, 0.6), ['(a) geometric tortuosity, 130 + 64 beds; ×: 24 non-percolating beds in the 130-bed set;', 'not transport tortuosity · (b) porosity, 130-bed set'])
    D[10] = 'sdcp_gain_prereg_v2 (inference on hold: CL-33/34/41); phase_a_rep100 README; lee_abs_sigma_e_prereg §5 (R 2.074, Q 0.840); release v1.2 README + CSV (figure F_dataset redrawn: tortuosity_SE_wall, porosity_union_exact_pct vs se_of_solid_vol); codex gen-2 reverify 7'
    notes(s, 'Sources: docs/reviews/sdcp_gain_prereg_v2_20260816.md — the eight-origin (voxel 0.15 µm) study: numerical ratios and the h1 inference are on hold in the claims ledger (CL-33/34/41), so no result is reported; '
             'docs/data/phase_a_rep100_20260930/README.md (104 arms; 72/72; replay 0.00000 %); docs/reviews/lee_abs_sigma_e_prereg_20260923.md §5 (R = 2.074 ± 0.004, Tier 1 "order-of-magnitude", not "within band"; '
             'Q = 0.840 = registered model-sensitivity prediction met, not an experimental hit; comparison with Lee 2025 [12]); docs/data/lhs_release_20261006_v12/README.md (130 × 132 and 64 × 131; physical-target HOLD 59/130 and 62/64 — released as model descriptors, not validated physical targets; '
             'figure F_dataset drawn from the released CSV columns se_of_solid_vol, tortuosity_SE_wall, porosity_union_exact_pct); docs/reviews/codex_review_gen2_network_reverify7_20261012.md (seven software re-verification rounds completed; the updated 194-bed production run, v1.3 and the 20-bed extension are not yet launched or approved). '
             '"Origin" = voxel-grid origin offsets of the same bed, not independent replicate beds. In the 130-bed set: SE percolation 106/130; geometric shortest-path tortuosity 1.3–4.2 (percolating beds; not transport tortuosity); 64-bed set 1.3–1.6. '
             'Loading is not a free design axis (derived from thickness and composition). The collaborator porosity model (OOF R² 0.99) and our five-knob fit (0.98) use unverified and different CV protocols; no comparison is claimed.')

    # ── 11 Mixing 1 — layered campaign ─────────────────────────────────
    s, ft = body_slide('Mixing', 'Drum mixing of a layered NCM–SE bed: no segregation detected',
                       [[('Rotating-drum DEM with a cohesive soft-contact model: ', None), ('ten runs from a layered initial state, eight revolutions', R), ('; Lacey index M on a 16×16×4 bin grid vs a random reference.', None)],
                        [('No NCM–NCM cohesion vs reference cohesion: ΔM within seed-to-seed scatter = the ', None), ('registered "no segregation" category', R), (' (not "no effect"), direction and size unresolved.', None)]],
                       src='This work (left) · figure: presenter — bed view, layered start and after eight revolutions (right)')
    fh = FZ_BOTTOM - ft - 0.6
    place_pic(s, figs / 'F_mix1.png', (0.45, ft, 4.75, fh), name='Fig mixing index')
    note_line(s, (0.45, ft + fh + 0.04, 4.75, 0.6), ['M = 1: fully mixed; error bars = within-run SD; all ten runs reach M ≈ 0.96–1.00, so this bin grid cannot separate the arms;',
                                                     'the registered 1 % overlap condition of the contact model was not met (diagnostic)'])
    D[11] = 'mixer_layered_prereg_20260921 §2 · §5 · §6 (registered no-segregation category: Δ +0.0012, SE 0.0073 at 16×16×4; negative control, floor and flatness checks passed); report_20261021 G7 = mixer_final10_20260930 (figure F_mix1 redrawn)'
    notes(s, 'Sources: docs/reviews/mixer_layered_prereg_20260921.md §2 (verdict line), §5 (report wording), §6 (readout 2026-09-30); docs/data/mixer_final10_20260930/ (ten runs; figure F_mix1 redrawn from the registered final values, report_20261021 G7). '
             'Registered estimand: paired contrast of the final Lacey index between the arm without NCM–NCM cohesion and the reference-cohesion arm, three seeds, 16×16×4 bins: Δ = +0.0012, seed-to-seed SE 0.0073 ⇒ |Δ| ≤ SE ⇒ registered "no segregation" category '
             '("direction and size not resolved" — in this model, this drum, the tested inputs and this bin convention; not "no effect", not "equivalent"). Coarser grids (report only): 12×12×3 Δ −0.0156, SE 0.0058 (ambiguous); 8×8×2 Δ +0.0081, SE 0.0144. '
             'Negative control (no cohesion): M ≥ 0.8 in all three grids; floor check 10/10; flatness 30/30; QC 30/30. '
             'The two cohesion matrices differ in NCM–NCM and NCM–wall adhesion together; the comparison is not a reproduction of real coating or of powder flowability. '
             'Contract qualifier (Codex wording, 2026-09-28): the result is a comparison of mixing indices after eight revolutions in the specified soft-contact model; in the reference state the maximum stored-frame overlap was 4.7–5.8 % particle–particle and 4.5–5.7 % at the wall, '
             'so the registered 1 % condition was not met; the effect of the overlap error on the mixing index has not been verified. Conditions: three phases, 100,000 particles, drum radius 13.1 mm, 75 rpm. '
             'Presenter figure (right, blank): bed view of one run, layered start and after eight revolutions (webapp bed view; display only).')

    # ── 12 Mixing 2 — cohesion exploration ─────────────────────────────
    s, ft = body_slide('Mixing', 'Cohesion exploration: stronger cohesion lowers mixing',
                       [[('Exploratory single-seed runs, two revolutions: NCM–NCM cohesion raised ', None), ('10× and 30×', R), (', or all contact pairs raised ', None), ('10× and 14×', R), (', vs the run without NCM–NCM cohesion.', None)],
                        [('ΔM grows with the Bond number in every bin grid, but each arm is ', None), ('one seed chosen after earlier runs', R), (', and wall adhesion changes with the pair — ', None), ('no verdict', R), (' is drawn.', None)]],
                       src='This work (left) · figure: presenter — bed snapshots, non-cohesive vs cohesive NCM, two revolutions (right)')
    fh = FZ_BOTTOM - ft - 0.6
    place_pic(s, figs / 'F_mix2.png', (0.45, ft, 4.75, fh), name='Fig cohesion exploration')
    note_line(s, (0.45, ft + fh + 0.04, 4.75, 0.6), ['hatched: bin grid failed the volume-coverage check; ΔM mixes the start-state difference', 'with the post-rotation difference; M < 0 is outside the Lacey range'])
    D[12] = 'mixer_highbo_devrot_prelim_20261002 (reference contrast +0.010 / −0.006 / +0.008); mixer_highbo_devbo_prelim_20261003 (×10: +0.83 / +0.93 / +0.99; ×30: +1.88 / +2.43 / +3.16); mixer_highbo_devu_prelim_20261006 (uniform ×10: +0.24 / +0.17 / +0.10; ×14: +0.87 (QC) / +0.48 / +0.27) — figure F_mix2'
    notes(s, 'Sources: docs/data/mixer_highbo_devrot_prelim_20261002/README.md, mixer_highbo_devbo_prelim_20261003/README.md, mixer_highbo_devu_prelim_20261006/README.md (development data; one seed; two revolutions; no pass/fail line; not used for any verdict; not mixed with the confirmatory seeds). '
             'ΔM = M(arm without NCM–NCM cohesion) − M(arm), bins 16×16×4 · 12×12×3 · 8×8×2: reference cohesion +0.010 / −0.006 / +0.008; NCM–NCM cohesion ×10 +0.83 / +0.93 / +0.99; ×30 +1.88 / +2.43 / +3.16; '
             'all pairs ×10 +0.24 / +0.17 / +0.10; all pairs ×14 +0.87 (16×16×4 failed the volume-coverage check) / +0.48 / +0.27. The 16×16×4 cells of the ×10 and ×30 NCM runs also failed that check. '
             'Negative M values of the strongly cohesive arms lie outside the Lacey range [0, 1] (their start-state variance is smaller); ΔM therefore mixes a start-state difference with the post-rotation difference. '
             'The ×10 / ×30 levels were chosen after seeing the first exploratory contrast (post-hoc); NCM–NCM and NCM–wall adhesion change together (joint intervention); the uniform runs scale all nine cohesion matrix elements and include SE stiffening. '
             'No sentence about "real coating effect" or "real powder" is made. Presenter figure (right, blank): bed snapshots of the arm without NCM–NCM cohesion vs a cohesive arm after two revolutions (webapp bed view).')

    # ── 13 Mixing 3 — surface energy and SE–SE adhesion ────────────────
    s, ft = body_slide('Mixing', 'Surface energy and coating: SE–SE adhesion is the open term',
                       [[('Two NCM–SE contacts → NCM–NCM + SE–SE frees ', None), ('ΔW = 2(√γ_A − √γ_S)² ≥ 0', R), ('; the small SE phase keeps like–like contacts best, so uncoated beds can retain SE–SE clusters.', None)],
                        [('Uniformly cohesive runs keep most SE in SE–SE clusters after two revolutions, unlike the run without NCM–NCM cohesion; the ', None), ('confirmatory 18-run block stays on hold', R), ('.', None)]],
                       src='This work — exploratory runs (left) · ratified thesis argument (right)')
    fh = FZ_BOTTOM - ft - 0.55
    place_pic(s, figs / 'F_mix3.png', (0.45, ft, 3.45, fh), name='Fig SE clusters')
    place_pic(s, figs / 'F_mix3_scheme.png', (4.15, ft, 5.3, fh), name='Fig adhesion scheme')
    note_line(s, (0.45, ft + fh + 0.04, 9.0, 0.5), ['Left: one seed, averaged over the second revolution; whiskers = dump-resolution range. Right: work-of-adhesion argument (γ = surface energy,',
                                                    'W = work of adhesion, R* = reduced radius, ρ = density of the smaller particle); the uncoated NCM / SE surface-energy ratio is not known.'])
    D[13] = 'hetero_thesis.md §2 (ΔW = 2(√γ_A − √γ_S)²; SE–SE ≈ 194× AM_P–AM_P at equal γ, arithmetic estimate; γ_A/γ_S < 0.31 unknown); mixer_highbo_devu_prelim_20261006 §3 (SE–SE cluster fraction 0.401 [0.143, 0.542] · 0.954 [0.875, 0.957] · 0.988 [0.981, 0.988]); confirmatory block HOLD (Codex rounds 7–10)'
    notes(s, 'Sources: docs/hetero_thesis.md (ratified 2026-10-05) §2-1 · §2-2 · §2-8: W_ij ≈ 2√(γ_i γ_j); ΔW = W_AA + W_SS − 2 W_AS = 2(√γ_A − √γ_S)² ≥ 0; pair Bond number ∝ W·R*/(ρ g R³) of the smaller particle — '
             'at equal surface energy the SE–SE pair Bond number is about 194× that of the large-NCM pair and SE→NCM adhesion is 1.8× SE–SE (arithmetic estimates); for SE–SE clusters to survive in an uncoated bed the ratio γ_NCM/γ_SE would have to be below about 0.31 — that ratio is not known (open decision). '
             'Predictions P1–P3 of the thesis: uncoated beds keep SE–SE clustering above random; coating removes it; the mixing difference carries into the compacted SE pathway. The present cohesion model does not yet contain that mechanism (NCM–NCM cohesion was varied, not SE–SE). '
             'Left figure: docs/data/mixer_highbo_devu_prelim_20261006/README.md §3 — fraction of SE in SE–SE clusters (bin 1, after two revolutions): 0.401 [0.143, 0.542] without NCM–NCM cohesion; 0.954 [0.875, 0.957] all pairs ×10; 0.988 [0.981, 0.988] all pairs ×14 '
             '(ranges = dump-resolution [certain, possible]; difference only where ranges do not overlap; one seed; development data). '
             'Confirmatory 18-run block: registered, on hold after the Codex rounds (docs/reviews/mixer_highbo_prereg_20260927.md); no "real coating effect" sentence.')

    # ── 14 Plan ────────────────────────────────────────────────────────
    s = two_column('Plan', 'Tests for unresolved model limits',
                   'Grid sensitivity remains unresolved. Diagnostic results exist, while final hypothesis verdicts remain pending.',
                   'Electronic transport',
                   ['Interface resistance: film model implemented, production on hold',
                    'Carbon conductivity: comparison with DC polarization',
                    'NCM conductivity: literature-based input review'],
                   'Compaction and mixing',
                   ['AM load sharing: final hypothesis verdict pending',
                    'Unloading: springback and measured thickness',
                    'Mixing: an adhesion parameter for SE–SE contacts'],
                   bottom=[('Pending: ', BLUE), ('separate launch approvals for the updated 194-bed run and the 20-bed extension.', DARK)])
    D[14] = 'Rint v3.1; fam_platen_prereg §12; dem_unloading_ps45_prereg §13; hetero_thesis §4; gen-2 reverify 7 (launch approval separate)'
    notes(s, 'Sources: docs/reviews/contact_resistance_pipeline_draft_v3_1_20261007.md; docs/reviews/fam_platen_prereg_20260812.md §12; docs/reviews/dem_unloading_ps45_prereg_20261009.md §12–13; docs/hetero_thesis.md §4. '
             'Grid sensitivity: no new grid campaign is promised here. The updated 194-bed production run and the 20-bed extension have separate launch approvals pending (software review complete).')

    # ── 15 Summary ─────────────────────────────────────────────────────
    s = dup_slide(prs, P['summary'])
    set_runs_text(shape(s, 'TextBox 20').text_frame, 'Summary')
    set_body(shape(s, 'Rectangle 4'), 'Current findings and limitations',
             [[('DEM and contact networks: ', BLUE), ('composition trends from one bed per composition', None)],
              [('MPM: ', BLUE), ('plastic deformation, with compaction agreement dependent on initialization and loading protocol', None)],
              [('Voxel transport: ', BLUE), ('ideal-interface calculations with grid dependence and unresolved absolute conductivity', None)],
              [('Mixing: ', BLUE), ('no segregation detected at the registered resolution; cohesion exploration without a verdict', None)],
              [('Final model-validation verdicts remain pending', None)]], line=1.3, after=8)
    place(shape(s, 'Rectangle 4'), (0.61, 1.08, 8.84, 4.2))
    drop(s, 'Conclusion')
    D[15] = 'summary of slides 3–14 (Codex SEM-01/06 wording + mixing line)'
    notes(s, 'Summary; no new numbers. No universal guarantee about registration or verification is claimed (SEM-06).')

    # ── 16 References (16 · 불변) ───────────────────────────────────────
    s = dup_slide(prs, P['refs'])
    set_runs_text(shape(s, 'TextBox 20').text_frame, 'References')
    left = [('h', 'Contact mechanics and DEM'),
            '[1] S. Luding, Granul. Matter 10, 235–246 (2008)',
            '[2] C. Thornton, Z. Ning, Powder Technol. 99, 154–162 (1998)',
            '[3] M. Paulick, M. Morgeneyer, A. Kwade, Powder Technol. 283, 66–76 (2015)',
            '[4] C.J. Coetzee, Powder Technol. 310, 104–142 (2017)',
            '[5] K. Giannis et al., Granul. Matter 23, 17 (2021)',
            '[6] W. Zunker, K. Kamrin, J. Mech. Phys. Solids 183, 105492 (2024)',
            '[7] LIGGGHTS-PUBLIC 3.8.0 (commit 3d5c00f, 2024), github.com/CFDEMproject/LIGGGHTS-PUBLIC',
            '[15] A. Sakuda, A. Hayashi, M. Tatsumisago, Sci. Rep. 3, 2261 (2013)']
    right = [('h', 'Solid-state battery electrodes'),
             '[8] A.M. Bazzoun et al., J. Power Sources 661, 238682 (2026)',
             '[9] M. So et al., J. Power Sources 508, 230344 (2021)',
             '[10] C.A. Varkey et al., Adv. Powder Technol. 37, 105338 (2026)',
             '[11] T. Shi et al., Adv. Energy Mater. 10, 1902881 (2020)',
             '[12] D.J. Lee et al., Nat. Commun. 16, 4200 (2025)',
             '[13] P. Minnmann et al., J. Electrochem. Soc. 168, 040537 (2021)',
             '[14] V. Miß, F. Lange, S. Staubitz, B. Roling, ACS Appl. Mater. Interfaces 18, 32830–32836 (2026)',
             '[16] S. Wang, M. Yan, Y. Li, C. Vinado, J. Yang, J. Power Sources 393, 75–82 (2018)']
    set_items(shape(s, 'Refs left'), left, size=12, line=1.1, before=3)
    set_items(shape(s, 'Refs right'), right, size=12, line=1.1, before=3)
    place(shape(s, 'Refs left'), (0.50, 1.00, 4.45, 5.85))
    place(shape(s, 'Refs right'), (5.05, 1.00, 4.45, 5.85))
    D[16] = 'references from canonical literature cards only (unchanged from v2.1; [15] sakuda2013_sulfide_mechanical_property, [16] wang2018_lco_nmc_electronic_ionic_conductivity_vs_ni)'
    notes(s, 'All entries are taken from the canonical literature cards (original PDFs); no bibliographic field was filled from memory. '
             'The constriction law (Holm/Maxwell) and the MLS-MPM numerical method have no canonical card yet and are cited as the models used, without a bibliographic entry. '
             'Literature crops on slides 4–6 are reproduced from the papers of [8], [1] and [9] (canonical figure crops).')

    # ── 17 Thank you ──────────────────────────────────────────────────
    s = dup_slide(prs, P['thanks'])
    D[17] = 'closing'

    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)[:n_orig]
    for sld in ids:
        rId = sld.get('{%s}id' % NS['r'])
        prs.part.drop_rel(rId)
        sldIdLst.remove(sld)
    prs.save(out)
    Path(str(out) + '.sources.json').write_text(json.dumps(D, ensure_ascii=False, indent=1), encoding='utf-8')
    return D


# ───────────────────────────── 검사 ─────────────────────────────

FOOT = {'HANYANG UNIVERSITY', 'Battery Materials Lab.'}
#: Codex 10-13 가 지적한 문구 (v2.1 목록 그대로) — 어느 장 본문에도 남으면 안 된다
FORBIDDEN = ['Large–large', 'large–large', 'follow no power law', 'changes sign', 'never to one another', 'cross-validation',
             'was retracted', 'Every production number', 'nothing below has a result', '30.7', 'Pressed honestly', 'contact-free branch',
             'Hertz constriction', 'single-crystal LPSCl', 'irreducible', 'closing the three problems', 'gain follows the represented',
             'fixed-platen hold', 'σy 0.05–0.30', 'Geometric tortuosity vs SE fraction (130 beds)', 'real coating effect', 'Figure pending']
INTERNAL = r'\b(?:CL-\d+|SELF-\d+|SEM-\d+|SEMV2-\d+|G2RR\d*-\d+|RGL[R0-9]*-\d+|UL-\d+|DEMP-\d+|RB-\d+|MV-\d+|PS45E-\d+|HBR\d*-\d+|stoic-knuth|pqwtv8|kgy|ibb|uma|WSL)\b'
#: 믹서 팔 코드 (1저자 10-13 — 덱 어디에도 쓰지 않는다 · notes 포함)
ARM_CODES = r'\b(?:L0|LA|LB[123]|LC|LH|LHx\d*|LU\d*|E0)\b'
NOT_COUNTED = {'TextBox 69', 'Credit', 'Lit tag', 'TextBox 20'}   # 출처 줄 · 크롭 출처 · 꼬리표 · 절 표지 = 작은 글씨 · 단어 수에 안 센다
BLANK_SLIDES = (5, 7, 11, 12)                                       # 빈 자리 (1저자 캡처) 가 있는 장 — 출처 줄에 'figure: presenter' 필수
ZONE_NAMES = ('Lit tag', 'Caption', 'Credit')


def deck_text(path):
    prs = Presentation(path)
    out = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if sh.has_table:
                out.extend(c.text for r in sh.table.rows for c in r.cells)
        if s.has_notes_slide:
            out.append(s.notes_slide.notes_text_frame.text)
    return '\n'.join(out)


def slide_texts(prs, with_tables=True):
    out = []
    for s in prs.slides:
        parts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        if with_tables:
            parts += [c.text for sh in s.shapes if sh.has_table for r in sh.table.rows for c in r.cells]
        out.append('\n'.join(parts))
    return out


def slide_stats(prs):
    """장별 (본문 단어 수 (출처 줄 · 크롭 출처 · 꼬리표 · 표 제외) · bullet pt 집합 · 해석 줄 pt 최소)."""
    out = []
    for s in prs.slides:
        words, bullet_pt, cap_pt = 0, set(), []
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if not t or t in FOOT or re.fullmatch(r'\d+', t) or sh.name in NOT_COUNTED:
                continue
            words += len(t.split())
            if sh.name == 'Rectangle 4':
                for p in sh.text_frame.paragraphs[1:]:
                    for r in p.runs:
                        if r.font.size:
                            bullet_pt.add(r.font.size.pt)
            if sh.name == 'Caption':
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            cap_pt.append(r.font.size.pt)
        out.append((words, bullet_pt, min(cap_pt) if cap_pt else None))
    return out


def overlap_check(prs):
    """본문 (Rectangle 4) 바닥 추정값 ≤ 그림 띠 (그림 · 표 · 흐름도 · 꼬리표 · 해석 줄) 윗선 — 장별 (추정 바닥, 윗선)."""
    bad = []
    for i, s in enumerate(prs.slides, 1):
        body = [sh for sh in s.shapes if sh.name == 'Rectangle 4']
        if not body or i == 1:
            continue
        paras = [p.text for p in body[0].text_frame.paragraphs if p.text.strip()]
        if len(paras) < 2:
            continue
        head, bullets = paras[0], paras[1:]
        est = est_body_bottom(head, bullets)
        tops = [sh.top / 914400 for sh in s.shapes
                if sh.shape_type == 13 or sh.has_table or sh.name.startswith('Flow') or sh.name in ZONE_NAMES]
        if not tops:
            continue
        if len(head) > CHARS_HEAD or est_lines(head, HEAD_PT, HEAD_W) != 1 or est > min(tops) + 0.02:
            bad.append((i, round(est, 2), round(min(tops), 2), len(head)))
    return bad


def selftest(out, repo, figs=None, word_cap=100, wide_cap=120, wide=(6, 7, 8, 10, 11, 12, 13, 14)):
    txt = deck_text(out)
    ban = json.loads((Path(repo) / 'docs/reviews/claims.json').read_text(encoding='utf-8')).get('quotation_ban', [])
    pats = [q.get('pattern') if isinstance(q, dict) else q for q in ban]
    hits = [p for p in pats if p and p in txt]
    assert not hits, f'banned strings in deck: {hits}'
    prs = Presentation(out)
    texts = slide_texts(prs)
    slide_txt = '\n'.join(texts)
    leak = re.findall(INTERNAL, slide_txt)
    assert not leak, f'internal identifiers on slides: {leak}'
    codes = re.findall(ARM_CODES, txt)
    assert not codes, f'mixer arm codes in deck (slides or notes): {codes}'
    bad = [f for f in FORBIDDEN if f in slide_txt]
    assert not bad, f'phrases flagged by the Codex review still present: {bad}'
    n = len(prs.slides)
    assert n == N_SLIDES, n
    angle = [i for i, t in enumerate(texts, 1) if re.search(r'(^|\n)< .* >($|\n)', t)]
    assert not angle, f'angle-bracket captions (old grammar) on slides: {angle}'
    for i in BLANK_SLIDES:
        src = [sh.text_frame.text for sh in prs.slides[i - 1].shapes if sh.name == 'TextBox 69']
        assert src and 'figure: presenter' in src[0], (i, src)
    long_src = [(i, len(sh.text_frame.text)) for i, s in enumerate(prs.slides, 1) for sh in s.shapes if sh.name == 'TextBox 69' and len(sh.text_frame.text) > 110]
    assert not long_src, f'source lines over 110 characters: {long_src}'
    ov = overlap_check(prs)
    assert not ov, f'body text would reach the figure zone (slide, est. body bottom, zone top, head chars): {ov}'
    caps10 = '\n'.join(sh.text_frame.text for sh in prs.slides[9].shapes if sh.name == 'Caption')
    assert '130 + 64' in caps10 and 'tortuosity' in caps10 and '130-bed set' in caps10 and 'porosity' in caps10, caps10
    assert 'registered "no segregation" category' in texts[10] and 'not "no effect"' in texts[10], texts[10][:300]
    for i in (11, 12, 13):
        assert 'cohesi' in texts[i - 1], i
    stats = slide_stats(prs)
    heavy = [(i, w) for i, (w, _, _) in enumerate(stats, 1) if 2 <= i <= 15 and w > (wide_cap if i in wide else word_cap)]
    assert not heavy, f'slides over {word_cap}/{wide_cap} words (small print and tables excluded): {heavy}'
    small = [(i, bp) for i, (_, bp, _) in enumerate(stats, 1) if bp and min(bp) < BODY_PT]
    assert not small, f'body bullets below {BODY_PT} pt: {small}'
    cap_small = [(i, c) for i, (_, _, c) in enumerate(stats, 1) if c is not None and c < CAP_PT]
    assert not cap_small, f'note lines below {CAP_PT} pt: {cap_small}'
    pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)
    assert pics >= 12, pics
    if figs:
        svg = (Path(figs) / 'F_limits.svg').read_text(encoding='utf-8')
        labels = [k for k in ('0:10', '3:7', '5:5', '7:3', '10:0') if k in svg]
        assert len(labels) == 5, f'F_limits composition labels present in svg: {labels}'
    print(f'selftest ok — {n} slides · banned 0 · internal 0 · arm codes 0 · flagged phrases 0 · angle captions 0 · overlap 0 · blanks {len(BLANK_SLIDES)} · '
          f'max words {max(w for w, _, _ in stats[1:15])} · bullets ≥ {BODY_PT} pt · note lines ≥ {CAP_PT} pt · pictures {pics}')


def dump_text(deck, out):
    prs = Presentation(deck)
    lines = ['# 연구세미나 덱 v3 "보는 판" — 장별 글 전문 (생성기 출력에서 추출 · 표 · 출처 줄 · notes 포함)', '']
    for i, s in enumerate(prs.slides, 1):
        lines.append(f'## p{i:02d}')
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                for p in sh.text_frame.paragraphs:
                    if p.text.strip():
                        lines.append(f'- [{sh.name}] {p.text}')
            if sh.has_table:
                for r in sh.table.rows:
                    lines.append('- [' + sh.name + '] ' + ' | '.join(c.text for c in r.cells))
            if sh.shape_type == 13:
                lines.append(f'- [picture] {sh.name}')
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            lines.append('- notes: ' + s.notes_slide.notes_text_frame.text)
        lines.append('')
    Path(out).write_text('\n'.join(lines), encoding='utf-8')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--template')
    ap.add_argument('--repo')
    ap.add_argument('--figs')
    ap.add_argument('--lit', help='문헌 크롭 폴더 (없으면 <out 폴더>/lit · 비어 있으면 litdb 정본 브랜치에서 받는다)')
    ap.add_argument('--out')
    ap.add_argument('--selftest-only', help='기존 덱 검사만 (생성 없음)')
    ap.add_argument('--dump', help='덱의 장별 글 전문을 --out (md) 로')
    a = ap.parse_args()
    if a.dump:
        dump_text(a.dump, a.out)
        return
    if a.selftest_only:
        selftest(a.selftest_only, a.repo, a.figs)
        return
    build(a.template, a.repo, a.figs, a.lit, a.out)
    selftest(a.out, a.repo, a.figs)
    print('written', a.out)


if __name__ == '__main__':
    main()
