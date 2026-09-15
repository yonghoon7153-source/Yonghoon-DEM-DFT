from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import contextlib
import io
import json
import subprocess
import sys
import unittest

PACKAGE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(PACKAGE))
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.pipeline import Pipeline
from config import Config,DataError,EvaluationError,config_dict,VERSION
from demo import make_demo
from extract_capacity import (CYCLE_COLUMNS,build_features,convert_mpt,find_cv_start,
                              load_raw,read_cycle,validate_frame)
from export_for_origin import main,run
from predict import predict_frame
from ridge_capacity import classification_metrics,nested_evaluate,select_model

def small_config(**changes):
    values=dict(frequency_grid_hz=[.1,1.,10.,1000.,10000.,100000.],
                feature_sets=['single_1e3'],ridge_alphas=[.1,10.],logistic_cs=[.1,1.],
                inner_splits=2,use_ocv=False)
    values.update(changes)
    return Config(**values).validate()

ROWS=[[0,0,1,10,0,3.7,0], [1,0,1,20,.001,3.8,.01],
      [2,1,1,20,.015,4.0,.01], [3,2,1,20,.018,4.2,.005],
      [4,3,1,20,.020,4.2,.002], [5,0,1,30,.025,4.1,-.01]]

def cycle_file(path,units=True,rows=None,metadata_lines=16):
    text='metadata\n'*metadata_lines+'\t'.join(CYCLE_COLUMNS)+'\n'
    if units:
        text+='s\ts\tindex\tindex\tAh\tV\tA\n'
    text+=''.join('\t'.join(map(str,row))+'\n' for row in (rows if rows is not None else ROWS))
    path.write_text(text,encoding='ascii')
    return path

def eis_file(path,cfg):
    rows=['freq/Hz\tRe(Z)/Ohm\t-Im(Z)/Ohm\t|Z|/Ohm\tPhase(Z)/deg\t<Ewe>/V']
    for f in cfg.frequency_grid_hz:
        z=10+20/(1+1j*f/300)
        rows.append(f'{f}\t{z.real}\t{-z.imag}\t{abs(z)}\t{np.angle(z,deg=True)}\t3.7')
    path.write_text('EC-Lab ASCII FILE\n'+'\n'.join(rows)+'\n',encoding='ascii')
    return path

class TemporaryCase(unittest.TestCase):
    def setUp(self):
        self.parent=Path.cwd().resolve()
        self.tmp=TemporaryDirectory(prefix='ml_validation_',dir=self.parent)
        self.folder=Path(self.tmp.name).resolve()
        self.assertTrue(self.folder.is_relative_to(self.parent))

    def tearDown(self):
        self.assertTrue(Path(self.tmp.name).resolve().is_relative_to(self.parent))
        self.tmp.cleanup()

class InputTests(TemporaryCase):
    def test_config_rejects_ambiguous_booleans_and_fractional_splits(self):
        for setting in [{'use_ocv':'false'},{'inner_splits':2.5},{'target_step_i':-1}]:
            with self.subTest(setting=setting),self.assertRaises(DataError):
                small_config(**setting)

    def test_invalid_optional_target_is_not_silently_dropped(self):
        cfg=small_config(); data=make_demo(cfg,8)
        data['cc_frac']=data['cc_frac'].astype(object); data.loc[0,'cc_frac']='corrupted'
        with self.assertRaisesRegex(DataError,'nonnumeric'):
            validate_frame(data,cfg,[])

    def test_malformed_feature_name_rejected(self):
        cfg=small_config(); data=make_demo(cfg,8).rename(columns={'mag@1000Hz':'mag@1kHz'})
        with self.assertRaisesRegex(DataError,'Invalid EIS feature name'):
            validate_frame(data,cfg,[])

    def test_failed_raw_input_keeps_file_provenance(self):
        cfg=small_config(); cycle_dir=self.folder/'cycle'; eis_dir=self.folder/'eis'
        cycle_dir.mkdir(); eis_dir.mkdir()
        cycle_file(cycle_dir/'cell41.txt'); (eis_dir/'cell41.mpt').write_text('invalid',encoding='ascii')
        events=[]
        with self.assertRaisesRegex(DataError,'Cell 41'):
            load_raw(cycle_dir,eis_dir,cfg,events)
        self.assertEqual(events[-1]['status'],'failed')
        self.assertEqual(len(events[-1]['eis_sha256']),64)
        self.assertEqual(len(events[-1]['cycle_sha256']),64)

    def test_units_row_and_no_units_keep_same_observations(self):
        a,_=read_cycle(cycle_file(self.folder/'cell41_a.txt',units=True,metadata_lines=3))
        b,_=read_cycle(cycle_file(self.folder/'cell41_b.txt',units=False,metadata_lines=27))
        self.assertEqual(a,b)
        self.assertAlmostEqual(a['cap'],.02)
        self.assertAlmostEqual(a['cap_cc'],.018)
        self.assertAlmostEqual(a['cap_cv'],.002)

    def test_cell_only_name_and_optional_pressure(self):
        result,_=read_cycle(cycle_file(self.folder/'cell41_cycle.txt'))
        self.assertEqual(result['cell_number'],41)
        self.assertIsNone(result['pressure_mpa'])

    def test_missing_target_segment_rejected(self):
        with self.assertRaisesRegex(DataError,'segment missing'):
            read_cycle(cycle_file(self.folder/'cell41.txt',rows=ROWS[:1]))

    def test_time_reversal_rejected(self):
        with self.assertRaisesRegex(DataError,'time is not monotonic'):
            read_cycle(cycle_file(self.folder/'cell41.txt',rows=[ROWS[0],ROWS[1],ROWS[3],ROWS[2],ROWS[4]]))

    def test_decreasing_capacity_rejected(self):
        rows=[r.copy() for r in ROWS]
        rows[3][4]=.01
        with self.assertRaisesRegex(DataError,'monotonic'):
            read_cycle(cycle_file(self.folder/'cell41.txt',rows=rows))

    def test_zero_initial_current_rejected(self):
        with self.assertRaisesRegex(DataError,'nonzero'):
            find_cv_start(pd.DataFrame({'current':[0,1,.5]}))

    def test_wrong_current_sign_rejected(self):
        rows=[r.copy() for r in ROWS]
        rows[2][6]=-.01
        with self.assertRaisesRegex(DataError,'current sign'):
            read_cycle(cycle_file(self.folder/'cell41.txt',rows=rows))

    def test_target_raw_step_supported(self):
        result,detail=read_cycle(cycle_file(self.folder/'cell41.txt'),small_config(target_raw_step=20,target_step_i=0))
        self.assertEqual(detail['raw_step'],20)
        self.assertAlmostEqual(result['cap'],.02)

    def test_nan_capacity_never_becomes_normal(self):
        cfg=small_config(); data=make_demo(cfg,8); data.loc[0,'cap']=np.nan; events=[]
        with self.assertRaisesRegex(DataError,'Invalid/missing capacity'):
            validate_frame(data,cfg,events)
        self.assertEqual(events[-1]['status'],'unlabeled_invalid')
        self.assertNotIn('true_label',data)

    def test_inf_features_rejected(self):
        cfg=small_config(); data=make_demo(cfg,8); data.loc[0,'mag@1000Hz']=np.inf
        with self.assertRaisesRegex(DataError,'NaN/Inf'):
            validate_frame(data,cfg,[])

    def test_duplicate_cell_rejected(self):
        cfg=small_config(); data=make_demo(cfg,8); data.loc[1,'cell_number']=data.loc[0,'cell_number']
        with self.assertRaisesRegex(DataError,'Duplicate cell'):
            validate_frame(data,cfg,[])

    def test_missing_ocv_does_not_silently_shrink_cohort(self):
        cfg=small_config(use_ocv=True); data=make_demo(cfg,8); data.loc[0,'ocv']=np.nan
        with self.assertRaisesRegex(DataError,'NaN/Inf'):
            validate_frame(data,cfg,[])
        cfg.use_ocv=False
        clean,features=validate_frame(data,cfg,[])
        self.assertEqual(len(clean),8)
        self.assertNotIn('ocv',features)

    def test_excluded_cell_does_not_define_grid(self):
        cfg=small_config(exclude_cells={'60':'predeclared measurement QC'})
        cycle_dir=self.folder/'cycle'; eis_dir=self.folder/'eis'; cycle_dir.mkdir(); eis_dir.mkdir()
        cycle_file(cycle_dir/'cell41.txt'); eis_file(eis_dir/'cell41.mpt',cfg)
        # Excluded malformed files are recorded and never influence the fixed grid.
        (cycle_dir/'cell60.txt').write_text('invalid',encoding='ascii')
        (eis_dir/'cell60.mpt').write_text('invalid',encoding='ascii')
        events=[]; frame=load_raw(cycle_dir,eis_dir,cfg,events)
        self.assertEqual(frame['cell_number'].tolist(),[41])
        self.assertIn('mag@100000Hz',frame)
        self.assertTrue(any(e.get('cell_number')==60 and e['status']=='excluded' for e in events))

    def test_duplicate_eis_requires_manifest(self):
        cfg=small_config(); cycle_dir=self.folder/'cycle'; eis_dir=self.folder/'eis'; cycle_dir.mkdir(); eis_dir.mkdir()
        cycle_file(cycle_dir/'cell41.txt'); eis_file(eis_dir/'cell41_a.mpt',cfg); eis_file(eis_dir/'cell41_b.mpt',cfg)
        with self.assertRaisesRegex(DataError,'Duplicate cell'):
            load_raw(cycle_dir,eis_dir,cfg,[])
        manifest=self.folder/'selection.csv'
        pd.DataFrame([{'cell_number':41,'cycle_file':'cell41.txt','eis_file':'cell41_b.mpt','batch_id':'B1'}]).to_csv(manifest,index=False)
        frame=load_raw(cycle_dir,eis_dir,cfg,[],manifest)
        self.assertEqual(frame['batch_id'].tolist(),['B1'])

    def test_empty_inputs_raise_clear_error(self):
        with self.assertRaisesRegex(DataError,'No .mpt'):
            load_raw(None,self.folder,small_config(),[])

    def test_mpt_imaginary_sign(self):
        cfg=small_config(use_ocv=True)
        df,_=convert_mpt(eis_file(self.folder/'cell41.mpt',cfg),cfg)
        self.assertTrue((df['im']<0).all())

    def test_no_implicit_extrapolation(self):
        cfg=small_config(frequency_grid_hz=[1.,10.,1000.])
        spec=pd.DataFrame({'f':[1.04,990.],'re':[4.,9.],'im':[0.,0.]})
        with self.assertRaisesRegex(DataError,'exceed measured support'):
            build_features(spec,cfg)
        cfg.boundary_policy='clamp'
        result,info=build_features(spec,cfg)
        self.assertEqual(result['re@1Hz'],4.)
        self.assertEqual(info['clamped_frequencies_hz'],[1.,1000.])

    def test_duplicate_frequency_policy(self):
        cfg=small_config(frequency_grid_hz=[1.,10.,100.])
        df=pd.DataFrame({'f':[1.,10.,10.,100.],'re':[1.,2.,20.,100.],'im':[0.,0.,0.,0.]})
        with self.assertRaisesRegex(DataError,'Duplicate EIS'):
            build_features(df,cfg)
        cfg.duplicate_frequency_policy='mean'
        result,_=build_features(df,cfg)
        self.assertEqual(result['re@10Hz'],11.)

    def test_complex_components_remain_consistent(self):
        cfg=small_config(frequency_grid_hz=[10.])
        df=pd.DataFrame({'f':[1.,100.],'re':[1.,0.],'im':[0.,-1.]})
        result,_=build_features(df,cfg)
        self.assertAlmostEqual(result['mag@10Hz'],np.hypot(result['re@10Hz'],result['im@10Hz']))
        self.assertAlmostEqual(result['phase@10Hz'],-45.)

class EvaluationTests(unittest.TestCase):
    def test_outer_target_cannot_change_its_own_training_selection(self):
        cfg=small_config(); df=make_demo(cfg,8); clean,features=validate_frame(df,cfg,[])
        X=clean[features]; y=clean['cap'].to_numpy(); ids=clean['cell_number'].to_numpy()
        pred,a=nested_evaluate(X,y,ids,cfg)
        changed=y.copy(); changed[0]=10.
        pred2,b=nested_evaluate(X,changed,ids,cfg)
        self.assertEqual(a[0]['selected'],b[0]['selected'])
        self.assertEqual(a[0]['best_inner_score'],b[0]['best_inner_score'])
        self.assertAlmostEqual(pred[0],pred2[0])

    def test_scaler_and_inner_folds_use_outer_training_only(self):
        cfg=small_config(); df=make_demo(cfg,8); clean,features=validate_frame(df,cfg,[])
        X=clean[features]; ids=clean['cell_number'].to_numpy()
        _,folds=nested_evaluate(X,clean['cap'].to_numpy(),ids,cfg)
        for fold in folds:
            train=set(fold['training_cells']); test=set(fold['test_cells'])
            self.assertFalse(train&test)
            columns=fold['selected']['features']
            expected=X.loc[clean['cell_number'].isin(train),columns].mean().to_numpy()
            np.testing.assert_allclose(fold['scaler_mean'],expected)
            for inner in fold['inner_folds']:
                a=set(inner['train_cells']); b=set(inner['validation_cells'])
                self.assertFalse(a&b); self.assertEqual(a|b,train); self.assertFalse((a|b)&test)

    def test_nested_group_separation(self):
        cfg=small_config(outer_cv='leave_one_group_out'); df=make_demo(cfg,12); clean,features=validate_frame(df,cfg,[])
        _,folds=nested_evaluate(clean[features],clean['cap'].to_numpy(),clean['cell_number'].to_numpy(),cfg,groups=clean['batch_id'].to_numpy())
        self.assertEqual(len(folds),3)
        for fold in folds:
            for inner in fold['inner_folds']:
                a=set(inner['train_groups']); b=set(inner['validation_groups']); test=set(fold['test_groups'])
                self.assertFalse(a&b); self.assertFalse((a|b)&test)

    def test_logistic_uses_finite_log_loss_not_one_cell_auc(self):
        cfg=small_config(); df=make_demo(cfg,8); clean,features=validate_frame(df,cfg,[])
        y=(clean['cap'].to_numpy()<=cfg.threshold_ah).astype(int)
        pred,folds=nested_evaluate(clean[features],y,clean['cell_number'].to_numpy(),cfg,kind='classification')
        self.assertTrue(np.isfinite(pred).all())
        self.assertTrue(all(f['inner_metric']=='negative_log_loss' and np.isfinite(f['best_inner_score']) for f in folds))

    def test_one_class_auc_is_unavailable(self):
        result=classification_metrics([0,0],[0,0],[.1,.2],probabilities=True)
        self.assertIsNone(result['AUC_pooled_OOF'])
        self.assertIsNone(result['sensitivity'])

    def test_no_group_breaking_fallback_for_unbalanced_groups(self):
        cfg=small_config(outer_cv='leave_one_group_out'); df=make_demo(cfg,12); clean,features=validate_frame(df,cfg,[])
        groups=clean['batch_id'].to_numpy(); y=(groups=='demo_batch_1').astype(int)
        with self.assertRaises(EvaluationError):
            nested_evaluate(clean[features],y,clean['cell_number'].to_numpy(),cfg,kind='classification',groups=groups)

class IntegrationTests(TemporaryCase):
    def test_cc_missing_rows_keep_cell_alignment_and_workbook(self):
        cfg=small_config(); df=make_demo(cfg,8); df.loc[0,'cc_frac']=np.nan
        out=self.folder/'run'; out.mkdir()
        result=run(df,cfg,out,[],{'data_kind':'synthetic_demo'})
        self.assertEqual(result['status'],'completed')
        actual=pd.read_csv(out/'oof_predictions.csv').set_index('cell_number')
        self.assertTrue(pd.isna(actual.loc[41,'pred_cc_frac']))
        self.assertEqual(int(actual['pred_cc_frac'].notna().sum()),7)
        np.testing.assert_allclose(actual.loc[42:,'measured_cc_frac'],df.set_index('cell_number').loc[42:,'cc_frac'])
        wb=load_workbook(out/'results_for_origin.xlsx',read_only=True,data_only=True)
        self.assertEqual(len(wb.sheetnames),11)
        self.assertEqual(wb['07_clf_cellwise'].max_row,9)
        wb.close()

    def test_prediction_uses_saved_fit_without_refitting(self):
        cfg=small_config(); df=make_demo(cfg,8); clean,features=validate_frame(df,cfg,[])
        X=clean[features]; model,_=select_model(X,clean['cap'].to_numpy(),clean['cell_number'].to_numpy(),cfg,'regression')
        bundle={'pipeline':model,'config':config_dict(cfg),'version':VERSION,'input_features':features,
                'training_cells':clean['cell_number'].tolist(),'training_min':X.min().to_dict(),'training_max':X.max().to_dict()}
        new=clean.drop(columns=['cap','cc_frac']).copy(); new['cell_number']+=100
        with patch.object(Pipeline,'fit',side_effect=AssertionError('Prediction must never fit')):
            result=predict_frame(new,{'capacity':bundle})
        self.assertEqual(len(result),8)
        self.assertFalse(result['capacity_seen_in_training'].any())

    def test_failed_run_writes_audit_and_nonzero_exit(self):
        cfg=small_config(); df=make_demo(cfg,8); df.loc[0,'cap']=np.nan
        csv=self.folder/'bad.csv'; df.to_csv(csv,index=False)
        config=self.folder/'config.json'; config.write_text(json.dumps(config_dict(cfg)),encoding='utf-8')
        out=self.folder/'run'
        with contextlib.redirect_stderr(io.StringIO()):
            code=main(['--features',str(csv),'--config',str(config),'--output',str(out)])
        self.assertEqual(code,2)
        self.assertEqual(json.loads((out/'run.json').read_text())['status'],'failed')
        events=json.loads((out/'data_audit.json').read_text())
        self.assertTrue(any(e['status']=='unlabeled_invalid' for e in events))

    def test_output_overwrite_refused(self):
        out=self.folder/'existing'; out.mkdir(); (out/'sentinel.txt').write_text('keep')
        with contextlib.redirect_stderr(io.StringIO()),self.assertRaises(SystemExit) as e:
            main(['--demo','--output',str(out)])
        self.assertEqual(e.exception.code,2)
        self.assertEqual((out/'sentinel.txt').read_text(),'keep')

    def test_imports_do_not_create_analysis_outputs(self):
        script=f'import sys; sys.path.insert(0,{str(PACKAGE)!r}); import extract_capacity, ridge_capacity, export_for_origin, feature_band_test, predict'
        result=subprocess.run([sys.executable,'-B','-c',script],cwd=self.folder,capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertEqual(list(self.folder.iterdir()),[])

if __name__=='__main__':
    unittest.main(verbosity=2)
