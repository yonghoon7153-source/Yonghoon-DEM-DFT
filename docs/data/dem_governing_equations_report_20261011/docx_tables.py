"""pandoc docx 표에 가로줄 (위 · 아래 · 행 사이) · 머리 행 음영 · 9 pt 를 입힌다."""
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
d = Document(sys.argv[1])
for tbl in d.tables:
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for side, val in (('top', 'single'), ('bottom', 'single'), ('insideH', 'single'), ('left', 'nil'), ('right', 'nil'), ('insideV', 'nil')):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), val)
        if val != 'nil':
            e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '888888')
        b.append(e)
    tblPr.append(b)
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW'); tblPr.append(w)
    w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), '5000')
    for ri, row in enumerate(tbl.rows):
        for cell in row.cells:
            if ri == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'EFEFEF')
                tcPr.append(sh)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
d.save(sys.argv[2])
print('tables', len(d.tables))
