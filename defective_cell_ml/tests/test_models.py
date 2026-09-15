"""Model selection, chance-level reporting and export bookkeeping (review fixes F05-F74)."""
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import contextlib
import io
import json
import platform
import sys
import unittest
import warnings

PACKAGE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(PACKAGE))
import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.exceptions import ConvergenceWarning
from sklearn.linear_model import LogisticRegression
from config import Config,EvaluationError,config_dict
from demo import make_demo
import export_for_origin
import ridge_capacity
from export_for_origin import main,run,versions,write_workbook
from extract_capacity import validate_frame
from ridge_capacity import (classification_metrics,select_model,selection_counts,selection_diagnostics,
                            skipped_candidates,uses_eis)

def small_config(**changes):
    values=dict(frequency_grid_hz=[.1,1.,10.,1000.,10000.,100000.],
                feature_sets=['single_1e3'],ridge_alphas=[.1,10.],logistic_cs=[.1,1.],
                inner_splits=2,use_ocv=False)
    values.update(changes)
    return Config(**values).validate()

def warn_on_fit(original):
    """Make every logistic fit emit a ConvergenceWarning, as an unscaled real cohort would."""
    def fit(self,X,y=None,**kwargs):
        warnings.warn('lbfgs failed to converge (status=1)',ConvergenceWarning)
        return original(self,X,y,**kwargs)
    return fit

class TemporaryCase(unittest.TestCase):
    """Outputs go to the system temporary directory; nothing is written inside the package."""
    def setUp(self):
        self.tmp=TemporaryDirectory(prefix='ml_models_')
        self.folder=Path(self.tmp.name).resolve()

    def tearDown(self):
        self.tmp.cleanup()

    def execute(self,cfg,frame=None,folder='run'):
        out=self.folder/folder
        out.mkdir()
        result=run(make_demo(cfg,8) if frame is None else frame,cfg,out,[],{'data_kind':'synthetic_demo'})
        return result,out

class SelectionTests(unittest.TestCase):
    def test_nonconverged_inner_fit_is_an_evaluation_error_not_a_crash(self):
        # F16: the warning stays promoted, but as EvaluationError so only this target is lost.
        cfg=small_config(); clean,features=validate_frame(make_demo(cfg,8),cfg,[])
        y=(clean['cap'].to_numpy()<=cfg.threshold_ah).astype(int)
        with patch.object(LogisticRegression,'fit',warn_on_fit(LogisticRegression.fit)):
            with self.assertRaisesRegex(EvaluationError,'did not converge'):
                select_model(clean[features],y,clean['cell_number'].to_numpy(),cfg,'classification')

    def test_inner_metric_and_folds_come_from_the_search_object(self):
        # F29: folds.json must report what GridSearchCV received, not a restated literal.
        cfg=small_config(); clean,features=validate_frame(make_demo(cfg,8),cfg,[])
        ids=clean['cell_number'].to_numpy(); captured={}
        original=ridge_capacity.GridSearchCV
        def spy(*args,**kwargs):
            captured.update(cv=kwargs['cv'],scoring=kwargs['scoring'])
            return original(*args,**kwargs)
        with patch.object(ridge_capacity,'GridSearchCV',spy):
            _,trace=select_model(clean[features],clean['cap'].to_numpy(),ids,cfg,'regression')
        self.assertEqual(len(captured['cv']),len(trace['inner_folds']))
        for (train,validation),recorded in zip(captured['cv'],trace['inner_folds']):
            self.assertEqual([int(ids[i]) for i in train],recorded['train_cells'])
            self.assertEqual([int(ids[i]) for i in validation],recorded['validation_cells'])
        self.assertEqual(trace['inner_metric'],ridge_capacity.INNER_METRICS[captured['scoring']])

    def test_inner_metric_follows_a_changed_scoring(self):
        # F29: a different scoring must change the recorded name instead of being reported as log loss.
        cfg=small_config(); clean,features=validate_frame(make_demo(cfg,8),cfg,[])
        y=(clean['cap'].to_numpy()<=cfg.threshold_ah).astype(int)
        original=ridge_capacity.GridSearchCV
        def forced(*args,**kwargs):
            return original(*args,**{**kwargs,'scoring':'accuracy'})
        with patch.dict(ridge_capacity.INNER_METRICS,{'accuracy':'accuracy_probe'}), \
             patch.object(ridge_capacity,'GridSearchCV',forced):
            _,trace=select_model(clean[features],y,clean['cell_number'].to_numpy(),cfg,'classification')
        self.assertEqual(trace['inner_metric'],'accuracy_probe')

    def test_missing_columns_name_the_skipped_candidates(self):
        # F10: single_1e3 and the |Z|@1 kHz baseline disappear silently without this bookkeeping.
        cfg=small_config(frequency_grid_hz=[.1,1.,10.,10000.,100000.],feature_sets=['single_1e3','band_low'])
        clean,features=validate_frame(make_demo(cfg,8),cfg,[])
        self.assertEqual(skipped_candidates(clean[features],cfg),
                         {'feature_sets':['single_1e3'],'baselines':['mag@1000Hz']})
        full=small_config()
        clean,features=validate_frame(make_demo(full,8),full,[])
        self.assertEqual(skipped_candidates(clean[features],full),{'feature_sets':[],'baselines':[]})

    def test_prior_baseline_anchors_log_loss_and_brier(self):
        # F31: class-weighted probabilities are only readable next to the constant prior.
        result=classification_metrics([0,0,0,1],[0,0,0,1],[.1,.2,.3,.9],probabilities=True)
        prior=.25
        self.assertAlmostEqual(result['log_loss_prior_baseline'],-(3*np.log(1-prior)+np.log(prior))/4)
        self.assertAlmostEqual(result['brier_prior_baseline'],(3*prior**2+(1-prior)**2)/4)
        self.assertLess(result['log_loss'],result['log_loss_prior_baseline'])

    def test_degenerate_and_boundary_selection_are_reported(self):
        # F24: a constant 0.5 model and a grid-edge C must not look like a fitted classifier.
        cfg=small_config(logistic_cs=[.1,1.,10.])
        traces=[{'selected':{'C':.1}} for _ in range(4)]
        notes=selection_diagnostics(np.full(6,.5),traces,cfg)
        self.assertTrue(any(note.startswith('degenerate') for note in notes))
        self.assertTrue(any(note.startswith('boundary') for note in notes))
        self.assertEqual(selection_diagnostics(np.array([.1,.9,.2,.8]),[{'selected':{'C':1.}}],cfg),[])

    def test_selection_counts_and_eis_usage(self):
        # F17: the summary must be able to say which estimator won and whether EIS was used at all.
        traces=[{'selected':{'estimator':'Ridge','features':['mag@1000Hz'],'alpha':1.,'C':None}},
                {'selected':{'estimator':'Ridge','features':['mag@1000Hz'],'alpha':1.,'C':None}},
                {'selected':{'estimator':'LinearRegression','features':['ocv'],'alpha':None,'C':None}}]
        self.assertEqual(selection_counts(traces),
                         [('Ridge, n_features=1, alpha/C=1',2),('LinearRegression, n_features=1',1)])
        self.assertFalse(uses_eis(['ocv']))
        self.assertTrue(uses_eis(['ocv','mag@1000Hz']))

class ExportTests(TemporaryCase):
    def test_literal_strings_keep_their_value_and_never_become_formulas(self):
        # F70: the old guard stored an apostrophe inside the exported value.
        table=pd.DataFrame([{'text':'=SUM(A1)','other':'-note'},{'text':'@handle','other':'plain'}])
        path=self.folder/'guard.xlsx'
        write_workbook(path,{'01_guard':table})
        sheet=load_workbook(path)['01_guard']
        self.assertEqual([c.value for c in sheet[2]],['=SUM(A1)','-note'])
        self.assertEqual([c.value for c in sheet[3]],['@handle','plain'])
        self.assertTrue(all(cell.data_type!='f' for row in sheet for cell in row))
        self.assertTrue(sheet.cell(row=2,column=1).quotePrefix)

    def test_versions_record_platform_and_blas(self):
        # F74: cross-platform reproduction cannot be judged from library versions alone.
        recorded=versions()
        self.assertEqual(recorded['platform'],platform.platform())
        self.assertTrue(recorded['machine'])
        self.assertTrue(all(set(entry)<= {'user_api','internal_api','prefix','version','num_threads'}
                            for entry in recorded['blas']))

    def test_cc_fraction_reason_names_the_missing_values(self):
        # F45: 'at least 4 cells' was reported even when the column was entirely empty.
        cfg=small_config(); frame=make_demo(cfg,8); frame['cc_frac']=np.nan
        result,out=self.execute(cfg,frame)
        self.assertEqual(result['status'],'partial')
        self.assertIn('every value is missing',result['unavailable']['cc_fraction'])
        frame=make_demo(cfg,8); frame.loc[3:,'cc_frac']=np.nan
        result,_=self.execute(cfg,frame,'run_few')
        self.assertIn('only 3 of 8 cells have cc_frac',result['unavailable']['cc_fraction'])

    def test_deployment_fit_failure_keeps_the_evaluation_it_did_not_touch(self):
        # F07: metrics and unavailable['logistic'] used to contradict each other.
        cfg=small_config()
        original=ridge_capacity.select_model
        def deployment_fails(x,y,ids,local_cfg,kind,groups=None):
            if kind=='classification':
                raise EvaluationError('no valid inner split on the full cohort')
            return original(x,y,ids,local_cfg,kind,groups)
        with patch.object(export_for_origin,'select_model',deployment_fails):
            result,out=self.execute(cfg)
        self.assertEqual(result['status'],'partial')
        self.assertNotIn('logistic',result['unavailable'])
        self.assertIn('deployment fit is unavailable',result['unavailable']['logistic_deployment_model'])
        metrics=json.loads((out/'metrics.json').read_text(encoding='utf-8'))
        self.assertIn('logistic_nested_selection',[row['model'] for row in metrics['classification']])
        self.assertFalse((out/'models'/'logistic.joblib').exists())
        self.assertTrue((out/'models'/'capacity.joblib').exists())
        self.assertIn('전체 자료 배포 모델만 적합 불가',(out/'run_summary.md').read_text(encoding='utf-8'))

    def test_nonconverged_classification_leaves_the_capacity_results_on_disk(self):
        # F16: one warning used to fail the whole run and delete nothing but leave no results either.
        cfg=small_config()
        with patch.object(LogisticRegression,'fit',warn_on_fit(LogisticRegression.fit)):
            result,out=self.execute(cfg)
        self.assertEqual(result['status'],'partial')
        self.assertIn('did not converge',result['unavailable']['logistic'])
        self.assertTrue((out/'models'/'capacity.joblib').exists())
        self.assertIn('capacity_nested_selection',pd.read_csv(out/'regression_metrics.csv')['model'].tolist())

    def test_skipped_candidates_are_reported_without_changing_the_status(self):
        # F10: a missing 1 kHz column must be visible but must not turn a good run into 'partial'.
        cfg=small_config(frequency_grid_hz=[.1,1.,10.,10000.,100000.],feature_sets=['single_1e3','band_low'])
        result,out=self.execute(cfg)
        self.assertEqual(result['status'],'completed')
        metrics=json.loads((out/'metrics.json').read_text(encoding='utf-8'))
        self.assertEqual(metrics['skipped_candidates'],{'feature_sets':['single_1e3'],'baselines':['mag@1000Hz']})
        self.assertIn('## 제외된 후보·기준선',(out/'run_summary.md').read_text(encoding='utf-8'))

class ReportingTests(TemporaryCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg=small_config()
        cls.shared=TemporaryDirectory(prefix='ml_models_shared_')
        cls.out=Path(cls.shared.name)/'run'
        cls.out.mkdir()
        cls.result=run(make_demo(cls.cfg,8),cls.cfg,cls.out,[],{'data_kind':'synthetic_demo'})
        cls.metrics=json.loads((cls.out/'metrics.json').read_text(encoding='utf-8'))
        cls.summary=(cls.out/'run_summary.md').read_text(encoding='utf-8')

    @classmethod
    def tearDownClass(cls):
        cls.shared.cleanup()

    def rows(self,kind):
        return {row['model']:row for row in self.metrics[kind]}

    def test_baseline_rows_show_the_leave_one_out_artefacts(self):
        # F11: Q2 of the LOO mean is a constant and its threshold AUC is 0 whatever the data says.
        n=self.result['n_cells']
        self.assertAlmostEqual(self.rows('regression')['baseline_mean']['Q2'],1-(n/(n-1))**2)
        self.assertEqual(self.rows('classification')['baseline_mean_threshold']['AUC_pooled_OOF'],0.)
        self.assertIn('1-(n/(n-1))^2',self.summary)

    def test_classification_table_has_a_chance_level_row(self):
        # F24: accuracy of the fitted classifier is unreadable without the prior-only row.
        prior=self.rows('classification')['baseline_prior']
        self.assertEqual(prior['n_samples'],self.result['n_cells'])
        self.assertIsNotNone(prior['log_loss_prior_baseline'])
        exported=pd.read_csv(self.out/'classification_metrics.csv')
        self.assertIn('baseline_prior',exported['model'].tolist())
        workbook=load_workbook(self.out/'results_for_origin.xlsx',read_only=True)
        self.assertIn('baseline_prior',{row[0] for row in workbook['10_confusion_matrix'].iter_rows(min_row=2,values_only=True)})
        workbook.close()

    def test_summary_reports_metrics_selection_and_chance_level(self):
        # F17/F05: numbers and the selected model per fold were only in folds.json before.
        self.assertIn('## 지표 요약 (바깥 OOF)',self.summary)
        self.assertIn('## fold별 선택과 최종 적합',self.summary)
        self.assertIn('전체 자료 적합:',self.summary)
        self.assertIn('우연 수준이 정확히 0.5가 아니고',self.summary)
        self.assertIn('capacity Q2=',self.result['headline'])
        self.assertTrue(self.metrics['notes'])

    def test_workbook_columns_follow_the_documented_order(self):
        # F46: residuals lacked pred_cc_frac, coefficients mixed model kinds, metrics started with n_samples.
        wb=load_workbook(self.out/'results_for_origin.xlsx')
        self.assertEqual(len(wb.sheetnames),11)
        self.assertIn('pred_cc_frac',[c.value for c in wb['04_reg_residual'][1]])
        self.assertEqual([c.value for c in wb['02_reg_metrics'][1]][:2],['model','target'])
        self.assertIn('model_kind',[c.value for c in wb['03_reg_coefs'][1]])
        kinds={row[1] for row in wb['03_reg_coefs'].iter_rows(min_row=2,values_only=True)}
        self.assertTrue(kinds<= {'regression','classification'})
        self.assertIsNone([c.value for c in wb['05_clf_roc'][2]][3])
        wb.close()

class CommandTests(TemporaryCase):
    def test_cli_records_sorted_code_hashes_and_prints_the_headline(self):
        # F64 (glob order) and F17 (console numbers).
        cfg=small_config()
        features=self.folder/'features.csv'; make_demo(cfg,8).to_csv(features,index=False)
        config=self.folder/'config.json'; config.write_text(json.dumps(config_dict(cfg)),encoding='utf-8')
        out=self.folder/'cli'
        stdout=io.StringIO()
        with contextlib.redirect_stdout(stdout):
            code=main(['--features',str(features),'--config',str(config),'--synthetic-data','--output',str(out)])
        self.assertEqual(code,0)
        state=json.loads((out/'run.json').read_text(encoding='utf-8'))
        self.assertEqual(list(state['code_sha256']),sorted(state['code_sha256']))
        self.assertIn('capacity Q2=',stdout.getvalue())
        self.assertIn('baseline_*',stdout.getvalue())
        bundle=out/'models'/'capacity.joblib'
        self.assertTrue(bundle.exists())

    def test_saved_bundle_carries_selected_features_and_versions(self):
        # Training-side half of the predict checks: the bundle must say what it actually uses.
        cfg=small_config()
        _,out=self.execute(cfg)
        bundle=joblib.load(out/'models'/'capacity.joblib')
        self.assertEqual(bundle['selected_features'],list(bundle['pipeline'].named_steps['select'].columns))
        self.assertTrue(set(bundle['selected_features'])<= set(bundle['input_features']))
        self.assertEqual(bundle['versions']['python'],versions()['python'])

if __name__=='__main__':
    unittest.main(verbosity=2)
